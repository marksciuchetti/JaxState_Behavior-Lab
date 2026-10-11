#!/usr/bin/env python3
"""
make_excel_workbook.py
Builds SIMULATED_fea_analysis.xlsx from the SIMULATED CSV files with openpyxl, and writes
analysis_excel_walkthrough.md from the SAME cell registry, so that every cell reference in the
walkthrough is the address of the formula in the workbook (the walkthrough is generated, so it cannot
drift from the workbook).

The workbook mirrors analysis.R for students who work in Excel:

  README   what the workbook is, and that every number is SIMULATED
  Data     the FEA Metrics Export as imported (NA kept as text), plus two helper columns
           (Keep, Present)
  Summary  table by expression and interval built with COUNTIFS, AVERAGEIFS, and array formulas
           (the hand built PivotTable); Engagement as a raw mean
  H1       Joy Percent of time present, Commercial versus Baseline: one row per kept respondent
           (SUMIFS pulls the two values from Data), the differences, paired t test with
           T.TEST(...,2,1), hand computed t, d_z, CI, viewers with any Joy, exact McNemar, and a hand
           computed Wilcoxon signed rank summary (ranks, V, normal approximation) and sign test
  Bins     H2: Joy by bin per respondent, the proportion of viewers with Joy present per bin
           (COUNTIFS), a repeated measures ANOVA by hand (its F equals the mixed model F test in this
           balanced design), a Friedman test by hand, and two contrasts
  H3       Positive Valence Percent of time present (Commercial) and Ad_liking_1to7: CORREL (Pearson),
           Spearman by ranks, t, p, Fisher intervals
  Check    live Excel results next to the values analysis.R produced, with a status

Sheet names H1 and H3 look like cell addresses, so formulas that point at them quote the name,
for example 'H1'!L10. The walkthrough writes the same addresses without the quotes (H1!L10); type
the quotes in a formula.

All formulas are ordinary Excel formulas. The STDEV.S(IF(...)) and MEDIAN(IF(...)) formulas on the
Summary sheet are array formulas; Excel 365 evaluates them with Enter, older Excel needs
Ctrl+Shift+Enter (openpyxl writes them as array formulas already).

Run: python3 make_excel_workbook.py   (after make_simulated_data.py and analysis.R)
"""

import csv
import os
import re
import statistics as st

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.data_source import NumDataSource, NumRef
from openpyxl.chart.error_bar import ErrorBars
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.formula import ArrayFormula

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "SIMULATED_fea_analysis.xlsx")
MD_OUT = os.path.join(HERE, "analysis_excel_walkthrough.md")

NAVY = "1F3864"
GRAY = "8C8C8C"
ACCENT = "A6561A"
LIGHT = "F2F2F2"

hdr_font = Font(bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor=NAVY)
sub_fill = PatternFill("solid", fgColor=LIGHT)
bold = Font(bold=True)
note_font = Font(italic=True, color="555555")
code_font = Font(name="Consolas", size=9, color="555555")
thin = Side(style="thin", color="BFBFBF")
box = Border(top=thin, bottom=thin, left=thin, right=thin)

EXPRS = ["Joy", "Surprise", "Brow Furrow", "Positive Valence", "Negative Valence"]
INTERVALS = ["Baseline", "Commercial", "Recovery"] + [f"Bin_{j:02d}" for j in range(1, 7)]
BINS = INTERVALS[3:]
DQ_RULE = 80


def read_csv(name):
    with open(os.path.join(HERE, name), encoding="utf-8") as fh:
        raw = fh.readlines()
    banner = raw[0].strip()
    rows = list(csv.reader([ln for ln in raw if not ln.startswith("#")]))
    return banner, rows[0], rows[1:]


def to_cell(v):
    """Numbers become numbers; NA stays the text NA (as iMotions writes it)."""
    if v == "NA" or v == "":
        return v
    try:
        f = float(v)
        return int(f) if f.is_integer() and "." not in v else f
    except ValueError:
        return v


def style_header(ws, row=1, ncol=None, first_col=1):
    ncol = ncol or ws.max_column
    for c in range(first_col, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = box


def autowidth(ws, min_w=8, max_w=44):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            L = len(str(cell.value)) if not str(cell.value).startswith("=") else 12
            widths[cell.column_letter] = max(widths.get(cell.column_letter, 0), L)
    for col, w in widths.items():
        ws.column_dimensions[col].width = max(min_w, min(max_w, w + 2))


def show_formula(cell, formula):
    """Write the formula text into a cell as TEXT (not evaluated), without the _xlfn. prefixes."""
    cell.value = formula.replace("_xlfn.", "")
    cell.data_type = "s"
    cell.font = code_font


REG = {}


def reg(sheet, key, cell, formula, label=""):
    REG[(sheet, key)] = {"cell": cell, "formula": formula.replace("_xlfn.", ""), "label": label}


banner, header, rows = read_csv("SIMULATED_fea_metrics.csv")
SEED = re.search(r"seed (\d+)", banner).group(1)
idx = {h: i for i, h in enumerate(header)}
n_data = len(rows)
FIRST, LAST = 2, n_data + 1
COL = {h: get_column_letter(i + 1) for i, h in enumerate(header)}
HELP = ["Keep", "Present"]
HCOL = {h: get_column_letter(len(header) + 1 + i) for i, h in enumerate(HELP)}


def rng(col):
    return f"Data!${col}${FIRST}:${col}${LAST}"


R_RESP, R_ITYPE, R_INT, R_EXPR = (rng(COL["Respondent Name"]), rng(COL["Interval type"]),
                                  rng(COL["Interval label"]), rng(COL["Expression"]))
R_PCT, R_SEG, R_RAW, R_DQ = (rng(COL["Percent of time present"]), rng(COL["Number of segments"]),
                             rng(COL["Mean raw score (0-100)"]), rng(COL["Data quality (%)"]))
R_KEEP, R_PRES = rng(HCOL["Keep"]), rng(HCOL["Present"])

dq_of = {}
for r in rows:
    dq_of[r[idx["Respondent Name"]]] = float(r[idx["Data quality (%)"]])
kept = sorted(k for k, v in dq_of.items() if v >= DQ_RULE)
excluded = sorted(k for k, v in dq_of.items() if v < DQ_RULE)
N = len(kept)

wb = Workbook()

# ---------------------------------------------------------------- README
ws = wb.active
ws.title = "README"
readme = [
    "SIMULATED facial expression analysis workbook (teaching material)",
    "",
    f"Every number in this workbook is SIMULATED. It was generated by make_simulated_data.py (seed {SEED}) to resemble",
    "the FEA Metrics Export of the iMotions FEA thresholding R Notebook for one video commercial: a 60 s Baseline,",
    "a 60 s Commercial cut into six 10 s bins (Bin_01 to Bin_06), and a 15 s Recovery, recorded with the lab webcam",
    "and processed with Affectiva AFFDEX. No real viewer was recorded.",
    "",
    "Sheets:",
    f"  Data     FEA Metrics Export as imported (rows {FIRST} to {LAST}). NA is kept as text, as iMotions writes it.",
    f"           Helper columns {HCOL['Keep']} and {HCOL['Present']}: Keep (Data quality >= {DQ_RULE}), Present (1 when Percent of time present > 0).",
    "  Summary  Mean, SD, median, viewers with the expression present, and segments by expression and interval,",
    "           built with COUNTIFS, AVERAGEIFS, and array formulas. Engagement is a raw mean score.",
    "  H1       Joy, Commercial versus Baseline. One row per kept respondent; paired t test (T.TEST type 1), d_z,",
    "           viewers with any Joy, exact McNemar test, and a hand computed Wilcoxon signed rank summary.",
    "  Bins     H2. Joy by bin, the proportion of viewers with Joy present per bin (COUNTIFS), repeated measures",
    "           ANOVA by hand, Friedman test by hand, contrasts.",
    "  H3       Positive Valence during the Commercial and Ad_liking_1to7: Pearson (CORREL) and Spearman.",
    "  Check    Live Excel results next to the values analysis.R produced, with an OK or CHECK status.",
    "",
    "Design: ONE viewing by every respondent, in a fixed order. Baseline Joy is the comparison; there is no",
    "counterbalancing. Every H1 difference is Commercial MINUS Baseline.",
    "Column names are modeled on the manual's description of the FEA Metrics Export (percent of time present and number of",
    "segments per respondent and time interval). The manuals print no header row: verify the headers on the lab machine.",
    f"Exclusion rule (the lab's own rule, fixed before collection; the manuals state none): Data quality below {DQ_RULE} percent",
    f"(respondents {', '.join(excluded)}).",
    "Engagement is not thresholded; it is a raw mean score (0 to 100) per interval.",
    "A viewer has an expression 'present' in an interval when Percent of time present is above 0.",
    "",
    "Companion files: analysis_excel_walkthrough.md (step by step), analysis.R (same analysis in R),",
    "analysis_output.txt (its console output), results_text.md (model paragraphs).",
    "Sheets named H1 and H3 look like cell addresses, so formulas that point at them quote the name: 'H1'!L10.",
]
for i, line in enumerate(readme, start=1):
    ws.cell(row=i, column=1, value=line)
ws["A1"].font = Font(bold=True, size=14, color=NAVY)
ws.column_dimensions["A"].width = 122

# ---------------------------------------------------------------- Data
ws = wb.create_sheet("Data")
ws.append(header + HELP)
style_header(ws)
cK, cP = HCOL["Keep"], HCOL["Present"]
Q_, H_ = COL["Data quality (%)"], COL["Percent of time present"]
for i, r in enumerate(rows, start=FIRST):
    ws.append([to_cell(v) for v in r])
    ws[f"{cK}{i}"] = f'=IF({Q_}{i}>={DQ_RULE},"Yes","No")'
    ws[f"{cP}{i}"] = f'=IF(ISNUMBER({H_}{i}),IF({H_}{i}>0,1,0),"")'
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(header) + len(HELP))}{LAST}"
autowidth(ws, max_w=24)
for h in HELP:
    ws[f"{HCOL[h]}1"].fill = PatternFill("solid", fgColor=ACCENT)
reg("Data", "keep", f"{cK}2", f'=IF({Q_}2>={DQ_RULE},"Yes","No")', "Keep")
reg("Data", "present", f"{cP}2", f'=IF(ISNUMBER({H_}2),IF({H_}2>0,1,0),"")', "Present")

# ---------------------------------------------------------------- Summary
ws = wb.create_sheet("Summary")
ws["A1"] = "Descriptive table by expression and interval (kept respondents only). SIMULATED data."
ws["A1"].font = Font(bold=True, size=12, color=NAVY)
ws["A2"] = ("n = kept respondents with a row; Mean, SD, and Median are of Percent of time present; Viewers present = respondents whose "
            "Percent of time present is above 0; Mean segments = Number of segments. Engagement (second table) is a raw mean score.")
ws["A2"].font = note_font
cols = ["Expression", "Interval", "n", "Mean time present (%)", "SD (%)", "Median (%)", "Viewers present (n)",
        "Viewers present (%)", "Mean segments"]
ws.append([])
ws.append(cols)
style_header(ws, row=4, ncol=len(cols))
r = 5
SUM_FIRST = r
summary_rows = {}
for e in EXPRS:
    for iv in INTERVALS:
        ws.cell(row=r, column=1, value=e)
        ws.cell(row=r, column=2, value=iv)
        crit = f'{R_EXPR},$A{r},{R_INT},$B{r},{R_KEEP},"Yes"'
        f_n = f"=COUNTIFS({crit})"
        f_m = f"=AVERAGEIFS({R_PCT},{crit})"
        f_sd = f'=_xlfn.STDEV.S(IF(({R_EXPR}=$A{r})*({R_INT}=$B{r})*({R_KEEP}="Yes"),{R_PCT}))'
        f_md = f'=MEDIAN(IF(({R_EXPR}=$A{r})*({R_INT}=$B{r})*({R_KEEP}="Yes"),{R_PCT}))'
        f_np = f'=COUNTIFS({crit},{R_PCT},">0")'
        f_pp = f"=100*G{r}/C{r}"
        f_sg = f"=AVERAGEIFS({R_SEG},{crit})"
        ws.cell(row=r, column=3, value=f_n)
        ws.cell(row=r, column=4, value=f_m)
        ws.cell(row=r, column=5).value = ArrayFormula(f"E{r}", f_sd)
        ws.cell(row=r, column=6).value = ArrayFormula(f"F{r}", f_md)
        ws.cell(row=r, column=7, value=f_np)
        ws.cell(row=r, column=8, value=f_pp)
        ws.cell(row=r, column=9, value=f_sg)
        for c in range(1, 10):
            ws.cell(row=r, column=c).border = box
            if c >= 4:
                ws.cell(row=r, column=c).number_format = "0.00"
        if r == SUM_FIRST:
            for key, col, f in (("n", "C", f_n), ("mean", "D", f_m), ("sd", "E", f_sd), ("median", "F", f_md),
                                ("n_present", "G", f_np), ("pct_present", "H", f_pp), ("mean_seg", "I", f_sg)):
                reg("Summary", key, f"{col}{r}", f, cols[ord(col) - ord("A")])
        summary_rows[(e, iv)] = r
        r += 1
SUM_LAST = r - 1
# Engagement block
ENG_HDR = SUM_LAST + 3
ws.cell(row=ENG_HDR - 1, column=1, value="Engagement: raw mean score (0 to 100), kept respondents. Not thresholded.").font = bold
ecols = ["Expression", "Interval", "n", "Mean raw score", "SD", "Median"]
for c, h in enumerate(ecols, start=1):
    ws.cell(row=ENG_HDR, column=c, value=h)
style_header(ws, row=ENG_HDR, ncol=len(ecols))
r = ENG_HDR + 1
ENG_FIRST = r
for iv in INTERVALS:
    ws.cell(row=r, column=1, value="Engagement")
    ws.cell(row=r, column=2, value=iv)
    crit = f'{R_EXPR},$A{r},{R_INT},$B{r},{R_KEEP},"Yes"'
    f_n = f'=COUNTIFS({crit},{R_RAW},">=0")'
    f_m = f"=AVERAGEIFS({R_RAW},{crit})"
    f_sd = f'=_xlfn.STDEV.S(IF(({R_EXPR}=$A{r})*({R_INT}=$B{r})*({R_KEEP}="Yes"),{R_RAW}))'
    f_md = f'=MEDIAN(IF(({R_EXPR}=$A{r})*({R_INT}=$B{r})*({R_KEEP}="Yes"),{R_RAW}))'
    ws.cell(row=r, column=3, value=f_n)
    ws.cell(row=r, column=4, value=f_m)
    ws.cell(row=r, column=5).value = ArrayFormula(f"E{r}", f_sd)
    ws.cell(row=r, column=6).value = ArrayFormula(f"F{r}", f_md)
    for c in range(1, 7):
        ws.cell(row=r, column=c).border = box
        if c >= 4:
            ws.cell(row=r, column=c).number_format = "0.00"
    if r == ENG_FIRST:
        for key, col, f in (("e_n", "C", f_n), ("e_mean", "D", f_m), ("e_sd", "E", f_sd), ("e_median", "F", f_md)):
            reg("Summary", key, f"{col}{r}", f, ecols[ord(col) - ord("A")])
    summary_rows[("Engagement", iv)] = r
    r += 1
ENG_LAST = r - 1
ws[f"A{ENG_LAST + 2}"] = ("Array formulas (SD and Median columns): Excel 365 evaluates them with Enter; "
                           "Excel 2019 or older needs Ctrl+Shift+Enter after editing.")
ws[f"A{ENG_LAST + 2}"].font = note_font
autowidth(ws, max_w=22)
ws.column_dimensions["A"].width = 20
ws.column_dimensions["B"].width = 12
# chart table and column chart of mean Joy time percent by interval
ct = ENG_LAST + 6
ws.cell(row=ct - 1, column=1, value="Chart table: mean Joy Percent of time present by interval").font = bold
ws.cell(row=ct, column=1, value="Interval")
ws.cell(row=ct, column=2, value="Mean Joy (%)")
for i, iv in enumerate(INTERVALS, start=1):
    ws.cell(row=ct + i, column=1, value=iv)
    ws.cell(row=ct + i, column=2, value=f"=D{summary_rows[('Joy', iv)]}").number_format = "0.0"
chart = BarChart()
chart.type = "col"
chart.style = 10
chart.title = "Mean Joy Percent of time present by interval (SIMULATED)"
chart.y_axis.title = "Percent of interval time (%)"
chart.add_data(Reference(ws, min_col=2, min_row=ct, max_row=ct + len(INTERVALS)), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=1, min_row=ct + 1, max_row=ct + len(INTERVALS)))
chart.series[0].graphicalProperties.solidFill = NAVY
chart.legend = None
chart.height = 8
chart.width = 16
ws.add_chart(chart, f"D{ct - 1}")


# ---------------------------------------------------------------- shared block writer
class Block:
    """Writes label, value, formula text, and note columns for one sheet and records every cell in REG."""

    def __init__(self, ws, sheet, label_col, start_row):
        self.ws, self.sheet, self.lc = ws, sheet, label_col
        self.row = start_row
        self.ref = {}
        for k, name in enumerate(("Statistic", "Value", "Formula shown", "Note")):
            c = ws.cell(row=start_row - 1, column=label_col + k, value=name)
            c.font = hdr_font
            c.fill = hdr_fill

    def title(self, text):
        c = self.ws.cell(row=self.row, column=self.lc, value=text)
        c.font = bold
        c.fill = sub_fill
        self.row += 1

    def put(self, key, label, formula, note="", fmt="0.0000"):
        f = formula.format(**{k: f"{get_column_letter(self.lc + 1)}{v}" for k, v in self.ref.items()})
        i = self.row
        self.ws.cell(row=i, column=self.lc, value=label)
        vc = self.ws.cell(row=i, column=self.lc + 1, value=f)
        vc.number_format = fmt
        show_formula(self.ws.cell(row=i, column=self.lc + 2), f)
        if note:
            self.ws.cell(row=i, column=self.lc + 3, value=note).font = note_font
        self.ref[key] = i
        reg(self.sheet, key, f"{get_column_letter(self.lc + 1)}{i}", f, label)
        self.row += 1

    def blank(self):
        self.row += 1

    def addr(self, key):
        return f"{get_column_letter(self.lc + 1)}{self.ref[key]}"


def paired_block(blk, title, ra, rb, rd, unit, name_a, name_b, note_n="", note_diff=""):
    """Paired t test, d_z, and CIs. ra, rb, rd are the first, second, and difference ranges (same sheet)."""
    blk.title(title)
    blk.put("m_b", f"{name_a} mean ({unit})", f"=AVERAGE({ra})", f"AVERAGE of the {name_a} column")
    blk.put("s_b", f"{name_a} SD ({unit})", f"=_xlfn.STDEV.S({ra})", "STDEV.S = sample standard deviation")
    blk.put("m_p", f"{name_b} mean ({unit})", f"=AVERAGE({rb})")
    blk.put("s_p", f"{name_b} SD ({unit})", f"=_xlfn.STDEV.S({rb})")
    blk.put("n", "n pairs", f"=COUNT({rd})", note_n, fmt="0")
    blk.put("md", f"Mean difference, {name_b} minus {name_a} ({unit})", f"=AVERAGE({rd})", note_diff)
    blk.put("sdd", f"SD of the differences ({unit})", f"=_xlfn.STDEV.S({rd})", "the SD that matters in a paired test")
    blk.put("se", "SE of the mean difference", "={sdd}/SQRT({n})", "SD of differences / sqrt(n)")
    blk.put("t", "Paired t", "={md}/{se}", "t = mean difference / SE")
    blk.put("df", "df", "={n}-1", "n pairs minus 1", fmt="0")
    blk.put("p_tt", "p, two tailed (T.TEST type 1)", f"=_xlfn.T.TEST({ra},{rb},2,1)", "tails = 2, type = 1 means paired",
            fmt="0.000000")
    blk.put("p_hand", "p from the hand computed t", "=_xlfn.T.DIST.2T(ABS({t}),{df})", "should match the T.TEST cell", fmt="0.000000")
    blk.put("tcrit", "Critical t (95%)", "=_xlfn.T.INV.2T(0.05,{df})", "two tailed, alpha = .05")
    blk.put("ci_lo", "Mean difference 95% CI lower", "={md}-{tcrit}*{se}")
    blk.put("ci_hi", "Mean difference 95% CI upper", "={md}+{tcrit}*{se}")
    blk.put("dz", "Cohen's d_z", "={md}/{sdd}", "d_z = mean difference / SD of differences")
    blk.put("se_dz", "SE of d_z (approximation)", "=SQRT(1/{n}+{dz}^2/(2*{n}))", "sqrt(1/n + d_z^2 / (2 n))")
    blk.put("dz_lo", "d_z 95% CI lower", "={dz}-1.96*{se_dz}", "large sample approximation; R uses the noncentral t interval")
    blk.put("dz_hi", "d_z 95% CI upper", "={dz}+1.96*{se_dz}")
    blk.put("r", "Correlation of the two columns", f"=CORREL({ra},{rb})", "within respondent correlation")
    blk.blank()


# ---------------------------------------------------------------- H1
ws = wb.create_sheet("H1")
SH = "H1"
ws["A1"] = "H1. Joy Percent of time present, Commercial versus Baseline, kept respondents. SIMULATED data."
ws["A1"].font = Font(bold=True, size=12, color=NAVY)
ws["A2"] = ("One row per kept respondent. Columns C and D pull the two values from the Data sheet with SUMIFS (each respondent has exactly one "
            "Joy row per interval). Difference is Commercial minus Baseline, rounded to 2 decimals so tied differences tie.")
ws["A2"].font = note_font
h1_cols = ["Respondent Name", "Data quality (%)", "Baseline Joy (%)", "Commercial Joy (%)", "Difference (Commercial minus Baseline)",
           "Absolute difference", "Rank of absolute difference (zeros left out)", "Sign of difference", "Tie term (c^2 - 1)"]
ws.append([])
ws.append(h1_cols)
style_header(ws, row=4, ncol=len(h1_cols))
ws.row_dimensions[4].height = 48
P1 = 5
P2 = P1 + N - 1
RC, RD_, RDF = f"C{P1}:C{P2}", f"D{P1}:D{P2}", f"E{P1}:E{P2}"
RAB, RRK, RSG, RTI = f"F{P1}:F{P2}", f"G{P1}:G{P2}", f"H{P1}:H{P2}", f"I{P1}:I{P2}"
for k, rid in enumerate(kept):
    r = P1 + k
    ws.cell(row=r, column=1, value=rid)
    ws.cell(row=r, column=2, value=f"=INDEX({R_DQ},MATCH($A{r},{R_RESP},0))")
    ws.cell(row=r, column=3, value=f'=SUMIFS({R_PCT},{R_RESP},$A{r},{R_EXPR},"Joy",{R_INT},"Baseline")')
    ws.cell(row=r, column=4, value=f'=SUMIFS({R_PCT},{R_RESP},$A{r},{R_EXPR},"Joy",{R_INT},"Commercial")')
    ws.cell(row=r, column=5, value=f"=ROUND(D{r}-C{r},2)")
    ws.cell(row=r, column=6, value=f"=ABS(E{r})")
    ws.cell(row=r, column=7, value=f'=IF(E{r}=0,"",_xlfn.RANK.AVG(F{r},$F${P1}:$F${P2},1)-COUNTIF($E${P1}:$E${P2},0))')
    ws.cell(row=r, column=8, value=f"=SIGN(E{r})")
    ws.cell(row=r, column=9, value=f"=IF(E{r}=0,0,COUNTIF($F${P1}:$F${P2},F{r})^2-1)")
    for c in range(1, 10):
        ws.cell(row=r, column=c).border = box
        if c in (3, 4, 5, 6):
            ws.cell(row=r, column=c).number_format = "0.00"
regs = [("dq", "B", f"=INDEX({R_DQ},MATCH($A{P1},{R_RESP},0))", "Data quality (%)"),
        ("base", "C", ws.cell(row=P1, column=3).value, "Baseline Joy (%)"),
        ("comm", "D", ws.cell(row=P1, column=4).value, "Commercial Joy (%)"),
        ("diff", "E", f"=ROUND(D{P1}-C{P1},2)", "Difference (Commercial minus Baseline)"),
        ("abs", "F", f"=ABS(E{P1})", "Absolute difference"),
        ("rank", "G", ws.cell(row=P1, column=7).value, "Rank of absolute difference"),
        ("sign", "H", f"=SIGN(E{P1})", "Sign of difference"),
        ("tie", "I", ws.cell(row=P1, column=9).value, "Tie term (c^2 - 1)")]
for key, col, f, lab in regs:
    reg(SH, "row_" + key, f"{col}{P1}", f, lab)
blk = Block(ws, SH, 11, 5)          # label column K, value L, formula M, note N
paired_block(blk, "Paired t test and d_z: Joy Percent of time present", RC, RD_, RDF, "%", "Baseline", "Commercial",
             "COUNT of the difference column; must equal the check row below")
blk.title("Viewers with any Joy present (Percent of time present above 0)")
blk.put("any_b", "Viewers with Joy, Baseline", f'=COUNTIF({RC},">0")', "COUNTIF counts the numbers above 0", fmt="0")
blk.put("any_c", "Viewers with Joy, Commercial", f'=COUNTIF({RD_},">0")', fmt="0")
blk.put("pct_b", "Percent of viewers with Joy, Baseline", "=100*{any_b}/{n}", fmt="0.0")
blk.put("pct_c", "Percent of viewers with Joy, Commercial", "=100*{any_c}/{n}", fmt="0.0")
blk.put("disc_b", "Joy at Baseline only (Joy at Baseline, none in the Commercial)", f"=SUMPRODUCT(({RC}>0)*({RD_}=0))",
        "discordant pair type 1", fmt="0")
blk.put("disc_c", "Joy in the Commercial only", f"=SUMPRODUCT(({RC}=0)*({RD_}>0))", "discordant pair type 2", fmt="0")
blk.put("mcn_p", "Exact McNemar p, two tailed", "=MIN(1,2*_xlfn.BINOM.DIST(MIN({disc_b},{disc_c}),{disc_b}+{disc_c},0.5,TRUE))",
        "binomial test of the discordant pairs", fmt="0.000000")
blk.blank()
blk.title("Wilcoxon signed rank test by hand (normal approximation with tie and continuity correction, as in R)")
blk.put("n_pos", "Differences above 0 (count)", f'=COUNTIF({RDF},">0")', fmt="0")
blk.put("n_neg", "Differences below 0 (count)", f'=COUNTIF({RDF},"<0")', fmt="0")
blk.put("n_zero", "Differences equal to 0 (count)", f"=COUNTIF({RDF},0)", "zero differences are dropped", fmt="0")
blk.put("n_nz", "Nonzero differences", "={n_pos}+{n_neg}", fmt="0")
blk.put("w_pos", "V = sum of ranks of positive differences", f"=SUMIFS({RRK},{RSG},1)", "this is the V that R prints", fmt="0.0")
blk.put("w_neg", "Sum of ranks of negative differences", f"=SUMIFS({RRK},{RSG},-1)", fmt="0.0")
blk.put("w_chk", "V plus negative rank sum minus n(n+1)/2", "={w_pos}+{w_neg}-{n_nz}*({n_nz}+1)/2", "must be 0", fmt="0.0")
blk.put("v_mean", "Expected V under H0", "={n_nz}*({n_nz}+1)/4", fmt="0.00")
blk.put("tie3", "Sum of (t^3 - t) over tie groups", f"=SUM({RTI})", "from the tie term column", fmt="0")
blk.put("sigma", "SD of V under H0", "=SQRT({n_nz}*({n_nz}+1)*(2*{n_nz}+1)/24-{tie3}/48)", "includes the tie correction", fmt="0.0000")
blk.put("z", "z with continuity correction", "=({w_pos}-{v_mean}-0.5*SIGN({w_pos}-{v_mean}))/{sigma}", fmt="0.0000")
blk.put("p_w", "Wilcoxon p, two tailed", "=2*(1-_xlfn.NORM.S.DIST(ABS({z}),TRUE))", "matches wilcox.test(exact = FALSE)", fmt="0.000000")
blk.put("r_rb", "Matched pairs rank biserial correlation", "=({w_pos}-{w_neg})/({w_pos}+{w_neg})", "(positive ranks - negative ranks) / all ranks",
        fmt="0.0000")
blk.put("p_sign", "Sign test p, two tailed", "=MIN(1,2*_xlfn.BINOM.DIST(MIN({n_pos},{n_neg}),{n_pos}+{n_neg},0.5,TRUE))",
        "ignores the size of the differences", fmt="0.000000")
blk.blank()
blk.title("Checks")
blk.put("chk_n", "Kept respondents with a Joy Baseline row in Data",
        f'=COUNTIFS({R_EXPR},"Joy",{R_INT},"Baseline",{R_KEEP},"Yes")',
        "must equal n pairs; if not, the roster in column A is out of date", fmt="0")
blk.put("chk_dq", "Smallest Data quality in column B", f"=MIN(B{P1}:B{P2})", f"must be at least {DQ_RULE}", fmt="0.0")
blk.blank()
for col, w in (("A", 17), ("B", 12), ("C", 14), ("D", 16), ("E", 22), ("F", 14), ("G", 20), ("H", 12), ("I", 12), ("J", 3),
               ("K", 62), ("L", 12), ("M", 76), ("N", 60)):
    ws.column_dimensions[col].width = w
cr = blk.row + 1
ws.cell(row=cr, column=11, value="Chart table: mean Joy percent with 95% CI half width").font = bold
ws.cell(row=cr + 1, column=11, value="Interval")
ws.cell(row=cr + 1, column=12, value="Mean Joy (%)")
ws.cell(row=cr + 1, column=13, value="CI half width")
ws.cell(row=cr + 2, column=11, value="Baseline")
ws.cell(row=cr + 2, column=12, value=f"={blk.addr('m_b')}")
ws.cell(row=cr + 2, column=13, value=f"=_xlfn.CONFIDENCE.T(0.05,{blk.addr('s_b')},{blk.addr('n')})")
ws.cell(row=cr + 3, column=11, value="Commercial")
ws.cell(row=cr + 3, column=12, value=f"={blk.addr('m_p')}")
ws.cell(row=cr + 3, column=13, value=f"=_xlfn.CONFIDENCE.T(0.05,{blk.addr('s_p')},{blk.addr('n')})")
for rr_ in (cr + 2, cr + 3):
    ws.cell(row=rr_, column=12).number_format = "0.00"
    ws.cell(row=rr_, column=13).number_format = "0.00"
reg(SH, "chart_ci_b", f"M{cr + 2}", f"=_xlfn.CONFIDENCE.T(0.05,{blk.addr('s_b')},{blk.addr('n')})", "CI half width Baseline")
reg(SH, "chart_ci_p", f"M{cr + 3}", f"=_xlfn.CONFIDENCE.T(0.05,{blk.addr('s_p')},{blk.addr('n')})", "CI half width Commercial")
chart2 = BarChart()
chart2.type = "col"
chart2.style = 10
chart2.title = "Mean Joy by stimulus (SIMULATED)"
chart2.y_axis.title = "Percent of stimulus time (%)"
chart2.add_data(Reference(ws, min_col=12, min_row=cr + 1, max_row=cr + 3), titles_from_data=True)
chart2.set_categories(Reference(ws, min_col=11, min_row=cr + 2, max_row=cr + 3))
chart2.series[0].graphicalProperties.solidFill = NAVY
chart2.series[0].errBars = ErrorBars(errDir="y", errBarType="both", errValType="cust",
                                      plus=NumDataSource(numRef=NumRef(f=f"'H1'!$M${cr + 2}:$M${cr + 3}")),
                                      minus=NumDataSource(numRef=NumRef(f=f"'H1'!$M${cr + 2}:$M${cr + 3}")))
chart2.legend = None
chart2.height = 8
chart2.width = 12
ws.add_chart(chart2, f"N{cr}")
H1B = blk
H1_CHART_ROW = cr

# ---------------------------------------------------------------- Bins
ws = wb.create_sheet("Bins")
SH = "Bins"
ws["A1"] = "H2. Joy Percent of time present across the six 10 s bins, kept respondents. SIMULATED data."
ws["A1"].font = Font(bold=True, size=12, color=NAVY)
ws["A2"] = ("Columns B to G pull each respondent's Joy value for the bin with SUMIFS. The rows under the roster give the proportion of viewers with "
            "Joy present (COUNTIFS on Data), the mean, and the Friedman rank sums. The block on the right tests the bin factor.")
ws["A2"].font = note_font
b_cols = ["Respondent Name"] + BINS + ["Commercial Joy (%)", "Respondent mean over the bins"] + [f"Rank {b}" for b in BINS] + ["Tie term"]
ws.append([])
ws.append(b_cols)
style_header(ws, row=4, ncol=len(b_cols))
ws.row_dimensions[4].height = 32
B1 = 5
B2 = B1 + N - 1
for k, rid in enumerate(kept):
    r = B1 + k
    ws.cell(row=r, column=1, value=rid)
    for j, b in enumerate(BINS):
        ws.cell(row=r, column=2 + j,
                value=f'=SUMIFS({R_PCT},{R_RESP},$A{r},{R_EXPR},"Joy",{R_INT},{get_column_letter(2 + j)}$4)').number_format = "0.00"
    ws.cell(row=r, column=8, value=f'=SUMIFS({R_PCT},{R_RESP},$A{r},{R_EXPR},"Joy",{R_INT},"Commercial")').number_format = "0.00"
    ws.cell(row=r, column=9, value=f"=AVERAGE(B{r}:G{r})").number_format = "0.00"
    for j in range(6):
        ws.cell(row=r, column=10 + j, value=f"=_xlfn.RANK.AVG({get_column_letter(2 + j)}{r},$B{r}:$G{r},1)").number_format = "0.0"
    ws.cell(row=r, column=16, value=f"=SUMPRODUCT(COUNTIF($B{r}:$G{r},$B{r}:$G{r})^2-1)")
    for c in range(1, 17):
        ws.cell(row=r, column=c).border = box
reg(SH, "row_bin", f"B{B1}", ws.cell(row=B1, column=2).value, "Bin_01 Joy (%)")
reg(SH, "row_comm", f"H{B1}", ws.cell(row=B1, column=8).value, "Commercial Joy (%)")
reg(SH, "row_resp_mean", f"I{B1}", f"=AVERAGE(B{B1}:G{B1})", "Respondent mean over the bins")
reg(SH, "row_rank", f"J{B1}", ws.cell(row=B1, column=10).value, "Rank of Bin_01")
reg(SH, "row_tie", f"P{B1}", ws.cell(row=B1, column=16).value, "Tie term")
RN, RNJ, RPR, RPC, RMN, RSD, RCI, RRS = (B2 + 3 + i for i in range(8))
labels = {RN: "Viewers (n)", RNJ: "Viewers with Joy present", RPR: "Proportion with Joy present", RPC: "Percent of viewers with Joy present",
          RMN: "Mean Joy (%)", RSD: "SD (%)", RCI: "95% CI half width", RRS: "Friedman rank sum"}
for rw, lab in labels.items():
    ws.cell(row=rw, column=1, value=lab).font = bold
for j, b in enumerate(BINS):
    col = get_column_letter(2 + j)
    rkc = get_column_letter(10 + j)
    f_n = f'=COUNTIFS({R_EXPR},"Joy",{R_INT},{col}$4,{R_KEEP},"Yes")'
    f_nj = f'=COUNTIFS({R_EXPR},"Joy",{R_INT},{col}$4,{R_KEEP},"Yes",{R_PCT},">0")'
    f_pr = f"={col}{RNJ}/{col}{RN}"
    f_pc = f"=100*{col}{RPR}"
    f_mn = f"=AVERAGE({col}{B1}:{col}{B2})"
    f_sd = f"=_xlfn.STDEV.S({col}{B1}:{col}{B2})"
    f_ci = f"=_xlfn.CONFIDENCE.T(0.05,{col}{RSD},{col}{RN})"
    f_rs = f"=SUM({rkc}{B1}:{rkc}{B2})"
    for rw, f, fm in ((RN, f_n, "0"), (RNJ, f_nj, "0"), (RPR, f_pr, "0.000"), (RPC, f_pc, "0.0"), (RMN, f_mn, "0.00"),
                      (RSD, f_sd, "0.00"), (RCI, f_ci, "0.00"), (RRS, f_rs, "0.0")):
        c = ws.cell(row=rw, column=2 + j, value=f)
        c.number_format = fm
        c.border = box
    if j == 0:
        for key, rw, f in (("b_n", RN, f_n), ("b_nj", RNJ, f_nj), ("b_prop", RPR, f_pr), ("b_pct", RPC, f_pc), ("b_mean", RMN, f_mn),
                           ("b_sd", RSD, f_sd), ("b_ci", RCI, f_ci), ("b_rs", RRS, f_rs)):
            reg(SH, key, f"B{rw}", f, labels[rw])
blk = Block(ws, SH, 18, 5)          # label column R, value S, formula T, note U
blk.title("Repeated measures ANOVA by hand (its F equals the mixed model F test for this balanced design)")
blk.put("n_r", "Respondents (n)", f"=COUNT(I{B1}:I{B2})", fmt="0")
blk.put("k", "Bins (k)", "=COUNTA(B4:G4)", fmt="0")
blk.put("grand", "Grand mean Joy (%)", f"=AVERAGE(B{B1}:G{B2})", fmt="0.0000")
blk.put("ss_tot", "SS total", f"=DEVSQ(B{B1}:G{B2})", "sum of squared deviations from the grand mean", fmt="0.00")
blk.put("ss_bin", "SS bins", f"={{n_r}}*DEVSQ(B{RMN}:G{RMN})", "n x squared deviations of the six bin means", fmt="0.00")
blk.put("ss_resp", "SS respondents", f"={{k}}*DEVSQ(I{B1}:I{B2})", "k x squared deviations of the respondent means", fmt="0.00")
blk.put("ss_err", "SS error", "={ss_tot}-{ss_bin}-{ss_resp}", fmt="0.00")
blk.put("df_bin", "df bins", "={k}-1", fmt="0")
blk.put("df_err", "df error", "=({n_r}-1)*({k}-1)", "equals the Satterthwaite df of the mixed model here", fmt="0")
blk.put("ms_bin", "MS bins", "={ss_bin}/{df_bin}", fmt="0.0000")
blk.put("ms_err", "MS error", "={ss_err}/{df_err}", "equals the residual variance of the mixed model", fmt="0.0000")
blk.put("F", "F", "={ms_bin}/{ms_err}", fmt="0.0000")
blk.put("p_F", "p (F.DIST.RT)", "=_xlfn.F.DIST.RT({F},{df_bin},{df_err})", fmt="0.000000")
blk.put("var_resp", "Respondent variance (ANOVA estimate)", "=((({ss_resp})/({n_r}-1))-{ms_err})/{k}", "equals the random intercept variance of the model",
        fmt="0.0000")
blk.put("icc", "ICC (adjusted)", "={var_resp}/({var_resp}+{ms_err})", "share of the remaining variance that lies between viewers", fmt="0.0000")
blk.blank()
blk.title("Friedman test by hand (nonparametric check, ties corrected as in R)")
blk.put("fr_num", "12 x sum of squared deviations of the rank sums",
        f"=12*SUMPRODUCT((B{RRS}:G{RRS}-{{n_r}}*({{k}}+1)/2)^2)", "rank sums are in the Friedman rank sum row", fmt="0.00")
blk.put("fr_den", "Denominator with tie correction", f"={{n_r}}*{{k}}*({{k}}+1)-SUM(P{B1}:P{B2})/({{k}}-1)", "n k (k + 1) minus the tie sum / (k - 1)",
        fmt="0.00")
blk.put("fr_chi", "Friedman chi-squared", "={fr_num}/{fr_den}", fmt="0.0000")
blk.put("fr_p", "p (CHISQ.DIST.RT)", "=_xlfn.CHISQ.DIST.RT({fr_chi},{k}-1)", fmt="0.000000")
blk.put("fr_w", "Kendall's W", "={fr_chi}/({n_r}*({k}-1))", "effect size for the Friedman test: chi-squared / (n (k - 1))", fmt="0.0000")
blk.blank()
blk.title("Contrasts: a late bin against the mean of Bin_01 to Bin_04 (no multiplicity adjustment here)")
blk.put("c5_est", "Bin_05 minus mean(Bin_01..Bin_04) (percentage points)", f"=F{RMN}-AVERAGE(B{RMN}:E{RMN})", "difference of the bin means", fmt="0.0000")
blk.put("c5_se", "SE of the contrast", "=SQRT({ms_err}*1.25/{n_r})", "sqrt(MS error x sum of squared weights / n); the weights are 1 and four of -1/4, sum of squares 1.25",
        fmt="0.0000")
blk.put("c5_t", "t", "={c5_est}/{c5_se}", fmt="0.0000")
blk.put("c5_p", "p, two tailed", "=_xlfn.T.DIST.2T(ABS({c5_t}),{df_err})", fmt="0.000000")
blk.put("c6_est", "Bin_06 minus mean(Bin_01..Bin_04) (percentage points)", f"=G{RMN}-AVERAGE(B{RMN}:E{RMN})", fmt="0.0000")
blk.put("c6_se", "SE of the contrast", "=SQRT({ms_err}*1.25/{n_r})", fmt="0.0000")
blk.put("c6_t", "t", "={c6_est}/{c6_se}", fmt="0.0000")
blk.put("c6_p", "p, two tailed", "=_xlfn.T.DIST.2T(ABS({c6_t}),{df_err})", fmt="0.000000")
blk.blank()
blk.title("Checks")
blk.put("chk_comm", "Largest gap between Commercial Joy and the mean of the six bins (percentage points)",
        f"=SUMPRODUCT(MAX(ABS(H{B1}:H{B2}-I{B1}:I{B2})))", "the bins tile the commercial; any gap is rounding only", fmt="0.0000")
blk.put("chk_rank", "Rank sums add up to n k (k + 1) / 2", f"=SUM(B{RRS}:G{RRS})-{{n_r}}*{{k}}*({{k}}+1)/2", "must be 0", fmt="0.0")
blk.blank()
for col, w in (("A", 34), ("B", 11), ("C", 11), ("D", 11), ("E", 11), ("F", 11), ("G", 11), ("H", 13), ("I", 14), ("J", 9), ("K", 9),
               ("L", 9), ("M", 9), ("N", 9), ("O", 9), ("P", 9), ("Q", 3), ("R", 70), ("S", 12), ("T", 70), ("U", 70)):
    ws.column_dimensions[col].width = w
BN = blk

# ---------------------------------------------------------------- H3
sb, sh, srows = read_csv("SIMULATED_fea_survey.csv")
ws = wb.create_sheet("H3")
SH = "H3"
ws["A1"] = "H3. Positive Valence Percent of time present (Commercial) and Ad_liking_1to7. SIMULATED data."
ws["A1"].font = Font(bold=True, size=12, color=NAVY)
ws.append([])
sv_cols = sh + ["Keep", "Positive Valence during the Commercial (%)", "Rank of Ad_liking_1to7", "Rank of Positive Valence (%)"]
ws.append(sv_cols)
style_header(ws, row=3, ncol=len(sv_cols))
ws.row_dimensions[3].height = 48
sidx = {h: i for i, h in enumerate(sh)}
srows.sort(key=lambda rw: (rw[sidx["Respondent Name"]] in excluded, rw[sidx["Respondent Name"]]))
S1 = 4
S2 = S1 + N - 1
S_last_all = S1 + len(srows) - 1
for k, rw in enumerate(srows):
    r = S1 + k
    for c, v in enumerate(rw, start=1):
        ws.cell(row=r, column=c, value=to_cell(v))
    ws.cell(row=r, column=7, value=f'=IF(INDEX({R_DQ},MATCH($A{r},{R_RESP},0))>={DQ_RULE},"Yes","No")')
    ws.cell(row=r, column=8, value=f'=SUMIFS({R_PCT},{R_RESP},$A{r},{R_EXPR},"Positive Valence",{R_INT},"Commercial")').number_format = "0.00"
    if k < N:
        ws.cell(row=r, column=9, value=f"=_xlfn.RANK.AVG(D{r},$D${S1}:$D${S2},1)")
        ws.cell(row=r, column=10, value=f"=_xlfn.RANK.AVG(H{r},$H${S1}:$H${S2},1)")
for key, col, f, lab in (("keep", "G", ws.cell(row=S1, column=7).value, "Keep"),
                         ("pv", "H", ws.cell(row=S1, column=8).value, sv_cols[7]),
                         ("rk_like", "I", f"=_xlfn.RANK.AVG(D{S1},$D${S1}:$D${S2},1)", sv_cols[8]),
                         ("rk_pv", "J", f"=_xlfn.RANK.AVG(H{S1},$H${S1}:$H${S2},1)", sv_cols[9])):
    reg(SH, "row_" + key, f"{col}{S1}", f, lab)
ws.cell(row=S_last_all + 2, column=1, value=(
    f"Rows {S1} to {S2} are the {N} kept respondents; the excluded respondents ({', '.join(excluded)}) are listed last "
    f"(rows {S2 + 1} to {S_last_all}) and left out of every test range and rank range.")).font = note_font
RPV, RLK, RRL, RRP = f"H{S1}:H{S2}", f"D{S1}:D{S2}", f"I{S1}:I{S2}", f"J{S1}:J{S2}"
blk = Block(ws, SH, 12, 4)          # label column L, value M, formula N, note O
blk.title("Descriptives (kept respondents)")
blk.put("n", "n (kept respondents)", f"=COUNT({RPV})", fmt="0")
blk.put("m_pv", "Positive Valence mean (%)", f"=AVERAGE({RPV})")
blk.put("s_pv", "Positive Valence SD (%)", f"=_xlfn.STDEV.S({RPV})")
blk.put("m_lk", "Ad_liking_1to7 mean", f"=AVERAGE({RLK})")
blk.put("s_lk", "Ad_liking_1to7 SD", f"=_xlfn.STDEV.S({RLK})")
blk.blank()
blk.title("Pearson correlation")
blk.put("r", "Pearson r (CORREL)", f"=CORREL({RPV},{RLK})", "Pearson r of the two columns")
blk.put("t_r", "t for r", "={r}*SQRT(({n}-2)/(1-{r}^2))", "t = r sqrt((n - 2) / (1 - r^2))")
blk.put("df_r", "df", "={n}-2", fmt="0")
blk.put("p_r", "p, two tailed", "=_xlfn.T.DIST.2T(ABS({t_r}),{df_r})", fmt="0.000000")
blk.put("r_lo", "r 95% CI lower (Fisher z)", "=TANH(ATANH({r})-_xlfn.NORM.S.INV(0.975)/SQRT({n}-3))")
blk.put("r_hi", "r 95% CI upper (Fisher z)", "=TANH(ATANH({r})+_xlfn.NORM.S.INV(0.975)/SQRT({n}-3))")
blk.put("slope", "Least squares slope (percentage points per liking point)", f"=SLOPE({RPV},{RLK})")
blk.put("icpt", "Least squares intercept (%)", f"=INTERCEPT({RPV},{RLK})")
blk.put("rsq", "R squared", f"=RSQ({RPV},{RLK})")
blk.blank()
blk.title("Spearman correlation (Pearson r of the average ranks)")
blk.put("rho", "Spearman rho", f"=CORREL({RRL},{RRP})", "RANK.AVG gives tied values the average rank, as R does")
blk.put("t_rho", "t for rho", "={rho}*SQRT(({n}-2)/(1-{rho}^2))", "the t approximation that R uses when there are ties")
blk.put("p_rho", "p, two tailed", "=_xlfn.T.DIST.2T(ABS({t_rho}),{n}-2)", fmt="0.000000")
blk.put("rho_lo", "rho 95% CI lower (Bonett and Wright)", "=TANH(ATANH({rho})-_xlfn.NORM.S.INV(0.975)*SQRT((1+{rho}^2/2)/({n}-3)))")
blk.put("rho_hi", "rho 95% CI upper (Bonett and Wright)", "=TANH(ATANH({rho})+_xlfn.NORM.S.INV(0.975)*SQRT((1+{rho}^2/2)/({n}-3)))")
blk.blank()
blk.title("Checks")
blk.put("chk_keep", "Kept respondents in the Keep column", f'=COUNTIF(G{S1}:G{S_last_all},"Yes")', "must equal n", fmt="0")
blk.put("chk_dq", "Rows listed in the roster flagged Yes", f'=COUNTIF(G{S1}:G{S2},"Yes")', "must equal n; the excluded rows sit below the roster", fmt="0")
blk.blank()
for col, w in (("A", 17), ("B", 9), ("C", 7), ("D", 14), ("E", 16), ("F", 14), ("G", 8), ("H", 20), ("I", 14), ("J", 16), ("K", 3),
               ("L", 56), ("M", 12), ("N", 80), ("O", 60)):
    ws.column_dimensions[col].width = w
H3B = blk

# ---------------------------------------------------------------- Check
ws = wb.create_sheet("Check")
ws["A1"] = "Live Excel results next to the values analysis.R produced (SIMULATED data). Status is OK when they agree within the tolerance."
ws["A1"].font = Font(bold=True, size=12, color=NAVY)
ws.append([])
ck_cols = ["Quantity", "Excel (live formula)", "analysis.R", "Difference", "Tolerance", "Status"]
ws.append(ck_cols)
style_header(ws, row=3, ncol=len(ck_cols))
_, th, trows = read_csv("SIMULATED_test_results.csv")
T = {(r_[0], r_[1], r_[2]): dict(zip(th, r_)) for r_ in trows}


def tv(h, a, q, col="estimate"):
    return float(T[(h, a, q)][col])


ck_rows = []


def add_check(label, sheet, blkobj, key, r_value, tol, note=""):
    ck_rows.append((label, f"='{sheet}'!{blkobj.addr(key)}", float(r_value), tol, note))


def ptol(p):
    return 5e-4 * float(p) + 1e-9


H1T = "H1 threshold 50"
add_check("H1 Joy: mean Baseline (%)", "H1", H1B, "m_b", tv(H1T, "Descriptive", "Baseline mean Percent of time present"), 1e-4)
add_check("H1 Joy: mean Commercial (%)", "H1", H1B, "m_p", tv(H1T, "Descriptive", "Commercial mean Percent of time present"), 1e-4)
add_check("H1 Joy: n pairs", "H1", H1B, "n", T[(H1T, "Paired t test", "Mean difference (Commercial minus Baseline)")]["n"], 1e-9)
q = "Mean difference (Commercial minus Baseline)"
add_check("H1 Joy: mean difference", "H1", H1B, "md", tv(H1T, "Paired t test", q), 1e-4)
add_check("H1 Joy: mean difference CI lower", "H1", H1B, "ci_lo", tv(H1T, "Paired t test", q, "ci_low"), 1e-4)
add_check("H1 Joy: mean difference CI upper", "H1", H1B, "ci_hi", tv(H1T, "Paired t test", q, "ci_high"), 1e-4)
add_check("H1 Joy: paired t", "H1", H1B, "t", tv(H1T, "Paired t test", q, "statistic"), 1e-4)
add_check("H1 Joy: df", "H1", H1B, "df", tv(H1T, "Paired t test", q, "df"), 1e-9)
add_check("H1 Joy: p (T.TEST type 1)", "H1", H1B, "p_tt", tv(H1T, "Paired t test", q, "p"), ptol(tv(H1T, "Paired t test", q, "p")))
add_check("H1 Joy: Cohen's d_z", "H1", H1B, "dz", tv(H1T, "Paired t test", "Cohen's d_z"), 1e-4)
add_check("H1 Joy: d_z CI lower (method differs)", "H1", H1B, "dz_lo", tv(H1T, "Paired t test", "Cohen's d_z", "ci_low"), 0.10,
          "Excel: large sample approximation; R: noncentral t")
add_check("H1 Joy: d_z CI upper (method differs)", "H1", H1B, "dz_hi", tv(H1T, "Paired t test", "Cohen's d_z", "ci_high"), 0.10,
          "Excel: large sample approximation; R: noncentral t")
add_check("H1 Joy: viewers with Joy, Baseline", "H1", H1B, "any_b", tv(H1T, "Descriptive", "Viewers with any Joy, Baseline"), 1e-9)
add_check("H1 Joy: viewers with Joy, Commercial", "H1", H1B, "any_c", tv(H1T, "Descriptive", "Viewers with any Joy, Commercial"), 1e-9)
mq = "Joy present only at Baseline versus only during the Commercial"
if T[(H1T, "Exact McNemar test", mq)]["p"] not in ("", "NA"):
    add_check("H1 Joy: exact McNemar p", "H1", H1B, "mcn_p", tv(H1T, "Exact McNemar test", mq, "p"), ptol(tv(H1T, "Exact McNemar test", mq, "p")))
add_check("H1 Joy: Wilcoxon V", "H1", H1B, "w_pos", tv(H1T, "Wilcoxon signed rank (primary)", "V (sum of positive ranks)", "statistic"), 1e-6)
add_check("H1 Joy: Wilcoxon p", "H1", H1B, "p_w", tv(H1T, "Wilcoxon signed rank (primary)", "V (sum of positive ranks)", "p"),
          ptol(tv(H1T, "Wilcoxon signed rank (primary)", "V (sum of positive ranks)", "p")))
add_check("H1 Joy: rank biserial correlation", "H1", H1B, "r_rb", tv(H1T, "Wilcoxon signed rank (primary)", "Matched pairs rank biserial correlation"), 1e-4)
hb = "H2"
fq = "Bin factor F test"
add_check("H2 Joy by bin: F", "Bins", BN, "F", tv(hb, "Mixed model pct ~ bin + (1 | respondent)", fq, "statistic"), 1e-4)
add_check("H2 Joy by bin: denominator df", "Bins", BN, "df_err", tv(hb, "Mixed model pct ~ bin + (1 | respondent)", fq, "df"), 1e-6)
add_check("H2 Joy by bin: p", "Bins", BN, "p_F", tv(hb, "Mixed model pct ~ bin + (1 | respondent)", fq, "p"),
          ptol(tv(hb, "Mixed model pct ~ bin + (1 | respondent)", fq, "p")))
add_check("H2 Joy by bin: ICC adjusted", "Bins", BN, "icc", tv(hb, "Mixed model pct ~ bin + (1 | respondent)", "ICC adjusted"), 1e-4)
add_check("H2 Friedman chi-squared", "Bins", BN, "fr_chi", tv(hb, "Friedman test", "Chi-squared", "statistic"), 1e-4)
add_check("H2 Friedman p", "Bins", BN, "fr_p", tv(hb, "Friedman test", "Chi-squared", "p"), ptol(tv(hb, "Friedman test", "Chi-squared", "p")))
add_check("H2 Kendall's W", "Bins", BN, "fr_w", tv(hb, "Friedman test", "Kendall's W"), 1e-4)
c5q, c6q = "Bin_05 minus mean(Bin_01..Bin_04)", "Bin_06 minus mean(Bin_01..Bin_04)"
add_check("H2 contrast Bin_05 minus Bin_01..04: estimate", "Bins", BN, "c5_est", tv(hb, "Contrast (Satterthwaite)", c5q), 1e-4)
add_check("H2 contrast Bin_05 minus Bin_01..04: t", "Bins", BN, "c5_t", tv(hb, "Contrast (Satterthwaite)", c5q, "statistic"), 1e-3)
add_check("H2 contrast Bin_06 minus Bin_01..04: estimate", "Bins", BN, "c6_est", tv(hb, "Contrast (Satterthwaite)", c6q), 1e-4)
add_check("H2 contrast Bin_06 minus Bin_01..04: t", "Bins", BN, "c6_t", tv(hb, "Contrast (Satterthwaite)", c6q, "statistic"), 1e-3)
for j, b in enumerate(BINS):
    col = get_column_letter(2 + j)
    ck_rows.append((f"H2 {b}: proportion of viewers with Joy present", f"=Bins!{col}{RPR}",
                    tv(hb, "Proportion of viewers with Joy present", f"{b}: proportion with Joy present"), 1e-6, ""))
for j, b in enumerate(BINS):
    col = get_column_letter(2 + j)
    ck_rows.append((f"H2 {b}: mean Joy (%)", f"=Bins!{col}{RMN}", tv(hb, "Mean Percent of time present", f"{b}: mean Joy percent"), 1e-4, ""))
rq = "r (Positive Valence % with Ad_liking_1to7)"
add_check("H3 Pearson r", "H3", H3B, "r", tv("H3", "Pearson", rq), 1e-4)
add_check("H3 Pearson: t", "H3", H3B, "t_r", tv("H3", "Pearson", rq, "statistic"), 1e-3)
add_check("H3 Pearson: p", "H3", H3B, "p_r", tv("H3", "Pearson", rq, "p"), ptol(tv("H3", "Pearson", rq, "p")))
add_check("H3 Pearson: CI lower", "H3", H3B, "r_lo", tv("H3", "Pearson", rq, "ci_low"), 1e-4)
add_check("H3 Pearson: CI upper", "H3", H3B, "r_hi", tv("H3", "Pearson", rq, "ci_high"), 1e-4)
sq = "rho (Positive Valence % with Ad_liking_1to7)"
add_check("H3 Spearman rho", "H3", H3B, "rho", tv("H3", "Spearman", sq), 1e-4)
add_check("H3 Spearman: t", "H3", H3B, "t_rho", tv("H3", "Spearman", sq, "statistic"), 1e-3)
add_check("H3 Spearman: p", "H3", H3B, "p_rho", tv("H3", "Spearman", sq, "p"), ptol(tv("H3", "Spearman", sq, "p")))
add_check("H3 Spearman: CI lower (Bonett and Wright)", "H3", H3B, "rho_lo", tv("H3", "Spearman", sq, "ci_low"), 1e-4)
add_check("H3 Spearman: CI upper (Bonett and Wright)", "H3", H3B, "rho_hi", tv("H3", "Spearman", sq, "ci_high"), 1e-4)
add_check("H3 least squares slope", "H3", H3B, "slope", tv("H3", "Least squares line", "Slope (percentage points per liking point)"), 1e-4)
CK1 = 4
for k, (lab, f, rv, tol, note) in enumerate(ck_rows):
    r = CK1 + k
    ws.cell(row=r, column=1, value=lab)
    ws.cell(row=r, column=2, value=f).number_format = "0.000000"
    ws.cell(row=r, column=3, value=rv).number_format = "0.000000"
    ws.cell(row=r, column=4, value=f"=B{r}-C{r}").number_format = "0.000000"
    ws.cell(row=r, column=5, value=tol).number_format = "0.000000"
    ws.cell(row=r, column=6, value=f'=IF(ABS(D{r})<=E{r},"OK","CHECK")')
    if note:
        ws.cell(row=r, column=7, value=note).font = note_font
CK2 = CK1 + len(ck_rows) - 1
ws.cell(row=CK2 + 2, column=1, value="Rows with status CHECK").font = bold
ws.cell(row=CK2 + 2, column=2, value=f'=COUNTIF(F{CK1}:F{CK2},"CHECK")')
reg("Check", "n_check", f"B{CK2 + 2}", f'=COUNTIF(F{CK1}:F{CK2},"CHECK")', "Rows with status CHECK")
reg("Check", "status_row", f"F{CK1}", f'=IF(ABS(D{CK1})<=E{CK1},"OK","CHECK")', "Status")
ws.cell(row=CK2 + 3, column=1, value=("Not reproducible in Excel: the mixed model's random intercept estimates and Satterthwaite output (the F test above is its "
                                      "balanced design equivalent), the Hodges-Lehmann estimate, the noncentral t interval for d_z, the exact "
                                      "Clopper-Pearson intervals for the proportions, the square root model, and the threshold 25 and 75 sensitivity "
                                      "runs. See analysis_output.txt.")).font = note_font
r = CK2 + 6
ws.cell(row=r - 1, column=1, value="analysis.R: tests (SIMULATED_test_results.csv)").font = bold
for c, h in enumerate(th, start=1):
    ws.cell(row=r, column=c, value=h)
style_header(ws, row=r, ncol=len(th))
for rw in trows:
    r += 1
    for c, v in enumerate(rw, start=1):
        ws.cell(row=r, column=c, value=to_cell(v))
r += 3
_, dh, drows = read_csv("SIMULATED_descriptives_by_interval.csv")
ws.cell(row=r - 1, column=1, value="analysis.R: descriptives (SIMULATED_descriptives_by_interval.csv)").font = bold
for c, h in enumerate(dh, start=1):
    ws.cell(row=r, column=c, value=h)
style_header(ws, row=r, ncol=len(dh))
for rw in drows:
    r += 1
    for c, v in enumerate(rw, start=1):
        ws.cell(row=r, column=c, value=to_cell(v))
autowidth(ws, max_w=34)
ws.column_dimensions["A"].width = 62
for col in "BCDEF":
    ws.column_dimensions[col].width = 20

wb.calculation = CalcProperties(fullCalcOnLoad=True)    # Excel computes every formula when the file opens
wb.save(OUT)
print("wrote", OUT)

# Recalculate a scratch copy with LibreOffice (when installed) so the walkthrough can quote the value of every
# formula it lists. Without LibreOffice the walkthrough leaves those result cells empty; the workbook is unaffected.
import shutil
import subprocess
import tempfile
from openpyxl import load_workbook


def recalculated():
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return None
    tmp = tempfile.mkdtemp()
    try:
        subprocess.run([exe, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, OUT],
                       capture_output=True, timeout=240, check=True)
        return load_workbook(os.path.join(tmp, os.path.basename(OUT)), data_only=True)
    except (subprocess.SubprocessError, OSError):
        return None


RC_WB = recalculated()


def xv(sheet, key, nd=2, across=False):
    """Value of a registered formula cell after recalculation; across=True returns the six bin cells (B to G) of that row."""
    if RC_WB is None:
        return ""
    cell = REG[(sheet, key)]["cell"]
    wsx = RC_WB[sheet]
    if across:
        rown = int(re.sub(r"[A-Z]+", "", cell))
        vals = [wsx.cell(row=rown, column=2 + j).value for j in range(6)]
        return ", ".join(f"{float(v):.{nd}f}" for v in vals)
    v = wsx[cell].value
    return "" if v is None else f"{float(v):.{nd}f}"

# =============================================================================
# The walkthrough, generated from REG so that every cell reference is the real one
# =============================================================================


def cl(sheet, key):
    return REG[(sheet, key)]["cell"]


def fx(sheet, key):
    return REG[(sheet, key)]["formula"]


def row(label, sheet, key, result=""):
    return f"| {label} | `{sheet}!{cl(sheet, key)}` | `{fx(sheet, key)}` | {result} |"


def row3(label, sheet, key):
    return f"| {label} | `{sheet}!{cl(sheet, key)}` | `{fx(sheet, key)}` |"


def rcell(sheet, key):
    return f"{sheet}!{cl(sheet, key)}"


def fnum(x, nd=2):
    return f"{float(x):.{nd}f}"


def pnum(p):
    p = float(p)
    return f"{p:.6f}" if p >= 1e-4 else f"{p:.2e}"


A = lambda a, q, col="estimate": T[(H1T, a, q)][col]   # H1 values at threshold 50
h1d = T[(H1T, "Paired t test", q)]
m_b, m_c = tv(H1T, "Descriptive", "Baseline mean Percent of time present"), tv(H1T, "Descriptive", "Commercial mean Percent of time present")
any_b, any_c = int(tv(H1T, "Descriptive", "Viewers with any Joy, Baseline")), int(tv(H1T, "Descriptive", "Viewers with any Joy, Commercial"))
wxq = "V (sum of positive ranks)"
wV, wp = tv(H1T, "Wilcoxon signed rank (primary)", wxq, "statistic"), tv(H1T, "Wilcoxon signed rank (primary)", wxq, "p")
wn = int(tv(H1T, "Wilcoxon signed rank (primary)", wxq, "n"))
r_rb = tv(H1T, "Wilcoxon signed rank (primary)", "Matched pairs rank biserial correlation")
pcs = {}
for iv in ("Baseline", "Commercial"):
    pcs[iv] = [float(r_[idx["Percent of time present"]]) for r_ in rows
               if r_[idx["Expression"]] == "Joy" and r_[idx["Interval label"]] == iv and r_[idx["Respondent Name"]] in kept]
sd_b, sd_c = st.stdev(pcs["Baseline"]), st.stdev(pcs["Commercial"])
diffs = [round(c - b, 2) for b, c in zip(pcs["Baseline"], pcs["Commercial"])]
sd_diff = st.stdev(diffs)
n_pos, n_neg, n_zero = sum(d > 0 for d in diffs), sum(d < 0 for d in diffs), sum(d == 0 for d in diffs)
se_dz = (1 / N + (st.mean(diffs) / sd_diff) ** 2 / (2 * N)) ** 0.5
dz = st.mean(diffs) / sd_diff
tcrit = (float(h1d["ci_high"]) - float(h1d["ci_low"])) / 2 / (sd_diff / N ** 0.5)
d_row = {(r_[0], r_[1], r_[2]): dict(zip(dh, r_)) for r_ in drows}   # descriptives keyed (expression, type, interval)
dd = lambda e, iv: next(v for k, v in d_row.items() if k[0] == e and k[2] == iv)
Ffac = tv(hb, "Mixed model pct ~ bin + (1 | respondent)", fq, "statistic")
Fp = tv(hb, "Mixed model pct ~ bin + (1 | respondent)", fq, "p")
Fdf = tv(hb, "Mixed model pct ~ bin + (1 | respondent)", fq, "df")
frc, frp = tv(hb, "Friedman test", "Chi-squared", "statistic"), tv(hb, "Friedman test", "Chi-squared", "p")
frw = tv(hb, "Friedman test", "Kendall's W")
prop = {b: 100 * tv(hb, "Proportion of viewers with Joy present", f"{b}: proportion with Joy present") for b in BINS}   # percent
bmean = {b: tv(hb, "Mean Percent of time present", f"{b}: mean Joy percent") for b in BINS}
r_p, r_lo, r_hi = tv("H3", "Pearson", rq), tv("H3", "Pearson", rq, "ci_low"), tv("H3", "Pearson", rq, "ci_high")
r_pp = tv("H3", "Pearson", rq, "p")
rho_v, rho_pp = tv("H3", "Spearman", sq), tv("H3", "Spearman", sq, "p")
excl_txt = " and ".join(excluded)
dq_excl = " and ".join(f"{k} ({dq_of[k]:.1f} percent)" for k in excluded)
d_vals = [dq_of[k] for k in dq_of]
mean_pv = st.mean(float(r_[idx["Percent of time present"]]) for r_ in rows
                  if r_[idx["Expression"]] == "Positive Valence" and r_[idx["Interval label"]] == "Commercial" and r_[idx["Respondent Name"]] in kept)

md = f"""# Facial expression analysis in Excel: step by step walkthrough

**SIMULATED DATA.** Every number in this walkthrough comes from `SIMULATED_fea_metrics.csv`, which `make_simulated_data.py` generated (seed {SEED}). No real viewer was recorded. The design follows the structure of the lab's library study "Nike Campaign Ad Evaluation" (a 60 s Baseline, a 60 s commercial, a 15 s Recovery, a short questionnaire, the respondent camera recorded for AFFDEX); no data from that study were used, and nothing is said about what the commercial shows.

This is the same analysis as `analysis.R`, done in Excel for students who do not use R. The finished workbook is `SIMULATED_fea_analysis.xlsx`; the sheet names below refer to it. You can follow along in that workbook or rebuild it from the CSV files. **Every cell address in this walkthrough is written as `Sheet!Cell` and was generated from the workbook itself** (`make_excel_workbook.py` writes both), so the addresses match the file. The sheets `H1` and `H3` have names that look like cell addresses, so when you type a formula that points at them, put the name in single quotes (`'H1'!L10`); the walkthrough leaves the quotes out.

Companion files: `analysis.R` (R version), `analysis_output.txt` (its console output), `results_text.md` (model paragraphs), `figures/` (PNG figures).

---

## 0. What you are analyzing

**Design.** One viewing. Each of {len(dq_of)} respondents (V01 to V{len(dq_of):02d}) watched the same sequence once on the lab monitor while the lab webcam recorded the face: a 60 s Baseline, a 60 s commercial, a 15 s Recovery, and a short questionnaire. The video was post processed with Affectiva AFFDEX (every frame), then the **FEA thresholding** R Notebook turned each expression into present (1) or absent (0) frame by frame at its default thresholds (Emotion 50, Valence +50 and -50, Action Units 50). A binning stimulus annotation cut the commercial into six 10 s bins, `Bin_01` to `Bin_06`. Because everyone saw the Baseline first and only once, the Baseline is a comparison period and not a counterbalanced control.

**Export.** The **FEA Metrics Export** (analysis > Export > R Processing > FEA Metrics Export; choose all or selected expressions) reports, for each respondent and each time interval (Stimulus, Annotation, or Scene), how often and how long an expression was present. The file here has one row per respondent by interval by expression: {N + len(excluded)} respondents x 9 intervals x 6 expressions = {n_data} data rows. **The manuals print no header row for this export, so the column names below are modeled on their description and must be checked against a real export on the lab machine.** The columns used:

| Column | Meaning | Used for |
|---|---|---|
| `Data quality (%)` | Percentage of frames successfully analyzed after post processing (the Quality Metrics value); not a measure of accuracy | Exclusion rule |
| `Interval type`, `Interval label` | Stimulus (`Baseline`, `Commercial`, `Recovery`) or Annotation (`Bin_01` to `Bin_06`) | Selecting rows |
| `Expression` | Joy, Surprise, Brow Furrow, Positive Valence, Negative Valence, Engagement | Selecting rows |
| `Threshold` | The threshold the notebook applied (50; -50 for Negative Valence) | Documentation |
| `Percent of time present` | Percent of the interval in which the expression was present | H1, H2, H3 |
| `Number of segments` | How many separate stretches of presence the interval holds (how often it was expressed) | Descriptives |
| `Mean raw score (0-100)` | Engagement only: mean AFFDEX score in the interval | Descriptives |

Engagement is not thresholded: the notebook article lists default thresholds for Emotion, Valence, Action Units, and the head position channels, but does not name Engagement, so it is reported as a raw mean (the kind of statistic the AFFDEX Raw Data Statistics Export gives). iMotions eye tracking exports write **NA** when a metric cannot be calculated; the simulated file uses NA the same way (the threshold, percent, and segment columns are NA on Engagement rows, and the raw mean is NA on all other rows). Whether the FEA export does this is an item to verify.

**Hypotheses.**
- H1: the share of time with Joy present is larger during the commercial than during the Baseline. Joy is zero inflated (most viewers show none at Baseline), so the primary test is the **Wilcoxon signed rank test**; the paired t test and d_z are reported beside it.
- H2: Joy `Percent of time present` differs across the six bins; the plan is a concentration in `Bin_05` and `Bin_06`. Tested with a repeated measures ANOVA (the mixed model F test in R), the proportion of viewers with Joy present per bin, and a Friedman test.
- H3: viewers with more Positive Valence time during the commercial like the ad more (Pearson and Spearman correlation with `Ad_liking_1to7`).

**Decision rules fixed before looking at the data.** Exclude a respondent whose `Data quality (%)` is below {DQ_RULE} percent. The manuals state no minimum; {DQ_RULE} percent is the lab's own rule. A viewer has an expression **present** in an interval when `Percent of time present` is above 0 (that is the quantity the thresholding Aggregation counts per time window). The H1 difference is **Commercial minus Baseline**. Two tailed tests, alpha = .05.

---

## 1. Import the export into Excel

iMotions exports are comma separated with a dot decimal. Double clicking the file uses your regional defaults, which can put everything in one column (Opening Data Export Files in Excel article). Use the import route instead.

**Route A, Get Data (Excel 2016 and later):**
1. **Data > From Text/CSV**, select `SIMULATED_fea_metrics.csv`, click **Import**.
2. In the preview set **Delimiter: Comma** and **File Origin: 65001 Unicode (UTF-8)**. **You should see** {len(header)} columns with `Study Name` in the first column.
3. Click **Transform Data** if the first row shows the `# SIMULATED ...` comment line: in Power Query choose **Home > Remove Rows > Remove Top Rows > 1**, then **Home > Use First Row as Headers**, then **Close & Load**. (A real iMotions export has no comment line; the simulated file carries one so nobody mistakes it for real data.)
4. **You should see** a table with {n_data} data rows.

**Route B, Text Import Wizard (any version):** **File > Open > Browse**, set the file type to **All Files**, open the .csv, choose **Delimited**, start at row 2 (to skip the comment line), tick **Comma**, click **Advanced** and set the decimal symbol to a dot, then **Finish**.

**Common mistake.** The NA cells must stay as the text `NA` (or be blank). Do not replace NA with 0: a 0 in `Percent of time present` would mean "the expression was never present", which is a measurement, while NA means "no value exists". Step 2 handles both correctly.

Save as **.xlsx** right away; a .csv does not keep formulas or extra sheets. Keep the original .csv read only.

In the workbook this is the **Data** sheet, header in row 1 and data in rows {FIRST} to {LAST}. Column letters there: `{COL['Respondent Name']}` = Respondent Name, `{COL['Interval type']}` = Interval type, `{COL['Interval label']}` = Interval label, `{COL['Expression']}` = Expression, `{COL['Percent of time present']}` = Percent of time present, `{COL['Number of segments']}` = Number of segments, `{COL['Mean raw score (0-100)']}` = Mean raw score, `{COL['Data quality (%)']}` = Data quality (%).

---

## 2. Clean with helper columns

Add two columns to the right of the data (`{HCOL['Keep']}` and `{HCOL['Present']}` in the workbook, headers in orange) and fill them down to the last row. The formulas below are the ones in row 2.

| Header | Cell | Formula | Meaning |
|---|---|---|---|
| Keep | `Data!{cl('Data','keep')}` | `{fx('Data','keep')}` | Exclusion rule on Data quality |
| Present | `Data!{cl('Data','present')}` | `{fx('Data','present')}` | 1 when the expression was present in the interval; blank for NA |

**You should see** `Keep = No` on all 54 rows for each excluded respondent, {dq_excl}; every other respondent is `Yes`. {excl_txt} were recorded in full and post processed; the session log explains each exclusion (glare on glasses lenses for V07; looking down at a phone during the commercial for V19). The same ({N}) respondents appear in every analysis below.

`ISNUMBER` is the key in the second formula: it is TRUE for a number and FALSE for the text NA or a blank cell, so one formula handles both ways an export can mark a missing metric.

---

## 3. Descriptive table by expression and interval

### 3a. With a PivotTable (the iMotions recommended route)

The Creating Excel Pivot Tables article gives the generic steps: click the first cell of the data, **Insert > PivotTable > OK**, drag a category to **Rows**, drag a metric to **Values**, right click the value, **Value Field Settings**, change **Sum** to **Average**, **OK**, then **PivotChart** if you want a quick picture. For this study:

1. Click cell A1 of the Data sheet, **Insert > PivotTable**, **New Worksheet**, **OK**.
2. Drag **Keep** to **Filters** and set it to **Yes**.
3. Drag **Expression** to **Filters** too and choose **Joy** (or put it in **Rows** to see all six).
4. Drag **Interval label** to **Rows**.
5. Drag **Percent of time present** to **Values** twice: **Average** and **StdDev** (the sample standard deviation, the same as STDEV.S).
6. Drag **Present** to **Values** and set it to **Sum** for the number of viewers with the expression present, or **Average** and multiply by 100 for the percent of viewers.
7. Drag **Number of segments** to **Values** and set it to **Average**.

**You should see** for `Joy`: `Baseline` average {dd('Joy', 'Baseline')['mean_pct']}, `Commercial` average {dd('Joy', 'Commercial')['mean_pct']}, and `Bin_05` average {dd('Joy', 'Bin_05')['mean_pct']}.

**Common mistake.** Leaving **Keep** out of the filter. {excl_txt} would then add rows of thin data (many missing frames, so low percentages).

### 3b. With formulas (the Summary sheet)

A formula table updates when the data change and shows exactly what is being averaged. The expression name is in column A and the interval in column B of each row (row {SUM_FIRST} is the first). The formulas in row {SUM_FIRST} (`Joy`, `Baseline`):

| Statistic | Cell | Formula |
|---|---|---|
""" + "\n".join(
    f"| {REG[('Summary', k)]['label']} | `Summary!{cl('Summary', k)}` | `{fx('Summary', k)}` |"
    for k in ("n", "mean", "sd", "median", "n_present", "pct_present", "mean_seg")) + f"""

The two array formulas (`STDEV.S(IF(...))` and `MEDIAN(IF(...))`, columns E and F) are needed because Excel has no STDEVIFS or MEDIANIFS. In Excel 365 type them and press **Enter**. In Excel 2019 or older press **Ctrl+Shift+Enter**; Excel then shows the formula in braces. `AVERAGEIFS` and `COUNTIFS` ignore the text NA.

**You should see** the same numbers as the PivotTable, in rows {SUM_FIRST} to {SUM_LAST} (nine rows per expression: the three stimuli, then the six bins) and a column chart of mean Joy by interval under the Engagement table. The **Engagement** table starts in row {ENG_FIRST} (`Summary!{REG[('Summary','e_mean')]['cell']}` is the mean raw score for the first interval):

| Statistic | Cell | Formula |
|---|---|---|
""" + "\n".join(
    f"| {REG[('Summary', k)]['label']} | `Summary!{cl('Summary', k)}` | `{fx('Summary', k)}` |"
    for k in ("e_n", "e_mean", "e_sd", "e_median")) + f"""

---

## 4. H1: Joy, Commercial versus Baseline

A paired test needs the two measurements of the same respondent side by side, so first reshape the long export into one row per respondent. The **H1** sheet does this for the {N} kept respondents (listed in column A, rows {P1} to {P2}); each of columns C and D is a `SUMIFS` that picks the one matching Joy row from the Data sheet:

| Column | Cell (first respondent) | Formula |
|---|---|---|
""" + "\n".join(
    f"| {REG[('H1', k)]['label']} | `H1!{REG[('H1', k)]['cell']}` | `{REG[('H1', k)]['formula']}` |"
    for k in ("row_dq", "row_base", "row_comm", "row_diff", "row_abs", "row_rank", "row_sign", "row_tie")) + f"""

(By hand you would get the same table from the PivotTable of section 3a: put **Respondent Name** in Rows, **Interval label** in Columns, **Percent of time present** in Values as Sum, and filter **Expression** to `Joy` and **Keep** to `Yes`.)

The difference is rounded to two decimals (`ROUND`) because the export has two decimals; without it, floating point noise can make two equal differences look unequal, which breaks the tie handling in the rank columns.

Now the test. The statistics sit in column L of the H1 sheet with the formula text shown beside them in column M:

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('Baseline mean (%)', 'H1', 'm_b', fnum(m_b))}
{row('Baseline SD (%)', 'H1', 's_b', fnum(sd_b))}
{row('Commercial mean (%)', 'H1', 'm_p', fnum(m_c))}
{row('Commercial SD (%)', 'H1', 's_p', fnum(sd_c))}
{row('n pairs', 'H1', 'n', N)}
{row('Mean difference, Commercial minus Baseline (%)', 'H1', 'md', fnum(h1d['estimate']))}
{row('SD of the differences (%)', 'H1', 'sdd', fnum(sd_diff))}
{row('**p, two tailed (T.TEST type 1)**', 'H1', 'p_tt', '**' + pnum(h1d['p']) + '**')}

`T.TEST(array1, array2, tails, type)`: `tails = 2` for a two tailed test; `type = 1` for **paired** data, which is the type for this within subjects design. `type = 2` and `type = 3` treat the two columns as different people and give the wrong answer for paired data. `T.TEST` returns only the p value, so compute t and df by hand with the paired formulas (a paired t test is a one sample t test on the differences):

| Statistic | Cell | Formula | Result |
|---|---|---|---|
{row('SE of the mean difference', 'H1', 'se', fnum(sd_diff / N ** 0.5, 3))}
{row('Paired t', 'H1', 't', fnum(h1d['statistic']))}
{row('df', 'H1', 'df', int(float(h1d['df'])))}
{row('p from the hand computed t', 'H1', 'p_hand', pnum(h1d['p']) + ', matches T.TEST')}
{row('Critical t (95%)', 'H1', 'tcrit', fnum(tcrit, 3))}
{row('Mean difference 95% CI lower', 'H1', 'ci_lo', fnum(h1d['ci_low']))}
{row('Mean difference 95% CI upper', 'H1', 'ci_hi', fnum(h1d['ci_high']))}

**You should see** the hand computed p equal the `T.TEST` p. If they differ, the `type` argument or a column range is wrong.

### Cohen's d_z by hand, with a 95 percent confidence interval

For paired data the effect size is **d_z**: the mean difference in units of the SD of the *differences*. It is linked to t by t = d_z x sqrt(n), so t, n, and d_z always agree.

| Step | Cell | Formula | Result |
|---|---|---|---|
{row(chr(67) + "ohen's d_z", 'H1', 'dz', fnum(dz))}
{row('SE of d_z (approximation)', 'H1', 'se_dz', fnum(se_dz, 3))}
{row('d_z 95% CI lower', 'H1', 'dz_lo', fnum(dz - 1.96 * se_dz))}
{row('d_z 95% CI upper', 'H1', 'dz_hi', fnum(dz + 1.96 * se_dz))}
{row('Correlation of the two columns', 'H1', 'r', fnum(st.correlation(pcs['Baseline'], pcs['Commercial'])))}

In words: d_z = mean difference / SD of differences; SE(d_z) = sqrt(1/n + d_z^2 / (2 n)). The interval is the usual large sample approximation. `analysis.R` reports the interval from the noncentral t distribution (effectsize package), which is more exact at this sample size; d_z itself matches, and the limits differ slightly ([{fnum(tv(H1T, 'Paired t test', "Cohen's d_z", 'ci_low'))}, {fnum(tv(H1T, 'Paired t test', "Cohen's d_z", 'ci_high'))}] in R). Say which method you used. Rule of thumb for reading d_z: 0.2 small, 0.5 medium, 0.8 large.

### Viewers with any Joy, and the exact McNemar test

Because most Baseline values are 0, the count of viewers who showed any Joy is part of the result.

| Statistic | Cell | Formula | Result |
|---|---|---|---|
{row('Viewers with Joy, Baseline', 'H1', 'any_b', any_b)}
{row('Viewers with Joy, Commercial', 'H1', 'any_c', any_c)}
{row('Percent of viewers with Joy, Baseline', 'H1', 'pct_b', fnum(100 * any_b / N, 1))}
{row('Percent of viewers with Joy, Commercial', 'H1', 'pct_c', fnum(100 * any_c / N, 1))}
{row('Joy at Baseline only', 'H1', 'disc_b', int(sum(1 for b, c in zip(pcs['Baseline'], pcs['Commercial']) if b > 0 and c == 0)))}
{row('Joy in the Commercial only', 'H1', 'disc_c', int(sum(1 for b, c in zip(pcs['Baseline'], pcs['Commercial']) if b == 0 and c > 0)))}
{row('Exact McNemar p, two tailed', 'H1', 'mcn_p', pnum(tv(H1T, 'Exact McNemar test', mq, 'p')))}

The McNemar test looks only at the viewers whose status changed between the two periods (Joy in one, none in the other) and asks whether the changes favor one direction; with the exact binomial the p value is `MIN(1, 2 x BINOM.DIST(smaller count, total changed, 0.5, TRUE))`.

### The Wilcoxon signed rank test by hand (the primary test)

Excel has no signed rank function, but the test is a short calculation on the difference column. The steps: drop the zero differences; rank the absolute differences from smallest to largest (tied values share the average rank, `RANK.AVG`); add up the ranks of the positive differences (that sum is the V that R prints); compare V with its expected value under no difference, using the normal approximation with the tie and continuity corrections that R applies when there are ties (`wilcox.test(exact = FALSE)`).

| Statistic | Cell | Formula | Result |
|---|---|---|---|
{row('Differences above 0', 'H1', 'n_pos', n_pos)}
{row('Differences below 0', 'H1', 'n_neg', n_neg)}
{row('Differences equal to 0', 'H1', 'n_zero', n_zero)}
{row('Nonzero differences', 'H1', 'n_nz', wn)}
{row('V = sum of ranks of positive differences', 'H1', 'w_pos', fnum(wV, 1))}
{row('Sum of ranks of negative differences', 'H1', 'w_neg', fnum(wn * (wn + 1) / 2 - wV, 1))}
{row('V plus negative rank sum minus n(n+1)/2', 'H1', 'w_chk', '0.0')}
{row('Expected V under H0', 'H1', 'v_mean', fnum(wn * (wn + 1) / 4))}
{row('Sum of (t^3 - t) over tie groups', 'H1', 'tie3', xv('H1', 'tie3', 0))}
{row('SD of V under H0', 'H1', 'sigma', xv('H1', 'sigma', 3))}
{row('z with continuity correction', 'H1', 'z', xv('H1', 'z', 3))}
{row('**Wilcoxon p, two tailed**', 'H1', 'p_w', '**' + pnum(wp) + '**')}
{row('Matched pairs rank biserial correlation', 'H1', 'r_rb', fnum(r_rb))}
{row('Sign test p, two tailed', 'H1', 'p_sign', xv('H1', 'p_sign', 6))}

**You should see** V equal to {fnum(wV, 0)} and the p value equal to the one in `analysis_output.txt` ({pnum(wp)}). The rank biserial correlation is the effect size that goes with the signed rank test: +1 means every nonzero difference is positive and large relative to the negative ones. The sign test uses only the direction of each difference and is the simplest check of all.

The two check rows at the bottom of the block:

| Statistic | Cell | Formula |
|---|---|---|
{row3('Kept respondents with a Joy Baseline row in Data', 'H1', 'chk_n')}
{row3('Smallest Data quality in column B', 'H1', 'chk_dq')}

**You should see** {N} in the first (it must equal n pairs) and a number of at least {DQ_RULE} in the second.

---

## 5. H2: Joy across the six bins

The **Bins** sheet lists the {N} kept respondents (rows {B1} to {B2}) with one column per bin (`B` to `G`), each a `SUMIFS` on the Joy rows of the Data sheet; the bin names in row 4 are the lookup keys, so do not rename them.

| Column | Cell (first respondent) | Formula |
|---|---|---|
""" + "\n".join(
    f"| {REG[('Bins', k)]['label']} | `Bins!{REG[('Bins', k)]['cell']}` | `{REG[('Bins', k)]['formula']}` |"
    for k in ("row_bin", "row_comm", "row_resp_mean", "row_rank", "row_tie")) + f"""

### The proportion of viewers with Joy present per bin

This is the quantity the FEA thresholding **Aggregation** shows in the Aggregate Replay (how many respondents showed an expression in a time window). Rows {RN} to {RRS} under the roster hold, for each bin, the count of viewers, the count with Joy present (`COUNTIFS` with the criterion `">0"` on Percent of time present), the proportion, and the mean. The formulas in column B (`Bin_01`):

| Statistic | Cell | Formula | Results across the bins (simulated) |
|---|---|---|---|
{row('Viewers (n)', 'Bins', 'b_n', N)}
{row('Viewers with Joy present', 'Bins', 'b_nj', ', '.join(str(round(prop[b] / 100 * N)) for b in BINS))}
{row('Proportion with Joy present', 'Bins', 'b_prop', ', '.join(fnum(prop[b] / 100, 3) for b in BINS))}
{row('Percent of viewers with Joy present', 'Bins', 'b_pct', ', '.join(fnum(prop[b], 1) for b in BINS))}
{row('Mean Joy (%)', 'Bins', 'b_mean', ', '.join(fnum(bmean[b]) for b in BINS))}
{row('SD (%)', 'Bins', 'b_sd', xv('Bins', 'b_sd', 2, True))}
{row('95% CI half width', 'Bins', 'b_ci', xv('Bins', 'b_ci', 2, True))}
{row('Friedman rank sum', 'Bins', 'b_rs', xv('Bins', 'b_rs', 1, True))}

**You should see** the proportion and the mean both highest in `Bin_05` (column F) and next highest in `Bin_06` (column G).

### Repeated measures ANOVA by hand

`analysis.R` fits `pct ~ bin + (1 | respondent)` with lme4 and lmerTest. With every respondent measured in every bin (a balanced design) and a positive respondent variance, that model's F test for the bin factor is the classic repeated measures ANOVA, which Excel can compute from sums of squares. The block sits in column S (formula text in column T):

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('Respondents (n)', 'Bins', 'n_r', N)}
{row('Bins (k)', 'Bins', 'k', 6)}
{row('SS total', 'Bins', 'ss_tot', xv('Bins', 'ss_tot', 1))}
{row('SS bins', 'Bins', 'ss_bin', xv('Bins', 'ss_bin', 1))}
{row('SS respondents', 'Bins', 'ss_resp', xv('Bins', 'ss_resp', 1))}
{row('SS error', 'Bins', 'ss_err', xv('Bins', 'ss_err', 1))}
{row('df error', 'Bins', 'df_err', int(Fdf))}
{row('MS bins', 'Bins', 'ms_bin', xv('Bins', 'ms_bin', 2))}
{row('MS error', 'Bins', 'ms_err', xv('Bins', 'ms_err', 2))}
{row('**F**', 'Bins', 'F', '**' + fnum(Ffac) + '**')}
{row('p (F.DIST.RT)', 'Bins', 'p_F', pnum(Fp))}
{row('Respondent variance (ANOVA estimate)', 'Bins', 'var_resp', xv('Bins', 'var_resp', 2))}
{row('ICC (adjusted)', 'Bins', 'icc', fnum(tv(hb, 'Mixed model pct ~ bin + (1 | respondent)', 'ICC adjusted'), 3))}

`DEVSQ` returns the sum of squared deviations from the mean. The error term uses (n - 1)(k - 1) = {int(Fdf)} degrees of freedom, which is also the Satterthwaite denominator df that lmerTest reports for this balanced design. The respondent variance and ICC are the ANOVA estimates, which equal the REML estimates of the mixed model here.

### Friedman test by hand (the nonparametric check)

Joy is zero inflated and right skewed, so the F test's normality assumption is doubtful (see the assumption plots from `analysis.R`). The Friedman test ranks each respondent's six bin values from 1 to 6 (ties share an average rank) and asks whether the bins differ in their rank sums. The tie correction matters here because many bins hold 0.

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('12 x sum of squared deviations of the rank sums', 'Bins', 'fr_num', xv('Bins', 'fr_num', 2))}
{row('Denominator with tie correction', 'Bins', 'fr_den', xv('Bins', 'fr_den', 2))}
{row('**Friedman chi-squared**', 'Bins', 'fr_chi', '**' + fnum(frc) + '**')}
{row('p (CHISQ.DIST.RT)', 'Bins', 'fr_p', pnum(frp))}
{row("Kendall's W", 'Bins', 'fr_w', fnum(frw, 3))}

Kendall's W is the effect size for the Friedman test (0 = no agreement among respondents about the order of the bins, 1 = complete agreement).

### Contrasts: where is the Joy?

An omnibus test says the bins differ, not where. Two descriptive contrasts compare a late bin with the mean of `Bin_01` to `Bin_04`; the standard error uses the error mean square and the squared weights (1 and four times -1/4 give 1.25):

| Statistic | Cell | Formula |
|---|---|---|
""" + "\n".join(row3(REG[('Bins', k)]['label'], 'Bins', k) for k in ("c5_est", "c5_se", "c5_t", "c5_p", "c6_est", "c6_se", "c6_t", "c6_p")) + f"""

R reports the same contrasts with Satterthwaite degrees of freedom and a Holm adjustment for the two contrasts; the unadjusted Excel p values match the unadjusted R p values.

The two check rows at the bottom of the block:

| Statistic | Cell | Formula |
|---|---|---|
{row3('Largest gap between Commercial Joy and the mean of the six bins (percentage points)', 'Bins', 'chk_comm')}
{row3('Rank sums add up to n k (k + 1) / 2', 'Bins', 'chk_rank')}

**You should see** a gap of at most 0.01 percentage points (the bins tile the commercial, so the stimulus value is the mean of the bin values, up to rounding) and 0 for the rank sum check.

---

## 6. H3: Positive Valence and ad liking

The **H3** sheet holds `SIMULATED_fea_survey.csv` (one row per respondent, rows {S1} to {S_last_all}), with the kept respondents first (rows {S1} to {S2}) and the excluded respondents last (rows {S2 + 1} to {S_last_all}), so the test ranges below leave them out. Column G flags Keep from the Data sheet, column H pulls Positive Valence `Percent of time present` during the commercial, and columns I and J rank the two variables for the Spearman correlation:

| Column | Cell (first row) | Formula |
|---|---|---|
""" + "\n".join(
    f"| {REG[('H3', k)]['label']} | `H3!{REG[('H3', k)]['cell']}` | `{REG[('H3', k)]['formula']}` |"
    for k in ("row_keep", "row_pv", "row_rk_like", "row_rk_pv")) + f"""

Positive Valence is the Valence channel above +50. The statistics sit in column M (formula text in column N):

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('n (kept respondents)', 'H3', 'n', N)}
{row('Positive Valence mean (%)', 'H3', 'm_pv', fnum(mean_pv))}
{row('**Pearson r (CORREL)**', 'H3', 'r', '**' + fnum(r_p) + '**')}
{row('t for r', 'H3', 't_r', fnum(tv('H3', 'Pearson', rq, 'statistic')))}
{row('df', 'H3', 'df_r', N - 2)}
{row('p, two tailed', 'H3', 'p_r', pnum(r_pp))}
{row('r 95% CI lower (Fisher z)', 'H3', 'r_lo', fnum(r_lo))}
{row('r 95% CI upper (Fisher z)', 'H3', 'r_hi', fnum(r_hi))}
{row('Least squares slope (percentage points per liking point)', 'H3', 'slope', fnum(tv('H3', 'Least squares line', 'Slope (percentage points per liking point)')))}
{row('**Spearman rho**', 'H3', 'rho', '**' + fnum(rho_v) + '**')}
{row('t for rho', 'H3', 't_rho', fnum(tv('H3', 'Spearman', sq, 'statistic')))}
{row('p, two tailed', 'H3', 'p_rho', pnum(rho_pp))}
{row('rho 95% CI lower (Bonett and Wright)', 'H3', 'rho_lo', fnum(tv('H3', 'Spearman', sq, 'ci_low')))}
{row('rho 95% CI upper (Bonett and Wright)', 'H3', 'rho_hi', fnum(tv('H3', 'Spearman', sq, 'ci_high')))}

A Pearson correlation uses the raw values, so a few viewers with a lot of Positive Valence can pull it; Spearman uses ranks (`RANK.AVG` gives tied values the average rank, and ratings on a 7 point scale tie often), so it is the check against that. `Ad_liking_1to7` is a whole number rating, and it is correlated with a time percentage: a correlation says the two move together, not that one causes the other, and AFFDEX scores show what the face expressed, not what the viewer felt.

---

## 7. Chart with 95 percent confidence intervals

1. Make a small chart table: two rows (Baseline, Commercial), columns **Mean Joy (%)** (point to the mean cells) and **CI half width** = `CONFIDENCE.T(0.05, SD cell, n cell)`. `CONFIDENCE.T` returns half the width of a t based 95 percent interval around a mean. The workbook has this table on the H1 sheet at `H1!K{H1_CHART_ROW + 2}:M{H1_CHART_ROW + 3}`; the half widths are in `{rcell('H1', 'chart_ci_b')}` and `{rcell('H1', 'chart_ci_p')}`.
2. Select the interval names and the means, **Insert > Column chart > Clustered Column**.
3. Click the bars, **Chart Design > Add Chart Element > Error Bars > More Error Bars Options**, choose **Custom > Specify Value**, and select the two CI half width cells for both **Positive** and **Negative**.
4. Set the bar fill to navy (`#1F3864`) and lighten the gridlines.

**Common mistake.** Using the default error bars (**Standard Error** or a fixed 5 percent). Those are not 95 percent confidence intervals. And remember that error bars around the two means are *not* the interval that goes with the paired test: the paired test uses the interval of the *differences* (`{rcell('H1', 'ci_lo')}` and `{rcell('H1', 'ci_hi')}`), which is much narrower because it removes the between respondent variation. With zero inflated data, a bar chart of means also hides the shape: show the individual values (the R figure does).

---

## 8. Check your numbers and know what Excel cannot do

The **Check** sheet puts each live Excel result next to the value `analysis.R` produced and prints OK or CHECK (status in column F, the count of CHECK rows in `{rcell('Check', 'n_check')}`). Your Excel results should match to rounding; the d_z limits differ by method and have a looser tolerance. Several parts of the R analysis have no Excel equivalent:

- **Mixed model estimates** (random intercept variance as fitted by REML, Satterthwaite degrees of freedom for unbalanced or missing data, the model on the square root of the percent).
- **Hodges-Lehmann estimate** and its interval for the median difference.
- **Exact Clopper-Pearson intervals** for the proportions of viewers per bin.
- **The sensitivity analysis at thresholds of 25 and 75.** In iMotions you rerun the thresholding notebook with a batch threshold of 25 and then 75 and export the FEA Metrics Export again. In Excel, import `SIMULATED_fea_metrics_threshold25.csv` into a copy of the Data sheet (same columns, same helper formulas) and the H1 sheet recalculates for that threshold. Expect the mean Joy percent to rise at 25 and fall at 75, because a lower threshold counts weaker expressions, and longer stretches of them, as present.

If you have the **Analysis ToolPak** add in (**File > Options > Add-ins > Manage Excel Add-ins > Go > Analysis ToolPak**), **Data > Data Analysis > t-Test: Paired Two Sample for Means** prints t, df, and p in one table and is a good cross check of section 4; **Anova: Two-Factor Without Replication** gives the repeated measures table of section 5 (rows = respondents, columns = bins).

---

## 9. Common mistakes, collected

1. Opening the .csv by double clicking and getting one column: use **Data > From Text/CSV** or the Text Import Wizard.
2. Replacing NA with 0. NA means no value exists (for example, the thresholded columns on Engagement rows).
3. Forgetting the exclusion filter (Keep = Yes) in the PivotTable or in COUNTIFS, or leaving {excl_txt} in a roster.
4. `T.TEST` with `type = 2` or `type = 3` on a within subjects design. Use `type = 1` and make sure the two columns are in the same respondent order.
5. Reporting the paired t test alone for a zero inflated outcome. Report the signed rank test, the count of viewers with any Joy, and d_z with its interval.
6. Treating {n_data} rows as {n_data} viewers. Each respondent contributes 54 rows; tests use one value per respondent and interval.
7. Adding up bin percentages as if they were separate observations of 60 s: each bin is 10 s of the 60 s commercial, so the commercial value is the **mean** of the six bins, not the sum.
8. Reading the Baseline as a counterbalanced control. Everyone saw the Baseline first, so any difference between Baseline and Commercial also includes time in the chair.
9. Reading `Percent of time present` as the strength of an emotion. FEA reports whether the face showed an expression above the threshold, not how intensely the viewer felt it.

## Sources

- iMotions help center, R Notebook (FEA): Facial expression analysis thresholding and aggregation: https://help.imotions.com/lab/analysis/facial-expression-analysis/r-notebook-fea-facial-expression-analysis-thresholding-and-aggregation (local copy: `C:\\JaxState_Behavior Lab\\iMotion_Manual\\R Notebook FEA Facial expression analysis thresholding and aggregation.pdf`)
- iMotions help center, Facial Expression Analysis: Importing and Post Processing Videos: https://help.imotions.com/lab/import-and-post-processing/facial-expression-analysis-importing-and-post-processing-videos (local copy: `C:\\JaxState_Behavior Lab\\iMotion_Manual\\Facial Expression Analysis Post Processing face recordings.pdf`)
- iMotions help center, Affectiva (AFFDEX) Signals Overview: https://help.imotions.com/lab/analysis/facial-expression-analysis/affdex/affectiva-affdex-signals-overview (local copy: `C:\\JaxState_Behavior Lab\\iMotion_Manual\\AFFDEX SDK 52 Signals.pdf`)
- iMotions help center, R Notebooks (FEA): AFFDEX Raw Data Statistics Export: https://help.imotions.com/lab/analysis/facial-expression-analysis/affdex/r-notebooks-fea-affdex-raw-data-statistics-export (local copy: `C:\\JaxState_Behavior Lab\\iMotion_Manual\\R Notebooks FEA AFFDEX Raw Data Statistics Export.pdf`)
- iMotions help center, Facial Expression Analysis: Why does AFFDEX show low QA in subject replay?: https://help.imotions.com/lab/analysis/facial-expression-analysis/affdex/facial-expression-analysis-why-does-affdex-show-low-qa-in-subject-replay (local copy: `C:\\JaxState_Behavior Lab\\iMotion_Manual\\Facial Expression Analysis Why does AFFDEX show low QA in subject replay.pdf`)
- iMotions help center, Demo study: Sony Bravia (Testing of video-commercials) - 9.0 (binning annotation): https://help.imotions.com/lab/getting-started/demo-studies/9-0-demo-studies/demo-study-sony-bravia-testing-of-video-commercials-9-0 (local copy: `C:\\JaxState_Behavior Lab\\iMotion_Manual\\1Printed\\Demo study Sony Bravia Testing of videocommercials  90.pdf`)
- iMotions help center, Opening Data Export Files in Excel: https://help.imotions.com/lab/exports/additional-export-information/opening-data-export-files-in-excel (local copy: `iMotion_Manual\\Opening Data Export Files in Excel.pdf`)
- iMotions help center, Creating Excel Pivot Tables: https://help.imotions.com/lab/exports/additional-export-information/creating-excel-pivot-tables (local copy: `iMotion_Manual\\Making Excel Pivot Tables.pdf`)

## Items to verify on the lab computer

- The exact column header layout, file name, and file format of the FEA Metrics Export in iMotions Lab 11.1 (the manuals describe the content but print no header row; one article calls R exports .txt and others .csv). Adjust the column letters in sections 1 and 2 once a real export is open.
- Whether the real export writes `NA` or leaves cells blank where a metric does not apply; `ISNUMBER` handles both.
- Whether the denominator of the percent of time present is the whole interval or only the frames with a face found, and how a segment that crosses a bin boundary is counted (the simulated file divides by the whole interval and counts the segment in both bins).
- Whether the thresholding notebook offers Engagement as a thresholded expression (the article lists Emotion, Valence, Action Units, and head position channels but does not name Engagement); until then Engagement stays a raw mean.
- What the thresholding Aggregation counts as "showed the expression" in a time window, and its default bin length and overlap (the manual states neither); the workbook uses Percent of time present above 0 in the 10 s bin.
- The default for the short or long classification time threshold (not stated), and the per channel threshold dialog.
- The {DQ_RULE} percent data quality exclusion rule is the lab's; no manual states a minimum. Decide the rule after viewing Quality Metrics on pilot data.
- Which AFFDEX version is installed in 11.1.13 and whether the rolling baseline setting changes the numbers.
"""
with open(MD_OUT, "w", encoding="utf-8") as fh:
    fh.write(md)
print("wrote", MD_OUT)
