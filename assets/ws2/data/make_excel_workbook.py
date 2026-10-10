#!/usr/bin/env python3
"""
make_excel_workbook.py
Builds SIMULATED_ad_test_analysis.xlsx from the three SIMULATED CSV files with openpyxl,
and writes analysis_excel_walkthrough.md from the SAME cell registry, so that every cell
reference in the walkthrough is the address of the formula in the workbook (Workshop 1's
walkthrough drifted from its workbook; this one cannot, because it is generated).

The workbook mirrors analysis.R for students who work in Excel:

  README     what the workbook is, and that every number is SIMULATED
  Data       the Individual AOI metrics export as imported (NA kept as text), plus five
             helper columns (Keep, Position, Dwell_pct_clean, Fixated, Fix_count_clean)
  Summary    table by AOI and commercial built with COUNTIFS, AVERAGEIFS, and
             STDEV.S array formulas (the hand built PivotTable)
  ProductTV  H1: one row per kept respondent (SUMIFS pulls the two dwell shares from
             Data), paired t test with T.TEST(...,2,1), hand computed t, d_z, CI, and the
             simple order effect estimate
  BrandLogo  H2: the same paired recipe for TTFF AOI (ms), on respondents with a TTFF in
             both commercials, plus the respondent ratio (an extra sheet beyond the six
             the workshop asked for, because H2 is the paired test with missing values)
  Survey     H3: paired liking, preference counts, and the point biserial correlation
             between Brand_Logo dwell time and recall (CORREL)
  Check      live Excel results next to the values analysis.R produced, with a status

All formulas are ordinary Excel formulas. The STDEV.S(IF(...)) formulas on the Summary
sheet are array formulas; Excel 365 evaluates them with Enter, older Excel needs
Ctrl+Shift+Enter (openpyxl writes them as array formulas already).

Run: python3 make_excel_workbook.py   (after make_simulated_data.py and analysis.R)
"""

import csv
import os
import re
import statistics as st

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.workbook.properties import CalcProperties
from openpyxl.chart.data_source import NumDataSource, NumRef
from openpyxl.chart.error_bar import ErrorBars
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "SIMULATED_ad_test_analysis.xlsx")
MD_OUT = os.path.join(HERE, "analysis_excel_walkthrough.md")

NAVY = "1F3864"
GRAY = "8C8C8C"
ACCENT = "A6561A"
LIGHT = "F2F2F2"

hdr_font = Font(bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor=NAVY)
sub_fill = PatternFill("solid", fgColor=LIGHT)
bold = Font(bold=True)
note_font = Font(italic=True, color="555555")
code_font = Font(name="Consolas", size=9, color="555555")
thin = Side(style="thin", color="BFBFBF")
box = Border(top=thin, bottom=thin, left=thin, right=thin)

BALLS, PAINT = "Commercial_1_Balls", "Commercial_2_Paint"
AOIS = ["Product_TV", "Brand_Logo", "Tagline_Text", "Hero_Element"]


def read_csv(name):
    with open(os.path.join(HERE, name), encoding="utf-8") as fh:
        raw = fh.readlines()
    banner = raw[0].strip()
    rows = list(csv.reader([ln for ln in raw if not ln.startswith("#")]))
    return banner, rows[0], rows[1:]


def to_cell(v):
    """Numbers become numbers; NA stays the text NA (as iMotions writes it)."""
    if v == "NA" or v == "":
        return v
    try:
        f = float(v)
        return int(f) if f.is_integer() and "." not in v else f
    except ValueError:
        return v


def style_header(ws, row=1, ncol=None, first_col=1):
    ncol = ncol or ws.max_column
    for c in range(first_col, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = box


def autowidth(ws, min_w=8, max_w=44):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            L = len(str(cell.value)) if not str(cell.value).startswith("=") else 12
            widths[cell.column_letter] = max(widths.get(cell.column_letter, 0), L)
    for col, w in widths.items():
        ws.column_dimensions[col].width = max(min_w, min(max_w, w + 2))


def show_formula(cell, formula):
    """Write the formula text into a cell as TEXT (not evaluated), without the _xlfn. prefixes."""
    cell.value = formula.replace("_xlfn.", "")
    cell.data_type = "s"
    cell.font = code_font


# registry of documented cells: REG[(sheet, key)] = {"cell": "J12", "formula": "=...", "label": "..."}
REG = {}


def reg(sheet, key, cell, formula, label=""):
    REG[(sheet, key)] = {"cell": cell, "formula": formula.replace("_xlfn.", ""), "label": label}


banner, header, rows = read_csv("SIMULATED_ad_test_AOI_metrics.csv")
SEED = re.search(r"seed (\d+)", banner).group(1)
idx = {h: i for i, h in enumerate(header)}
n_data = len(rows)
FIRST, LAST = 2, n_data + 1                       # Data rows 2 to 241
COL = {h: get_column_letter(i + 1) for i, h in enumerate(header)}   # header -> Data column letter
HELP = ["Keep", "Position", "Dwell_pct_clean", "Fixated", "Fix_count_clean"]
HCOL = {h: get_column_letter(len(header) + 1 + i) for i, h in enumerate(HELP)}   # Z to AD


def rng(col):
    return f"Data!${col}${FIRST}:${col}${LAST}"


R_RESP, R_GROUP, R_STIM, R_AOI = rng(COL["Respondent Name"]), rng(COL["Group"]), rng(COL["Stimulus"]), rng(COL["AOI Name"])
R_VALID, R_TTFF, R_DWELL_MS = rng(COL["Valid data"]), rng(COL["TTFF AOI (ms)"]), rng(COL["Dwell time (ms)"])
R_KEEP, R_POS, R_DWELL, R_FIX, R_FIXC = (rng(HCOL["Keep"]), rng(HCOL["Position"]), rng(HCOL["Dwell_pct_clean"]),
                                         rng(HCOL["Fixated"]), rng(HCOL["Fix_count_clean"]))

# kept respondents (Valid data >= 70) and the per respondent values, computed here only to build
# the static rosters; every number on the analysis sheets is a live formula on the Data sheet
by_resp = {}
for r in rows:
    by_resp.setdefault(r[idx["Respondent Name"]], []).append(r)
valid = {k: float(v[0][idx["Valid data"]]) for k, v in by_resp.items()}
kept = sorted(k for k, v in valid.items() if v >= 70)
excluded = sorted(k for k, v in valid.items() if v < 70)


def value_of(resp, stim, aoi, col):
    for r in by_resp[resp]:
        if r[idx["Stimulus"]] == stim and r[idx["AOI Name"]] == aoi:
            return r[idx[col]]


both_logo = [k for k in kept if value_of(k, BALLS, "Brand_Logo", "TTFF AOI (ms)") != "NA"
             and value_of(k, PAINT, "Brand_Logo", "TTFF AOI (ms)") != "NA"]

wb = Workbook()

# ---------------------------------------------------------------- README
ws = wb.active
ws.title = "README"
readme = [
    "SIMULATED ad test workbook (teaching material)",
    "",
    f"Every number in this workbook is SIMULATED. It was generated by make_simulated_data.py (seed {SEED}) to",
    "resemble a two commercial test in the style of the iMotions demonstration study DEMO - SonyBravia:",
    "each respondent watched BOTH commercials (Commercial_1_Balls, Commercial_2_Paint, 60 s each) on a",
    "monitor with the Smart Eye Aurora tracker. No real viewer was recorded.",
    "",
    "Sheets:",
    f"  Data       Individual AOI metrics export as imported (rows {FIRST} to {LAST}). NA is kept as text, as iMotions writes it.",
    f"             Helper columns {HCOL['Keep']} to {HCOL['Fix_count_clean']}: Keep (Valid data >= 70), Position (first or second viewing),",
    "             Dwell_pct_clean (NA counts as 0), Fixated (1 or 0), Fix_count_clean (NA counts as 0).",
    "  Summary    Mean, SD, n by AOI and commercial, built with COUNTIFS, AVERAGEIFS, STDEV.S(IF(...)).",
    "  ProductTV  H1. One row per kept respondent; paired t test (T.TEST type 1), hand computed t, d_z, CI;",
    "             simple order effect estimate.",
    "  BrandLogo  H2. The same paired recipe for TTFF AOI (ms) on respondents with a TTFF in both commercials.",
    "  Survey     H3. Paired liking, preference counts, point biserial correlation of Brand_Logo dwell with recall.",
    "  Check      Live Excel results next to the values analysis.R produced, with an OK or CHECK status.",
    "",
    "Design: WITHIN subjects. The Group column is the ORDER (A = Balls first, B = Paint first), not a treatment.",
    "Every difference is Paint MINUS Balls. Column names follow the iMotions AOI Metrics help article; where",
    "the same metric name exists as a gaze based and a fixation based metric, the gaze based one is prefixed 'Gaze'.",
    f"Exclusion rule, fixed before data collection: Valid data below 70 percent (respondent {', '.join(excluded)}).",
    "Never fixated AOI: TTFF is not observed (left out of TTFF statistics); dwell time counts as 0.",
    "Dynamic AOIs: Dwell time (%) is relative to the time the AOI was active (AOI duration (ms) on the Data sheet).",
    "",
    "Companion files: analysis_excel_walkthrough.md (step by step), analysis.R (same analysis in R),",
    "analysis_output.txt (its console output), results_text.md (model paragraphs).",
]
for i, line in enumerate(readme, start=1):
    ws.cell(row=i, column=1, value=line)
ws["A1"].font = Font(bold=True, size=14, color=NAVY)
ws.column_dimensions["A"].width = 118

# ---------------------------------------------------------------- Data
ws = wb.create_sheet("Data")
ws.append(header + HELP)
style_header(ws)
cK, cP, cD, cF, cFC = (HCOL[h] for h in HELP)
L_, G_, E_, S_, W_, R_ = (COL["Valid data"], COL["Stimulus"], COL["Group"], COL["TTFF AOI (ms)"],
                          COL["Dwell time (%)"], COL["Fixation count"])
for i, r in enumerate(rows, start=FIRST):
    ws.append([to_cell(v) for v in r])
    ws[f"{cK}{i}"] = f'=IF({L_}{i}>=70,"Yes","No")'
    ws[f"{cP}{i}"] = (f'=IF(OR(AND({E_}{i}="A",{G_}{i}="{BALLS}"),AND({E_}{i}="B",{G_}{i}="{PAINT}")),"first","second")')
    ws[f"{cD}{i}"] = f"=IF(ISNUMBER({W_}{i}),{W_}{i},0)"
    ws[f"{cF}{i}"] = f"=IF(ISNUMBER({S_}{i}),1,0)"
    ws[f"{cFC}{i}"] = f"=IF(ISNUMBER({R_}{i}),{R_}{i},0)"
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(header) + len(HELP))}{LAST}"
autowidth(ws, max_w=24)
for h in HELP:
    ws[f"{HCOL[h]}1"].fill = PatternFill("solid", fgColor=ACCENT)
reg("Data", "keep", f"{cK}2", f'=IF({L_}2>=70,"Yes","No")', "Keep")
reg("Data", "position", f"{cP}2", ws[f"{cP}2"].value, "Position")
reg("Data", "dwell_clean", f"{cD}2", f"=IF(ISNUMBER({W_}2),{W_}2,0)", "Dwell_pct_clean")
reg("Data", "fixated", f"{cF}2", f"=IF(ISNUMBER({S_}2),1,0)", "Fixated")
reg("Data", "fixc", f"{cFC}2", f"=IF(ISNUMBER({R_}2),{R_}2,0)", "Fix_count_clean")

# ---------------------------------------------------------------- Summary
ws = wb.create_sheet("Summary")
ws["A1"] = "Descriptive table by AOI and commercial (kept respondents only). SIMULATED data."
ws["A1"].font = Font(bold=True, size=12, color=NAVY)
ws["A2"] = ("n = kept respondents with a row for the AOI; Respondent ratio = percent who fixated the AOI; "
            "TTFF statistics use fixators only; dwell and fixation count statistics count never fixated AOIs as 0.")
ws["A2"].font = note_font
cols = ["AOI", "Stimulus", "n", "Respondent ratio (%)", "n with TTFF", "TTFF AOI mean (ms)", "TTFF AOI SD (ms)",
        "Dwell time mean (%)", "Dwell time SD (%)", "Fixation count mean"]
ws.append([])
ws.append(cols)
style_header(ws, row=4, ncol=len(cols))
r = 5
SUM_FIRST = r
summary_rows = {}
for a in AOIS:
    for stim in (BALLS, PAINT):
        ws.cell(row=r, column=1, value=a)
        ws.cell(row=r, column=2, value=stim)
        crit = f'{R_STIM},$B{r},{R_AOI},$A{r},{R_KEEP},"Yes"'
        f_n = f"=COUNTIFS({crit})"
        f_rr = f"=100*AVERAGEIFS({R_FIX},{crit})"
        f_nt = f'=COUNTIFS({crit},{R_TTFF},">0")'
        f_tm = f"=AVERAGEIFS({R_TTFF},{crit})"
        f_ts = (f'=_xlfn.STDEV.S(IF(({R_STIM}=$B{r})*({R_AOI}=$A{r})*({R_KEEP}="Yes")*ISNUMBER({R_TTFF}),{R_TTFF}))')
        f_dm = f"=AVERAGEIFS({R_DWELL},{crit})"
        f_ds = f'=_xlfn.STDEV.S(IF(({R_STIM}=$B{r})*({R_AOI}=$A{r})*({R_KEEP}="Yes"),{R_DWELL}))'
        f_fm = f"=AVERAGEIFS({R_FIXC},{crit})"
        ws.cell(row=r, column=3, value=f_n)
        ws.cell(row=r, column=4, value=f_rr)
        ws.cell(row=r, column=5, value=f_nt)
        ws.cell(row=r, column=6, value=f_tm)
        ws.cell(row=r, column=7).value = ArrayFormula(f"G{r}", f_ts)
        ws.cell(row=r, column=8, value=f_dm)
        ws.cell(row=r, column=9).value = ArrayFormula(f"I{r}", f_ds)
        ws.cell(row=r, column=10, value=f_fm)
        for c in range(1, 11):
            ws.cell(row=r, column=c).border = box
            if c >= 3:
                ws.cell(row=r, column=c).number_format = "0.00"
        if r == SUM_FIRST:
            for key, col, f in (("n", "C", f_n), ("resp_ratio", "D", f_rr), ("n_ttff", "E", f_nt), ("ttff_mean", "F", f_tm),
                                ("ttff_sd", "G", f_ts), ("dwell_mean", "H", f_dm), ("dwell_sd", "I", f_ds),
                                ("fix_mean", "J", f_fm)):
                reg("Summary", key, f"{col}{r}", f, cols[ord(col) - ord("A")])
        summary_rows[(a, stim)] = r
        r += 1
SUM_LAST = r - 1
ws[f"A{SUM_LAST + 2}"] = ("Array formulas (columns G and I): Excel 365 evaluates them with Enter; "
                           "Excel 2019 or older needs Ctrl+Shift+Enter after editing.")
ws[f"A{SUM_LAST + 2}"].font = note_font
autowidth(ws, max_w=22)
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 22

# chart table and clustered bar chart of dwell share by AOI and commercial
ct = SUM_LAST + 5
ws.cell(row=ct - 1, column=1, value="Chart table: mean Dwell time (%) by AOI and commercial").font = bold
ws.cell(row=ct, column=1, value="AOI")
ws.cell(row=ct, column=2, value="Balls")
ws.cell(row=ct, column=3, value="Paint")
for i, a in enumerate(AOIS, start=1):
    ws.cell(row=ct + i, column=1, value=a)
    for j, stim in enumerate((BALLS, PAINT), start=2):
        ws.cell(row=ct + i, column=j, value=f"=H{summary_rows[(a, stim)]}").number_format = "0.0"
chart = BarChart()
chart.type = "bar"
chart.grouping = "clustered"
chart.style = 10
chart.title = "Mean dwell time share by AOI and commercial (SIMULATED)"
chart.x_axis.title = "AOI"
chart.y_axis.title = "Dwell time (% of AOI active time)"
chart.add_data(Reference(ws, min_col=2, max_col=3, min_row=ct, max_row=ct + len(AOIS)), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=1, min_row=ct + 1, max_row=ct + len(AOIS)))
chart.series[0].graphicalProperties.solidFill = GRAY
chart.series[1].graphicalProperties.solidFill = NAVY
chart.height = 9
chart.width = 18
ws.add_chart(chart, f"E{ct - 1}")


# ---------------------------------------------------------------- paired test block (shared recipe)
class Block:
    """Writes label, value, formula text, and note columns for one sheet and records every cell in REG."""

    def __init__(self, ws, sheet, label_col, start_row):
        self.ws, self.sheet, self.lc = ws, sheet, label_col
        self.row = start_row
        self.ref = {}
        for k, name in enumerate(("Statistic", "Value", "Formula shown", "Note")):
            c = ws.cell(row=start_row - 1, column=label_col + k, value=name)
            c.font = hdr_font
            c.fill = hdr_fill

    def title(self, text):
        c = self.ws.cell(row=self.row, column=self.lc, value=text)
        c.font = bold
        c.fill = sub_fill
        self.row += 1

    def put(self, key, label, formula, note="", fmt="0.0000"):
        f = formula.format(**{k: f"{get_column_letter(self.lc + 1)}{v}" for k, v in self.ref.items()})
        i = self.row
        self.ws.cell(row=i, column=self.lc, value=label)
        vc = self.ws.cell(row=i, column=self.lc + 1, value=f)
        vc.number_format = fmt
        show_formula(self.ws.cell(row=i, column=self.lc + 2), f)
        if note:
            self.ws.cell(row=i, column=self.lc + 3, value=note).font = note_font
        self.ref[key] = i
        reg(self.sheet, key, f"{get_column_letter(self.lc + 1)}{i}", f, label)
        self.row += 1

    def blank(self):
        self.row += 1

    def addr(self, key):
        return f"{get_column_letter(self.lc + 1)}{self.ref[key]}"


def paired_block(blk, title, rb, rp, rd, unit, note_n="", note_diff=""):
    """Paired t test, d_z, and CIs. rb, rp, rd are the Balls, Paint, and difference ranges (same sheet)."""
    blk.title(title)
    blk.put("m_b", f"Balls mean ({unit})", f"=AVERAGE({rb})", "AVERAGE of the Balls column")
    blk.put("s_b", f"Balls SD ({unit})", f"=_xlfn.STDEV.S({rb})", "STDEV.S = sample standard deviation")
    blk.put("m_p", f"Paint mean ({unit})", f"=AVERAGE({rp})")
    blk.put("s_p", f"Paint SD ({unit})", f"=_xlfn.STDEV.S({rp})")
    blk.put("n", "n pairs", f"=COUNT({rd})", note_n, fmt="0")
    blk.put("md", f"Mean difference, Paint minus Balls ({unit})", f"=AVERAGE({rd})", note_diff)
    blk.put("sdd", f"SD of the differences ({unit})", f"=_xlfn.STDEV.S({rd})", "the SD that matters in a paired test")
    blk.put("se", "SE of the mean difference", "={sdd}/SQRT({n})", "SD of differences / sqrt(n)")
    blk.put("t", "Paired t", "={md}/{se}", "t = mean difference / SE")
    blk.put("df", "df", "={n}-1", "n pairs minus 1", fmt="0")
    blk.put("p_tt", "p, two tailed (T.TEST type 1)", f"=_xlfn.T.TEST({rb},{rp},2,1)",
            "tails = 2, type = 1 means paired", fmt="0.000000")
    blk.put("p_hand", "p from the hand computed t", "=_xlfn.T.DIST.2T(ABS({t}),{df})", "should match the T.TEST cell",
            fmt="0.000000")
    blk.put("tcrit", "Critical t (95%)", "=_xlfn.T.INV.2T(0.05,{df})", "two tailed, alpha = .05")
    blk.put("ci_lo", "Mean difference 95% CI lower", "={md}-{tcrit}*{se}")
    blk.put("ci_hi", "Mean difference 95% CI upper", "={md}+{tcrit}*{se}")
    blk.put("dz", "Cohen's d_z", "={md}/{sdd}", "d_z = mean difference / SD of differences")
    blk.put("se_dz", "SE of d_z (approximation)", "=SQRT(1/{n}+{dz}^2/(2*{n}))", "sqrt(1/n + d_z^2 / (2 n))")
    blk.put("dz_lo", "d_z 95% CI lower", "={dz}-1.96*{se_dz}", "large sample approximation; R uses the noncentral t interval")
    blk.put("dz_hi", "d_z 95% CI upper", "={dz}+1.96*{se_dz}")
    blk.put("r", "Correlation of the two columns", f"=CORREL({rb},{rp})", "within respondent correlation")
    blk.blank()


# ---------------------------------------------------------------- ProductTV
ws = wb.create_sheet("ProductTV")
SH = "ProductTV"
ws["A1"] = "H1. Product_TV Dwell time (%), Balls versus Paint, kept respondents. SIMULATED data."
ws["A1"].font = Font(bold=True, size=12, color=NAVY)
ws["A2"] = ("One row per kept respondent. Columns D and E pull the two dwell shares from the Data sheet with SUMIFS "
            "(each respondent has exactly one Product_TV row per commercial). Difference is Paint minus Balls.")
ws["A2"].font = note_font
ptv_cols = ["Respondent Name", "Group (order)", "Valid data", "Balls Dwell time (%)", "Paint Dwell time (%)",
            "Difference (Paint minus Balls)", "Paint seen"]
ws.append([])
ws.append(ptv_cols)
style_header(ws, row=4, ncol=len(ptv_cols))
P1 = 5
for k, rid in enumerate(kept):
    r = P1 + k
    ws.cell(row=r, column=1, value=rid)
    ws.cell(row=r, column=2, value=f"=INDEX({R_GROUP},MATCH($A{r},{R_RESP},0))")
    ws.cell(row=r, column=3, value=f"=INDEX({R_VALID},MATCH($A{r},{R_RESP},0))")
    ws.cell(row=r, column=4, value=f'=SUMIFS({R_DWELL},{R_RESP},$A{r},{R_AOI},"Product_TV",{R_STIM},"{BALLS}")')
    ws.cell(row=r, column=5, value=f'=SUMIFS({R_DWELL},{R_RESP},$A{r},{R_AOI},"Product_TV",{R_STIM},"{PAINT}")')
    ws.cell(row=r, column=6, value=f"=E{r}-D{r}")
    ws.cell(row=r, column=7, value=f'=IF(B{r}="A","second","first")')
    for c in range(1, 8):
        ws.cell(row=r, column=c).border = box
        if c in (4, 5, 6):
            ws.cell(row=r, column=c).number_format = "0.00"
P2 = P1 + len(kept) - 1
regs = [("grp", "B", f"=INDEX({R_GROUP},MATCH($A{P1},{R_RESP},0))", "Group (order)"),
        ("valid", "C", f"=INDEX({R_VALID},MATCH($A{P1},{R_RESP},0))", "Valid data"),
        ("balls", "D", ws.cell(row=P1, column=4).value, "Balls Dwell time (%)"),
        ("paint", "E", ws.cell(row=P1, column=5).value, "Paint Dwell time (%)"),
        ("diff", "F", f"=E{P1}-D{P1}", "Difference (Paint minus Balls)"),
        ("seen", "G", f'=IF(B{P1}="A","second","first")', "Paint seen")]
for key, col, f, lab in regs:
    reg(SH, "row_" + key, f"{col}{P1}", f, lab)
RB, RP, RD, RG = f"D{P1}:D{P2}", f"E{P1}:E{P2}", f"F{P1}:F{P2}", f"B{P1}:B{P2}"
blk = Block(ws, SH, 9, 5)           # label column I, value J, formula K, note L
paired_block(blk, "Paired t test and d_z: Product_TV Dwell time (%)", RB, RP, RD, "%",
             "COUNT of the difference column; must equal the check row below")
blk.title("Order effect (simple within respondent estimate)")
blk.put("oe_A", "Mean difference, Group A (Paint seen second)", f'=AVERAGEIFS({RD},{RG},"A")',
        "difference = commercial effect + order effect")
blk.put("oe_B", "Mean difference, Group B (Paint seen first)", f'=AVERAGEIFS({RD},{RG},"B")',
        "difference = commercial effect - order effect")
blk.put("oe_nA", "n, Group A", f'=COUNTIF({RG},"A")', fmt="0")
blk.put("oe_nB", "n, Group B", f'=COUNTIF({RG},"B")', fmt="0")
blk.put("oe", "Order effect (second minus first)", "=({oe_A}-{oe_B})/2",
        "same as the lmer coefficient for order_position here; Excel cannot fit the mixed model")
blk.put("commercial", "Commercial effect (Paint minus Balls), order adjusted", "=({oe_A}+{oe_B})/2",
        "average of the two group differences")
blk.blank()
blk.title("Checks")
blk.put("chk_n", "Kept respondents with a Product_TV row for Balls in Data",
        f'=COUNTIFS({R_STIM},"{BALLS}",{R_AOI},"Product_TV",{R_KEEP},"Yes")',
        "must equal n pairs; if not, the roster in column A is out of date", fmt="0")
blk.put("chk_valid", "Smallest Valid data in column C", f"=MIN(C{P1}:C{P2})", "must be at least 70", fmt="0.0")
blk.blank()
for col, w in (("A", 17), ("B", 13), ("C", 11), ("D", 20), ("E", 20), ("F", 26), ("G", 11), ("H", 3),
               ("I", 54), ("J", 12), ("K", 72), ("L", 62)):
    ws.column_dimensions[col].width = w
# chart: mean Product_TV dwell share by commercial with 95% CI half widths
cr = blk.row + 1
ws.cell(row=cr, column=9, value="Chart table: mean dwell share with 95% CI half width").font = bold
ws.cell(row=cr + 1, column=9, value="Commercial")
ws.cell(row=cr + 1, column=10, value="Mean dwell (%)")
ws.cell(row=cr + 1, column=11, value="CI half width")
ws.cell(row=cr + 2, column=9, value="Balls")
ws.cell(row=cr + 2, column=10, value=f"={blk.addr('m_b')}")
ws.cell(row=cr + 2, column=11, value=f"=_xlfn.CONFIDENCE.T(0.05,{blk.addr('s_b')},{blk.addr('n')})")
ws.cell(row=cr + 3, column=9, value="Paint")
ws.cell(row=cr + 3, column=10, value=f"={blk.addr('m_p')}")
ws.cell(row=cr + 3, column=11, value=f"=_xlfn.CONFIDENCE.T(0.05,{blk.addr('s_p')},{blk.addr('n')})")
for rr in (cr + 2, cr + 3):
    ws.cell(row=rr, column=10).number_format = "0.00"
    ws.cell(row=rr, column=11).number_format = "0.00"
reg(SH, "chart_ci_b", f"K{cr + 2}", f"=_xlfn.CONFIDENCE.T(0.05,{blk.addr('s_b')},{blk.addr('n')})", "CI half width Balls")
reg(SH, "chart_ci_p", f"K{cr + 3}", f"=_xlfn.CONFIDENCE.T(0.05,{blk.addr('s_p')},{blk.addr('n')})", "CI half width Paint")
chart2 = BarChart()
chart2.type = "col"
chart2.style = 10
chart2.title = "Mean Product_TV dwell share by commercial (SIMULATED)"
chart2.y_axis.title = "Dwell time (% of AOI active time)"
chart2.add_data(Reference(ws, min_col=10, min_row=cr + 1, max_row=cr + 3), titles_from_data=True)
chart2.set_categories(Reference(ws, min_col=9, min_row=cr + 2, max_row=cr + 3))
chart2.series[0].graphicalProperties.solidFill = NAVY
chart2.series[0].errBars = ErrorBars(errDir="y", errBarType="both", errValType="cust",
                                      plus=NumDataSource(numRef=NumRef(f=f"ProductTV!$K${cr + 2}:$K${cr + 3}")),
                                      minus=NumDataSource(numRef=NumRef(f=f"ProductTV!$K${cr + 2}:$K${cr + 3}")))
chart2.legend = None
chart2.height = 8
chart2.width = 12
ws.add_chart(chart2, f"M{cr}")
PT = blk
PT_CHART_ROW = cr

# ---------------------------------------------------------------- BrandLogo
ws = wb.create_sheet("BrandLogo")
SH = "BrandLogo"
ws["A1"] = "H2. Brand_Logo TTFF AOI (ms), Balls versus Paint, respondents with a TTFF in both commercials. SIMULATED data."
ws["A1"].font = Font(bold=True, size=12, color=NAVY)
ws["A2"] = ("A respondent who never fixated the logo has NA, which cannot enter a paired test, so the roster below lists only "
            "kept respondents with a number in BOTH commercials. TTFF AOI counts from the start of the end card. "
            "Difference is Paint minus Balls.")
ws["A2"].font = note_font
bl_cols = ["Respondent Name", "Group (order)", "Balls TTFF AOI (ms)", "Paint TTFF AOI (ms)", "Difference (Paint minus Balls)"]
ws.append([])
ws.append(bl_cols)
style_header(ws, row=4, ncol=len(bl_cols))
B1 = 5
for k, rid in enumerate(both_logo):
    r = B1 + k
    ws.cell(row=r, column=1, value=rid)
    ws.cell(row=r, column=2, value=f"=INDEX({R_GROUP},MATCH($A{r},{R_RESP},0))")
    ws.cell(row=r, column=3, value=f'=SUMIFS({R_TTFF},{R_RESP},$A{r},{R_AOI},"Brand_Logo",{R_STIM},"{BALLS}")')
    ws.cell(row=r, column=4, value=f'=SUMIFS({R_TTFF},{R_RESP},$A{r},{R_AOI},"Brand_Logo",{R_STIM},"{PAINT}")')
    ws.cell(row=r, column=5, value=f"=D{r}-C{r}")
    for c in range(1, 6):
        ws.cell(row=r, column=c).border = box
        if c >= 3:
            ws.cell(row=r, column=c).number_format = "0"
B2 = B1 + len(both_logo) - 1
for key, col, f, lab in (("balls", "C", ws.cell(row=B1, column=3).value, bl_cols[2]),
                         ("paint", "D", ws.cell(row=B1, column=4).value, bl_cols[3]),
                         ("diff", "E", f"=D{B1}-C{B1}", bl_cols[4])):
    reg(SH, "row_" + key, f"{col}{B1}", f, lab)
blk = Block(ws, SH, 7, 5)           # label column G, value H, formula I, note J
paired_block(blk, "Paired t test and d_z: Brand_Logo TTFF AOI (ms)", f"C{B1}:C{B2}", f"D{B1}:D{B2}", f"E{B1}:E{B2}", "ms",
             "pairs only: respondents with a TTFF in both commercials")
blk.title("Respondent ratio (kept respondents who fixated the logo)")
crit_l = f'{R_AOI},"Brand_Logo",{R_KEEP},"Yes"'
blk.put("kept_b", "Kept respondents with a Brand_Logo row, Balls", f'=COUNTIFS({crit_l},{R_STIM},"{BALLS}")', fmt="0")
blk.put("fix_b", "Fixated the logo, Balls", f'=COUNTIFS({crit_l},{R_STIM},"{BALLS}",{R_FIX},1)', fmt="0")
blk.put("ratio_b", "Respondent ratio, Balls (%)", "=100*{fix_b}/{kept_b}", "fixation based Respondent ratio (%)", fmt="0.0")
blk.put("kept_p", "Kept respondents with a Brand_Logo row, Paint", f'=COUNTIFS({crit_l},{R_STIM},"{PAINT}")', fmt="0")
blk.put("fix_p", "Fixated the logo, Paint", f'=COUNTIFS({crit_l},{R_STIM},"{PAINT}",{R_FIX},1)', fmt="0")
blk.put("ratio_p", "Respondent ratio, Paint (%)", "=100*{fix_p}/{kept_p}", fmt="0.0")
blk.blank()
blk.title("Checks")
blk.put("chk_both", "Kept respondents with a TTFF in both commercials (counted from Data)",
        f'=SUMPRODUCT(--(COUNTIFS({R_RESP},ProductTV!$A${P1}:$A${P2},{R_AOI},"Brand_Logo",{R_STIM},"{BALLS}",{R_FIX},1)=1),'
        f'--(COUNTIFS({R_RESP},ProductTV!$A${P1}:$A${P2},{R_AOI},"Brand_Logo",{R_STIM},"{PAINT}",{R_FIX},1)=1))',
        "must equal n pairs; checks the roster against every kept respondent listed on the ProductTV sheet", fmt="0")
blk.blank()
for col, w in (("A", 17), ("B", 13), ("C", 20), ("D", 20), ("E", 26), ("F", 3), ("G", 54), ("H", 12), ("I", 72), ("J", 62)):
    ws.column_dimensions[col].width = w
BL = blk

# ---------------------------------------------------------------- Survey
sb, sh, srows = read_csv("SIMULATED_ad_test_survey.csv")
ws = wb.create_sheet("Survey")
SH = "Survey"
ws["A1"] = "H3. Liking after each commercial, joined to Brand_Logo dwell time. SIMULATED data."
ws["A1"].font = Font(bold=True, size=12, color=NAVY)
ws.append([])
sv_cols = sh + ["Keep", "Difference (Paint minus Balls)", "Brand_Logo Dwell time (ms), both commercials", "Recall (1 = Yes)"]
ws.append(sv_cols)
style_header(ws, row=3, ncol=len(sv_cols))
sidx = {h: i for i, h in enumerate(sh)}
srows.sort(key=lambda rw: (rw[sidx["Respondent Name"]] in excluded, rw[sidx["Respondent Name"]]))   # excluded last
S1 = 4
for k, rw in enumerate(srows):
    r = S1 + k
    for c, v in enumerate(rw, start=1):
        ws.cell(row=r, column=c, value=to_cell(v))
    ws.cell(row=r, column=8, value=f'=IF(INDEX({R_VALID},MATCH($B{r},{R_RESP},0))>=70,"Yes","No")')
    ws.cell(row=r, column=9, value=f"=E{r}-D{r}")
    ws.cell(row=r, column=10, value=f'=SUMIFS({R_DWELL_MS},{R_RESP},$B{r},{R_AOI},"Brand_Logo")')
    ws.cell(row=r, column=11, value=f'=IF(F{r}="Yes",1,0)')
S_last_all = S1 + len(srows) - 1
S2 = S1 + len(kept) - 1            # last kept row (excluded respondents are sorted below it)
for key, col, f, lab in (("keep", "H", ws.cell(row=S1, column=8).value, "Keep"),
                         ("diff", "I", f"=E{S1}-D{S1}", "Difference (Paint minus Balls)"),
                         ("logo", "J", ws.cell(row=S1, column=10).value, sv_cols[9]),
                         ("recall01", "K", f'=IF(F{S1}="Yes",1,0)', "Recall (1 = Yes)")):
    reg(SH, "row_" + key, f"{col}{S1}", f, lab)
ws.cell(row=S_last_all + 2, column=1, value=(
    f"Rows {S1} to {S2} are the {len(kept)} kept respondents; the excluded respondent ({', '.join(excluded)}) is listed "
    f"last (row {S_last_all}) and left out of every test range.")).font = note_font
sblk = Block(ws, SH, 13, 4)          # label column M, value N, formula O, note P
LB, LP, LD = f"D{S1}:D{S2}", f"E{S1}:E{S2}", f"I{S1}:I{S2}"
paired_block(sblk, "Paired t test and d_z: Liking (1 to 7)", LB, LP, LD, "points", "kept respondents only")
sblk.title("Preference and ratings (kept respondents)")
sblk.put("pref_b", "Ad_preferred = Balls (count)", f'=COUNTIF(G{S1}:G{S2},"Balls")', fmt="0")
sblk.put("pref_p", "Ad_preferred = Paint (count)", f'=COUNTIF(G{S1}:G{S2},"Paint")', fmt="0")
sblk.put("lik_b_hi", "Rated Balls higher (count)", f"=SUMPRODUCT(--({LB}>{LP}))", fmt="0")
sblk.put("lik_p_hi", "Rated Paint higher (count)", f"=SUMPRODUCT(--({LP}>{LB}))", fmt="0")
sblk.put("lik_tie", "Equal ratings (count)", f"=SUMPRODUCT(--({LB}={LP}))", fmt="0")
sblk.blank()
sblk.title("Brand recall and Brand_Logo dwell time (point biserial)")
sblk.put("rec_yes", "Brand_recall = Yes (count)", f'=COUNTIF(F{S1}:F{S2},"Yes")', fmt="0")
sblk.put("rec_n", "Kept respondents", f"=COUNTA(B{S1}:B{S2})", fmt="0")
sblk.put("rec_pct", "Percent Yes", "=100*{rec_yes}/{rec_n}", fmt="0.0")
sblk.put("mean_yes", "Mean total logo dwell (ms), recall Yes", f'=AVERAGEIFS(J{S1}:J{S2},F{S1}:F{S2},"Yes")', fmt="0")
sblk.put("mean_no", "Mean total logo dwell (ms), recall No", f'=AVERAGEIFS(J{S1}:J{S2},F{S1}:F{S2},"No")', fmt="0")
sblk.put("rpb", "Point biserial r (CORREL)", f"=CORREL(J{S1}:J{S2},K{S1}:K{S2})",
         "Pearson r of a continuous variable with a 0/1 variable")
sblk.put("rpb_t", "t for r", "={rpb}*SQRT(({rec_n}-2)/(1-{rpb}^2))", "t = r sqrt((n - 2) / (1 - r^2))")
sblk.put("rpb_p", "p, two tailed", "=_xlfn.T.DIST.2T(ABS({rpb_t}),{rec_n}-2)", "df = n - 2", fmt="0.000000")
sblk.put("rpb_lo", "r 95% CI lower (Fisher z)", "=TANH(ATANH({rpb})-_xlfn.NORM.S.INV(0.975)/SQRT({rec_n}-3))")
sblk.put("rpb_hi", "r 95% CI upper (Fisher z)", "=TANH(ATANH({rpb})+_xlfn.NORM.S.INV(0.975)/SQRT({rec_n}-3))")
sblk.blank()
for col, w in (("A", 34), ("B", 15), ("C", 8), ("D", 18), ("E", 18), ("F", 13), ("G", 14), ("H", 8), ("I", 18), ("J", 22),
               ("K", 14), ("L", 3), ("M", 54), ("N", 12), ("O", 72), ("P", 52)):
    ws.column_dimensions[col].width = w
ws.row_dimensions[3].height = 48
SV = sblk

# ---------------------------------------------------------------- Check
ws = wb.create_sheet("Check")
ws["A1"] = "Live Excel results next to the values analysis.R produced (SIMULATED data). Status is OK when they agree within the tolerance."
ws["A1"].font = Font(bold=True, size=12, color=NAVY)
ws.append([])
ck_cols = ["Quantity", "Excel (live formula)", "analysis.R", "Difference", "Tolerance", "Status"]
ws.append(ck_cols)
style_header(ws, row=3, ncol=len(ck_cols))
_, th, trows = read_csv("SIMULATED_test_results.csv")
_, mh, mrows = read_csv("SIMULATED_model_results.csv")
T = {r[0].split(" ")[0]: dict(zip(th, r)) for r in trows if not r[0].startswith("H2b")}      # keys H1, H2, H3
M = {(r[0], r[1]): dict(zip(mh, r)) for r in mrows}
mkey = ("Mixed model Product_TV dwell_pct", "Order position (second minus first)")
pb = M[("Point biserial: total Brand_Logo dwell (ms) with recall", "r_pb")]
ck_rows = []


def add_check(label, sheet, blkobj, key, r_value, tol, note=""):
    ck_rows.append((label, f"={sheet}!{blkobj.addr(key)}", float(r_value), tol, note))


for hid, sheet, blkobj, nm in (("H1", "ProductTV", PT, "Product_TV dwell (%)"), ("H2", "BrandLogo", BL, "Brand_Logo TTFF (ms)"),
                               ("H3", "Survey", SV, "Liking")):
    t = T[hid]
    add_check(f"{hid} {nm}: mean Balls", sheet, blkobj, "m_b", t["m_balls"], 1e-4)
    add_check(f"{hid} {nm}: mean Paint", sheet, blkobj, "m_p", t["m_paint"], 1e-4)
    add_check(f"{hid} {nm}: mean difference", sheet, blkobj, "md", t["mean_diff"], 1e-4)
    add_check(f"{hid} {nm}: SD of differences", sheet, blkobj, "sdd", t["sd_diff"], 1e-4)
    add_check(f"{hid} {nm}: paired t", sheet, blkobj, "t", t["t"], 1e-4)
    add_check(f"{hid} {nm}: df", sheet, blkobj, "df", t["df"], 1e-9)
    add_check(f"{hid} {nm}: p (T.TEST type 1)", sheet, blkobj, "p_tt", t["p"], 5e-4 * float(t["p"]) + 1e-9)
    add_check(f"{hid} {nm}: mean difference CI lower", sheet, blkobj, "ci_lo", t["diff_ci_low"], 1e-4)
    add_check(f"{hid} {nm}: mean difference CI upper", sheet, blkobj, "ci_hi", t["diff_ci_high"], 1e-4)
    add_check(f"{hid} {nm}: Cohen's d_z", sheet, blkobj, "dz", t["dz"], 1e-4)
    add_check(f"{hid} {nm}: d_z CI lower (method differs)", sheet, blkobj, "dz_lo", t["dz_ci_low"], 0.10,
              "Excel: large sample approximation; R: noncentral t")
    add_check(f"{hid} {nm}: d_z CI upper (method differs)", sheet, blkobj, "dz_hi", t["dz_ci_high"], 0.10,
              "Excel: large sample approximation; R: noncentral t")
add_check("Order effect: simple estimate versus lmer order_position", "ProductTV", PT, "oe", M[mkey]["estimate"], 1e-3,
          "equal here; in general lmer weights groups differently")
add_check("Point biserial r (recall with total logo dwell)", "Survey", SV, "rpb", pb["estimate"], 1e-4)
add_check("Point biserial: t", "Survey", SV, "rpb_t", pb["statistic"], 1e-3)
add_check("Point biserial: p", "Survey", SV, "rpb_p", pb["p"], 5e-4 * float(pb["p"]) + 1e-9)
add_check("Point biserial: CI lower", "Survey", SV, "rpb_lo", pb["ci_low"], 1e-4)
add_check("Point biserial: CI upper", "Survey", SV, "rpb_hi", pb["ci_high"], 1e-4)
CK1 = 4
for k, (lab, f, rv, tol, note) in enumerate(ck_rows):
    r = CK1 + k
    ws.cell(row=r, column=1, value=lab)
    ws.cell(row=r, column=2, value=f).number_format = "0.000000"
    ws.cell(row=r, column=3, value=rv).number_format = "0.000000"
    ws.cell(row=r, column=4, value=f"=B{r}-C{r}").number_format = "0.000000"
    ws.cell(row=r, column=5, value=tol).number_format = "0.0000"
    ws.cell(row=r, column=6, value=f'=IF(ABS(D{r})<=E{r},"OK","CHECK")')
    if note:
        ws.cell(row=r, column=7, value=note).font = note_font
CK2 = CK1 + len(ck_rows) - 1
ws.cell(row=CK2 + 2, column=1, value="Rows with status CHECK").font = bold
ws.cell(row=CK2 + 2, column=2, value=f'=COUNTIF(F{CK1}:F{CK2},"CHECK")')
reg("Check", "n_check", f"B{CK2 + 2}", f'=COUNTIF(F{CK1}:F{CK2},"CHECK")', "Rows with status CHECK")
reg("Check", "status_row", f"F{CK1}", f'=IF(ABS(D{CK1})<=E{CK1},"OK","CHECK")', "Status")
ws.cell(row=CK2 + 3, column=1, value=("Not reproducible in Excel: the Wilcoxon signed rank tests, the sensitivity analysis with never fixators "
                                      "set to 5000 ms, the exact binomial and McNemar tests, the mixed model with its ICC, and the "
                                      "noncentral t interval for d_z. See analysis_output.txt.")).font = note_font
# static copies of the R tables
r = CK2 + 6
ws.cell(row=r - 1, column=1, value="analysis.R: paired tests (SIMULATED_test_results.csv)").font = bold
for c, h in enumerate(th, start=1):
    ws.cell(row=r, column=c, value=h)
style_header(ws, row=r, ncol=len(th))
for rw in trows:
    r += 1
    for c, v in enumerate(rw, start=1):
        ws.cell(row=r, column=c, value=to_cell(v))
r += 3
ws.cell(row=r - 1, column=1, value="analysis.R: mixed model, point biserial, exploratory (SIMULATED_model_results.csv)").font = bold
for c, h in enumerate(mh, start=1):
    ws.cell(row=r, column=c, value=h)
style_header(ws, row=r, ncol=len(mh))
for rw in mrows:
    r += 1
    for c, v in enumerate(rw, start=1):
        ws.cell(row=r, column=c, value=to_cell(v))
r += 3
_, dh, drows = read_csv("SIMULATED_descriptives_by_commercial_AOI.csv")
ws.cell(row=r - 1, column=1, value="analysis.R: descriptives (SIMULATED_descriptives_by_commercial_AOI.csv)").font = bold
for c, h in enumerate(dh, start=1):
    ws.cell(row=r, column=c, value=h)
style_header(ws, row=r, ncol=len(dh))
for rw in drows:
    r += 1
    for c, v in enumerate(rw, start=1):
        ws.cell(row=r, column=c, value=to_cell(v))
autowidth(ws, max_w=34)
ws.column_dimensions["A"].width = 58
for col in "BCDEF":
    ws.column_dimensions[col].width = 20

wb.calculation = CalcProperties(fullCalcOnLoad=True)    # Excel computes every formula when the file opens
wb.save(OUT)
print("wrote", OUT)

# =============================================================================
# The walkthrough, generated from REG so that every cell reference is the real one
# =============================================================================
_, _, tr = read_csv("SIMULATED_test_results.csv")


def cl(sheet, key):
    return REG[(sheet, key)]["cell"]


def fx(sheet, key):
    return REG[(sheet, key)]["formula"]


def row(label, sheet, key, result=""):
    return f"| {label} | `{sheet}!{cl(sheet, key)}` | `{fx(sheet, key)}` | {result} |"


def row3(label, sheet, key):
    return f"| {label} | `{sheet}!{cl(sheet, key)}` | `{fx(sheet, key)}` |"


def rcell(sheet, key):
    return f"{sheet}!{cl(sheet, key)}"


# numbers for the text, all from analysis.R's CSV output or computed from the data
def fnum(x, nd=2):
    return f"{float(x):.{nd}f}"


def sd_of(vals):
    return st.stdev(vals)


def paired_values(stim_col_values):
    return stim_col_values


h1, h2, h3 = T["H1"], T["H2"], T["H3"]
b_ptv = [float(value_of(k, BALLS, "Product_TV", "Dwell time (%)")) for k in kept]
p_ptv = [float(value_of(k, PAINT, "Product_TV", "Dwell time (%)")) for k in kept]
b_logo = [float(value_of(k, BALLS, "Brand_Logo", "TTFF AOI (ms)")) for k in both_logo]
p_logo = [float(value_of(k, PAINT, "Brand_Logo", "TTFF AOI (ms)")) for k in both_logo]
sv_by = {rw[sidx["Respondent Name"]]: rw for rw in srows}
b_lik = [float(sv_by[k][sidx["Liking_Balls_1to7"]]) for k in kept]
p_lik = [float(sv_by[k][sidx["Liking_Paint_1to7"]]) for k in kept]


def se_dz(n, dz):
    return (1 / n + dz ** 2 / (2 * n)) ** 0.5


def approx_ci(h):
    n, dz = int(h["n"]), float(h["dz"])
    s = se_dz(n, dz)
    return dz - 1.96 * s, dz + 1.96 * s, s


a1, a2, a3 = approx_ci(h1), approx_ci(h2), approx_ci(h3)
pT = lambda h: f"{float(h['p']):.6f}"
nrow_data = n_data
excl_txt = ", ".join(excluded)
mm_order = M[mkey]
mm_comm = M[("Mixed model Product_TV dwell_pct", "Commercial (Paint minus Balls)")]
icc_adj = M[("Mixed model Product_TV dwell_pct", "ICC adjusted (fitted model)")]["estimate"]
pb_r, pb_p = pb["estimate"], pb["p"]
sumrow_ptv_b = summary_rows[("Product_TV", BALLS)]
sumrow_ptv_p = summary_rows[("Product_TV", PAINT)]
sumrow_logo_p = summary_rows[("Brand_Logo", PAINT)]

# values on the Summary sheet for the "You should see" lines, read back from the R descriptives
drow = {(r_[0], r_[1]): dict(zip(dh, r_)) for r_ in drows}
d_ptv_b, d_ptv_p = drow[("Product_TV", "Balls")], drow[("Product_TV", "Paint")]
d_logo_b, d_logo_p = drow[("Brand_Logo", "Balls")], drow[("Brand_Logo", "Paint")]
n_fix_b = sum(1 for k in kept if value_of(k, BALLS, "Brand_Logo", "TTFF AOI (ms)") != "NA")
n_fix_p = sum(1 for k in kept if value_of(k, PAINT, "Brand_Logo", "TTFF AOI (ms)") != "NA")
grp_of = {k: by_resp[k][0][idx["Group"]] for k in kept}
dA = [pp - bb for k, bb, pp in zip(kept, b_ptv, p_ptv) if grp_of[k] == "A"]
dB = [pp - bb for k, bb, pp in zip(kept, b_ptv, p_ptv) if grp_of[k] == "B"]
n_ties = sum(1 for x, y in zip(b_lik, p_lik) if x == y)
tcrit_df = (float(h1["diff_ci_high"]) - float(h1["diff_ci_low"])) / 2 / (float(h1["sd_diff"]) / int(h1["n"]) ** 0.5)
h2_split = (float(h2["p"]) < 0.05) != (float(h2["wilcoxon_p"]) < 0.05)
h2_txt = (f"in this sample the signed rank test (p = {float(h2['wilcoxon_p']):.3f}) and the paired t test "
          f"(p = {float(h2['p']):.3f}) fall on opposite sides of .05, so run it in R before you draw a conclusion"
          if h2_split else
          f"in this sample the signed rank test (p = {float(h2['wilcoxon_p']):.3f}) agrees with the paired t test "
          f"(p = {float(h2['p']):.3f}), but run it in R to confirm")
DZ = "Cohen's d_z"

md = f"""# Ad test analysis in Excel: step by step walkthrough

**SIMULATED DATA.** Every number in this walkthrough comes from `SIMULATED_ad_test_AOI_metrics.csv`, which `make_simulated_data.py` generated (seed {SEED}). No real viewer was recorded. The design follows the structure of the iMotions demonstration study "DEMO - SonyBravia" (one brand, two television commercials, an eye tracker on a monitor); no data from that study were used.

This is the same analysis as `analysis.R`, done in Excel for students who do not use R. The finished workbook is `SIMULATED_ad_test_analysis.xlsx`; the sheet names below refer to it. You can follow along in that workbook or rebuild it from the CSV files. **Every cell address in this walkthrough is written as `Sheet!Cell` and was generated from the workbook itself** (`make_excel_workbook.py` writes both), so the addresses match the file.

Companion files: `analysis.R` (R version), `analysis_output.txt` (its console output), `results_text.md` (model paragraphs), `figures/` (PNG figures).

---

## 0. What you are analyzing

**Design.** Within subjects. Each respondent watched both television commercials for the same brand, `Commercial_1_Balls` and `Commercial_2_Paint` (60 s each in the simulation), on the lab monitor with the Smart Eye Aurora (60 Hz). A Baseline slide came before each commercial and a rating survey slide after each. The order was counterbalanced: Group A saw Balls first, Group B saw Paint first. 30 respondents, R01 to R30, 15 per order. In Workshop 1 the groups were different people (a Welch t test); here the two commercials are measured on the *same* people, so the analysis is a **paired** test and the order of viewing must be checked.

**Export.** One row per respondent by stimulus by AOI, the layout of the iMotions **Individual AOI metrics** export, which iMotions recommends for statistics because it keeps each respondent's values instead of averaging across the segment (Areas of Interest article, Export options). There are {nrow_data} data rows: 30 respondents x 2 commercials x 4 AOIs. The columns after the identifiers are iMotions AOI metrics with the names from the **AOI Metrics** help article. The ones used here:

| Column | iMotions definition (AOI Metrics article) | Used for |
|---|---|---|
| `Valid data` | Percentage of collected and interpolated samples; below 100 means missing data | Exclusion rule |
| `AOI duration (ms)` and `AOI duration (%)` | How long the (dynamic) AOI was active, and that time as a percentage of the stimulus duration | Interpreting Dwell time (%) |
| `TTFF AOI (ms)` | Time from the AOI's start until the first fixation on the AOI | H2 (Brand_Logo) |
| `TTFF parent (ms)` | The same, counted from the start of the AOI's parent (here the commercial) | Not used (see below) |
| `TTFF max. (ms)` | Counted from the start of the parent; a respondent who never fixated the AOI gets the parent's full presentation duration (60000 ms here) | Not used (see below) |
| `Dwell time (ms)` | Total time fixating on the AOI | Brand_Logo total dwell and recall |
| `Dwell time (%)` | Dwell time relative to the time during which the AOI was active | H1 (Product_TV) |
| `Fixation count` | Number of fixations inside the AOI | Descriptives |

Where the same metric name exists as a gaze based and a fixation based metric (Dwell time, Revisit count), the gaze based column carries the prefix `Gaze`. iMotions writes **NA** when a metric cannot be calculated, for example when the respondent never fixated the AOI. The simulated file does the same.

**Dynamic AOIs.** An AOI that follows an element on screen is active only while the element is shown (Areas of Interest article, dynamic AOIs). Product_TV (the television set in the product shots) is active for 14 s of the 60 s in Balls and 20 s in Paint; Brand_Logo for the last 5 s; Tagline_Text for the last 4 s; Hero_Element (the bouncing balls or the paint explosion) for 40 s in Balls and 32 s in Paint. `Dwell time (%)` divides by this active time, which is why it can be compared across the two commercials even though the AOIs are on screen for different lengths.

**Hypotheses.**
- H1: the share of dwell time on `Product_TV` differs between the commercials (the simulation built in a larger share for Paint, which shows the set for 20 s against 14 s).
- H2: `TTFF AOI (ms)` to `Brand_Logo` (from the start of the end card) differs between the commercials.
- H3: liking (1 to 7) differs between the commercials.
- Order effect: the commercial seen second gets extra Product_TV dwell share (carryover from the first viewing).

**Decision rules fixed before looking at the data.** Exclude a respondent whose `Valid data` is below 70 percent. For an AOI the respondent never fixated, TTFF is *not observed* (leave it out of TTFF statistics) and dwell time is *0* (keep it in dwell statistics). Every difference is **Paint minus Balls**. Two tailed tests, alpha = .05.

---

## 1. Import the export into Excel

iMotions exports are comma separated with a dot decimal. Double clicking the file uses your regional defaults, which can put everything in one column (Opening Data Export Files in Excel article). Use the import route instead.

**Route A, Get Data (Excel 2016 and later):**
1. **Data > From Text/CSV**, select `SIMULATED_ad_test_AOI_metrics.csv`, click **Import**.
2. In the preview set **Delimiter: Comma** and **File Origin: 65001 Unicode (UTF-8)**. **You should see** {len(header)} columns with `Study Name` in the first column.
3. Click **Transform Data** if the first row shows the `# SIMULATED ...` comment line: in Power Query choose **Home > Remove Rows > Remove Top Rows > 1**, then **Home > Use First Row as Headers**, then **Close & Load**. (A real iMotions export has no comment line; the simulated file carries one so nobody mistakes it for real data.)
4. **You should see** a table with {nrow_data} data rows.

**Route B, Text Import Wizard (any version):** **File > Open > Browse**, set the file type to **All Files**, open the .csv, choose **Delimited**, start at row 2 (to skip the comment line), tick **Comma**, click **Advanced** and set the decimal symbol to a dot, then **Finish**.

**Common mistake.** The NA cells must stay as the text `NA` (or be blank). Do not replace NA with 0 in a TTFF column: a zero TTFF would mean the respondent fixated the logo at the very first instant, which is false. Step 2 handles NA correctly.

Save as **.xlsx** right away; a .csv does not keep formulas or extra sheets. Keep the original .csv read only.

In the workbook this is the **Data** sheet, header in row 1 and data in rows {FIRST} to {LAST}. Column letters there: `{COL['Respondent Name']}` = Respondent Name, `{COL['Group']}` = Group (the order), `{COL['Stimulus']}` = Stimulus, `{COL['AOI Name']}` = AOI Name, `{COL['Valid data']}` = Valid data, `{COL['Fixation count']}` = Fixation count, `{COL['TTFF AOI (ms)']}` = TTFF AOI (ms), `{COL['Dwell time (ms)']}` = Dwell time (ms), `{COL['Dwell time (%)']}` = Dwell time (%).

---

## 2. Clean with helper columns

Add five columns to the right of the data (`{HCOL['Keep']}` to `{HCOL['Fix_count_clean']}` in the workbook, headers in orange) and fill them down to the last row. The formulas below are the ones in row 2.

| Header | Cell | Formula | Meaning |
|---|---|---|---|
| Keep | `Data!{cl('Data','keep')}` | `{fx('Data','keep')}` | Exclusion rule on Valid data |
| Position | `Data!{cl('Data','position')}` | `{fx('Data','position')}` | Was this commercial the first or the second one this respondent saw? Group A saw Balls first, Group B saw Paint first |
| Dwell_pct_clean | `Data!{cl('Data','dwell_clean')}` | `{fx('Data','dwell_clean')}` | Never fixated AOI counts as 0 percent |
| Fixated | `Data!{cl('Data','fixated')}` | `{fx('Data','fixated')}` | 1 when at least one fixation hit the AOI |
| Fix_count_clean | `Data!{cl('Data','fixc')}` | `{fx('Data','fixc')}` | Fixation count with NA counted as 0 |

**You should see** `Keep = No` on all eight rows of respondent {excl_txt} (Valid data 58 percent; the session log says calibration stayed Poor after two recalibrations and the session was recorded anyway). Every other respondent is `Yes`. Respondents R09 and R17 needed one recalibration (Poor, then Good) but have complete data and stay in; a recalibration is not by itself a reason to exclude.

`ISNUMBER` is the key: it is TRUE for a number and FALSE for the text NA or a blank cell, so one formula handles both ways an export can mark a missing metric.

---

## 3. Descriptive table by commercial and AOI

### 3a. With a PivotTable (the iMotions recommended route)

The Creating Excel Pivot Tables article gives the generic steps: click the first cell of the data, **Insert > PivotTable > OK**, drag a category to **Rows**, drag a metric to **Values**, right click the value, **Value Field Settings**, change **Sum** to **Average**, **OK**, then **PivotChart** if you want a quick picture. For this study:

1. Click cell A1 of the Data sheet, **Insert > PivotTable**, **New Worksheet**, **OK**.
2. Drag **Keep** to **Filters** and set it to **Yes**.
3. Drag **AOI Name** to **Rows** and **Stimulus** to **Columns**.
4. Drag **TTFF AOI (ms)** to **Values** three times. Set the first to **Average**, the second to **StdDev** (the sample standard deviation, the same as STDEV.S), the third to **Count**. The text NA is ignored by Average and StdDev, but **Count** counts every non blank cell including the text NA; use **Count Numbers** (Value Field Settings > Summarize Values By > Count Numbers) for the number of fixators, or the formula table in 3b.
5. Drag **Dwell_pct_clean** to **Values** twice: **Average** and **StdDev**.
6. Drag **Fixated** to **Values** and set it to **Average**; multiply by 100 beside the table. This is the fixation based **Respondent ratio (%)** from the iMotions AOI table.

**You should see** `Product_TV`: Balls average Dwell_pct_clean {d_ptv_b['dwell_mean_pct']} and Paint {d_ptv_p['dwell_mean_pct']}; `Brand_Logo`: respondent ratio {d_logo_b['resp_ratio_pct']} percent for Balls and {d_logo_p['resp_ratio_pct']} percent for Paint.

**Common mistake.** Leaving **Keep** out of the filter. {excl_txt} would then add eight rows of thin data.

### 3b. With formulas (the Summary sheet)

A formula table updates when the data change and shows exactly what is being averaged. The AOI name is in column A and the stimulus in column B of each row (row {SUM_FIRST} is the first). The formulas in row {SUM_FIRST} (`Product_TV`, `{BALLS}`):

| Statistic | Cell | Formula |
|---|---|---|
""" + "\n".join(
    f"| {REG[('Summary', k)]['label']} | `Summary!{cl('Summary', k)}` | `{fx('Summary', k)}` |"
    for k in ("n", "resp_ratio", "n_ttff", "ttff_mean", "ttff_sd", "dwell_mean", "dwell_sd", "fix_mean")) + f"""

The two `STDEV.S(IF(...))` formulas (columns G and I) are **array formulas** because Excel has no STDEVIFS. In Excel 365 type them and press **Enter**. In Excel 2019 or older press **Ctrl+Shift+Enter**; Excel then shows the formula in braces. `AVERAGEIFS` and `COUNTIFS` ignore the text NA, which is exactly the "TTFF not observed" rule.

**You should see** the same numbers as the PivotTable, in rows {SUM_FIRST} to {SUM_LAST} (two rows per AOI), and a clustered bar chart of mean dwell share by AOI and commercial under the table.

---

## 4. H1: Product_TV dwell share, paired t test

A paired test needs the two measurements of the same respondent side by side, so first reshape the long export into one row per respondent. The **ProductTV** sheet does this for the {len(kept)} kept respondents (listed in column A, rows {P1} to {P2}); each of columns D and E is a `SUMIFS` that picks the one matching Product_TV row from the Data sheet:

| Column | Cell (first respondent) | Formula |
|---|---|---|
""" + "\n".join(
    f"| {REG[(SH_, k)]['label']} | `ProductTV!{REG[(SH_, k)]['cell']}` | `{REG[(SH_, k)]['formula']}` |"
    for SH_, k in (("ProductTV", "row_grp"), ("ProductTV", "row_valid"), ("ProductTV", "row_balls"), ("ProductTV", "row_paint"),
                   ("ProductTV", "row_diff"), ("ProductTV", "row_seen"))) + f"""

(By hand you would get the same table from the PivotTable of section 3a: put **Respondent Name** in Rows, **Stimulus** in Columns, **Dwell_pct_clean** in Values as Sum, and filter **AOI Name** to `Product_TV` and **Keep** to `Yes`.)

Now the test. The statistics sit in column J of the ProductTV sheet with the formula text shown beside them in column K:

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('Balls mean (%)', 'ProductTV', 'm_b', fnum(h1['m_balls']))}
{row('Balls SD (%)', 'ProductTV', 's_b', fnum(sd_of(b_ptv)))}
{row('Paint mean (%)', 'ProductTV', 'm_p', fnum(h1['m_paint']))}
{row('Paint SD (%)', 'ProductTV', 's_p', fnum(sd_of(p_ptv)))}
{row('n pairs', 'ProductTV', 'n', h1['n'])}
{row('Mean difference, Paint minus Balls (%)', 'ProductTV', 'md', fnum(h1['mean_diff']))}
{row('SD of the differences (%)', 'ProductTV', 'sdd', fnum(h1['sd_diff']))}
{row('**p, two tailed (T.TEST type 1)**', 'ProductTV', 'p_tt', '**' + pT(h1) + '**')}

`T.TEST(array1, array2, tails, type)`: `tails = 2` for a two tailed test; `type = 1` for **paired** data, which is the type for this within subjects design. `type = 2` (equal variances) and `type = 3` (unequal variances, the Welch test of Workshop 1) treat the two columns as different people and give the wrong answer for paired data. `T.TEST` returns only the p value, so compute t and df by hand with the paired formulas (a paired t test is a one sample t test on the differences):

| Statistic | Cell | Formula | Result |
|---|---|---|---|
{row('SE of the mean difference', 'ProductTV', 'se', fnum(float(h1['sd_diff']) / int(h1['n']) ** 0.5, 3))}
{row('Paired t', 'ProductTV', 't', fnum(h1['t']))}
{row('df', 'ProductTV', 'df', h1['df'])}
{row('p from the hand computed t', 'ProductTV', 'p_hand', pT(h1) + ', matches T.TEST')}
{row('Critical t (95%)', 'ProductTV', 'tcrit', fnum(tcrit_df, 3))}
{row('Mean difference 95% CI lower', 'ProductTV', 'ci_lo', fnum(h1['diff_ci_low']))}
{row('Mean difference 95% CI upper', 'ProductTV', 'ci_hi', fnum(h1['diff_ci_high']))}

**You should see** the hand computed p equal the `T.TEST` p. If they differ, the `type` argument or a column range is wrong.

### Cohen's d_z by hand, with a 95 percent confidence interval

For paired data the effect size is **d_z**: the mean difference in units of the SD of the *differences*. It is linked to t by t = d_z x sqrt(n), so t, n, and d_z always agree.

| Step | Cell | Formula | Result |
|---|---|---|---|
{row(DZ, 'ProductTV', 'dz', fnum(h1['dz']))}
{row('SE of d_z (approximation)', 'ProductTV', 'se_dz', fnum(a1[2], 3))}
{row('d_z 95% CI lower', 'ProductTV', 'dz_lo', fnum(a1[0]))}
{row('d_z 95% CI upper', 'ProductTV', 'dz_hi', fnum(a1[1]))}
{row('Correlation of the two columns', 'ProductTV', 'r', fnum(st.correlation(b_ptv, p_ptv)))}

In words: d_z = mean difference / SD of differences; SE(d_z) = sqrt(1/n + d_z^2 / (2 n)). The interval is the usual large sample approximation. `analysis.R` reports the interval from the noncentral t distribution (effectsize package), which is more exact at this sample size; d_z itself matches, and the limits differ slightly ([{fnum(h1['dz_ci_low'])}, {fnum(h1['dz_ci_high'])}] in R). Say which method you used. Rule of thumb for reading d_z: 0.2 small, 0.5 medium, 0.8 large. The correlation in the last row is the within respondent correlation: respondents who looked at the television a lot in one commercial also did in the other, which is why a paired test is more powerful than a test of two independent groups.

---

## 5. The order effect

Because every respondent saw both commercials, the second commercial may have benefited from the first (carryover). The simulated viewers were built with about +3 percentage points of Product_TV dwell share for whichever commercial came second. The paired differences carry this: for Group A (Paint seen second) the difference is commercial effect + order effect, and for Group B (Paint seen first) it is commercial effect - order effect. So half the gap between the two group means is the order effect:

| Statistic | Cell | Formula | Result |
|---|---|---|---|
{row('Mean difference, Group A (Paint seen second)', 'ProductTV', 'oe_A', fnum(st.mean(dA)))}
{row('Mean difference, Group B (Paint seen first)', 'ProductTV', 'oe_B', fnum(st.mean(dB)))}
{row('n, Group A', 'ProductTV', 'oe_nA', len(dA))}
{row('n, Group B', 'ProductTV', 'oe_nB', len(dB))}
{row('Order effect (second minus first)', 'ProductTV', 'oe', fnum(mm_order['estimate']))}
{row('Commercial effect (Paint minus Balls), order adjusted', 'ProductTV', 'commercial', fnum(mm_comm['estimate']))}

`analysis.R` fits `dwell_pct ~ commercial * order_position + (1 | respondent)` with lme4 and lmerTest and gets the order effect b = {fnum(mm_order['estimate'])}, SE = {fnum(mm_order['se'])}, p = {fnum(mm_order['p'], 3)}, and the same commercial effect {fnum(mm_comm['estimate'])}; with 14 and 15 respondents in the two order groups the simple Excel estimate coincides with the model estimate. Excel cannot supply the standard error with the right degrees of freedom, the interaction test, or the **intraclass correlation** (ICC = {fnum(icc_adj, 2)} after the fixed effects): that is the share of the remaining variance that lies between respondents, the reason the two rows of a respondent are not independent. Use R for those.

The two check rows at the bottom of the block:

| Statistic | Cell | Formula |
|---|---|---|
{row3('Kept respondents with a Product_TV row for Balls in Data', 'ProductTV', 'chk_n')}
{row3('Smallest Valid data in column C', 'ProductTV', 'chk_valid')}

**You should see** {len(kept)} in the first (it must equal n pairs) and a number of at least 70 in the second.

---

## 6. H2: Brand_Logo time to first fixation

TTFF has a complication that dwell share does not: a respondent who never fixated the logo has **NA**, and a pair with an NA cannot enter a paired test. The **BrandLogo** sheet therefore lists only the {len(both_logo)} kept respondents who have a TTFF in **both** commercials (rows {B1} to {B2}); the other {len(kept) - len(both_logo)} kept respondents have an NA in at least one commercial and leave the paired test. Columns C and D use the same `SUMIFS` idea on `TTFF AOI (ms)` (SUMIFS ignores the text NA, and a roster respondent has a number in both rows). `TTFF AOI (ms)` counts from the start of the end card, which is what the question is about; `TTFF max. (ms)` and `TTFF parent (ms)` count from the start of the commercial, so they are not used here.

| Column | Cell (first respondent) | Formula |
|---|---|---|
""" + "\n".join(
    f"| {REG[('BrandLogo', k)]['label']} | `BrandLogo!{REG[('BrandLogo', k)]['cell']}` | `{REG[('BrandLogo', k)]['formula']}` |"
    for k in ("row_balls", "row_paint", "row_diff")) + f"""

The block in column H repeats the paired recipe of section 4 on this table:

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('Balls mean (ms)', 'BrandLogo', 'm_b', fnum(h2['m_balls']))}
{row('Paint mean (ms)', 'BrandLogo', 'm_p', fnum(h2['m_paint']))}
{row('n pairs', 'BrandLogo', 'n', h2['n'])}
{row('Mean difference, Paint minus Balls (ms)', 'BrandLogo', 'md', fnum(h2['mean_diff']))}
{row('SD of the differences (ms)', 'BrandLogo', 'sdd', fnum(h2['sd_diff']))}
{row('Paired t', 'BrandLogo', 't', fnum(h2['t']))}
{row('df', 'BrandLogo', 'df', h2['df'])}
{row('p, two tailed (T.TEST type 1)', 'BrandLogo', 'p_tt', pT(h2))}
{row(DZ, 'BrandLogo', 'dz', fnum(h2['dz']))}
{row('d_z 95% CI lower', 'BrandLogo', 'dz_lo', fnum(a2[0]))}
{row('d_z 95% CI upper', 'BrandLogo', 'dz_hi', fnum(a2[1]))}

The respondent ratio (the fixation based **Respondent ratio (%)**) is below the block:

| Statistic | Cell | Formula | Result |
|---|---|---|---|
{row('Fixated the logo, Balls', 'BrandLogo', 'fix_b', n_fix_b)}
{row('Respondent ratio, Balls (%)', 'BrandLogo', 'ratio_b', fnum(100 * n_fix_b / len(kept), 1))}
{row('Fixated the logo, Paint', 'BrandLogo', 'fix_p', n_fix_p)}
{row('Respondent ratio, Paint (%)', 'BrandLogo', 'ratio_p', fnum(100 * n_fix_p / len(kept), 1))}

**Common mistake.** Treating the NA as 0 ms: that would make the respondent the fastest viewer in the sample. The other tempting shortcut, giving non fixators the longest possible value (5000 ms, the length of the end card), answers a different question; `analysis.R` runs it as a sensitivity analysis (H2b) and the result is in `analysis_output.txt`. Excel has no built in Wilcoxon signed rank test either; {h2_txt}.

---

## 7. H3: Liking

The **Survey** sheet holds `SIMULATED_ad_test_survey.csv` (one row per respondent, rows {S1} to {S_last_all}), with the kept respondents first (rows {S1} to {S2}) and the excluded respondent last (row {S_last_all}), so the test ranges below leave it out. Column H flags Keep from the Data sheet, column I is the difference, column J sums the Brand_Logo dwell time over both commercials, and column K codes recall as 1 or 0:

| Column | Cell (first row) | Formula |
|---|---|---|
""" + "\n".join(
    f"| {REG[('Survey', k)]['label']} | `Survey!{REG[('Survey', k)]['cell']}` | `{REG[('Survey', k)]['formula']}` |"
    for k in ("row_keep", "row_diff", "row_logo", "row_recall01")) + f"""

The paired recipe sits in column N (formula text in column O):

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('Balls mean (points)', 'Survey', 'm_b', fnum(h3['m_balls']))}
{row('Paint mean (points)', 'Survey', 'm_p', fnum(h3['m_paint']))}
{row('n pairs', 'Survey', 'n', h3['n'])}
{row('Mean difference, Paint minus Balls (points)', 'Survey', 'md', fnum(h3['mean_diff']))}
{row('SD of the differences (points)', 'Survey', 'sdd', fnum(h3['sd_diff']))}
{row('Paired t', 'Survey', 't', fnum(h3['t']))}
{row('df', 'Survey', 'df', h3['df'])}
{row('p, two tailed (T.TEST type 1)', 'Survey', 'p_tt', pT(h3))}
{row(DZ, 'Survey', 'dz', fnum(h3['dz']))}
{row('d_z 95% CI lower', 'Survey', 'dz_lo', fnum(a3[0]))}
{row('d_z 95% CI upper', 'Survey', 'dz_hi', fnum(a3[1]))}

The negative sign means Paint was liked *less*. Balls is liked more even though Paint draws more attention to the television set: attention (what the eyes did) and attitude (what the respondent says) are different measures, which is the main teaching point of this study. The counts below the test block (`COUNTIF` of Ad_preferred, `SUMPRODUCT` of the ratings) show that {n_ties} of the {len(kept)} respondents gave the two commercials the same rating (a tied pair carries no information for a signed rank test, which drops it); the Wilcoxon test is in R.

### Brand recall and logo attention (point biserial correlation)

`Brand_recall` is one answer per respondent, so the logo attention measure is one number per respondent too: the total fixation based dwell time on Brand_Logo over both commercials (column J; the text NA is ignored by `SUMIFS`, so a never fixated logo counts as 0 ms). A Pearson correlation between a continuous variable and a 0/1 variable is the **point biserial** correlation:

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('Point biserial r (CORREL)', 'Survey', 'rpb', fnum(pb_r))}
{row('t for r', 'Survey', 'rpb_t', fnum(pb['statistic']))}
{row('p, two tailed', 'Survey', 'rpb_p', fnum(pb_p, 4))}
{row('r 95% CI lower (Fisher z)', 'Survey', 'rpb_lo', fnum(pb['ci_low']))}
{row('r 95% CI upper (Fisher z)', 'Survey', 'rpb_hi', fnum(pb['ci_high']))}

---

## 8. Chart with 95 percent confidence intervals

1. Make a small chart table: two rows (Balls, Paint), columns **Mean dwell (%)** (point to the mean cells) and **CI half width** = `CONFIDENCE.T(0.05, SD cell, n cell)`. `CONFIDENCE.T` returns half the width of a t based 95 percent interval around a mean. The workbook has this table on the ProductTV sheet at `ProductTV!I{PT_CHART_ROW + 2}:K{PT_CHART_ROW + 3}`; the half widths are in `{rcell('ProductTV', 'chart_ci_b')}` and `{rcell('ProductTV', 'chart_ci_p')}`.
2. Select the commercial names and the means, **Insert > Column chart > Clustered Column**.
3. Click the bars, **Chart Design > Add Chart Element > Error Bars > More Error Bars Options**, choose **Custom > Specify Value**, and select the two CI half width cells for both **Positive** and **Negative**.
4. Set the bar fill to navy (`#1F3864`) and lighten the gridlines.

**Common mistake.** Using the default error bars (**Standard Error** or a fixed 5 percent). Those are not 95 percent confidence intervals. And remember that error bars around the two means are *not* the interval that goes with the paired test: the paired test uses the interval of the *differences* (`{rcell('ProductTV', 'ci_lo')}` and `{rcell('ProductTV', 'ci_hi')}`), which is much narrower because it removes the between respondent variation.

---

## 9. Check your numbers and know what Excel cannot do

The **Check** sheet puts each live Excel result next to the value `analysis.R` produced and prints OK or CHECK (status in column F, the count of CHECK rows in `{rcell('Check', 'n_check')}`). Your Excel results should match to rounding; the d_z limits differ by method and have a looser tolerance. Several parts of the R analysis have no Excel equivalent:

- **Wilcoxon signed rank tests** (a check that the conclusion does not depend on normality). In R: H1 V = {h1['wilcoxon_V']}, H2 V = {h2['wilcoxon_V']}, H3 V = {h3['wilcoxon_V']}.
- **Mixed effects model** with a random intercept per respondent and the **ICC**: `lme4::lmer(dwell_pct ~ commercial * order_position + (1 | respondent))`.
- **Sensitivity analysis** for never fixators (H2b), **exact McNemar** test of the two respondent ratios, and **exact binomial** test of the preference split.

If you have the **Analysis ToolPak** add in (**File > Options > Add-ins > Manage Excel Add-ins > Go > Analysis ToolPak**), **Data > Data Analysis > t-Test: Paired Two Sample for Means** prints t, df, and p in one table and is a good cross check of sections 4, 6, and 7.

---

## 10. Common mistakes, collected

1. Opening the .csv by double clicking and getting one column: use **Data > From Text/CSV** or the Text Import Wizard.
2. Replacing NA with 0 in a TTFF column. NA means not observed.
3. Forgetting the exclusion filter (Keep = Yes) in the PivotTable or in COUNTIFS, or leaving {excl_txt} in a roster.
4. `T.TEST` with `type = 2` or `type = 3` on a within subjects design. Use `type = 1` and make sure the two columns are in the same respondent order.
5. Reporting `T.TEST` alone. Always give M, SD, n, t, df, p, and d_z with its interval.
6. Pairing the wrong rows: the two ranges must list the respondents in the same order, and a respondent with an NA has to leave the roster entirely.
7. Treating {nrow_data} rows as {nrow_data} viewers. Each respondent contributes eight rows; per AOI tests use one value per respondent and commercial, and the all rows analysis needs the mixed model.
8. Ignoring the order of viewing. The commercial seen second gains extra dwell share; counterbalancing makes the commercial comparison fair, but the order effect still adds noise and should be checked.

## Sources

- iMotions help center, AOI Metrics: https://help.imotions.com/lab/analysis/eye-tracking/areas-of-interest/aoi-metrics (local copy: `C:\\JaxState_Behavior Lab\\Eye Tracking Glasses (ETG)\\AOI Metrics.pdf`)
- iMotions help center, Areas of Interest (AOIs), including dynamic AOIs and export options: https://help.imotions.com/lab/analysis/eye-tracking/areas-of-interest/areas-of-interest-aois (local copy: `C:\\JaxState_Behavior Lab\\Eye Tracking Glasses (ETG)\\Areas of Interest AOIs.pdf`)
- iMotions help center, Opening Data Export Files in Excel: https://help.imotions.com/lab/exports/additional-export-information/opening-data-export-files-in-excel (local copy: `iMotion_Manual\\Opening Data Export Files in Excel.pdf`)
- iMotions help center, Creating Excel Pivot Tables: https://help.imotions.com/lab/exports/additional-export-information/creating-excel-pivot-tables (local copy: `iMotion_Manual\\Making Excel Pivot Tables.pdf`)

## Items to verify on the lab laptop

- The exact column header layout of the Individual AOI metrics export in iMotions Lab 11.1 (the manuals name the metrics but do not print the header row). Adjust the column letters in sections 1 and 2 once a real export is open.
- Whether the real export writes `NA` or leaves cells blank for never fixated AOIs; `ISNUMBER` handles both.
- How the export reports `AOI duration (ms)`, `TTFF AOI (ms)`, and `Dwell time (%)` for a dynamic AOI that is active in several separate intervals (Product_TV and Hero_Element here). The AOI Metrics article defines the start time and the active time but does not show a multi interval example.
- How `TTFF max. (ms)` is written for a dynamic AOI. The AOI Metrics article says it counts from the start of the parent and uses the parent's presentation duration when there was no fixation; the simulated file follows that wording (60000 ms).
- Whether the dwell classification threshold (default 100 ms) and the first fixation exclusion were left at their defaults when the export was made; both change the numbers (Areas of Interest article, Export options).
"""
with open(MD_OUT, "w", encoding="utf-8") as fh:
    fh.write(md)
print("wrote", MD_OUT)
