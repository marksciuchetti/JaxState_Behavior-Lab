#!/usr/bin/env python3
"""
make_excel_workbook.py
Builds SIMULATED_gsr_analysis.xlsx from the SIMULATED CSV files with openpyxl, and writes
analysis_excel_walkthrough.md from the SAME cell registry, so that every cell reference in the
walkthrough is the address of the formula in the workbook. After saving, the script reads the
workbook back with openpyxl and checks that every registered cell holds the registered formula
(the verification step; it stops with an error list if anything differs).

The workbook mirrors analysis.R for students who work in Excel:

  README   what the workbook is, and that every number is SIMULATED
  Data     the GSR Peak Metrics export as exported (columns A to O, 180 rows), the Signal Quality,
           heart rate, and survey values matched to each row (columns P to T, pasted values), and
           live helper columns U to AA (Phase, Brand, Position, Keep, Nonresponder, PPM_recalc,
           PPM_if_Duration)
  Summary  table by stimulus built with COUNTIFS, AVERAGEIFS, and STDEV.S array formulas (the hand
           built PivotTable), the Signal Duration check, and the device check
  H1       arousal response: one row per kept respondent (SUMIFS pulls Peaks Per Minute and the
           self reported arousal from Data); paired t tests for Nike and e.l.f. against their own
           baselines, the sensitivity analysis without nonresponders, and the exploratory
           correlation of self reported arousal with Peaks Per Minute
  H2       Nike versus e.l.f. Peaks Per Minute, the simple order effect estimate, and the
           sensitivity analysis without nonresponders
  H3       heart rate change from baseline for each commercial and the difference between them
  Check    live Excel results next to the values analysis.R produced, with a status

All formulas are ordinary Excel formulas. The STDEV.S(IF(...)) formulas on the Summary sheet are
array formulas; Excel 365 evaluates them with Enter, older Excel needs Ctrl+Shift+Enter
(openpyxl writes them as array formulas already).

The heart rate column and the Signal Quality layout are MODELED (see make_simulated_data.py).

Run: python3 make_excel_workbook.py   (after make_simulated_data.py and analysis.R)
"""

import csv
import os
import re
import statistics as st
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.data_source import NumDataSource, NumRef
from openpyxl.chart.error_bar import ErrorBars
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.formula import ArrayFormula

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "SIMULATED_gsr_analysis.xlsx")
MD_OUT = os.path.join(HERE, "analysis_excel_walkthrough.md")

NAVY = "1F3864"
TEAL = "1B7F79"
GRAY = "8C8C8C"
ACCENT = "A6561A"
LIGHT = "F2F2F2"

hdr_font = Font(bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor=NAVY)
sub_fill = PatternFill("solid", fgColor=LIGHT)
bold = Font(bold=True)
note_font = Font(italic=True, color="555555")
code_font = Font(name="Consolas", size=9, color="555555")
thin = Side(style="thin", color="BFBFBF")
box = Border(top=thin, bottom=thin, left=thin, right=thin)

LABELS = ["Baseline_Nike", "Nike_Commercial", "Recovery_Nike", "Baseline_elf", "elf_Commercial", "Recovery_elf"]
NIKE_B, NIKE_C, ELF_B, ELF_C = "Baseline_Nike", "Nike_Commercial", "Baseline_elf", "elf_Commercial"


def read_csv(name):
    with open(os.path.join(HERE, name), encoding="utf-8") as fh:
        raw = fh.readlines()
    banner = raw[0].strip()
    rows = list(csv.reader([ln for ln in raw if not ln.startswith("#")]))
    return banner, rows[0], rows[1:]


def to_cell(v):
    """Numbers become numbers; NA stays the text NA (as iMotions writes it)."""
    if v == "NA" or v == "":
        return v
    try:
        f = float(v)
        return int(f) if f.is_integer() and "." not in v else f
    except ValueError:
        return v


def style_header(ws, row=1, ncol=None, first_col=1):
    ncol = ncol or ws.max_column
    for c in range(first_col, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = box


def autowidth(ws, min_w=8, max_w=44):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            L = len(str(cell.value)) if not str(cell.value).startswith("=") else 12
            widths[cell.column_letter] = max(widths.get(cell.column_letter, 0), L)
    for col, w in widths.items():
        ws.column_dimensions[col].width = max(min_w, min(max_w, w + 2))


def show_formula(cell, formula):
    """Write the formula text into a cell as TEXT (not evaluated), without the _xlfn. prefixes."""
    cell.value = formula.replace("_xlfn.", "")
    cell.data_type = "s"
    cell.font = code_font


# registry of documented cells: REG[(sheet, key)] = {"cell": "J12", "formula": "=...", "label": "..."}
REG = {}


def reg(sheet, key, cell, formula, label=""):
    REG[(sheet, key)] = {"cell": cell, "formula": formula.replace("_xlfn.", ""), "label": label}


# ---------------------------------------------------------------- read the files
banner, header, rows = read_csv("SIMULATED_gsr_peak_metrics.csv")
SEED = re.search(r"seed (\d+)", banner).group(1)
idx = {h: i for i, h in enumerate(header)}
n_data = len(rows)
FIRST, LAST = 2, n_data + 1                       # Data rows 2 to 181
_, qh, qrows = read_csv("SIMULATED_gsr_signal_quality.csv")
qidx = {h: i for i, h in enumerate(qh)}
quality = {(r[qidx["Respondent Name"]], r[qidx["Stimulus"]]): r for r in qrows}
_, hh, hrows = read_csv("SIMULATED_heart_rate.csv")
hidx = {h: i for i, h in enumerate(hh)}
heart = {(r[hidx["Respondent Name"]], r[hidx["Stimulus"]]): r for r in hrows}
sb, sh, srows = read_csv("SIMULATED_gsr_survey.csv")
sidx = {h: i for i, h in enumerate(sh)}
survey = {r[sidx["Respondent Name"]]: r for r in srows}

PM_COLS = len(header)                              # 15
EXTRA = ["Device", "Signal-to-Noise ratio (dB)", "Signal missing (%)", "Mean PPG Heart Rate (BPM)", "Arousal_1to9 (survey)"]
HELP = ["Phase", "Brand", "Position", "Keep", "Nonresponder", "PPM_recalc", "PPM_if_Duration"]
ALLH = header + EXTRA + HELP
COL = {h: get_column_letter(i + 1) for i, h in enumerate(ALLH)}     # header -> Data column letter


def rng(col):
    return f"Data!${col}${FIRST}:${col}${LAST}"


R_RESP, R_GROUP, R_LABEL = rng(COL["Respondent Name"]), rng(COL["Group"]), rng(COL["Label"])
R_DUR, R_SIG = rng(COL["Duration (ms)"]), rng(COL["Signal Duration (ms)"])
R_HAS, R_COUNT, R_PPM, R_AMP = (rng(COL["Has Peaks"]), rng(COL["Peak Count"]), rng(COL["Peaks Per Minute"]),
                               rng(COL["Average Peak Amplitude (microSiemens)"]))
R_DEV, R_SNR, R_MISS, R_HR, R_AR = (rng(COL["Device"]), rng(COL["Signal-to-Noise ratio (dB)"]), rng(COL["Signal missing (%)"]),
                                    rng(COL["Mean PPG Heart Rate (BPM)"]), rng(COL["Arousal_1to9 (survey)"]))
R_PHASE, R_BRAND, R_POS, R_KEEP, R_NR, R_RECALC, R_IFDUR = (rng(COL[h]) for h in HELP)

# rosters (static, like the WS2 workbook; the Check sheet verifies them against Data with live formulas)
by_resp = {}
for r in rows:
    by_resp.setdefault(r[idx["Respondent Name"]], []).append(r)
ids = sorted(by_resp)


def phase_of(label):
    return "Baseline" if label.startswith("Baseline") else "Recovery" if label.startswith("Recovery") else "Commercial"


excluded = sorted(rid for rid in ids if any(
    phase_of(lab) == "Commercial" and (float(quality[(rid, lab)][qidx["Signal-to-Noise ratio (dB)"]]) < 0
                                       or float(quality[(rid, lab)][qidx["Signal missing (%)"]]) > 20) for lab in LABELS))
kept_ids = [i for i in ids if i not in excluded]
n_zero = {rid: sum(1 for r in by_resp[rid] if r[idx["Has Peaks"]] == "0") for rid in ids}
nonresp = sorted(i for i in kept_ids if n_zero[i] >= 4)
resp_ids = [i for i in kept_ids if i not in nonresp]
roster = resp_ids + nonresp                        # responders first, nonresponders last
N_KEPT, N_RESP = len(roster), len(resp_ids)


def value_of(resp, label, col):
    for r in by_resp[resp]:
        if r[idx["Label"]] == label:
            return r[idx[col]]


wb = Workbook()

# ---------------------------------------------------------------- README
ws = wb.active
ws.title = "README"
readme = [
    "SIMULATED GSR and heart rate workbook (teaching material)",
    "",
    f"Every number in this workbook is SIMULATED. It was generated by make_simulated_data.py (seed {SEED}) to",
    "resemble a two commercial arousal test in the style of the lab's own library studies Nike Campaign Ad Evaluation and",
    "e.l.f. Campaign Ad Evaluation: a 60 s Baseline, the commercial (Nike 60 s, e.l.f. 30 s), a 15 s Recovery, and a short",
    "questionnaire, recorded with two Shimmer3 GSR+ units (GSR plus PPG heart rate). No real person was recorded.",
    "",
    "Sheets:",
    f"  Data       GSR Peak Metrics export as exported (columns A to O, rows {FIRST} to {LAST}). Columns P to T hold the Device,",
    "             Signal-to-Noise ratio, Signal missing, mean PPG heart rate, and self reported arousal matched to each row (pasted values).",
    f"             Helper columns {COL['Phase']} to {COL['PPM_if_Duration']} are live formulas: Phase, Brand, Position (first or second sequence), Keep (exclusion rule),",
    "             Nonresponder (Has Peaks = 0 in at least 4 of 6 stimuli), PPM_recalc, PPM_if_Duration.",
    "  Summary    Mean, SD, n by stimulus (COUNTIFS, AVERAGEIFS, STDEV.S(IF(...))), the Signal Duration check, the device check.",
    "  H1         Arousal response. Paired t tests of each commercial against its own baseline (Peaks Per Minute), the sensitivity analysis",
    "             without nonresponders, and the exploratory correlation of self reported arousal with Peaks Per Minute.",
    "  H2         Nike versus e.l.f. Peaks Per Minute in the commercials, the simple order (carryover) estimate, sensitivity analysis.",
    "  H3         Heart rate change from baseline (commercial minus baseline) for each commercial and the difference between them.",
    "  Check      Live Excel results next to the values analysis.R produced, with an OK or CHECK status.",
    "",
    "Design: WITHIN subjects, 30 respondents, two sequences each, 2 minute rest between them. The Group column is the ORDER",
    "(A = Nike first, B = e.l.f. first), not a treatment. Respondents were run two at a time at Seat 1 (Shimmer3-8956) and Seat 2",
    "(Shimmer3-B, a placeholder name to verify); the analysis pools both devices.",
    "H1 and H3 differences are commercial MINUS baseline. H2 is Nike MINUS e.l.f.",
    f"Exclusion rule, written before collection (the lab's rule, built from the GSR Signal Quality article): SNR below 0 dB or Signal missing above 20 percent in any commercial (respondents {', '.join(excluded)}).",
    f"Nonresponders (kept in the main analysis, removed in the sensitivity analysis): {', '.join(nonresp)}.",
    "Peaks Per Minute = Peak Count / (Signal Duration / 60000): gaps longer than 4 s are left out of Signal Duration.",
    "MODELED, verify on the lab machine: the heart rate column name (the manuals name none), the Signal Quality file layout, Type, Parent Stimulus, and Device Shimmer3-B.",
    "",
    "Companion files: analysis_excel_walkthrough.md (step by step), analysis.R (same analysis in R),",
    "analysis_output.txt (its console output), results_text.md (model paragraphs).",
]
for i, line in enumerate(readme, start=1):
    ws.cell(row=i, column=1, value=line)
ws["A1"].font = Font(bold=True, size=14, color=NAVY)
ws.column_dimensions["A"].width = 150

# ---------------------------------------------------------------- Data
ws = wb.create_sheet("Data")
ws.append(ALLH)
style_header(ws)
cU, cV, cW, cX, cY, cZ, cAA = (COL[h] for h in HELP)
cL, cQ, cR, cE, cG, cM, cJ, cI, cB = (COL["Has Peaks"], COL["Signal-to-Noise ratio (dB)"], COL["Signal missing (%)"],
                                      COL["Group"], COL["Label"], COL["Peak Count"], COL["Signal Duration (ms)"],
                                      COL["Duration (ms)"], COL["Respondent Name"])
for i, r in enumerate(rows, start=FIRST):
    rid, lab = r[idx["Respondent Name"]], r[idx["Label"]]
    q, h = quality[(rid, lab)], heart[(rid, lab)]
    ar = "NA"
    if phase_of(lab) == "Commercial":
        ar = survey[rid][sidx["Arousal_Nike_1to9" if "Nike" in lab else "Arousal_elf_1to9"]]
    vals = [to_cell(v) for v in r] + [q[qidx["Device"]], to_cell(q[qidx["Signal-to-Noise ratio (dB)"]]),
                                      to_cell(q[qidx["Signal missing (%)"]]), to_cell(h[hidx["Mean PPG Heart Rate (BPM)"]]),
                                      to_cell(ar)]
    ws.append(vals)
    ws[f"{cU}{i}"] = f'=IF(LEFT({cG}{i},8)="Baseline","Baseline",IF(LEFT({cG}{i},8)="Recovery","Recovery","Commercial"))'
    ws[f"{cV}{i}"] = f'=IF(ISNUMBER(SEARCH("Nike",{cG}{i})),"Nike","elf")'
    ws[f"{cW}{i}"] = f'=IF(OR(AND({cE}{i}="A",{cV}{i}="Nike"),AND({cE}{i}="B",{cV}{i}="elf")),"first","second")'
    ws[f"{cX}{i}"] = (f'=IF(COUNTIFS({R_RESP},{cB}{i},{R_PHASE},"Commercial",{R_SNR},"<0")'
                      f'+COUNTIFS({R_RESP},{cB}{i},{R_PHASE},"Commercial",{R_MISS},">20")>0,"No","Yes")')
    ws[f"{cY}{i}"] = f"=IF(COUNTIFS({R_RESP},{cB}{i},{R_HAS},0)>=4,1,0)"
    ws[f"{cZ}{i}"] = f"={cM}{i}/({cJ}{i}/60000)"
    ws[f"{cAA}{i}"] = f"={cM}{i}/({cI}{i}/60000)"
    ws[f"{cZ}{i}"].number_format = "0.000"
    ws[f"{cAA}{i}"].number_format = "0.000"
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(ALLH))}{LAST}"
autowidth(ws, max_w=22)
for h in EXTRA:
    ws[f"{COL[h]}1"].fill = PatternFill("solid", fgColor=GRAY)
for h in HELP:
    ws[f"{COL[h]}1"].fill = PatternFill("solid", fgColor=ACCENT)
r0 = FIRST
reg("Data", "phase", f"{cU}{r0}", ws[f"{cU}{r0}"].value, "Phase")
reg("Data", "brand", f"{cV}{r0}", ws[f"{cV}{r0}"].value, "Brand")
reg("Data", "position", f"{cW}{r0}", ws[f"{cW}{r0}"].value, "Position")
reg("Data", "keep", f"{cX}{r0}", ws[f"{cX}{r0}"].value, "Keep")
reg("Data", "nonresp", f"{cY}{r0}", ws[f"{cY}{r0}"].value, "Nonresponder")
reg("Data", "recalc", f"{cZ}{r0}", ws[f"{cZ}{r0}"].value, "PPM_recalc")
reg("Data", "ifdur", f"{cAA}{r0}", ws[f"{cAA}{r0}"].value, "PPM_if_Duration")

# ---------------------------------------------------------------- Summary
ws = wb.create_sheet("Summary")
ws["A1"] = "Descriptive table by stimulus (kept respondents only). SIMULATED data."
ws["A1"].font = Font(bold=True, size=12, color=NAVY)
ws["A2"] = ("n = kept respondents. Peaks Per Minute statistics use every kept respondent (nonresponders included). "
            "Average Peak Amplitude averages the respondent values that exist (the text NA, a stimulus without a peak, is ignored).")
ws["A2"].font = note_font
cols = ["Stimulus", "n", "PPM mean", "PPM SD", "Has Peaks (%)", "Peak Count mean", "Peak amplitude mean (microSiemens)",
        "n with amplitude", "Heart rate mean (BPM)", "Heart rate SD (BPM)", "SNR mean (dB)", "Signal missing mean (%)"]
ws.append([])
ws.append(cols)
style_header(ws, row=4, ncol=len(cols))
r = 5
SUM_FIRST = r
summary_rows = {}
for lab in LABELS:
    ws.cell(row=r, column=1, value=lab)
    crit = f'{R_LABEL},$A{r},{R_KEEP},"Yes"'
    f_n = f"=COUNTIFS({crit})"
    f_pm = f"=AVERAGEIFS({R_PPM},{crit})"
    f_ps = f'=_xlfn.STDEV.S(IF(({R_LABEL}=$A{r})*({R_KEEP}="Yes"),{R_PPM}))'
    f_hp = f"=100*AVERAGEIFS({R_HAS},{crit})"
    f_pc = f"=AVERAGEIFS({R_COUNT},{crit})"
    f_am = f"=AVERAGEIFS({R_AMP},{crit})"
    f_na = f'=COUNTIFS({crit},{R_AMP},">0")'
    f_hm = f"=AVERAGEIFS({R_HR},{crit})"
    f_hs = f'=_xlfn.STDEV.S(IF(({R_LABEL}=$A{r})*({R_KEEP}="Yes"),{R_HR}))'
    f_sn = f"=AVERAGEIFS({R_SNR},{crit})"
    f_ms = f"=AVERAGEIFS({R_MISS},{crit})"
    ws.cell(row=r, column=2, value=f_n)
    ws.cell(row=r, column=3, value=f_pm)
    ws.cell(row=r, column=4).value = ArrayFormula(f"D{r}", f_ps)
    ws.cell(row=r, column=5, value=f_hp)
    ws.cell(row=r, column=6, value=f_pc)
    ws.cell(row=r, column=7, value=f_am)
    ws.cell(row=r, column=8, value=f_na)
    ws.cell(row=r, column=9, value=f_hm)
    ws.cell(row=r, column=10).value = ArrayFormula(f"J{r}", f_hs)
    ws.cell(row=r, column=11, value=f_sn)
    ws.cell(row=r, column=12, value=f_ms)
    for c in range(1, 13):
        ws.cell(row=r, column=c).border = box
        if c >= 3:
            ws.cell(row=r, column=c).number_format = "0.00"
    if r == SUM_FIRST:
        for key, col, f in (("n", "B", f_n), ("ppm_mean", "C", f_pm), ("ppm_sd", "D", f_ps), ("has_pct", "E", f_hp),
                            ("count_mean", "F", f_pc), ("amp_mean", "G", f_am), ("n_amp", "H", f_na), ("hr_mean", "I", f_hm),
                            ("hr_sd", "J", f_hs), ("snr_mean", "K", f_sn), ("miss_mean", "L", f_ms)):
            reg("Summary", key, f"{col}{r}", f, cols[ord(col) - ord("A")])
    summary_rows[lab] = r
    r += 1
SUM_LAST = r - 1
ws[f"A{SUM_LAST + 2}"] = ("Array formulas (columns D and J): Excel 365 evaluates them with Enter; "
                          "Excel 2019 or older needs Ctrl+Shift+Enter after editing.")
ws[f"A{SUM_LAST + 2}"].font = note_font

# Signal Duration check and device check
sd0 = SUM_LAST + 4
ws.cell(row=sd0 - 1, column=1, value="Why Peaks Per Minute uses Signal Duration (kept respondents)").font = bold
f_gap = f'=SUMPRODUCT(--({rng(COL["Signal Duration (ms)"])}<{rng(COL["Duration (ms)"])}),--({R_KEEP}="Yes"))'
f_under = f'=SUMPRODUCT(MAX(({R_KEEP}="Yes")*({R_RECALC}-{R_IFDUR})))'
f_maxd = f"=SUMPRODUCT(MAX(ABS({rng(COL['Peaks Per Minute'])}-{R_RECALC})))"
ws.cell(row=sd0, column=1, value="Rows with Signal Duration shorter than Duration")
ws.cell(row=sd0, column=2, value=f_gap).number_format = "0"
ws.cell(row=sd0 + 1, column=1, value="Largest understatement if Duration were used (peaks per minute)")
ws.cell(row=sd0 + 1, column=2, value=f_under).number_format = "0.000"
ws.cell(row=sd0 + 2, column=1, value="Largest gap between the exported Peaks Per Minute and PPM_recalc (all rows)")
ws.cell(row=sd0 + 2, column=2, value=f_maxd).number_format = "0.0000"
reg("Summary", "gap_rows", f"B{sd0}", f_gap, "Rows with Signal Duration shorter than Duration")
reg("Summary", "understate", f"B{sd0 + 1}", f_under, "Largest understatement if Duration were used")
reg("Summary", "max_gap", f"B{sd0 + 2}", f_maxd, "Largest gap between exported and recalculated Peaks Per Minute")
dv0 = sd0 + 5
ws.cell(row=dv0 - 2, column=1, value="Device check (kept respondents; the analysis pools the two devices)").font = bold
for c, h in enumerate(["Device", "Rows", "SNR mean (dB)", "PPM mean (baselines and commercials)"], start=1):
    ws.cell(row=dv0 - 1, column=c, value=h)
style_header(ws, row=dv0 - 1, ncol=4)
dev_rows = {}
for k, dname in enumerate(sorted({q[qidx["Device"]] for q in qrows})):
    rr = dv0 + k
    ws.cell(row=rr, column=1, value=dname)
    f_dn = f'=COUNTIFS({R_DEV},$A{rr},{R_KEEP},"Yes")'
    f_ds = f'=AVERAGEIFS({R_SNR},{R_DEV},$A{rr},{R_KEEP},"Yes")'
    f_dp = f'=AVERAGEIFS({R_PPM},{R_DEV},$A{rr},{R_KEEP},"Yes",{R_PHASE},"<>Recovery")'
    ws.cell(row=rr, column=2, value=f_dn)
    ws.cell(row=rr, column=3, value=f_ds).number_format = "0.00"
    ws.cell(row=rr, column=4, value=f_dp).number_format = "0.00"
    dev_rows[dname] = rr
    if k == 0:
        reg("Summary", "dev_n", f"B{rr}", f_dn, "Rows by device")
        reg("Summary", "dev_snr", f"C{rr}", f_ds, "SNR mean by device")
        reg("Summary", "dev_ppm", f"D{rr}", f_dp, "PPM mean by device")
autowidth(ws, max_w=24)
ws.column_dimensions["A"].width = 62
# chart table and bar chart of mean Peaks Per Minute by stimulus
ct = dv0 + 5
ws.cell(row=ct - 1, column=1, value="Chart table: mean Peaks Per Minute by stimulus").font = bold
ws.cell(row=ct, column=1, value="Stimulus")
ws.cell(row=ct, column=2, value="Mean PPM")
for i, lab in enumerate(LABELS, start=1):
    ws.cell(row=ct + i, column=1, value=lab)
    ws.cell(row=ct + i, column=2, value=f"=C{summary_rows[lab]}").number_format = "0.00"
chart = BarChart()
chart.type = "col"
chart.style = 10
chart.title = "Mean peaks per minute by stimulus (SIMULATED)"
chart.y_axis.title = "GSR peaks per minute"
chart.add_data(Reference(ws, min_col=2, min_row=ct, max_row=ct + len(LABELS)), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=1, min_row=ct + 1, max_row=ct + len(LABELS)))
chart.series[0].graphicalProperties.solidFill = NAVY
chart.legend = None
chart.height = 8
chart.width = 16
ws.add_chart(chart, f"D{ct - 1}")


# ---------------------------------------------------------------- paired and correlation blocks
class Block:
    """Writes label, value, formula text, and note columns for one block and records every cell in REG."""

    def __init__(self, ws, sheet, label_col, start_row, prefix):
        self.ws, self.sheet, self.lc, self.prefix = ws, sheet, label_col, prefix
        self.row = start_row
        self.ref = {}
        for k, name in enumerate(("Statistic", "Value", "Formula shown", "Note")):
            c = ws.cell(row=start_row - 1, column=label_col + k, value=name)
            c.font = hdr_font
            c.fill = hdr_fill

    def title(self, text):
        c = self.ws.cell(row=self.row, column=self.lc, value=text)
        c.font = bold
        c.fill = sub_fill
        self.row += 1

    def put(self, key, label, formula, note="", fmt="0.0000"):
        f = formula.format(**{k: f"{get_column_letter(self.lc + 1)}{v}" for k, v in self.ref.items()})
        i = self.row
        self.ws.cell(row=i, column=self.lc, value=label)
        vc = self.ws.cell(row=i, column=self.lc + 1, value=f)
        vc.number_format = fmt
        show_formula(self.ws.cell(row=i, column=self.lc + 2), f)
        if note:
            self.ws.cell(row=i, column=self.lc + 3, value=note).font = note_font
        self.ref[key] = i
        reg(self.sheet, f"{self.prefix}_{key}", f"{get_column_letter(self.lc + 1)}{i}", f, label)
        self.row += 1

    def blank(self):
        self.row += 1

    def addr(self, key):
        return f"{get_column_letter(self.lc + 1)}{self.ref[key]}"

    def next_start(self):
        return self.row + 2


def paired_block(blk, title, r1, r2, rd, unit, nm1, nm2, note_n="", note_diff=""):
    """Paired t test, d_z, and CIs. r1, r2, rd are the reference, comparison, and difference ranges (same sheet);
    difference = comparison minus reference."""
    blk.title(title)
    blk.put("m1", f"{nm1} mean ({unit})", f"=AVERAGE({r1})", "AVERAGE of the reference column")
    blk.put("s1", f"{nm1} SD ({unit})", f"=_xlfn.STDEV.S({r1})", "STDEV.S = sample standard deviation")
    blk.put("m2", f"{nm2} mean ({unit})", f"=AVERAGE({r2})")
    blk.put("s2", f"{nm2} SD ({unit})", f"=_xlfn.STDEV.S({r2})")
    blk.put("n", "n pairs", f"=COUNT({rd})", note_n, fmt="0")
    blk.put("md", f"Mean difference, {nm2} minus {nm1} ({unit})", f"=AVERAGE({rd})", note_diff)
    blk.put("sdd", f"SD of the differences ({unit})", f"=_xlfn.STDEV.S({rd})", "the SD that matters in a paired test")
    blk.put("se", "SE of the mean difference", "={sdd}/SQRT({n})", "SD of differences / sqrt(n)")
    blk.put("t", "Paired t", "={md}/{se}", "t = mean difference / SE")
    blk.put("df", "df", "={n}-1", "n pairs minus 1", fmt="0")
    blk.put("p_tt", "p, two tailed (T.TEST type 1)", f"=_xlfn.T.TEST({r1},{r2},2,1)",
            "tails = 2, type = 1 means paired", fmt="0.000000")
    blk.put("p_hand", "p from the hand computed t", "=_xlfn.T.DIST.2T(ABS({t}),{df})", "should match the T.TEST cell", fmt="0.000000")
    blk.put("tcrit", "Critical t (95%)", "=_xlfn.T.INV.2T(0.05,{df})", "two tailed, alpha = .05")
    blk.put("ci_lo", "Mean difference 95% CI lower", "={md}-{tcrit}*{se}")
    blk.put("ci_hi", "Mean difference 95% CI upper", "={md}+{tcrit}*{se}")
    blk.put("dz", "Cohen's d_z", "={md}/{sdd}", "d_z = mean difference / SD of differences")
    blk.put("se_dz", "SE of d_z (approximation)", "=SQRT(1/{n}+{dz}^2/(2*{n}))", "sqrt(1/n + d_z^2 / (2 n))")
    blk.put("dz_lo", "d_z 95% CI lower", "={dz}-1.96*{se_dz}", "large sample approximation; R uses the noncentral t interval")
    blk.put("dz_hi", "d_z 95% CI upper", "={dz}+1.96*{se_dz}")
    blk.put("r", "Correlation of the two columns", f"=CORREL({r1},{r2})", "within respondent correlation")
    blk.blank()


def corr_block(blk, title, rx, ry, note=""):
    blk.title(title)
    blk.put("n", "n", f"=COUNT({rx})", note, fmt="0")
    blk.put("r", "Pearson r (CORREL)", f"=CORREL({rx},{ry})")
    blk.put("t", "t for r", "={r}*SQRT(({n}-2)/(1-{r}^2))", "t = r sqrt((n - 2) / (1 - r^2))")
    blk.put("p", "p, two tailed", "=_xlfn.T.DIST.2T(ABS({t}),{n}-2)", "df = n - 2", fmt="0.000000")
    blk.put("lo", "r 95% CI lower (Fisher z)", "=TANH(ATANH({r})-_xlfn.NORM.S.INV(0.975)/SQRT({n}-3))")
    blk.put("hi", "r 95% CI upper (Fisher z)", "=TANH(ATANH({r})+_xlfn.NORM.S.INV(0.975)/SQRT({n}-3))")
    blk.blank()


def roster_formulas(ws, r, with_group=True):
    ws.cell(row=r, column=2, value=f"=INDEX({R_GROUP},MATCH($A{r},{R_RESP},0))")
    ws.cell(row=r, column=3, value=f"=INDEX({R_NR},MATCH($A{r},{R_RESP},0))")


def sumifs_ppm(r, label):
    return f'=SUMIFS({R_PPM},{R_RESP},$A{r},{R_LABEL},"{label}")'


def sumifs_hr(r, label):
    return f'=SUMIFS({R_HR},{R_RESP},$A{r},{R_LABEL},"{label}")'


def sumifs_ar(r, label):
    return f'=SUMIFS({R_AR},{R_RESP},$A{r},{R_LABEL},"{label}")'


P1 = 5
P2 = P1 + N_KEPT - 1                 # last roster row (all kept)
PR = P1 + N_RESP - 1                 # last responder row (nonresponders are listed below it)

# ---------------------------------------------------------------- H1
ws = wb.create_sheet("H1")
SH = "H1"
ws["A1"] = "H1. Peaks Per Minute in each commercial versus its own baseline, kept respondents. SIMULATED data."
ws["A1"].font = Font(bold=True, size=12, color=NAVY)
ws["A2"] = (f"One row per kept respondent; responders in rows {P1} to {PR}, the {len(nonresp)} nonresponders in rows {PR + 1} to {P2}. "
            "Columns D, E, G, H pull Peaks Per Minute from the Data sheet with SUMIFS (each respondent has exactly one row per stimulus); "
            "J and K pull the self reported arousal. Differences are commercial MINUS baseline.")
ws["A2"].font = note_font
h1_cols = ["Respondent Name", "Group (order)", "Nonresponder", "Baseline_Nike PPM", "Nike_Commercial PPM", "Difference Nike (commercial minus baseline)",
           "Baseline_elf PPM", "elf_Commercial PPM", "Difference e.l.f. (commercial minus baseline)", "Arousal_Nike_1to9", "Arousal_elf_1to9",
           "Mean arousal (both commercials)", "Mean commercial PPM (both)"]
ws.append([])
ws.append(h1_cols)
style_header(ws, row=4, ncol=len(h1_cols))
ws.row_dimensions[4].height = 48
for k, rid in enumerate(roster):
    r = P1 + k
    ws.cell(row=r, column=1, value=rid)
    roster_formulas(ws, r)
    ws.cell(row=r, column=4, value=sumifs_ppm(r, NIKE_B))
    ws.cell(row=r, column=5, value=sumifs_ppm(r, NIKE_C))
    ws.cell(row=r, column=6, value=f"=E{r}-D{r}")
    ws.cell(row=r, column=7, value=sumifs_ppm(r, ELF_B))
    ws.cell(row=r, column=8, value=sumifs_ppm(r, ELF_C))
    ws.cell(row=r, column=9, value=f"=H{r}-G{r}")
    ws.cell(row=r, column=10, value=sumifs_ar(r, NIKE_C))
    ws.cell(row=r, column=11, value=sumifs_ar(r, ELF_C))
    ws.cell(row=r, column=12, value=f"=(J{r}+K{r})/2")
    ws.cell(row=r, column=13, value=f"=(E{r}+H{r})/2")
    for c in range(1, 14):
        ws.cell(row=r, column=c).border = box
        if c in (4, 5, 6, 7, 8, 9, 13):
            ws.cell(row=r, column=c).number_format = "0.00"
for key, col, f, lab in (("row_grp", "B", f"=INDEX({R_GROUP},MATCH($A{P1},{R_RESP},0))", "Group (order)"),
                         ("row_nr", "C", f"=INDEX({R_NR},MATCH($A{P1},{R_RESP},0))", "Nonresponder"),
                         ("row_bn", "D", sumifs_ppm(P1, NIKE_B), "Baseline_Nike PPM"),
                         ("row_cn", "E", sumifs_ppm(P1, NIKE_C), "Nike_Commercial PPM"),
                         ("row_dn", "F", f"=E{P1}-D{P1}", "Difference Nike"),
                         ("row_be", "G", sumifs_ppm(P1, ELF_B), "Baseline_elf PPM"),
                         ("row_ce", "H", sumifs_ppm(P1, ELF_C), "elf_Commercial PPM"),
                         ("row_de", "I", f"=H{P1}-G{P1}", "Difference e.l.f."),
                         ("row_an", "J", sumifs_ar(P1, NIKE_C), "Arousal_Nike_1to9"),
                         ("row_ae", "K", sumifs_ar(P1, ELF_C), "Arousal_elf_1to9"),
                         ("row_am", "L", f"=(J{P1}+K{P1})/2", "Mean arousal"),
                         ("row_pm", "M", f"=(E{P1}+H{P1})/2", "Mean commercial PPM")):
    reg(SH, key, f"{col}{P1}", f, lab)
rA = lambda c, last: f"{c}{P1}:{c}{last}"
blk = Block(ws, SH, 15, 5, "nike")             # label column O, value P, formula Q, note R
paired_block(blk, "Nike: paired t test and d_z, all kept respondents (Peaks Per Minute)", rA("D", P2), rA("E", P2), rA("F", P2),
             "peaks per minute", "Baseline_Nike", "Nike_Commercial", "COUNT of the difference column; must equal the check below")
H1_NIKE = blk
blk = Block(ws, SH, 15, blk.next_start(), "elf")
paired_block(blk, "e.l.f.: paired t test and d_z, all kept respondents (Peaks Per Minute)", rA("G", P2), rA("H", P2), rA("I", P2),
             "peaks per minute", "Baseline_elf", "elf_Commercial")
H1_ELF = blk
blk = Block(ws, SH, 15, blk.next_start(), "snike")
paired_block(blk, f"Sensitivity, Nike: responders only (rows {P1} to {PR})", rA("D", PR), rA("E", PR), rA("F", PR),
             "peaks per minute", "Baseline_Nike", "Nike_Commercial", "nonresponders are listed below the responders and left out")
H1_SNIKE = blk
blk = Block(ws, SH, 15, blk.next_start(), "self")
paired_block(blk, f"Sensitivity, e.l.f.: responders only (rows {P1} to {PR})", rA("G", PR), rA("H", PR), rA("I", PR),
             "peaks per minute", "Baseline_elf", "elf_Commercial")
H1_SELF = blk
blk = Block(ws, SH, 15, blk.next_start(), "cn")
corr_block(blk, "Exploratory: Arousal_Nike_1to9 with Nike_Commercial PPM (all kept)", rA("J", P2), rA("E", P2))
H1_CN = blk
blk = Block(ws, SH, 15, blk.next_start(), "ce")
corr_block(blk, "Exploratory: Arousal_elf_1to9 with elf_Commercial PPM (all kept)", rA("K", P2), rA("H", P2))
H1_CE = blk
blk = Block(ws, SH, 15, blk.next_start(), "cm")
corr_block(blk, "Exploratory: respondent means over the two commercials (all kept)", rA("L", P2), rA("M", P2))
H1_CM = blk
blk = Block(ws, SH, 15, blk.next_start(), "scn")
corr_block(blk, "Sensitivity: Arousal_Nike_1to9 with Nike_Commercial PPM (responders only)", rA("J", PR), rA("E", PR))
H1_SCN = blk
blk = Block(ws, SH, 15, blk.next_start(), "sce")
corr_block(blk, "Sensitivity: Arousal_elf_1to9 with elf_Commercial PPM (responders only)", rA("K", PR), rA("H", PR))
H1_SCE = blk
blk = Block(ws, SH, 15, blk.next_start(), "chk")
blk.title("Checks")
blk.put("chk_n", "Kept respondents with a Nike_Commercial row in Data (Keep = Yes)",
        f'=COUNTIFS({R_LABEL},"{NIKE_C}",{R_KEEP},"Yes")', "must equal n pairs", fmt="0")
blk.put("chk_resp", "Kept responders with a Nike_Commercial row in Data (Nonresponder = 0)",
        f'=COUNTIFS({R_LABEL},"{NIKE_C}",{R_KEEP},"Yes",{R_NR},0)', f"must equal the n pairs of the sensitivity block (rows {P1} to {PR})", fmt="0")
blk.put("chk_nr", "Nonresponders listed in rows below the responders (column C = 1)", f"=SUM(C{P1}:C{P2})",
        "all nonresponders must sit below row " + str(PR) + "; this sum counts them", fmt="0")
blk.put("chk_nr_top", "Nonresponders inside the responder rows", f"=SUM(C{P1}:C{PR})", "must be 0", fmt="0")
H1_CHK = blk
blk.blank()
for col, w in (("A", 17), ("B", 13), ("C", 13), ("D", 14), ("E", 16), ("F", 22), ("G", 14), ("H", 16), ("I", 22), ("J", 14),
               ("K", 14), ("L", 16), ("M", 16), ("N", 3), ("O", 60), ("P", 12), ("Q", 76), ("R", 60)):
    ws.column_dimensions[col].width = w
# chart: mean PPM with 95% CI half widths (the four baseline and commercial cells)
cr = blk.row + 1
ws.cell(row=cr, column=15, value="Chart table: mean Peaks Per Minute with 95% CI half width").font = bold
ws.cell(row=cr + 1, column=15, value="Stimulus")
ws.cell(row=cr + 1, column=16, value="Mean PPM")
ws.cell(row=cr + 1, column=17, value="CI half width")
chart_rows = [("Baseline_Nike", H1_NIKE, "m1", "s1"), ("Nike_Commercial", H1_NIKE, "m2", "s2"),
              ("Baseline_elf", H1_ELF, "m1", "s1"), ("elf_Commercial", H1_ELF, "m2", "s2")]
for k, (lab, b, mk, sk) in enumerate(chart_rows):
    rr = cr + 2 + k
    ws.cell(row=rr, column=15, value=lab)
    ws.cell(row=rr, column=16, value=f"={b.addr(mk)}").number_format = "0.00"
    f_ci = f"=_xlfn.CONFIDENCE.T(0.05,{b.addr(sk)},{b.addr('n')})"
    ws.cell(row=rr, column=17, value=f_ci).number_format = "0.00"
    reg(SH, f"chart_ci_{k}", f"Q{rr}", f_ci, f"CI half width {lab}")
chart2 = BarChart()
chart2.type = "col"
chart2.style = 10
chart2.title = "Mean peaks per minute with 95% CI (SIMULATED)"
chart2.y_axis.title = "GSR peaks per minute"
chart2.add_data(Reference(ws, min_col=16, min_row=cr + 1, max_row=cr + 1 + len(chart_rows)), titles_from_data=True)
chart2.set_categories(Reference(ws, min_col=15, min_row=cr + 2, max_row=cr + 1 + len(chart_rows)))
chart2.series[0].graphicalProperties.solidFill = NAVY
chart2.series[0].errBars = ErrorBars(errDir="y", errBarType="both", errValType="cust",
                                      plus=NumDataSource(numRef=NumRef(f=f"H1!$Q${cr + 2}:$Q${cr + 1 + len(chart_rows)}")),
                                      minus=NumDataSource(numRef=NumRef(f=f"H1!$Q${cr + 2}:$Q${cr + 1 + len(chart_rows)}")))
chart2.legend = None
chart2.height = 8
chart2.width = 14
ws.add_chart(chart2, f"S{cr}")
H1_CHART_ROW = cr

# ---------------------------------------------------------------- H2
ws = wb.create_sheet("H2")
SH = "H2"
ws["A1"] = "H2. Peaks Per Minute in the commercials, Nike versus e.l.f., kept respondents. SIMULATED data."
ws["A1"].font = Font(bold=True, size=12, color=NAVY)
ws["A2"] = (f"Same roster order as H1 (responders in rows {P1} to {PR}, nonresponders below). Difference is Nike MINUS e.l.f. "
            "Group A saw Nike first, Group B saw e.l.f. first, so the simple order estimate is half the gap between the two group means.")
ws["A2"].font = note_font
h2_cols = ["Respondent Name", "Group (order)", "Nonresponder", "elf_Commercial PPM", "Nike_Commercial PPM", "Difference (Nike minus elf)", "Nike seen"]
ws.append([])
ws.append(h2_cols)
style_header(ws, row=4, ncol=len(h2_cols))
ws.row_dimensions[4].height = 36
for k, rid in enumerate(roster):
    r = P1 + k
    ws.cell(row=r, column=1, value=rid)
    roster_formulas(ws, r)
    ws.cell(row=r, column=4, value=sumifs_ppm(r, ELF_C))
    ws.cell(row=r, column=5, value=sumifs_ppm(r, NIKE_C))
    ws.cell(row=r, column=6, value=f"=E{r}-D{r}")
    ws.cell(row=r, column=7, value=f'=IF(B{r}="A","first","second")')
    for c in range(1, 8):
        ws.cell(row=r, column=c).border = box
        if c in (4, 5, 6):
            ws.cell(row=r, column=c).number_format = "0.00"
for key, col, f, lab in (("row_grp", "B", f"=INDEX({R_GROUP},MATCH($A{P1},{R_RESP},0))", "Group (order)"),
                         ("row_nr", "C", f"=INDEX({R_NR},MATCH($A{P1},{R_RESP},0))", "Nonresponder"),
                         ("row_elf", "D", sumifs_ppm(P1, ELF_C), "elf_Commercial PPM"),
                         ("row_nike", "E", sumifs_ppm(P1, NIKE_C), "Nike_Commercial PPM"),
                         ("row_diff", "F", f"=E{P1}-D{P1}", "Difference (Nike minus elf)"),
                         ("row_seen", "G", f'=IF(B{P1}="A","first","second")', "Nike seen")):
    reg(SH, key, f"{col}{P1}", f, lab)
RG = f"B{P1}:B{P2}"
RC = f"C{P1}:C{P2}"
RDall = f"F{P1}:F{P2}"
blk = Block(ws, SH, 9, 5, "h2")          # label column I, value J, formula K, note L
paired_block(blk, "Paired t test and d_z: Peaks Per Minute in the commercials, all kept respondents", rA("D", P2), rA("E", P2), RDall,
             "peaks per minute", "elf_Commercial", "Nike_Commercial", "COUNT of the difference column; must equal the check below")
H2_MAIN = blk
blk.title("Order effect (simple within respondent estimate)")
blk.put("oe_A", "Mean difference, Group A (Nike seen first)", f'=AVERAGEIFS({RDall},{RG},"A")', "difference = commercial effect - order effect")
blk.put("oe_B", "Mean difference, Group B (Nike seen second)", f'=AVERAGEIFS({RDall},{RG},"B")', "difference = commercial effect + order effect")
blk.put("oe_nA", "n, Group A", f'=COUNTIF({RG},"A")', fmt="0")
blk.put("oe_nB", "n, Group B", f'=COUNTIF({RG},"B")', fmt="0")
blk.put("oe", "Order effect (second minus first)", "=({oe_B}-{oe_A})/2",
        "same as the lmer coefficient for order_position; Excel cannot fit the mixed model")
blk.put("commercial", "Commercial effect (Nike minus elf), order adjusted", "=({oe_A}+{oe_B})/2", "average of the two group differences")
blk.blank()
blk.title("Checks")
blk.put("chk_n", "Kept respondents with a Nike_Commercial row in Data (Keep = Yes)", f'=COUNTIFS({R_LABEL},"{NIKE_C}",{R_KEEP},"Yes")',
        "must equal n pairs", fmt="0")
blk.blank()
blk = Block(ws, SH, 9, blk.next_start(), "h2s")
paired_block(blk, f"Sensitivity: responders only (rows {P1} to {PR})", rA("D", PR), rA("E", PR), f"F{P1}:F{PR}",
             "peaks per minute", "elf_Commercial", "Nike_Commercial", "nonresponders are listed below the responders and left out")
H2_SENS = blk
RGs, RDs, RCs = f"B{P1}:B{PR}", f"F{P1}:F{PR}", f"C{P1}:C{PR}"
blk.title("Order effect, responders only")
blk.put("oe_A", "Mean difference, Group A (responders)", f'=AVERAGEIFS({RDs},{RGs},"A")')
blk.put("oe_B", "Mean difference, Group B (responders)", f'=AVERAGEIFS({RDs},{RGs},"B")')
blk.put("oe_nA", "n, Group A (responders)", f'=COUNTIF({RGs},"A")', fmt="0")
blk.put("oe_nB", "n, Group B (responders)", f'=COUNTIF({RGs},"B")', fmt="0")
blk.put("oe", "Order effect (second minus first), responders only", "=({oe_B}-{oe_A})/2")
blk.blank()
H2_SENS_OE = blk
for col, w in (("A", 17), ("B", 13), ("C", 13), ("D", 20), ("E", 20), ("F", 26), ("G", 11), ("H", 3), ("I", 58), ("J", 12), ("K", 76), ("L", 62)):
    ws.column_dimensions[col].width = w

# ---------------------------------------------------------------- H3
ws = wb.create_sheet("H3")
SH = "H3"
ws["A1"] = "H3. Heart rate (BPM) change from baseline, kept respondents. SIMULATED data. The heart rate column name is MODELED; verify on the lab machine."
ws["A1"].font = Font(bold=True, size=12, color=NAVY)
ws["A2"] = ("Mean PPG heart rate per stimulus from the Data sheet (SUMIFS). Change = commercial mean minus the mean of the Baseline that preceded it. "
            "The difference of the changes is Nike change minus e.l.f. change. Same roster order as H1.")
ws["A2"].font = note_font
h3_cols = ["Respondent Name", "Group (order)", "Nonresponder", "Baseline_Nike HR", "Nike_Commercial HR", "Change Nike (commercial minus baseline)",
           "Baseline_elf HR", "elf_Commercial HR", "Change e.l.f. (commercial minus baseline)", "Difference of changes (Nike minus elf)"]
ws.append([])
ws.append(h3_cols)
style_header(ws, row=4, ncol=len(h3_cols))
ws.row_dimensions[4].height = 48
for k, rid in enumerate(roster):
    r = P1 + k
    ws.cell(row=r, column=1, value=rid)
    roster_formulas(ws, r)
    ws.cell(row=r, column=4, value=sumifs_hr(r, NIKE_B))
    ws.cell(row=r, column=5, value=sumifs_hr(r, NIKE_C))
    ws.cell(row=r, column=6, value=f"=E{r}-D{r}")
    ws.cell(row=r, column=7, value=sumifs_hr(r, ELF_B))
    ws.cell(row=r, column=8, value=sumifs_hr(r, ELF_C))
    ws.cell(row=r, column=9, value=f"=H{r}-G{r}")
    ws.cell(row=r, column=10, value=f"=F{r}-I{r}")
    for c in range(1, 11):
        ws.cell(row=r, column=c).border = box
        if c >= 4:
            ws.cell(row=r, column=c).number_format = "0.00"
for key, col, f, lab in (("row_bn", "D", sumifs_hr(P1, NIKE_B), "Baseline_Nike HR"), ("row_cn", "E", sumifs_hr(P1, NIKE_C), "Nike_Commercial HR"),
                         ("row_dn", "F", f"=E{P1}-D{P1}", "Change Nike"), ("row_be", "G", sumifs_hr(P1, ELF_B), "Baseline_elf HR"),
                         ("row_ce", "H", sumifs_hr(P1, ELF_C), "elf_Commercial HR"), ("row_de", "I", f"=H{P1}-G{P1}", "Change e.l.f."),
                         ("row_dd", "J", f"=F{P1}-I{P1}", "Difference of changes")):
    reg(SH, key, f"{col}{P1}", f, lab)
blk = Block(ws, SH, 12, 5, "hn")          # label column L, value M, formula N, note O
paired_block(blk, "Nike: heart rate in the commercial versus its baseline (change = mean difference)", rA("D", P2), rA("E", P2), rA("F", P2),
             "BPM", "Baseline_Nike", "Nike_Commercial", "COUNT of the change column; must equal the check below")
H3_NIKE = blk
blk = Block(ws, SH, 12, blk.next_start(), "he")
paired_block(blk, "e.l.f.: heart rate in the commercial versus its baseline (change = mean difference)", rA("G", P2), rA("H", P2), rA("I", P2),
             "BPM", "Baseline_elf", "elf_Commercial")
H3_ELF = blk
blk = Block(ws, SH, 12, blk.next_start(), "hd")
paired_block(blk, "Difference between the changes (Nike change minus e.l.f. change)", rA("I", P2), rA("F", P2), rA("J", P2),
             "BPM", "Change_elf", "Change_Nike", "the paired test on the two change columns")
H3_DIFF = blk
blk = Block(ws, SH, 12, blk.next_start(), "hchk")
blk.title("Checks")
blk.put("chk_n", "Kept respondents with a Nike_Commercial row in Data (Keep = Yes)", f'=COUNTIFS({R_LABEL},"{NIKE_C}",{R_KEEP},"Yes")',
        "must equal n pairs", fmt="0")
blk.put("chk_below_n", "Respondents whose heart rate fell during the Nike commercial", f'=COUNTIF(F{P1}:F{P2},"<0")', fmt="0")
blk.put("chk_below_e", "Respondents whose heart rate fell during the e.l.f. commercial", f'=COUNTIF(I{P1}:I{P2},"<0")', fmt="0")
H3_CHK = blk
for col, w in (("A", 17), ("B", 13), ("C", 13), ("D", 16), ("E", 18), ("F", 24), ("G", 16), ("H", 18), ("I", 24), ("J", 22), ("K", 3),
               ("L", 62), ("M", 12), ("N", 76), ("O", 60)):
    ws.column_dimensions[col].width = w

# ---------------------------------------------------------------- Check
ws = wb.create_sheet("Check")
ws["A1"] = "Live Excel results next to the values analysis.R produced (SIMULATED data). Status is OK when they agree within the tolerance."
ws["A1"].font = Font(bold=True, size=12, color=NAVY)
ws.append([])
ck_cols = ["Quantity", "Excel (live formula)", "analysis.R", "Difference", "Tolerance", "Status"]
ws.append(ck_cols)
style_header(ws, row=3, ncol=len(ck_cols))
_, th, trows = read_csv("SIMULATED_test_results.csv")
_, mh, mrows = read_csv("SIMULATED_model_results.csv")
_, dh, drows = read_csv("SIMULATED_descriptives_by_stimulus.csv")
T = {r[0]: dict(zip(th, r)) for r in trows}
M = {(r[0], r[1]): dict(zip(mh, r)) for r in mrows}
D = {r[0]: dict(zip(dh, r)) for r in drows}
AN = "Mixed model Peaks Per Minute in the commercials"
ck_rows = []


def add_check(label, ref, r_value, tol, note=""):
    ck_rows.append((label, f"={ref}", float(r_value), tol, note))


tests_map = [("H1_Nike", "H1", H1_NIKE, "H1 Nike PPM"), ("H1_elf", "H1", H1_ELF, "H1 e.l.f. PPM"), ("H2", "H2", H2_MAIN, "H2 Nike minus e.l.f. PPM"),
             ("H3_Nike", "H3", H3_NIKE, "H3 Nike HR"), ("H3_elf", "H3", H3_ELF, "H3 e.l.f. HR"), ("H3_diff", "H3", H3_DIFF, "H3 difference of changes"),
             ("S_H1_Nike", "H1", H1_SNIKE, "Sensitivity H1 Nike PPM (responders)"), ("S_H1_elf", "H1", H1_SELF, "Sensitivity H1 e.l.f. PPM (responders)"),
             ("S_H2", "H2", H2_SENS, "Sensitivity H2 PPM (responders)")]
for tid, sheet, b, nm in tests_map:
    t = T[tid]
    for key, rkey, tol, lab in (("m1", "m_ref", 1e-4, "mean, reference"), ("m2", "m_cmp", 1e-4, "mean, comparison"),
                                ("md", "mean_diff", 1e-4, "mean difference"), ("sdd", "sd_diff", 1e-4, "SD of differences"),
                                ("t", "t", 1e-4, "paired t"), ("df", "df", 1e-9, "df"), ("p_tt", "p", 5e-4 * float(t["p"]) + 1e-9, "p (T.TEST type 1)"),
                                ("ci_lo", "diff_ci_low", 1e-4, "mean difference CI lower"), ("ci_hi", "diff_ci_high", 1e-4, "mean difference CI upper"),
                                ("dz", "dz", 1e-4, "Cohen's d_z"), ("dz_lo", "dz_ci_low", 0.10, "d_z CI lower (method differs)"),
                                ("dz_hi", "dz_ci_high", 0.10, "d_z CI upper (method differs)")):
        add_check(f"{nm}: {lab}", f"{sheet}!{b.addr(key)}", t[rkey], tol, "Excel: large sample approximation; R: noncentral t" if key.startswith("dz_") else "")
mo, mc = M[(AN, "Order position (second minus first)")], M[(AN, "Commercial (Nike minus elf)")]
add_check("Order effect: simple estimate versus lmer order_position", f"H2!{H2_MAIN.addr('oe')}", mo["estimate"], 1e-4, "equal here because each group's mean vector is free in the saturated model")
add_check("Commercial effect: simple order adjusted estimate versus lmer", f"H2!{H2_MAIN.addr('commercial')}", mc["estimate"], 1e-4)
add_check("Order effect, responders only: simple estimate", f"H2!{H2_SENS_OE.addr('oe')}", M[("Sensitivity: responders only", "Simple order effect estimate ((Group B - Group A) / 2)")]["estimate"], 1e-4)
EX = "Exploratory: self reported arousal with Peaks Per Minute"
for key, term, bobj in (("cor", "r Nike", H1_CN), ("cor", "r elf", H1_CE), ("cor", "r respondent means", H1_CM)):
    m = M[(EX, term)]
    add_check(f"{term}: Pearson r", f"H1!{bobj.addr('r')}", m["estimate"], 1e-4)
    add_check(f"{term}: t", f"H1!{bobj.addr('t')}", m["statistic"], 1e-3)
    add_check(f"{term}: p", f"H1!{bobj.addr('p')}", m["p"], 5e-4 * float(m["p"]) + 1e-9)
    add_check(f"{term}: CI lower", f"H1!{bobj.addr('lo')}", m["ci_low"], 1e-4)
    add_check(f"{term}: CI upper", f"H1!{bobj.addr('hi')}", m["ci_high"], 1e-4)
for term, bobj in (("r Nike", H1_SCN), ("r elf", H1_SCE)):
    m = M[("Sensitivity: responders only", term)]
    add_check(f"Responders only, {term}: Pearson r", f"H1!{bobj.addr('r')}", m["estimate"], 1e-4)
    add_check(f"Responders only, {term}: p", f"H1!{bobj.addr('p')}", m["p"], 5e-4 * float(m["p"]) + 1e-9)
# descriptives
for lab in LABELS:
    d, rr = D[lab], summary_rows[lab]
    add_check(f"Summary {lab}: n", f"Summary!B{rr}", d["n"], 1e-9)
    add_check(f"Summary {lab}: PPM mean", f"Summary!C{rr}", d["ppm_mean"], 6e-4)
    add_check(f"Summary {lab}: PPM SD", f"Summary!D{rr}", d["ppm_sd"], 6e-4)
    add_check(f"Summary {lab}: Has Peaks (%)", f"Summary!E{rr}", d["pct_has_peaks"], 6e-2)
    add_check(f"Summary {lab}: peak amplitude mean", f"Summary!G{rr}", d["amp_mean"], 6e-4)
    add_check(f"Summary {lab}: heart rate mean", f"Summary!I{rr}", d["hr_mean_bpm"], 6e-4)
    add_check(f"Summary {lab}: heart rate SD", f"Summary!J{rr}", d["hr_sd_bpm"], 6e-4)
    add_check(f"Summary {lab}: SNR mean", f"Summary!K{rr}", d["snr_mean_db"], 6e-4)
# helper and consistency checks
n_h1, n_s = int(T["H1_Nike"]["n"]), int(T["S_H1_Nike"]["n"])
add_check("Data rows with Keep = No (R: 6 stimuli x excluded respondents)", f'COUNTIF({R_KEEP},"No")', 6 * (len(ids) - n_h1), 1e-9,
          "excluded respondents counted by the rule in R: 30 minus the kept n")
add_check("Kept nonresponder rows in Nike_Commercial (R: kept n minus responders n)", f'COUNTIFS({R_LABEL},"{NIKE_C}",{R_KEEP},"Yes",{R_NR},1)',
          n_h1 - n_s, 1e-9)
add_check("Largest gap between exported and recalculated Peaks Per Minute", f"Summary!B{sd0 + 2}", 0.0, 5.1e-4,
          "the export rounds Peaks Per Minute to 3 decimals; R reports the same gap in section 2")
with open(os.path.join(HERE, "analysis_output.txt"), encoding="utf-8") as fh:
    out_txt = fh.read()
m_gap = re.search(r"Among kept respondents: (\d+) rows with a gap, largest understatement ([0-9.]+) peaks per minute", out_txt)
assert m_gap, "section 9 line not found in analysis_output.txt"
add_check("Kept rows with Signal Duration shorter than Duration", f"Summary!B{sd0}", m_gap.group(1), 1e-9, "analysis_output.txt section 9")
add_check("Largest understatement if Duration were used (kept rows)", f"Summary!B{sd0 + 1}", m_gap.group(2), 1.1e-3, "analysis_output.txt section 9 (rounded to 3 decimals)")
m_dev = re.findall(r"^\s*(Shimmer3-\S+)\s+(\d+)\s+([0-9.]+)\s+([0-9.]+)\s+([0-9.]+)\s+([0-9.]+)\s*$", out_txt, flags=re.M)
for dname, nres, snr_m, snr_s, ppm_m, ppm_s in m_dev:
    add_check(f"Device {dname}: SNR mean", f"Summary!C{dev_rows[dname]}", snr_m, 6e-3, "analysis_output.txt section 9")
    add_check(f"Device {dname}: PPM mean (baselines and commercials)", f"Summary!D{dev_rows[dname]}", ppm_m, 6e-3, "analysis_output.txt section 9")
CK1 = 4
for k, (lab, f, rv, tol, note) in enumerate(ck_rows):
    r = CK1 + k
    ws.cell(row=r, column=1, value=lab)
    ws.cell(row=r, column=2, value=f).number_format = "0.000000"
    ws.cell(row=r, column=3, value=rv).number_format = "0.000000"
    ws.cell(row=r, column=4, value=f"=B{r}-C{r}").number_format = "0.000000"
    ws.cell(row=r, column=5, value=tol).number_format = "0.0000"
    ws.cell(row=r, column=6, value=f'=IF(ABS(D{r})<=E{r},"OK","CHECK")')
    if note:
        ws.cell(row=r, column=7, value=note).font = note_font
CK2 = CK1 + len(ck_rows) - 1
ws.cell(row=CK2 + 2, column=1, value="Rows with status CHECK").font = bold
ws.cell(row=CK2 + 2, column=2, value=f'=COUNTIF(F{CK1}:F{CK2},"CHECK")')
reg("Check", "n_check", f"B{CK2 + 2}", f'=COUNTIF(F{CK1}:F{CK2},"CHECK")', "Rows with status CHECK")
reg("Check", "status_row", f"F{CK1}", f'=IF(ABS(D{CK1})<=E{CK1},"OK","CHECK")', "Status")
ws.cell(row=CK2 + 3, column=1, value=("Not reproducible in Excel: the Wilcoxon signed rank tests, the Shapiro Wilk tests, the mixed model with its ICC, "
                                      "the noncentral t interval for d_z, the Spearman correlations, and the moment to moment bins (they need the GSR Peak Table). "
                                      "See analysis_output.txt.")).font = note_font
r = CK2 + 6
for title, hdr_, body in (("analysis.R: paired tests (SIMULATED_test_results.csv)", th, trows),
                          ("analysis.R: mixed model and correlations (SIMULATED_model_results.csv)", mh, mrows),
                          ("analysis.R: descriptives (SIMULATED_descriptives_by_stimulus.csv)", dh, drows)):
    ws.cell(row=r - 1, column=1, value=title).font = bold
    for c, h in enumerate(hdr_, start=1):
        ws.cell(row=r, column=c, value=h)
    style_header(ws, row=r, ncol=len(hdr_))
    for rw in body:
        r += 1
        for c, v in enumerate(rw, start=1):
            ws.cell(row=r, column=c, value=to_cell(v))
    r += 3
autowidth(ws, max_w=34)
ws.column_dimensions["A"].width = 66
for col in "BCDEF":
    ws.column_dimensions[col].width = 20

wb.calculation = CalcProperties(fullCalcOnLoad=True)    # Excel computes every formula when the file opens
wb.save(OUT)
print("wrote", OUT)

# =============================================================================
# Verification: read the workbook back with openpyxl and compare every registered cell
# =============================================================================
wb2 = load_workbook(OUT)
problems = []
for (sheet, key), info in REG.items():
    cell = wb2[sheet][info["cell"]]
    v = cell.value
    if isinstance(v, ArrayFormula):
        v = v.text
    v = str(v).replace("_xlfn.", "") if v is not None else None
    if v != info["formula"]:
        problems.append(f"{sheet}!{info['cell']} ({key}): workbook has {v!r}, registry has {info['formula']!r}")
for k, (lab, f, rv, tol, note) in enumerate(ck_rows):
    v = wb2["Check"][f"B{CK1 + k}"].value
    if v != f:
        problems.append(f"Check!B{CK1 + k}: {v!r} versus {f!r}")
if problems:
    print("VERIFICATION FAILED:")
    print("\n".join(problems[:40]))
    sys.exit(1)
print(f"verified: {len(REG)} registered cells and {len(ck_rows)} Check rows read back with openpyxl and match the registry")
n_formulas = sum(1 for wsx in wb2.worksheets for row in wsx.iter_rows() for c in row
                 if (isinstance(c.value, str) and c.value.startswith("=")) or isinstance(c.value, ArrayFormula))
print(f"formula cells in the workbook: {n_formulas}")

# =============================================================================
# The walkthrough, generated from REG so that every cell reference is the real one
# =============================================================================


def cl(sheet, key):
    return REG[(sheet, key)]["cell"]


def fx(sheet, key):
    return REG[(sheet, key)]["formula"]


def row(label, sheet, key, result=""):
    return f"| {label} | `{sheet}!{cl(sheet, key)}` | `{fx(sheet, key)}` | {result} |"


def row3(label, sheet, key):
    return f"| {label} | `{sheet}!{cl(sheet, key)}` | `{fx(sheet, key)}` |"


def rcell(sheet, key):
    return f"{sheet}!{cl(sheet, key)}"


def fnum(x, nd=2):
    return f"{float(x):.{nd}f}"


def approx_ci(t):
    n, dz = int(t["n"]), float(t["dz"])
    s = (1 / n + dz ** 2 / (2 * n)) ** 0.5
    return dz - 1.96 * s, dz + 1.96 * s, s


def pT(t):
    return f"{float(t['p']):.6f}"


def tcrit_of(t):
    return (float(t["diff_ci_high"]) - float(t["diff_ci_low"])) / 2 / (float(t["sd_diff"]) / int(t["n"]) ** 0.5)


h1n, h1e, h2, h3n, h3e, h3d = (T[k] for k in ("H1_Nike", "H1_elf", "H2", "H3_Nike", "H3_elf", "H3_diff"))
s1n, s1e, s2 = T["S_H1_Nike"], T["S_H1_elf"], T["S_H2"]
a_h1n, a_h1e, a_h2, a_h3n, a_h3e, a_h3d = (approx_ci(t) for t in (h1n, h1e, h2, h3n, h3e, h3d))


def _and(xs):
    xs = list(xs)
    return xs[0] if len(xs) == 1 else (" and ".join(xs) if len(xs) == 2 else ", ".join(xs[:-1]) + ", and " + xs[-1])


nr_txt = _and(nonresp)
excl_txt = _and(excluded)
cn, ce, cm = M[(EX, "r Nike")], M[(EX, "r elf")], M[(EX, "r respondent means")]
icc_adj = M[(AN, "ICC adjusted (fitted model)")]["estimate"]
grp_of = {k: by_resp[k][0][idx["Group"]] for k in kept_ids}
nA = sum(1 for k in kept_ids if grp_of[k] == "A")
nB = sum(1 for k in kept_ids if grp_of[k] == "B")
n_below_n = sum(1 for k in kept_ids if float(value_of(k, NIKE_C, "Peaks Per Minute")) > -1 and
                float(heart[(k, NIKE_C)][hidx["Mean PPG Heart Rate (BPM)"]]) < float(heart[(k, NIKE_B)][hidx["Mean PPG Heart Rate (BPM)"]]))
n_below_e = sum(1 for k in kept_ids if float(heart[(k, ELF_C)][hidx["Mean PPG Heart Rate (BPM)"]]) < float(heart[(k, ELF_B)][hidx["Mean PPG Heart Rate (BPM)"]]))
dd = {k: D[k] for k in D}


def _diff(k):
    return float(value_of(k, NIKE_C, "Peaks Per Minute")) - float(value_of(k, ELF_C, "Peaks Per Minute"))


gA = st.mean(_diff(k) for k in kept_ids if grp_of[k] == "A")
gB = st.mean(_diff(k) for k in kept_ids if grp_of[k] == "B")
gaps_kept = sorted(int(r[idx["Duration (ms)"]]) - int(r[idx["Signal Duration (ms)"]]) for r in rows
                   if r[idx["Respondent Name"]] in kept_ids and r[idx["Signal Duration (ms)"]] != r[idx["Duration (ms)"]])
gap_lo, gap_hi = gaps_kept[0] / 1000, gaps_kept[-1] / 1000
md = f"""# GSR and heart rate analysis in Excel: step by step walkthrough

**SIMULATED DATA.** Every number in this walkthrough comes from the SIMULATED files that `make_simulated_data.py` generated (seed {SEED}). No real person was recorded. The design follows the structure of the lab's own library studies "Nike Campaign Ad Evaluation" and "e.l.f. Campaign Ad Evaluation" (a 60 s Baseline, the commercial, a 15 s Recovery, and a short questionnaire; the Nike commercial is 60 s and the e.l.f. commercial 30 s); nothing about the content of either commercial was used.

This is the same analysis as `analysis.R`, done in Excel for students who do not use R. The finished workbook is `SIMULATED_gsr_analysis.xlsx`; the sheet names below refer to it. You can follow along in that workbook or rebuild it from the CSV files. **Every cell address in this walkthrough is written as `Sheet!Cell` and was generated from the workbook itself** (`make_excel_workbook.py` writes both and then reads the workbook back with openpyxl to confirm that every address holds the formula printed here), so the addresses match the file.

Companion files: `analysis.R` (R version), `analysis_output.txt` (its console output), `results_text.md` (model paragraphs), `figures/` (PNG figures).

---

## 0. What you are analyzing

**Design.** Within subjects. Each respondent (P01 to P30) watched two sequences on the lab monitor with a Shimmer3 GSR+ on the non-dominant hand: a Baseline (60 s), the commercial, a Recovery (15 s), and a short questionnaire, for Nike and for e.l.f., with a 2 minute rest between the two sequences. The order was counterbalanced: Group A saw Nike first, Group B saw e.l.f. first (15 each). Respondents were run two at a time at Seat 1 (unit `Shimmer3-8956`) and Seat 2 (unit `Shimmer3-B`, a placeholder name to verify on the lab machine). The analysis pools both devices, as the **Multiple device aggregation** option of the GSR R Notebooks does. The six stimulus labels are `Baseline_Nike`, `Nike_Commercial`, `Recovery_Nike`, `Baseline_elf`, `elf_Commercial`, and `Recovery_elf`.

**Exports.** Four files feed the workbook.

| File | iMotions export | Columns used here | Used for |
|---|---|---|---|
| `SIMULATED_gsr_peak_metrics.csv` | GSR Peak Metrics (Export: R Processed GSR Files), one row per respondent and stimulus | `Respondent Name`, `Group`, `Label`, `Duration (ms)`, `Signal Duration (ms)`, `Has Peaks`, `Peak Count`, `Peaks Per Minute`, `Average Peak Amplitude (microSiemens)` | H1, H2, descriptives |
| `SIMULATED_gsr_signal_quality.csv` | GSR Signal Quality export (the layout is modeled) | `Signal-to-Noise ratio (dB)`, `Signal missing (%)`, `Device` | Exclusion rule, device check |
| `SIMULATED_heart_rate.csv` | Mean of the PPG beats per minute channel per stimulus (the column name `Mean PPG Heart Rate (BPM)` is **modeled**; the manuals name no PPG heart rate column, so verify it on the lab machine) | `Mean PPG Heart Rate (BPM)` | H3 |
| `SIMULATED_gsr_survey.csv` | Questionnaire | `Arousal_Nike_1to9`, `Arousal_elf_1to9` | Exploratory correlation |

A fifth file, `SIMULATED_gsr_peak_table.csv` (one row per detected peak), is used in R for the moment to moment figure and for consistency checks; the workbook does not need it. `SIMULATED_gsr_session_log.csv` is the operator log (seat, device, electrode placement, skin note, warm up minutes, decision notes).

`Peaks Per Minute` is `Peak Count` divided by `Signal Duration` in minutes. `Signal Duration (ms)` leaves out gaps longer than 4 seconds, so a Bluetooth dropout shortens it (section 8). `Has Peaks` is 1 when the stimulus held at least one peak and 0 when it held none. `Average Peak Amplitude` is NA when there is no peak, because it cannot be calculated.

**Hypotheses.**
- H1 (arousal response): Peaks Per Minute is higher during each commercial than during its own baseline.
- H2: Peaks Per Minute differs between the Nike and the e.l.f. commercial, checked for an order (carryover) effect.
- H3: mean heart rate falls during each commercial relative to its baseline (a deceleration), and the change differs between the commercials.
- Exploratory: self reported arousal is related to Peaks Per Minute.
- Sensitivity: the H1 and H2 results without the nonresponders.

**Decision rules fixed before looking at the data.** Exclude a respondent whose Signal-to-Noise ratio is below 0 dB or whose Signal missing is above 20 percent in any commercial (the lab's rule, built from the GSR Signal Quality article, which calls an SNR below 0 dB poor and above 20 dB good). A **nonresponder** has `Has Peaks` = 0 in at least 4 of the 6 stimuli; nonresponders stay in the main analysis. H1 and H3 differences are **commercial minus baseline**; H2 is **Nike minus e.l.f.** Two tailed tests, alpha = .05.

---

## 1. Import the files and build the Data sheet

iMotions exports are comma separated with a dot decimal. Double clicking a file uses your regional defaults, which can put everything in one column (Opening Data Export Files in Excel article). Use the import route instead.

**Route A, Get Data (Excel 2016 and later):**
1. **Data > From Text/CSV**, select `SIMULATED_gsr_peak_metrics.csv`, click **Import**.
2. In the preview set **Delimiter: Comma** and **File Origin: 65001 Unicode (UTF-8)**. **You should see** {PM_COLS} columns with `Study Name` in the first column.
3. Click **Transform Data** if the first row shows the `# SIMULATED ...` comment line: in Power Query choose **Home > Remove Rows > Remove Top Rows > 1**, then **Home > Use First Row as Headers**, then **Close & Load**. (A real iMotions export has no comment line; the simulated files carry one so nobody mistakes them for real data. The Peak Table has a second comment line, the METADATA block, so remove 2 rows there.)
4. **You should see** a table with {n_data} data rows (30 respondents x 6 stimuli).

**Route B, Text Import Wizard (any version):** **File > Open > Browse**, set the file type to **All Files**, open the .csv, choose **Delimited**, start at row 2 (to skip the comment line), tick **Comma**, click **Advanced** and set the decimal symbol to a dot, then **Finish**.

**Joining the other files.** The Peak Metrics file has no Device, SNR, or heart rate. Each of those files has one row per respondent and stimulus, so add a key column to each (`=B2&"|"&G2` on the Peak Metrics sheet, where B is `Respondent Name` and G is `Label`; the key of the other files uses their own respondent and stimulus columns) and pull the values across with `INDEX` and `MATCH` (or `XLOOKUP` in Excel 365). In the finished workbook this is already done: columns {COL['Device']} to {COL['Arousal_1to9 (survey)']} of the **Data** sheet hold the matched values as pasted values (gray headers).

**Common mistake.** The NA cells in `Average Peak Amplitude` must stay as the text `NA` (or be blank). Do not replace NA with 0: a zero amplitude would pull the mean down for every respondent who had no peak.

Save as **.xlsx** right away; a .csv does not keep formulas or extra sheets. Keep the original .csv files read only.

In the workbook this is the **Data** sheet, header in row 1 and data in rows {FIRST} to {LAST}. Column letters there: `{COL['Respondent Name']}` = Respondent Name, `{COL['Group']}` = Group (the order), `{COL['Label']}` = Label, `{COL['Duration (ms)']}` = Duration (ms), `{COL['Signal Duration (ms)']}` = Signal Duration (ms), `{COL['Has Peaks']}` = Has Peaks, `{COL['Peak Count']}` = Peak Count, `{COL['Peaks Per Minute']}` = Peaks Per Minute, `{COL['Average Peak Amplitude (microSiemens)']}` = Average Peak Amplitude, `{COL['Device']}` = Device, `{COL['Signal-to-Noise ratio (dB)']}` = Signal-to-Noise ratio (dB), `{COL['Signal missing (%)']}` = Signal missing (%), `{COL['Mean PPG Heart Rate (BPM)']}` = Mean PPG Heart Rate (BPM), `{COL['Arousal_1to9 (survey)']}` = Arousal_1to9 (survey; filled on the commercial rows only).

---

## 2. Apply the rules with helper columns

Add seven columns to the right of the data (`{COL['Phase']}` to `{COL['PPM_if_Duration']}` in the workbook, headers in orange) and fill them down to the last row. The formulas below are the ones in row {FIRST}.

| Header | Cell | Formula | Meaning |
|---|---|---|---|
| Phase | `Data!{cl('Data','phase')}` | `{fx('Data','phase')}` | Baseline, Commercial, or Recovery, read from the label |
| Brand | `Data!{cl('Data','brand')}` | `{fx('Data','brand')}` | Nike or elf |
| Position | `Data!{cl('Data','position')}` | `{fx('Data','position')}` | Was this sequence the first or the second one this respondent saw? Group A saw Nike first, Group B saw e.l.f. first |
| Keep | `Data!{cl('Data','keep')}` | `{fx('Data','keep')}` | The exclusion rule: No when any commercial row of this respondent has SNR below 0 dB or Signal missing above 20 percent |
| Nonresponder | `Data!{cl('Data','nonresp')}` | `{fx('Data','nonresp')}` | 1 when `Has Peaks` is 0 in at least 4 of this respondent's 6 stimuli |
| PPM_recalc | `Data!{cl('Data','recalc')}` | `{fx('Data','recalc')}` | Peak Count divided by Signal Duration in minutes, to compare with the exported Peaks Per Minute |
| PPM_if_Duration | `Data!{cl('Data','ifdur')}` | `{fx('Data','ifdur')}` | Peak Count divided by Duration in minutes (the wrong denominator, for the check in section 8) |

**You should see** `Keep = No` on all {6 * len(excluded)} rows of respondents {excl_txt} and `Yes` on every other row. {excluded[0]} is excluded for an SNR below 0 dB in the Nike commercial (a loose electrode lead), and {excluded[1]} for Signal missing of 38.0 percent in the e.l.f. commercial (a Bluetooth dropout); the session log records both decisions. Respondents {nr_txt} have `Nonresponder = 1` on all six of their rows. {len(nonresp)} of the {N_KEPT} kept respondents ({100 * len(nonresp) / N_KEPT:.1f} percent) are nonresponders; they are kept and checked in the sensitivity analysis. {nA} kept respondents are in Group A and {nB} in Group B.

`COUNTIFS` inside `Keep` and `Nonresponder` counts rows of the same respondent on the whole Data range, so the rule is evaluated per respondent, not per row. A low SNR in a Baseline or Recovery row does not exclude anyone; the rule looks at the commercials only.

**Common mistake.** Applying the SNR rule to every stimulus. The rule was written for the commercials, because they carry the hypotheses.

---

## 3. Descriptive table by stimulus

### 3a. With a PivotTable (the iMotions recommended route)

The Creating Excel Pivot Tables article gives the generic steps: click the first cell of the data, **Insert > PivotTable > OK**, drag a category to **Rows**, drag a metric to **Values**, right click the value, **Value Field Settings**, change **Sum** to **Average**, **OK**. For this study:

1. Click cell A1 of the Data sheet, **Insert > PivotTable**, **New Worksheet**, **OK**.
2. Drag **Keep** to **Filters** and set it to **Yes**.
3. Drag **Label** to **Rows**.
4. Drag **Peaks Per Minute** to **Values** twice. Set the first to **Average** and the second to **StdDev** (the sample standard deviation, the same as `STDEV.S`).
5. Drag **Has Peaks** to **Values** and set it to **Average**; multiply by 100 beside the table to get the percent of respondents with at least one peak.

**You should see** `Nike_Commercial` with an average Peaks Per Minute of {dd['Nike_Commercial']['ppm_mean']} and `Baseline_Nike` with {dd['Baseline_Nike']['ppm_mean']}.

**Common mistake.** Leaving **Keep** out of the filter. {excl_txt} would then add twelve rows, and the loose electrode lead of {excluded[0]} would add nine spurious peaks to the Nike commercial.

### 3b. With formulas (the Summary sheet)

A formula table updates when the data change and shows exactly what is being averaged. The stimulus label is in column A of each row (row {SUM_FIRST} is the first, `Baseline_Nike`). The formulas in row {SUM_FIRST}:

| Statistic | Cell | Formula |
|---|---|---|
""" + "\n".join(
    f"| {REG[('Summary', k)]['label']} | `Summary!{cl('Summary', k)}` | `{fx('Summary', k)}` |"
    for k in ("n", "ppm_mean", "ppm_sd", "has_pct", "count_mean", "amp_mean", "n_amp", "hr_mean", "hr_sd", "snr_mean", "miss_mean")) + f"""

The two `STDEV.S(IF(...))` formulas (columns D and J) are **array formulas** because Excel has no `STDEVIFS`. In Excel 365 type them and press **Enter**. In Excel 2019 or older press **Ctrl+Shift+Enter**; Excel then shows the formula in braces. `AVERAGEIFS` and `COUNTIFS` ignore the text NA in the amplitude column, which is exactly the "amplitude not calculable" rule.

**You should see** the same numbers as the PivotTable, in rows {SUM_FIRST} to {SUM_LAST} (one row per stimulus), and a column chart of mean Peaks Per Minute by stimulus under the tables. In the simulation the mean Peaks Per Minute is {dd['Baseline_Nike']['ppm_mean']} and {dd['Baseline_elf']['ppm_mean']} in the baselines, {dd['Nike_Commercial']['ppm_mean']} in the Nike commercial, {dd['elf_Commercial']['ppm_mean']} in the e.l.f. commercial, and {dd['Recovery_Nike']['ppm_mean']} and {dd['Recovery_elf']['ppm_mean']} in the recoveries. The recoveries last only 15 s, so Peaks Per Minute there moves in steps of 4 (one peak is 4 peaks per minute) and is too coarse for a hypothesis test; the workshop reports it descriptively only.

---

## 4. H1: arousal response, each commercial against its own baseline

A paired test needs the two measurements of the same respondent side by side, so first reshape the long export into one row per respondent. The **H1** sheet does this for the {N_KEPT} kept respondents (listed in column A, rows {P1} to {P2}); the {N_RESP} responders come first (rows {P1} to {PR}) and the {len(nonresp)} nonresponders last (rows {PR + 1} to {P2}). Each of columns D, E, G, and H is a `SUMIFS` that picks the one matching Peaks Per Minute row from the Data sheet:

| Column | Cell (first respondent) | Formula |
|---|---|---|
""" + "\n".join(
    f"| {REG[('H1', k)]['label']} | `H1!{REG[('H1', k)]['cell']}` | `{REG[('H1', k)]['formula']}` |"
    for k in ("row_grp", "row_nr", "row_bn", "row_cn", "row_dn", "row_be", "row_ce", "row_de")) + f"""

(By hand you would get the same table from the PivotTable of section 3a: put **Respondent Name** in Rows, **Label** in Columns, **Peaks Per Minute** in Values as Sum, and filter **Keep** to `Yes`.)

Now the test for **Nike**. The statistics sit in column P of the H1 sheet with the formula text shown beside them in column Q:

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('Baseline_Nike mean (peaks per minute)', 'H1', 'nike_m1', fnum(h1n['m_ref']))}
{row('Baseline_Nike SD', 'H1', 'nike_s1', fnum(dd['Baseline_Nike']['ppm_sd']))}
{row('Nike_Commercial mean', 'H1', 'nike_m2', fnum(h1n['m_cmp']))}
{row('Nike_Commercial SD', 'H1', 'nike_s2', fnum(dd['Nike_Commercial']['ppm_sd']))}
{row('n pairs', 'H1', 'nike_n', h1n['n'])}
{row('Mean difference, commercial minus baseline', 'H1', 'nike_md', fnum(h1n['mean_diff']))}
{row('SD of the differences', 'H1', 'nike_sdd', fnum(h1n['sd_diff']))}
{row('**p, two tailed (T.TEST type 1)**', 'H1', 'nike_p_tt', '**' + pT(h1n) + '**')}

`T.TEST(array1, array2, tails, type)`: `tails = 2` for a two tailed test; `type = 1` for **paired** data, which is the type for this within subjects design. `type = 2` (equal variances) and `type = 3` (unequal variances) treat the two columns as different people and give the wrong answer for paired data. `T.TEST` returns only the p value, so compute t and df by hand with the paired formulas (a paired t test is a one sample t test on the differences):

| Statistic | Cell | Formula | Result |
|---|---|---|---|
{row('SE of the mean difference', 'H1', 'nike_se', fnum(float(h1n['sd_diff']) / int(h1n['n']) ** 0.5, 3))}
{row('Paired t', 'H1', 'nike_t', fnum(h1n['t']))}
{row('df', 'H1', 'nike_df', h1n['df'])}
{row('p from the hand computed t', 'H1', 'nike_p_hand', pT(h1n) + ', matches T.TEST')}
{row('Critical t (95%)', 'H1', 'nike_tcrit', fnum(tcrit_of(h1n), 3))}
{row('Mean difference 95% CI lower', 'H1', 'nike_ci_lo', fnum(h1n['diff_ci_low']))}
{row('Mean difference 95% CI upper', 'H1', 'nike_ci_hi', fnum(h1n['diff_ci_high']))}

**You should see** the hand computed p equal the `T.TEST` p. If they differ, the `type` argument or a column range is wrong.

### Cohen's d_z by hand, with a 95 percent confidence interval

For paired data the effect size is **d_z**: the mean difference in units of the SD of the *differences*. It is linked to t by t = d_z x sqrt(n), so t, n, and d_z always agree.

| Step | Cell | Formula | Result |
|---|---|---|---|
{row(chr(100) + "_z (Nike)", 'H1', 'nike_dz', fnum(h1n['dz']))}
{row('SE of d_z (approximation)', 'H1', 'nike_se_dz', fnum(a_h1n[2], 3))}
{row('d_z 95% CI lower', 'H1', 'nike_dz_lo', fnum(a_h1n[0]))}
{row('d_z 95% CI upper', 'H1', 'nike_dz_hi', fnum(a_h1n[1]))}
{row('Correlation of the two columns', 'H1', 'nike_r', fnum(st.correlation([float(value_of(k, NIKE_B, 'Peaks Per Minute')) for k in roster], [float(value_of(k, NIKE_C, 'Peaks Per Minute')) for k in roster])))}

In words: d_z = mean difference / SD of differences; SE(d_z) = sqrt(1/n + d_z^2 / (2 n)). The interval is the usual large sample approximation. `analysis.R` reports the interval from the noncentral t distribution (effectsize package), which is more exact at this sample size; d_z itself matches, and the limits differ slightly ([{fnum(h1n['dz_ci_low'])}, {fnum(h1n['dz_ci_high'])}] in R). Say which method you used. Rule of thumb for reading d_z: 0.2 small, 0.5 medium, 0.8 large.

### The same test for e.l.f.

The e.l.f. block repeats the recipe on columns G, H, and I (the baseline, the commercial, and their difference). The e.l.f. commercial lasts 30 s, so each peak adds 2 peaks per minute to a respondent's rate and the values move in steps of 2 (apart from the rows where a gap shortened `Signal Duration`).

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('Baseline_elf mean', 'H1', 'elf_m1', fnum(h1e['m_ref']))}
{row('elf_Commercial mean', 'H1', 'elf_m2', fnum(h1e['m_cmp']))}
{row('n pairs', 'H1', 'elf_n', h1e['n'])}
{row('Mean difference, commercial minus baseline', 'H1', 'elf_md', fnum(h1e['mean_diff']))}
{row('SD of the differences', 'H1', 'elf_sdd', fnum(h1e['sd_diff']))}
{row('Paired t', 'H1', 'elf_t', fnum(h1e['t']))}
{row('df', 'H1', 'elf_df', h1e['df'])}
{row('p, two tailed (T.TEST type 1)', 'H1', 'elf_p_tt', pT(h1e))}
{row(chr(100) + "_z (e.l.f.)", 'H1', 'elf_dz', fnum(h1e['dz']))}
{row('d_z 95% CI lower', 'H1', 'elf_dz_lo', fnum(a_h1e[0]))}
{row('d_z 95% CI upper', 'H1', 'elf_dz_hi', fnum(a_h1e[1]))}

The Wilcoxon signed rank tests, a check that the conclusion does not depend on normality, have no Excel function; in R: Nike V = {h1n['wilcoxon_V']}, e.l.f. V = {h1e['wilcoxon_V']}. The differences are built from whole peak counts, so they cannot be perfectly normal; the two tests agree in this sample.

The block below the e.l.f. blocks, with the checks, counts the kept respondents again from the Data sheet:

| Statistic | Cell | Formula |
|---|---|---|
{row3('Kept respondents with a Nike_Commercial row in Data', 'H1', 'chk_chk_n')}
{row3('Kept responders with a Nike_Commercial row in Data', 'H1', 'chk_chk_resp')}
{row3('Nonresponders listed in the roster', 'H1', 'chk_chk_nr')}
{row3('Nonresponders inside the responder rows', 'H1', 'chk_chk_nr_top')}

**You should see** {N_KEPT} in the first (it must equal n pairs), {N_RESP} in the second, {len(nonresp)} in the third, and 0 in the fourth. If the last one is not 0, a nonresponder is listed among the responders and the sensitivity blocks below are wrong.

**Common mistake.** Pairing the wrong rows. The two ranges must list the respondents in the same order, which is why every column of the H1 sheet is a lookup on the respondent name in column A rather than a pasted column.

---

## 5. H2: Nike versus e.l.f., and the order effect

The **H2** sheet lists the two commercial values per respondent (column D e.l.f., column E Nike, column F the difference **Nike minus e.l.f.**, column G whether Nike was the first or the second sequence) in the same roster order as H1. The paired recipe sits in column J (formula text in column K):

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('elf_Commercial mean (peaks per minute)', 'H2', 'h2_m1', fnum(h2['m_ref']))}
{row('Nike_Commercial mean', 'H2', 'h2_m2', fnum(h2['m_cmp']))}
{row('n pairs', 'H2', 'h2_n', h2['n'])}
{row('Mean difference, Nike minus elf', 'H2', 'h2_md', fnum(h2['mean_diff']))}
{row('SD of the differences', 'H2', 'h2_sdd', fnum(h2['sd_diff']))}
{row('Paired t', 'H2', 'h2_t', fnum(h2['t']))}
{row('df', 'H2', 'h2_df', h2['df'])}
{row('p, two tailed (T.TEST type 1)', 'H2', 'h2_p_tt', pT(h2))}
{row(chr(100) + "_z", 'H2', 'h2_dz', fnum(h2['dz']))}
{row('d_z 95% CI lower', 'H2', 'h2_dz_lo', fnum(a_h2[0]))}
{row('d_z 95% CI upper', 'H2', 'h2_dz_hi', fnum(a_h2[1]))}
{row('Correlation of the two columns', 'H2', 'h2_r', fnum(st.correlation([float(value_of(k, ELF_C, 'Peaks Per Minute')) for k in roster], [float(value_of(k, NIKE_C, 'Peaks Per Minute')) for k in roster])))}

Because every respondent saw both commercials, the second one may be affected by the first (carryover; for GSR the usual suspect is habituation, a smaller response to a repeated kind of stimulus). The simulated viewers were built with about 0.4 fewer peaks per minute for the commercial seen second and 0.4 more for the one seen first. The paired difference carries this: for Group A (Nike first) the difference Nike minus e.l.f. is commercial effect minus order effect, and for Group B (Nike second) it is commercial effect plus order effect. So half the gap between the two group means is the order effect:

| Statistic | Cell | Formula | Result |
|---|---|---|---|
{row('Mean difference, Group A (Nike seen first)', 'H2', 'h2_oe_A', fnum(gA))}
{row('Mean difference, Group B (Nike seen second)', 'H2', 'h2_oe_B', fnum(gB))}
{row('n, Group A', 'H2', 'h2_oe_nA', nA)}
{row('n, Group B', 'H2', 'h2_oe_nB', nB)}
{row('Order effect (second minus first)', 'H2', 'h2_oe', fnum(mo['estimate']))}
{row('Commercial effect (Nike minus elf), order adjusted', 'H2', 'h2_commercial', fnum(mc['estimate']))}

`analysis.R` fits `ppm ~ commercial * order_position + (1 | respondent)` with lme4 and lmerTest and gets the order effect b = {fnum(mo['estimate'])}, SE = {fnum(mo['se'])}, p = {fnum(mo['p'], 3)}, and the commercial effect {fnum(mc['estimate'])}; the simple Excel estimate coincides with the model coefficients because each order group's pair of means is free in the model. Excel cannot supply the standard error with the right degrees of freedom, the interaction test, or the **intraclass correlation** (ICC = {fnum(icc_adj, 2)} after the fixed effects): that is the share of the remaining variance that lies between respondents, the reason the two rows of a respondent are not independent. Use R for those.

The check at the end of the block:

| Statistic | Cell | Formula |
|---|---|---|
{row3('Kept respondents with a Nike_Commercial row in Data', 'H2', 'h2_chk_n')}

**You should see** {N_KEPT}, equal to n pairs.

### Sensitivity without the nonresponders

A second block on the H2 sheet (column J, below the first) repeats the test on rows {P1} to {PR} only, the {N_RESP} responders:

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('n pairs (responders)', 'H2', 'h2s_n', s2['n'])}
{row('Mean difference, Nike minus elf', 'H2', 'h2s_md', fnum(s2['mean_diff']))}
{row('Paired t', 'H2', 'h2s_t', fnum(s2['t']))}
{row('p, two tailed (T.TEST type 1)', 'H2', 'h2s_p_tt', pT(s2))}
{row(chr(100) + "_z", 'H2', 'h2s_dz', fnum(s2['dz']))}
{row('Order effect, responders only', 'H2', 'h2s_oe', fnum(M[('Sensitivity: responders only', 'Simple order effect estimate ((Group B - Group A) / 2)')]['estimate']))}

---

## 6. H3: heart rate change from baseline

The **H3** sheet holds the mean heart rate per stimulus (`SUMIFS` on the Data sheet) and the change (commercial minus the baseline that preceded it) in the same roster order:

| Column | Cell (first respondent) | Formula |
|---|---|---|
""" + "\n".join(
    f"| {REG[('H3', k)]['label']} | `H3!{REG[('H3', k)]['cell']}` | `{REG[('H3', k)]['formula']}` |"
    for k in ("row_bn", "row_cn", "row_dn", "row_be", "row_ce", "row_de", "row_dd")) + f"""

Three paired blocks sit in column M (formula text in column N): the Nike commercial against its baseline, the e.l.f. commercial against its baseline, and the **difference between the two changes** (Nike change minus e.l.f. change). In the first two, the mean difference **is** the heart rate change.

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('Nike: baseline mean (BPM)', 'H3', 'hn_m1', fnum(h3n['m_ref']))}
{row('Nike: commercial mean (BPM)', 'H3', 'hn_m2', fnum(h3n['m_cmp']))}
{row('Nike: mean change (BPM)', 'H3', 'hn_md', fnum(h3n['mean_diff']))}
{row('Nike: SD of the changes', 'H3', 'hn_sdd', fnum(h3n['sd_diff']))}
{row('Nike: paired t', 'H3', 'hn_t', fnum(h3n['t']))}
{row('Nike: p (T.TEST type 1)', 'H3', 'hn_p_tt', pT(h3n))}
{row('Nike: ' + chr(100) + '_z', 'H3', 'hn_dz', fnum(h3n['dz']))}
{row('e.l.f.: mean change (BPM)', 'H3', 'he_md', fnum(h3e['mean_diff']))}
{row('e.l.f.: paired t', 'H3', 'he_t', fnum(h3e['t']))}
{row('e.l.f.: p (T.TEST type 1)', 'H3', 'he_p_tt', pT(h3e))}
{row('e.l.f.: ' + chr(100) + '_z', 'H3', 'he_dz', fnum(h3e['dz']))}
{row('Difference of the changes, Nike minus e.l.f. (BPM)', 'H3', 'hd_md', fnum(h3d['mean_diff']))}
{row('Difference of the changes: SD', 'H3', 'hd_sdd', fnum(h3d['sd_diff']))}
{row('Difference of the changes: paired t', 'H3', 'hd_t', fnum(h3d['t']))}
{row('Difference of the changes: p (T.TEST type 1)', 'H3', 'hd_p_tt', pT(h3d))}
{row('Difference of the changes: ' + chr(100) + '_z', 'H3', 'hd_dz', fnum(h3d['dz']))}

A negative change means the heart rate was lower during the commercial than during its baseline. Heart rate comes from the PPG channel of the Shimmer3, a different sensor from the GSR electrodes, so a nonresponder in GSR can still show a normal heart rate change. The heart rate numbers rest on a modeled column name (see section 0); on real data first confirm in a pilot export that the beats per minute channel exists and what it contains.

The checks at the end of the H3 sheet:

| Statistic | Cell | Formula |
|---|---|---|
{row3('Kept respondents with a Nike_Commercial row in Data', 'H3', 'hchk_chk_n')}
{row3('Respondents whose heart rate fell during the Nike commercial', 'H3', 'hchk_chk_below_n')}
{row3('Respondents whose heart rate fell during the e.l.f. commercial', 'H3', 'hchk_chk_below_e')}

**You should see** {N_KEPT}, then {n_below_n}, then {n_below_e}.

---

## 7. Exploratory correlation and the nonresponder sensitivity analysis

### Self reported arousal and Peaks Per Minute

Self reported arousal (a 1 to 9 self assessment manikin style rating) is collected after each sequence. Columns J and K of the H1 sheet hold it for each respondent, and columns L and M hold the respondent means over the two commercials. The correlation blocks in column P use `CORREL`, then the usual t and Fisher z formulas:

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('Nike: Pearson r (arousal with Peaks Per Minute)', 'H1', 'cn_r', fnum(cn['estimate']))}
{row('Nike: t for r', 'H1', 'cn_t', fnum(cn['statistic']))}
{row('Nike: p, two tailed', 'H1', 'cn_p', fnum(cn['p'], 4))}
{row('Nike: r 95% CI lower (Fisher z)', 'H1', 'cn_lo', fnum(cn['ci_low']))}
{row('Nike: r 95% CI upper (Fisher z)', 'H1', 'cn_hi', fnum(cn['ci_high']))}
{row('e.l.f.: Pearson r', 'H1', 'ce_r', fnum(ce['estimate']))}
{row('Respondent means: Pearson r', 'H1', 'cm_r', fnum(cm['estimate']))}

The Spearman rank correlations and the check that no correlation hides behind the outlying nonresponders are in R. Because a rating and Peaks Per Minute are different kinds of measure (what the viewer says and what the skin did), a modest correlation is the expected result; GSR also measures only arousal, not whether the feeling is positive or negative.

### Sensitivity analysis: without the nonresponders

Nonresponders ({nr_txt}) show no peaks in most stimuli. They are a physiological finding (some people have very little phasic skin conductance activity), not a data fault, so the main analysis keeps them and the sensitivity analysis removes them. The roster puts them in the last {len(nonresp)} rows (rows {PR + 1} to {P2}), so the sensitivity blocks only change the row range from `{P1}:{P2}` to `{P1}:{PR}`:

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('Nike, responders only: n pairs', 'H1', 'snike_n', s1n['n'])}
{row('Nike, responders only: mean difference', 'H1', 'snike_md', fnum(s1n['mean_diff']))}
{row('Nike, responders only: paired t', 'H1', 'snike_t', fnum(s1n['t']))}
{row('Nike, responders only: p (T.TEST type 1)', 'H1', 'snike_p_tt', pT(s1n))}
{row('Nike, responders only: ' + chr(100) + '_z', 'H1', 'snike_dz', fnum(s1n['dz']))}
{row('e.l.f., responders only: mean difference', 'H1', 'self_md', fnum(s1e['mean_diff']))}
{row('e.l.f., responders only: paired t', 'H1', 'self_t', fnum(s1e['t']))}
{row('e.l.f., responders only: p (T.TEST type 1)', 'H1', 'self_p_tt', pT(s1e))}
{row('e.l.f., responders only: ' + chr(100) + '_z', 'H1', 'self_dz', fnum(s1e['dz']))}

Compare the effect sizes with the full sample: removing the nonresponders raises d_z for Nike from {fnum(h1n['dz'])} to {fnum(s1n['dz'])} and for e.l.f. from {fnum(h1e['dz'])} to {fnum(s1e['dz'])}. The conclusions do not change (every test stays on the same side of .05), and that is what a sensitivity analysis is for. The nonresponders also raise the correlation between a respondent's baseline and commercial values; in R, the correlations among responders only are lower.

---

## 8. Why Peaks Per Minute divides by Signal Duration

`Signal Duration (ms)` is the duration of the recorded signal without gaps longer than 4 seconds, and `Peaks Per Minute` divides `Peak Count` by it. A Bluetooth dropout hides a stretch of the stimulus from the sensor; no peak can be detected in the gap, so dividing by the full `Duration` would count the gap as time without arousal. The Summary sheet tests this:

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('Rows with Signal Duration shorter than Duration (kept respondents)', 'Summary', 'gap_rows', '3')}
{row('Largest understatement if Duration were used (peaks per minute)', 'Summary', 'understate', '0.415')}
{row('Largest gap between exported Peaks Per Minute and PPM_recalc', 'Summary', 'max_gap', '0.0004 or less (the export rounds to 3 decimals)')}

In this data set the biggest effect is on the excluded respondent {excluded[1]}, whose e.l.f. commercial row has a Signal Duration of 18600 ms out of 30000 ms: 2 peaks are 6.452 peaks per minute by Signal Duration and 4.000 by Duration. Among the kept respondents the gaps are shorter ({gap_lo:.1f} to {gap_hi:.1f} s of a 60 s stimulus), so the understatement is at most about 0.4 peaks per minute, but it would be a systematic error in the same direction for everyone who had a dropout.

**You should see** the exported `Peaks Per Minute` (column {COL['Peaks Per Minute']}) and `PPM_recalc` (column {COL['PPM_recalc']}) agree on every row.

**Common mistake.** Recomputing Peaks Per Minute yourself from `Peak Count` and `Duration (ms)`. Use `Signal Duration (ms)`.

The device check on the Summary sheet (a small table below the Signal Duration block, columns A to D) shows the mean SNR and the mean Peaks Per Minute of the baselines and commercials for each device; the two devices are close (R tests this in section 9 of `analysis_output.txt`), so pooling them is reasonable here.

| Statistic | Cell | Formula |
|---|---|---|
{row3('Rows by device (first device)', 'Summary', 'dev_n')}
{row3('SNR mean by device (first device)', 'Summary', 'dev_snr')}
{row3('PPM mean by device (first device)', 'Summary', 'dev_ppm')}

---

## 9. Chart with 95 percent confidence intervals

1. Make a small chart table: one row per baseline and commercial, columns **Mean PPM** (point to the mean cells of the paired blocks) and **CI half width** = `CONFIDENCE.T(0.05, SD cell, n cell)`. `CONFIDENCE.T` returns half the width of a t based 95 percent interval around a mean. The workbook has this table on the H1 sheet at `H1!O{H1_CHART_ROW + 2}:Q{H1_CHART_ROW + 5}`; the half widths are in `{rcell('H1', 'chart_ci_0')}`, `{rcell('H1', 'chart_ci_1')}`, `{rcell('H1', 'chart_ci_2')}`, and `{rcell('H1', 'chart_ci_3')}`.
2. Select the stimulus names and the means, **Insert > Column chart > Clustered Column**.
3. Click the bars, **Chart Design > Add Chart Element > Error Bars > More Error Bars Options**, choose **Custom > Specify Value**, and select the four CI half width cells for both **Positive** and **Negative**.
4. Set the bar fill to navy (`#1F3864`) and lighten the gridlines.

**Common mistake.** Using the default error bars (**Standard Error** or a fixed 5 percent). Those are not 95 percent confidence intervals. And remember that error bars around the two means are *not* the interval that goes with the paired test: the paired test uses the interval of the *differences* (`{rcell('H1', 'nike_ci_lo')}` and `{rcell('H1', 'nike_ci_hi')}` for Nike), which is much narrower because it removes the between respondent variation.

---

## 10. Check your numbers and know what Excel cannot do

The **Check** sheet puts each live Excel result next to the value `analysis.R` produced and prints OK or CHECK (status in column F, the count of CHECK rows in `{rcell('Check', 'n_check')}`). Your Excel results should match to rounding; the d_z limits differ by method and have a looser tolerance. Several parts of the R analysis have no Excel equivalent:

- **Wilcoxon signed rank tests** (a check that the conclusion does not depend on normality). In R: H1 Nike V = {h1n['wilcoxon_V']}, H1 e.l.f. V = {h1e['wilcoxon_V']}, H2 V = {h2['wilcoxon_V']}, H3 Nike V = {h3n['wilcoxon_V']}, H3 e.l.f. V = {h3e['wilcoxon_V']}, H3 difference V = {h3d['wilcoxon_V']}.
- **Mixed effects model** with a random intercept per respondent and the **ICC**: `lme4::lmer(ppm ~ commercial * order_position + (1 | respondent))`.
- **Shapiro Wilk** tests and QQ plots of the paired differences, **Spearman** correlations, and the **noncentral t** interval for d_z.
- **Moment to moment view** (percent of respondents with a peak in each 5 s bin with 1 s overlap): it needs the GSR Peak Table (or the Moment To Moment Statistics export) and is computed in R.

If you have the **Analysis ToolPak** add in (**File > Options > Add-ins > Manage Excel Add-ins > Go > Analysis ToolPak**), **Data > Data Analysis > t-Test: Paired Two Sample for Means** prints t, df, and p in one table and is a good cross check of sections 4 to 6.

---

## 11. Common mistakes, collected

1. Opening the .csv by double clicking and getting one column: use **Data > From Text/CSV** or the Text Import Wizard.
2. Replacing the NA in `Average Peak Amplitude` with 0. NA means no peak, so the amplitude does not exist.
3. Forgetting the exclusion filter (Keep = Yes) in the PivotTable or in `COUNTIFS`, or leaving {excl_txt} in a roster.
4. Applying the exclusion rule to the baselines and recoveries as well as the commercials.
5. `T.TEST` with `type = 2` or `type = 3` on a within subjects design. Use `type = 1` and make sure the two columns are in the same respondent order.
6. Reporting `T.TEST` alone. Always give M, SD, n, t, df, p, and d_z with its interval.
7. Treating {n_data} rows as {n_data} independent observations. Each respondent contributes six rows; tests use one value per respondent and stimulus, and an all rows analysis needs the mixed model.
8. Dividing Peak Count by Duration instead of Signal Duration, or reading a 15 s Recovery or a 30 s commercial Peaks Per Minute as if it had a one minute resolution.
9. Dropping the nonresponders without saying so, or keeping them without checking what they do to the result. Report both.
10. Reading a GSR peak as a positive or a negative feeling. GSR shows arousal (intensity), not valence.

## Sources

- iMotions help center, GSR R-Notebooks: Peak Detection and Aggregation: https://help.imotions.com/lab/analysis/eda-gsr/gsr-r-notebooks-peak-detection-and-aggregation (local copy: `C:\\JaxState_Behavior Lab\\iMotion_Manual\\GSR RNotebooks Peak Detection and Aggegation.pdf`)
- iMotions help center, GSR R-Notebooks: GSR Signal Quality: https://help.imotions.com/lab/analysis/eda-gsr/gsr-r-notebooks-gsr-signal-quality (local copy: `C:\\JaxState_Behavior Lab\\iMotion_Manual\\GSR RNotebooks GSR Signal Quality.pdf`)
- iMotions help center, Export: R Processed GSR Files: https://help.imotions.com/lab/exports/r-processed-exports/export-r-processed-gsr-files (local copy: `C:\\JaxState_Behavior Lab\\iMotion_Manual\\Export R Processed GSR Files.pdf`)
- iMotions help center, PPG, BVP, and Optical Heart Rate: https://help.imotions.com/lab/sensor-setup/shimmer/ppg-bvp-and-optical-heart-rate
- iMotions help center, Opening Data Export Files in Excel: https://help.imotions.com/lab/exports/additional-export-information/opening-data-export-files-in-excel (local copy: `iMotion_Manual\\Opening Data Export Files in Excel.pdf`)
- iMotions help center, Creating Excel Pivot Tables: https://help.imotions.com/lab/exports/additional-export-information/creating-excel-pivot-tables (local copy: `iMotion_Manual\\Making Excel Pivot Tables.pdf`)

## Items to verify on the lab laptop

- The name, calculation, and units of the PPG beats per minute column in a Sensor Data export, and whether any notebook accepts it. The manuals name no PPG heart rate notebook or column; the column `Mean PPG Heart Rate (BPM)` in the simulated file is a modeled name.
- The exact layout of the GSR Signal Quality export in Lab 11.1 (the manual names Signal-to-Noise ratio in decibels and Signal missing in percent per stimulus and respondent but prints no header row) and whether it is a .txt or a .csv file.
- The Type and Parent Stimulus values that the GSR Peak Metrics export writes for a plain stimulus row.
- The name of the second Shimmer unit (`Shimmer3-B` is a placeholder), the flow names, and whether the **Multiple device aggregation** box appears with two Shimmers in Lab 11.1.
- The peak detection parameter values used (gap interpolation, phasic filter length, low pass cutoff, onset, offset, amplitude threshold, minimum peak duration); the manuals do not state defaults, so record them from the parameter panel.
- Whether the real Moment To Moment Statistics file labels the bins the way the simulation does (the manual does not show the 1 s overlap in the file).
"""
with open(MD_OUT, "w", encoding="utf-8") as fh:
    fh.write(md)
print("wrote", MD_OUT)
