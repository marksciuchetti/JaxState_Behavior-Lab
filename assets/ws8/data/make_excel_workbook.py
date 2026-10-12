#!/usr/bin/env python3
"""
make_excel_workbook.py  (SIMULATED DATA, teaching only)
SIMULATED DATA. Builds SIMULATED_analysis_report.xlsx (the Excel route of Workshop 8) with openpyxl from
the Workshop 7 analysis ready file, and writes excel_walkthrough.md from the SAME cell registry, so that
every cell reference in the walkthrough is the address of a live formula in the workbook. No real person
was recorded; ALTAVO and BRENLOW are invented brands.

The Excel route covers:
  1. the analysis ready table (analysis_ready_long.csv) as the Text Import leaves it, plus a key column;
  2. a long AOI table (one row per participant, ad, and AOI) built with INDEX and MATCH, the shape that
     AVERAGEIFS and a PivotTable need;
  3. Descriptives: COUNTIFS, AVERAGEIFS, and STDEV.S, MEDIAN, QUARTILE.INC as array formulas
     (STDEV.S(IF(...))) by ad and AOI, plus the facial expressions, GSR, and liking by stimulus;
  4. a PivotTable style summary built with formulas (the walkthrough explains the real PivotTable);
  5. Q5: the three paired tests (Joy, Engagement, Brow Furrow, Ad_A versus Ad_B) with T.TEST, d_z, and
     its approximate 95% CI;
  6. a Holm correction table for the three Q5 tests;
  7. a Check sheet that compares every result with the value analysis_full.R produced.

Verification after the workbook is saved:
  1. openpyxl reads the workbook back and checks that every registered cell holds the registered formula.
  2. LibreOffice (headless) recalculates a copy in a temporary folder (the delivered workbook is not
     changed); every row of the Check sheet must say OK and no cell may show an error value.
  3. Every result printed in the walkthrough is the LibreOffice value of that cell, and it must agree
     with the R value at the number of decimals printed.

Run: python3 make_excel_workbook.py   (after Rscript analysis_full.R)
"""

import csv
import os
import re
import shutil
import subprocess
import sys
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.formula import ArrayFormula

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "SIMULATED_analysis_report.xlsx")
MD_OUT = os.path.join(HERE, "excel_walkthrough.md")

BLUE, ORANGE = "0072B2", "E69F00"
hdr_font = Font(bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor=BLUE)
help_fill = PatternFill("solid", fgColor=ORANGE)
yellow = PatternFill("solid", fgColor="FFF2CC")
bold = Font(bold=True)
note_font = Font(italic=True, color="555555")
thin = Side(style="thin", color="BFBFBF")
box = Border(top=thin, bottom=thin, left=thin, right=thin)


# ---------------------------------------------------------------- inputs
def ws7_dir():
    """The Workshop 7 folder if it is there, else the byte identical copies in ws7_inputs/ (as in analysis_full.R)."""
    for d in (os.path.normpath(os.path.join(HERE, "..", "..", "ws7", "data")), os.path.join(HERE, "ws7_inputs")):
        if os.path.exists(os.path.join(d, "analysis_ready_long.csv")):
            return d
    sys.exit("analysis_ready_long.csv not found in ../../ws7/data or ws7_inputs/")


def read_banner_csv(path):
    with open(path, encoding="utf-8") as fh:
        raw = fh.read().splitlines()
    start = 1 if raw[0].startswith("#") else 0
    rows = list(csv.reader(raw[start:]))
    return raw[0] if start else "", rows[0], rows[1:]


IN_DIR = ws7_dir()
LONG_PATH = os.path.join(IN_DIR, "analysis_ready_long.csv")
long_banner, L_HEAD, L_ROWS = read_banner_csv(LONG_PATH)
SEED = re.search(r"seed (\d+)", long_banner).group(1)
_, dh, drows = read_banner_csv(os.path.join(HERE, "SIMULATED_descriptives.csv"))
DS = {(r[dh.index("variable")], r[dh.index("stimulus")]): dict(zip(dh, r)) for r in drows}
_, th, trows = read_banner_csv(os.path.join(HERE, "SIMULATED_all_tests.csv"))
T = {r[th.index("id")]: dict(zip(th, r)) for r in trows}
_, eh, erows = read_banner_csv(os.path.join(HERE, "SIMULATED_q4_marginal_means.csv"))
EMM = {(r[0], r[1]): dict(zip(eh, r)) for r in erows}
with open(os.path.join(HERE, "analysis_output.txt"), encoding="utf-8") as fh:
    OUT_TXT = fh.read()

EXPR = [("Joy", "joy_pct", "Q5_JOY"), ("Engagement", "engagement_pct", "Q5_ENGAGEMENT"), ("Brow Furrow", "browfurrow_pct", "Q5_BROWFURROW")]


def q5_block_from_output(name):
    """SE, critical t, SE of d_z, and the counts that analysis_output.txt prints for one Q5 expression."""
    m = re.search(re.escape(name) + r" \(percent of time\), Ad_A minus Ad_B; at zero: Ad_A (\d+) of \d+, Ad_B (\d+) of \d+, both ads (\d+)\n"
                  r".*?\n.*?SE = ([0-9.]+), critical t = ([0-9.]+)\n.*?\(SE of d_z = ([0-9.]+)\)\n"
                  r".*?Ad_A higher: (\d+), Ad_B higher: (\d+), equal: (\d+)", OUT_TXT)
    assert m, f"Q5 block for {name} not found in analysis_output.txt"
    g = m.groups()
    return {"zero_a": int(g[0]), "zero_b": int(g[1]), "zero_both": int(g[2]), "se": float(g[3]), "tcrit": float(g[4]), "se_dz": float(g[5]),
            "n_pos": int(g[6]), "n_neg": int(g[7]), "n_tie": int(g[8])}


Q5OUT = {e: q5_block_from_output(e) for e, _, _ in EXPR}

# ---------------------------------------------------------------- registry of documented cells
REG, EXPECT = {}, {}


def reg(sheet, key, cell, formula, label="", array=False):
    REG[(sheet, key)] = {"cell": cell, "formula": formula.replace("_xlfn.", ""), "label": label, "array": array}


def put(ws, key, cell, formula, label="", fmt=None, expect=None, nd=3, array=False):
    ws[cell] = ArrayFormula(cell, formula) if array else formula
    if fmt:
        ws[cell].number_format = fmt
    reg(ws.title, key, cell, formula, label, array)
    if expect is not None:
        EXPECT[(ws.title, key)] = (expect, nd)


def style_header(ws, row=1, first=1, last=None, fill=None):
    last = last or ws.max_column
    for c in range(first, last + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill, cell.border = hdr_font, fill or hdr_fill, box
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def widths(ws, d):
    for col, w in d.items():
        ws.column_dimensions[col].width = w


def to_cell(v):
    """Numbers become numbers; NA stays the text NA (as Excel imports it); an empty cell stays empty."""
    if v == "":
        return None
    try:
        f = float(v)
        return int(f) if f.is_integer() and "." not in v else f
    except ValueError:
        return v


def f(x):
    return float(x)


wb = Workbook()
# ================================================================ README
ws = wb.active
ws.title = "README"
readme = [
    ("SIMULATED ANALYSIS REPORT WORKBOOK (Workshop 8)", bold),
    (f"SIMULATED DATA for teaching (seed {SEED}). No real person was recorded; ALTAVO and BRENLOW are invented brands. Built by make_excel_workbook.py "
     "from analysis_ready_long.csv (Workshop 7) and checked against analysis_full.R (Workshop 8).", None),
    ("Study: SIMULATED Multimodal Ad Study. 24 participants (P01 to P24), analysis sample 22 (include = 1). Baseline (60 s), Ad_A (ALTAVO) and "
     "Ad_B (BRENLOW), 30 s each, counterbalanced; AOIs Logo, Product, Headline; AFFDEX; Shimmer GSR; survey.", None),
    ("This workbook is the Excel route through the descriptive statistics, the Q5 paired tests, and the Holm correction of analysis_full.R.", None),
    ("", None),
    ("Sheets", bold),
    ("Data          analysis_ready_long.csv as the Text Import leaves it (header in row 1, 72 rows), plus the key column AD (orange)", None),
    ("AOI_long      one row per participant, ad, and AOI (144 rows): the dwell values stacked into ONE column with INDEX and MATCH", None),
    ("Descriptives  n, mean, SD, median, quartiles, IQR, and percent at zero by ad and AOI (and by stimulus for FEA, GSR, liking)", None),
    ("Pivot         the summary a PivotTable gives (Average of dwell by AOI and ad), built with AVERAGEIFS and COUNTIFS so that it can be checked", None),
    ("Q5            Joy, Engagement, and Brow Furrow, Ad_A versus Ad_B: paired T.TEST, d_z, and its approximate 95% CI", None),
    ("Holm          the Holm correction of the three Q5 p values, step by step, next to the Bonferroni correction", None),
    ("Check         every Excel result next to the value R produced, with OK or CHECK", None),
    ("", None),
    ("Colors: blue headers = data as exported; orange headers = columns you add; yellow cells are typed by you (the grid and the rosters).", None),
    ("Cells with braces in the formula bar ({=STDEV.S(IF(...))}) are array formulas: in Excel 365 press Enter; in Excel 2019 or older press "
     "Ctrl+Shift+Enter.", None),
    ("The walkthrough excel_walkthrough.md explains every formula with its cell address.", None),
]
for i, (txt, fnt) in enumerate(readme, start=1):
    c = ws.cell(row=i, column=1, value=txt)
    if fnt:
        c.font = fnt
ws.column_dimensions["A"].width = 150

# ================================================================ Data
ws = wb.create_sheet("Data")
for c, h in enumerate(L_HEAD, start=1):
    ws.cell(row=1, column=c, value=h)
for r, row in enumerate(L_ROWS, start=2):
    for c, v in enumerate(row, start=1):
        ws.cell(row=r, column=c, value=to_cell(v))
style_header(ws, last=len(L_HEAD))
D_LAST = len(L_ROWS) + 1
DC = {h: get_column_letter(i + 1) for i, h in enumerate(L_HEAD)}
KEY = get_column_letter(len(L_HEAD) + 1)
ws[f"{KEY}1"] = "key = participant|stimulus"
style_header(ws, first=len(L_HEAD) + 1, last=len(L_HEAD) + 1, fill=help_fill)
for r in range(2, D_LAST + 1):
    ws[f"{KEY}{r}"] = f'={DC["participant_id"]}{r}&"|"&{DC["stimulus"]}{r}'
reg("Data", "key", f"{KEY}2", f'={DC["participant_id"]}2&"|"&{DC["stimulus"]}2', "key of the first row")
ws.freeze_panes = "C2"
widths(ws, {get_column_letter(i): 11 for i in range(1, len(L_HEAD) + 2)})
ws.column_dimensions[DC["exclusion_reason"]].width = 30
ws.column_dimensions[KEY].width = 18


def DR(var):
    """Absolute range of one Data column, rows 2 to the last row."""
    col = DC[var] if var in DC else var
    return f"Data!${col}$2:${col}${D_LAST}"


R_KEY = DR(KEY)
IDS = sorted({row[0] for row in L_ROWS})
INCL = {row[0]: row[L_HEAD.index("include")] == "1" for row in L_ROWS}
ROSTER = [p for p in IDS if INCL[p]]

# ================================================================ AOI_long
ws = wb.create_sheet("AOI_long")
ahead = ["participant_id", "ad", "AOI", "include (from Data)", "gaze dwell (% of ad)"]
for c, h in enumerate(ahead, start=1):
    ws.cell(row=1, column=c, value=h)
style_header(ws, last=3)
style_header(ws, first=4, last=5, fill=help_fill)
AOIS = [("Logo", "logo_gaze_dwell_pct"), ("Product", "product_gaze_dwell_pct"), ("Headline", "headline_gaze_dwell_pct")]
ADS = ["Ad_A", "Ad_B"]
af = {
    "D": "=INDEX(" + DR("include") + ',MATCH(A{r}&"|"&B{r},' + R_KEY + ",0))",
    "E": '=IF(C{r}="Logo",INDEX(' + DR("logo_gaze_dwell_pct") + ',MATCH(A{r}&"|"&B{r},' + R_KEY + ',0)),IF(C{r}="Product",INDEX('
         + DR("product_gaze_dwell_pct") + ',MATCH(A{r}&"|"&B{r},' + R_KEY + ',0)),INDEX(' + DR("headline_gaze_dwell_pct") + ',MATCH(A{r}&"|"&B{r},'
         + R_KEY + ",0))))",
}
r = 2
for p in IDS:
    for ad in ADS:
        for aoi, _ in AOIS:
            ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"] = p, ad, aoi
            for col in "ABC":
                ws[f"{col}{r}"].fill = yellow
            for col, fm in af.items():
                ws[f"{col}{r}"] = fm.format(r=r)
            ws[f"E{r}"].number_format = "0.00"
            r += 1
A_LAST = r - 1
reg("AOI_long", "D", "D2", af["D"].format(r=2), "include")
reg("AOI_long", "E", "E2", af["E"].format(r=2), "gaze dwell")
ws["G2"] = "Checks"
ws["G2"].font = bold
ws["G3"] = "Rows (24 participants x 2 ads x 3 AOIs)"
put(ws, "n_rows", "H3", f"=COUNTA(A2:A{A_LAST})", "AOI_long rows", "0", 144, 0)
ws["G4"] = "Lookups that failed (#N/A)"
put(ws, "n_na", "H4", f"=SUMPRODUCT(--ISNA(D2:E{A_LAST}))", "AOI_long failed lookups", "0", 0, 0)
ws["G5"] = "Rows with include = 1 (22 x 6)"
put(ws, "n_incl", "H5", f"=COUNTIF(D2:D{A_LAST},1)", "AOI_long rows with include = 1", "0", 132, 0)
widths(ws, {"A": 14, "B": 8, "C": 10, "D": 12, "E": 14, "G": 40, "H": 8})
ws.freeze_panes = "A2"


def AR(col):
    return f"AOI_long!${col}$2:${col}${A_LAST}"


# ================================================================ Descriptives
ws = wb.create_sheet("Descriptives")
ws["A1"] = ("Descriptive statistics, analysis sample (include = 1). Gaze dwell from AOI_long; facial expressions, GSR (valid only, rule R3), "
            "and liking from Data. SD, median, and quartiles are array formulas.")
ws["A1"].font = bold
dhead = ["Measure", "Ad or stimulus", "AOI", "n", "Mean", "SD", "Median", "Q1", "Q3", "IQR", "% at zero"]
DESC_ROWS = {}


def desc_block(start_row, title, items):
    """items: (measure, stim, aoi or '', value range, [(range, criterion)...], variable, has_zero)."""
    ws.cell(row=start_row, column=1, value=title).font = bold
    for c, h in enumerate(dhead, start=1):
        ws.cell(row=start_row + 1, column=c, value=h)
    style_header(ws, row=start_row + 1, last=len(dhead))
    rr = start_row + 2
    for (meas, stim, aoi, vr, conds, var, zero) in items:
        ws[f"A{rr}"], ws[f"B{rr}"] = meas, stim
        if aoi:
            ws[f"C{rr}"] = aoi
        cif = ",".join(f"{rg},{cr.format(r=rr)}" for rg, cr in conds)
        arr = "*".join(f"({rg}={cr.format(r=rr)})" for rg, cr in conds)
        key = f"{var}|{stim}"
        ds = DS[(var, stim)]
        put(ws, f"n|{key}", f"D{rr}", f"=COUNTIFS({cif})", f"n, {meas}, {stim}", "0", int(ds["n"]), 0)
        put(ws, f"mean|{key}", f"E{rr}", f"=AVERAGEIFS({vr},{cif})", f"mean, {meas}, {stim}", "0.000", f(ds["mean"]), 3)
        put(ws, f"sd|{key}", f"F{rr}", f"=_xlfn.STDEV.S(IF({arr},{vr}))", f"SD, {meas}, {stim}", "0.000", f(ds["sd"]), 3, array=True)
        put(ws, f"median|{key}", f"G{rr}", f"=MEDIAN(IF({arr},{vr}))", f"median, {meas}, {stim}", "0.000", f(ds["median"]), 3, array=True)
        put(ws, f"q1|{key}", f"H{rr}", f"=_xlfn.QUARTILE.INC(IF({arr},{vr}),1)", f"Q1, {meas}, {stim}", "0.000", f(ds["q1"]), 3, array=True)
        put(ws, f"q3|{key}", f"I{rr}", f"=_xlfn.QUARTILE.INC(IF({arr},{vr}),3)", f"Q3, {meas}, {stim}", "0.000", f(ds["q3"]), 3, array=True)
        put(ws, f"iqr|{key}", f"J{rr}", f"=I{rr}-H{rr}", f"IQR, {meas}, {stim}", "0.000", f(ds["iqr"]), 3)
        if zero:
            put(ws, f"zero|{key}", f"K{rr}", f"=100*COUNTIFS({cif},{vr},0)/D{rr}", f"% at zero, {meas}, {stim}", "0.00", f(ds["pct_zero"]), 2)
        DESC_ROWS[key] = rr
        rr += 1
    return rr


dw_items = [("Gaze dwell (% of ad)", ad, aoi, AR("E"), [(AR("B"), "B{r}"), (AR("C"), "C{r}"), (AR("D"), "1")], var, True)
            for ad in ADS for aoi, var in AOIS]
nxt = desc_block(3, "A. Gaze dwell by ad and AOI (from AOI_long)", dw_items)
STIMS = ["Baseline", "Ad_A", "Ad_B"]
fea_items = [(f"{e} (% of time)", st, "", DR(var), [(DR("stimulus"), "B{r}"), (DR("include"), "1")], var, True) for e, var, _ in EXPR for st in STIMS]
nxt = desc_block(nxt + 1, "B. Facial expressions by stimulus (from Data)", fea_items)
gsr_items = [("GSR Peaks Per Minute", st, "", DR("gsr_peaks_per_min"), [(DR("stimulus"), "B{r}"), (DR("include"), "1"), (DR("gsr_valid"), "1")],
              "gsr_peaks_per_min", True) for st in STIMS]
nxt = desc_block(nxt + 1, "C. GSR Peaks Per Minute by stimulus, valid GSR only (rule R3) (from Data)", gsr_items)
lik_items = [("Liking (1 to 7)", ad, "", DR("liking"), [(DR("stimulus"), "B{r}"), (DR("include"), "1")], "liking", False) for ad in ADS]
nxt = desc_block(nxt + 1, "D. Liking by ad (from Data)", lik_items)
widths(ws, {"A": 24, "B": 14, "C": 10, "D": 6, "E": 10, "F": 10, "G": 10, "H": 10, "I": 10, "J": 10, "K": 10})

# ================================================================ Pivot
ws = wb.create_sheet("Pivot")
ws["A1"] = ("PivotTable style summary: Average of gaze dwell (% of ad) by AOI (rows) and ad (columns), include = 1. Built with AVERAGEIFS "
            "and COUNTIFS so that it can be checked; a real PivotTable gives the same numbers.")
ws["A1"].font = bold
for c, h in enumerate(["AOI", "Ad_A", "Ad_B", "Grand average", "Count Ad_A", "Count Ad_B"], start=1):
    ws.cell(row=3, column=c, value=h)
style_header(ws, row=3, last=6)
PV = {}
for i, (aoi, var) in enumerate(AOIS):
    rr = 4 + i
    ws[f"A{rr}"] = aoi
    for col, ad in (("B", "Ad_A"), ("C", "Ad_B")):
        put(ws, f"m_{aoi}_{ad}", f"{col}{rr}", f'=AVERAGEIFS({AR("E")},{AR("C")},$A{rr},{AR("B")},{col}$3,{AR("D")},1)', f"pivot mean {aoi} {ad}",
            "0.00", f(DS[(var, ad)]["mean"]), 3)
    put(ws, f"g_{aoi}", f"D{rr}", f'=AVERAGEIFS({AR("E")},{AR("C")},$A{rr},{AR("D")},1)', f"pivot grand average {aoi}", "0.00",
        f(EMM[("both ads", aoi)]["emm"]), 3)
    for col, ad in (("E", "Ad_A"), ("F", "Ad_B")):
        put(ws, f"n_{aoi}_{ad}", f"{col}{rr}", f'=COUNTIFS({AR("C")},$A{rr},{AR("B")},"{ad}",{AR("D")},1)', f"pivot count {aoi} {ad}", "0",
            int(DS[(var, ad)]["n"]), 0)
    PV[aoi] = rr
ws["A7"] = "Grand average"
ws["A7"].font = bold
for col, ad in (("B", "Ad_A"), ("C", "Ad_B")):
    put(ws, f"g_{ad}", f"{col}7", f'=AVERAGEIFS({AR("E")},{AR("B")},{col}$3,{AR("D")},1)', f"pivot grand average {ad}", "0.00",
        f(EMM[(ad, "all AOIs")]["emm"]), 3)
ws["A9"] = ("To build the real PivotTable: click any cell of AOI_long, Insert > PivotTable > New Worksheet > OK; drag AOI to Rows, ad to Columns, "
            "gaze dwell (% of ad) to Values; right click a value, Value Field Settings, change Sum to Average; drag include (from Data) to "
            "Filters and keep only 1. Excel sorts the rows A to Z (Headline, Logo, Product).")
ws["A9"].font = note_font
widths(ws, {"A": 16, "B": 10, "C": 10, "D": 15, "E": 11, "F": 11})


# ================================================================ Q5
class Block:
    """A column of labeled results; {key} in a formula becomes the address of that key's cell."""

    def __init__(self, ws, row, lc):
        self.ws, self.row, self.lc, self.ref = ws, row, lc, {}

    def title(self, text):
        self.ws.cell(row=self.row, column=self.lc, value=text).font = bold
        self.row += 1

    def put(self, key, label, formula, fmt="0.0000", expect=None, nd=3, pref=""):
        fm = formula
        for k, rr in self.ref.items():
            fm = fm.replace("{" + k + "}", f"{get_column_letter(self.lc + 1)}{rr}")
        self.ws.cell(row=self.row, column=self.lc, value=label)
        cell = f"{get_column_letter(self.lc + 1)}{self.row}"
        put(self.ws, pref + key, cell, fm, label, fmt, expect, nd)
        self.ref[key] = self.row
        self.row += 1


ws = wb.create_sheet("Q5")
ws["A1"] = ("Q5: facial expressions, Ad_A versus Ad_B (paired). Roster: the 22 participants with include = 1 (static list, checked on the "
            "Check sheet). Values from Data with INDEX and MATCH on the key column.")
ws["A1"].font = bold
q5head = ["participant_id"]
for e, _, _ in EXPR:
    q5head += [f"{e} Ad_A", f"{e} Ad_B", f"{e} Ad_A minus Ad_B"]
for c, h in enumerate(q5head, start=1):
    ws.cell(row=3, column=c, value=h)
style_header(ws, row=3, last=1)
style_header(ws, row=3, first=2, last=len(q5head), fill=help_fill)
Q5C = {}
for k, (e, var, _) in enumerate(EXPR):
    ca, cb, cd = (get_column_letter(2 + 3 * k + j) for j in range(3))
    Q5C[e] = (ca, cb, cd)
R0 = 4
for i, p in enumerate(ROSTER):
    rr = R0 + i
    ws[f"A{rr}"] = p
    ws[f"A{rr}"].fill = yellow
    for e, var, _ in EXPR:
        ca, cb, cd = Q5C[e]
        ws[f"{ca}{rr}"] = f'=INDEX({DR(var)},MATCH($A{rr}&"|Ad_A",{R_KEY},0))'
        ws[f"{cb}{rr}"] = f'=INDEX({DR(var)},MATCH($A{rr}&"|Ad_B",{R_KEY},0))'
        ws[f"{cd}{rr}"] = f"={ca}{rr}-{cb}{rr}"
Q5_LAST = R0 + len(ROSTER) - 1
for e, var, _ in EXPR:
    ca, cb, cd = Q5C[e]
    reg("Q5", f"{e}_a", f"{ca}{R0}", ws[f"{ca}{R0}"].value, f"{e} Ad_A, first participant")
    reg("Q5", f"{e}_b", f"{cb}{R0}", ws[f"{cb}{R0}"].value, f"{e} Ad_B, first participant")
    reg("Q5", f"{e}_d", f"{cd}{R0}", ws[f"{cd}{R0}"].value, f"{e} difference, first participant")
Q5B = {}
for k, (e, var, tid) in enumerate(EXPR):
    ca, cb, cd = Q5C[e]
    ra, rb, rd = (f"{c}{R0}:{c}{Q5_LAST}" for c in (ca, cb, cd))
    t, o = T[tid], Q5OUT[e]
    blk = Block(ws, Q5_LAST + 3, 1 + 4 * k)
    pf_ = f"{e}_"
    blk.title(f"{e}: paired t test, Ad_A minus Ad_B")
    blk.put("m1", "Ad_A mean", f"=AVERAGE({ra})", expect=f(t["m1"]), pref=pf_)
    blk.put("s1", "Ad_A SD", f"=_xlfn.STDEV.S({ra})", expect=f(t["sd1"]), pref=pf_)
    blk.put("m2", "Ad_B mean", f"=AVERAGE({rb})", expect=f(t["m2"]), pref=pf_)
    blk.put("s2", "Ad_B SD", f"=_xlfn.STDEV.S({rb})", expect=f(t["sd2"]), pref=pf_)
    blk.put("n", "n (participants)", f"=COUNT({rd})", fmt="0", expect=int(f(t["n"])), nd=0, pref=pf_)
    blk.put("md", "Mean difference", f"=AVERAGE({rd})", expect=f(t["mean_diff"]), pref=pf_)
    blk.put("sdd", "SD of the differences", f"=_xlfn.STDEV.S({rd})", expect=f(t["sd_diff"]), pref=pf_)
    blk.put("se", "SE of the mean difference", "={sdd}/SQRT({n})", expect=o["se"], nd=4, pref=pf_)
    blk.put("t", "t", "={md}/{se}", expect=f(t["statistic"]), pref=pf_)
    blk.put("df", "df", "={n}-1", fmt="0", expect=int(f(t["df1"])), nd=0, pref=pf_)
    blk.put("p_tt", "p, two tailed (T.TEST, paired)", f"=_xlfn.T.TEST({ra},{rb},2,1)", fmt="0.00000", expect=f(t["p"]), nd=5, pref=pf_)
    blk.put("p_hand", "p from t (T.DIST.2T)", "=_xlfn.T.DIST.2T(ABS({t}),{df})", fmt="0.00000", expect=f(t["p"]), nd=5, pref=pf_)
    blk.put("tcrit", "Critical t (95%)", "=_xlfn.T.INV.2T(0.05,{df})", expect=o["tcrit"], nd=4, pref=pf_)
    blk.put("ci_lo", "Mean difference, 95% CI lower", "={md}-{tcrit}*{se}", expect=f(t["diff_ci_low"]), pref=pf_)
    blk.put("ci_hi", "Mean difference, 95% CI upper", "={md}+{tcrit}*{se}", expect=f(t["diff_ci_high"]), pref=pf_)
    blk.put("dz", "Cohen's d_z", "={md}/{sdd}", expect=f(t["es"]), pref=pf_)
    blk.put("se_dz", "SE of d_z (approximation)", "=SQRT(1/{n}+{dz}^2/(2*{n}))", expect=o["se_dz"], nd=4, pref=pf_)
    blk.put("dz_lo", "d_z 95% CI lower (approximation)", "={dz}-1.96*{se_dz}", expect=f(t["es_ci_low_approx"]), pref=pf_)
    blk.put("dz_hi", "d_z 95% CI upper (approximation)", "={dz}+1.96*{se_dz}", expect=f(t["es_ci_high_approx"]), pref=pf_)
    blk.put("npos", "Ad_A higher (count)", f'=COUNTIF({rd},">0")', fmt="0", expect=o["n_pos"], nd=0, pref=pf_)
    blk.put("nneg", "Ad_B higher (count)", f'=COUNTIF({rd},"<0")', fmt="0", expect=o["n_neg"], nd=0, pref=pf_)
    blk.put("ntie", "No difference (count)", f"=COUNTIF({rd},0)", fmt="0", expect=o["n_tie"], nd=0, pref=pf_)
    blk.put("nzero", "Zero in both ads (count)", f"=COUNTIFS({ra},0,{rb},0)", fmt="0", expect=o["zero_both"], nd=0, pref=pf_)
    Q5B[e] = blk
widths(ws, {get_column_letter(i): 13 for i in range(1, len(q5head) + 1)})
for k in range(3):
    ws.column_dimensions[get_column_letter(1 + 4 * k)].width = 31
ws.freeze_panes = "B4"

# ================================================================ Holm
ws = wb.create_sheet("Holm")
ws["A1"] = ("Holm correction of the three Q5 p values (family: Joy, Engagement, Brow Furrow). Holm keeps the chance of at least one false "
            "positive in the family at .05 and is never less powerful than Bonferroni.")
ws["A1"].font = bold
hh = ["Test", "p (from Q5)", "Rank (smallest = 1)", "Multiplier (m - rank + 1)", "p x multiplier", "Holm adjusted p", "Bonferroni p", "Decision at .05"]
for c, h in enumerate(hh, start=1):
    ws.cell(row=3, column=c, value=h)
style_header(ws, row=3, last=1)
style_header(ws, row=3, first=2, last=len(hh), fill=help_fill)
H0, HN = 4, 4 + len(EXPR) - 1
ws["J3"] = "m (tests in the family)"
ws["J3"].font = bold
put(ws, "m", "K3", f"=COUNT(B{H0}:B{HN})", "m, number of tests", "0", 3, 0)
# sorted table
ws["A9"] = "Step by step, from the smallest p to the largest"
ws["A9"].font = bold
sh = ["Step k", "k-th smallest p (SMALL)", "Multiplier (m - k + 1)", "p x multiplier", "Holm adjusted (running maximum, at most 1)"]
for c, h in enumerate(sh, start=1):
    ws.cell(row=10, column=c, value=h)
style_header(ws, row=10, last=len(sh), fill=help_fill)
S0 = 11
for k in range(1, len(EXPR) + 1):
    rr = S0 + k - 1
    ws[f"A{rr}"] = k
    put(ws, f"s_p{k}", f"B{rr}", f"=SMALL($B${H0}:$B${HN},A{rr})", f"sorted p, step {k}", "0.00000")
    put(ws, f"s_m{k}", f"C{rr}", f"=$K$3-A{rr}+1", f"multiplier, step {k}", "0")
    put(ws, f"s_x{k}", f"D{rr}", f"=B{rr}*C{rr}", f"p x multiplier, step {k}", "0.00000")
    fm = f"=MIN(1,D{rr})" if k == 1 else f"=MIN(1,MAX(E{rr - 1},D{rr}))"
    put(ws, f"s_a{k}", f"E{rr}", fm, f"Holm adjusted, step {k}", "0.00000")
SN = S0 + len(EXPR) - 1
for i, (e, var, tid) in enumerate(EXPR):
    rr = H0 + i
    ws[f"A{rr}"] = e
    blk = Q5B[e]
    pcell = f"Q5!{get_column_letter(blk.lc + 1)}{blk.ref['p_tt']}"
    put(ws, f"p_{e}", f"B{rr}", f"={pcell}", f"p, {e}", "0.00000", f(T[tid]["p"]), 5)
    put(ws, f"rank_{e}", f"C{rr}", f"=_xlfn.RANK.EQ(B{rr},$B${H0}:$B${HN},1)", f"rank, {e}", "0")
    put(ws, f"mult_{e}", f"D{rr}", f"=$K$3-C{rr}+1", f"multiplier, {e}", "0")
    put(ws, f"px_{e}", f"E{rr}", f"=B{rr}*D{rr}", f"p x multiplier, {e}", "0.00000")
    put(ws, f"holm_{e}", f"F{rr}", f"=INDEX($E${S0}:$E${SN},C{rr})", f"Holm p, {e}", "0.00000", f(T[tid]["p_holm"]), 5)
    put(ws, f"bonf_{e}", f"G{rr}", f"=MIN(1,$K$3*B{rr})", f"Bonferroni p, {e}", "0.00000")
    put(ws, f"dec_{e}", f"H{rr}", f'=IF(F{rr}<0.05,"significant","not significant")', f"decision, {e}")
ws["A15"] = ("Read the table: sort the p values from smallest to largest; multiply the smallest by m, the next by m - 1, and so on; each "
             "adjusted value is at least as large as the one before it (the running maximum) and never above 1. Compare the adjusted p with .05.")
ws["A15"].font = note_font
_ord = sorted(((f(T[tid]["p"]), e) for e, _, tid in EXPR))
_m = len(_ord)
_prod = [(e, p_ * (_m - i)) for i, (p_, e) in enumerate(_ord)]
_raised = [(e, x, _prod[i - 1][1]) for i, (e, x) in enumerate(_prod) if i > 0 and x < max(y for _, y in _prod[:i])]
HOLM_ORDER = ", then ".join(f"{e} (p = {p_:.5f}, multiplied by {_m - i})" for i, (p_, e) in enumerate(_ord))
HOLM_RAISED = ("; ".join(f"{e}'s product {x:.5f} is raised to {max(y for _, y in _prod[:[q for q, _ in _prod].index(e)]):.5f}, the adjusted value before it"
                         for e, x, _ in _raised)) or "no product needed raising here"
ws["A16"] = ("p x multiplier (column E, rows 4 to 6) is not yet the Holm p: the running maximum in the step table can raise it (" + HOLM_RAISED + ").")
ws["A16"].font = note_font
widths(ws, {"A": 16, "B": 16, "C": 16, "D": 18, "E": 18, "F": 18, "G": 14, "H": 16, "J": 24, "K": 6})

# ================================================================ Check
ws = wb.create_sheet("Check")
ws["A1"] = "Excel results next to the values that analysis_full.R produced (SIMULATED data). Tolerance: half a unit in the last printed decimal."
ws["A1"].font = bold
for c, h in enumerate(["Result", "Excel", "R", "Difference", "Tolerance", "Status"], start=1):
    ws.cell(row=3, column=c, value=h)
style_header(ws, row=3, last=6)
ck = []
for (sheet, key), (val, nd) in EXPECT.items():
    tol = 1e-9 if nd == 0 else 0.5 * 10 ** (-nd)
    ck.append((f"{sheet}: {REG[(sheet, key)]['label']}", f"={sheet}!{REG[(sheet, key)]['cell']}", float(val), tol))
# the static Q5 roster: every listed participant must have include = 1 in Data, and the roster must hold all 22 of them
ck.append(("Q5 roster: listed participants whose Baseline row in Data has include = 1",
           f'=SUMPRODUCT(COUNTIFS({DR("participant_id")},Q5!A{R0}:A{Q5_LAST},{DR("stimulus")},"Baseline",{DR("include")},1))', float(len(ROSTER)), 1e-9))
ck.append(("Participants with include = 1 in Data (Baseline rows)", f'=COUNTIFS({DR("stimulus")},"Baseline",{DR("include")},1)', float(len(ROSTER)), 1e-9))
CK1 = 4
for k, (lab, fm, rv, tol) in enumerate(ck):
    rr = CK1 + k
    ws.cell(row=rr, column=1, value=lab)
    ws.cell(row=rr, column=2, value=fm).number_format = "0.000000"
    ws.cell(row=rr, column=3, value=rv).number_format = "0.000000"
    ws.cell(row=rr, column=4, value=f"=B{rr}-C{rr}").number_format = "0.000000"
    ws.cell(row=rr, column=5, value=tol).number_format = "0.000000"
    ws.cell(row=rr, column=6, value=f'=IF(ABS(D{rr})<=E{rr},"OK","CHECK")')
CK2 = CK1 + len(ck) - 1
ws.cell(row=CK2 + 2, column=1, value="Rows with status CHECK").font = bold
put(ws, "n_check", f"B{CK2 + 2}", f'=COUNTIF(F{CK1}:F{CK2},"CHECK")', "Rows with status CHECK", "0")
ws.cell(row=CK2 + 3, column=1, value=("Not reproduced in Excel: H1 to H3 (see Workshop 7), the Q4 mixed model and repeated measures ANOVA, the Q4 "
                                      "and Q6 paired tests, the order checks, the Wilcoxon checks, the exact noncentral t intervals of d_z, the "
                                      "Spearman correlations, and the power analysis. See analysis_output.txt.")).font = note_font
widths(ws, {"A": 80, "B": 14, "C": 14, "D": 13, "E": 11, "F": 9})

wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
print("wrote", OUT)

# =============================================================================
# Verification 1: read back and compare every registered cell
# =============================================================================
wb2 = load_workbook(OUT)
problems = []
for (sheet, key), info in REG.items():
    v = wb2[sheet][info["cell"]].value
    if isinstance(v, ArrayFormula):
        v = v.text
    v = str(v).replace("_xlfn.", "") if v is not None else None
    if v != info["formula"]:
        problems.append(f"{sheet}!{info['cell']} ({key}): workbook {v!r} registry {info['formula']!r}")
for k, (lab, fm, rv, tol) in enumerate(ck):
    if wb2["Check"][f"B{CK1 + k}"].value != fm:
        problems.append(f"Check!B{CK1 + k} differs")
if problems:
    print("VERIFICATION FAILED:\n" + "\n".join(problems[:40]))
    sys.exit(1)
n_formulas = sum(1 for w in wb2.worksheets for row in w.iter_rows() for c in row
                 if isinstance(c.value, ArrayFormula) or (isinstance(c.value, str) and c.value.startswith("=")))
n_arrays = sum(1 for w in wb2.worksheets for row in w.iter_rows() for c in row if isinstance(c.value, ArrayFormula))
print(f"verified (openpyxl): {len(REG)} registered cells and {len(ck)} Check rows hold the registered formulas; "
      f"formula cells in the workbook: {n_formulas} (array formulas: {n_arrays})")

# =============================================================================
# Verification 2 and 3: LibreOffice recalculation
# =============================================================================
LO = {}
soffice = shutil.which("soffice") or shutil.which("libreoffice")
lo_note = "LibreOffice was not found, so the formula values were not recalculated here; the printed values are the R values."
LO_SUMMARY = None


def same_at(got, val, nd):
    got, val = float(got), float(val)
    return f"{got:.{nd}f}" == f"{val:.{nd}f}" or abs(got - val) <= 1e-9 * max(1.0, abs(val))


if soffice:
    tmp = tempfile.mkdtemp(prefix="ws8_lo_")
    shutil.copy(OUT, os.path.join(tmp, "in.xlsx"))
    subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", os.path.join(tmp, "out"),
                    os.path.join(tmp, "in.xlsx")], check=True, capture_output=True, timeout=600)
    wb3 = load_workbook(os.path.join(tmp, "out", "in.xlsx"), data_only=True)
    for (sheet, key), info in REG.items():
        LO[(sheet, key)] = wb3[sheet][info["cell"]].value
    bad = [(ck[k][0], wb3["Check"][f"B{CK1 + k}"].value, ck[k][2]) for k in range(len(ck)) if wb3["Check"][f"F{CK1 + k}"].value != "OK"]
    errs = [(w.title, c.coordinate, c.value) for w in wb3.worksheets for row in w.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("#") and c.value[:4] in ("#N/A", "#VAL", "#REF", "#NAM", "#DIV", "#NUM")]
    if bad or errs:
        print("LIBREOFFICE CHECK FAILED:", bad[:20], errs[:20])
        sys.exit(1)
    mism = [(s, k, LO[(s, k)], v) for (s, k), (v, nd) in EXPECT.items() if LO[(s, k)] is None or isinstance(LO[(s, k)], str)
            or not same_at(LO[(s, k)], v, nd)]
    if mism:
        print("VALUES DIFFER FROM R AT THE PRINTED DECIMALS:", mism)
        sys.exit(1)
    maxdiff = max(abs(float(wb3["Check"][f"D{CK1 + k}"].value)) for k in range(len(ck)))
    lo_ver = subprocess.run([soffice, "--version"], capture_output=True, text=True).stdout.strip()
    n_check_cell = wb3["Check"][REG[("Check", "n_check")]["cell"]].value
    LO_SUMMARY = (len(ck), n_check_cell, maxdiff, lo_ver)
    lo_note = (f"A copy was recalculated with {lo_ver} (headless): all {len(ck)} Check rows are OK (`Check!{REG[('Check', 'n_check')]['cell']}` "
               f"= {int(n_check_cell)}), no cell shows an error value, the largest absolute difference from R is {maxdiff:.1e}, and the "
               f"{len(EXPECT)} results printed here agree with R at the printed decimals.")
    print(f"verified (LibreOffice): {len(ck)} Check rows OK; rows with CHECK = {int(n_check_cell)}; no error values; largest |Excel - R| = {maxdiff:.2e}; "
          f"{len(EXPECT)} printed results agree with R")
    shutil.rmtree(tmp, ignore_errors=True)
else:
    print(lo_note)


# =============================================================================
# The walkthrough, generated from REG
# =============================================================================
def cl(sheet, key):
    return REG[(sheet, key)]["cell"]


def fx(sheet, key):
    fm = REG[(sheet, key)]["formula"]
    return "{" + fm + "}" if REG[(sheet, key)]["array"] else fm


def val(sheet, key, nd=None):
    v, nd0 = EXPECT.get((sheet, key), (None, 3))
    nd = nd0 if nd is None else nd
    got = LO.get((sheet, key))
    if got is not None and not isinstance(got, str):
        v = got
    if v is None:
        return str(got) if got is not None else ""
    if isinstance(v, bool):
        return str(v).upper()
    return f"{float(v):.{nd}f}"


def lov(sheet, key, nd=None):
    v = LO.get((sheet, key))
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)) and nd is not None:
        return f"{float(v):.{nd}f}"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return "" if v is None else str(v)


def row_(label, sheet, key, nd=None):
    return f"| {label} | `{sheet}!{cl(sheet, key)}` | `{fx(sheet, key)}` | {val(sheet, key, nd)} |"


def p3(x):
    x = float(x)
    return "< .001" if x < 0.001 else "= " + f"{x:.3f}".lstrip("0")


dw_first = "logo_gaze_dwell_pct|Ad_A"
desc_keys = [("n", "n", 0), ("mean", "Mean (AVERAGEIFS)", 3), ("sd", "SD (STDEV.S, array)", 3), ("median", "Median (array)", 3), ("q1", "Q1 (QUARTILE.INC, array)", 3),
             ("q3", "Q3 (QUARTILE.INC, array)", 3), ("iqr", "IQR", 3), ("zero", "% at zero", 2)]


def desc_table(keys_list):
    out = ["| Measure | Ad or stimulus | n | Mean | SD | Median | Q1 | Q3 | IQR | % at zero |", "|---|---|---|---|---|---|---|---|---|---|"]
    for meas, st, var in keys_list:
        key = f"{var}|{st}"
        cells = []
        for k, _, nd in desc_keys:
            cells.append(val("Descriptives", f"{k}|{key}", nd) if ("Descriptives", f"{k}|{key}") in REG else "")
        out.append(f"| {meas} | {st} | " + " | ".join(cells) + " |")
    return "\n".join(out)


def q5_table():
    keys = [("m1", "Ad_A mean", 3), ("s1", "Ad_A SD", 3), ("m2", "Ad_B mean", 3), ("s2", "Ad_B SD", 3), ("n", "n", 0), ("md", "Mean difference", 3),
            ("sdd", "SD of the differences", 3), ("se", "SE of the mean difference", 4), ("t", "t", 3), ("df", "df", 0),
            ("p_tt", "p (T.TEST, paired)", 5), ("p_hand", "p from t (T.DIST.2T)", 5), ("tcrit", "Critical t", 4), ("ci_lo", "95% CI lower", 3),
            ("ci_hi", "95% CI upper", 3), ("dz", "Cohen's d_z", 3), ("se_dz", "SE of d_z", 4), ("dz_lo", "d_z CI lower (approx.)", 3),
            ("dz_hi", "d_z CI upper (approx.)", 3), ("npos", "Ad_A higher", 0), ("nneg", "Ad_B higher", 0), ("ntie", "No difference", 0),
            ("nzero", "Zero in both ads", 0)]
    out = ["| Result | Joy cell | Joy | Engagement cell | Engagement | Brow Furrow cell | Brow Furrow |", "|---|---|---|---|---|---|---|"]
    for k, lab, nd in keys:
        cells = []
        for e, _, _ in EXPR:
            cells += [f"`Q5!{cl('Q5', e + '_' + k)}`", val("Q5", f"{e}_{k}", nd)]
        out.append(f"| {lab} | " + " | ".join(cells) + " |")
    return "\n".join(out)


def q5_formulas():
    keys = ["m1", "s1", "m2", "s2", "n", "md", "sdd", "se", "t", "df", "p_tt", "p_hand", "tcrit", "ci_lo", "ci_hi", "dz", "se_dz", "dz_lo", "dz_hi",
            "npos", "nneg", "ntie", "nzero"]
    out = ["| Result | Cell | Formula |", "|---|---|---|"]
    for k in keys:
        out.append(f"| {REG[('Q5', 'Joy_' + k)]['label']} | `Q5!{cl('Q5', 'Joy_' + k)}` | `{fx('Q5', 'Joy_' + k)}` |")
    return "\n".join(out)


holm_rows = "\n".join(
    f"| {e} | {val('Holm', 'p_' + e, 5)} | {lov('Holm', 'rank_' + e)} | {lov('Holm', 'mult_' + e)} | {lov('Holm', 'px_' + e, 5)} | "
    f"{val('Holm', 'holm_' + e, 5)} | {lov('Holm', 'bonf_' + e, 5)} | {lov('Holm', 'dec_' + e)} |" for e, _, _ in EXPR)
step_rows = "\n".join(
    f"| {k} | `Holm!{cl('Holm', f's_p{k}')}` | `{fx('Holm', f's_p{k}')}` | {lov('Holm', f's_p{k}', 5)} | `{fx('Holm', f's_m{k}')}` | "
    f"`{fx('Holm', f's_x{k}')}` | {lov('Holm', f's_x{k}', 5)} | `{fx('Holm', f's_a{k}')}` | {lov('Holm', f's_a{k}', 5)} |" for k in range(1, len(EXPR) + 1))
tj, te, tb = T["Q5_JOY"], T["Q5_ENGAGEMENT"], T["Q5_BROWFURROW"]
check_line = (f"`Check!{cl('Check', 'n_check')}` counts the rows with status CHECK and must read 0" +
              (f"; in the LibreOffice recalculation it reads {int(LO_SUMMARY[1])} for {LO_SUMMARY[0]} rows, and the largest absolute difference "
               f"between Excel and R is {LO_SUMMARY[2]:.1e}" if LO_SUMMARY else ""))

md = f"""# Excel walkthrough: descriptive statistics, paired tests, and the Holm correction

**SIMULATED DATA.** Every number in this walkthrough comes from the SIMULATED Multimodal Ad Study of Workshop 7 (`analysis_ready_long.csv`, seed {SEED}). No real person was recorded, ALTAVO and BRENLOW are invented brands, and the numbers must never be quoted as findings.

This is the Excel route through three parts of `analysis_full.R`: the descriptive statistics by ad and AOI (STEP 3), the Q5 paired tests of the facial expressions (STEP 6), and the Holm correction. The finished workbook is `SIMULATED_analysis_report.xlsx`. **Every cell address below is written as `Sheet!Cell` and was generated from the workbook itself:** `make_excel_workbook.py` writes the workbook and this file from one list of cells, reads the workbook back with openpyxl to confirm that each address holds the formula printed here, and recalculates a copy to confirm the results. {lo_note}

Companion files: `analysis_full.R` and `analysis_output.txt` (the R analysis and its full output), `SIMULATED_descriptives.csv` and `SIMULATED_all_tests.csv` (the R values the `Check` sheet uses), `results_text.md` (model paragraphs), `README.md`.

---

## 0. What you are working with

| Sheet | What it holds | Rows |
|---|---|---|
| `Data` | `analysis_ready_long.csv` (Workshop 7) as the Text Import leaves it; helper column `{KEY}` (orange) | {len(L_ROWS)} |
| `AOI_long` | one row per participant, ad, and AOI; the dwell values stacked into one column | {A_LAST - 1} |
| `Descriptives` | n, mean, SD, median, quartiles, IQR, percent at zero | {len(DESC_ROWS)} |
| `Pivot` | Average of dwell by AOI and ad (the PivotTable, as formulas) | 4 |
| `Q5` | Joy, Engagement, Brow Furrow: Ad_A versus Ad_B, paired | {len(ROSTER)} participants |
| `Holm` | the Holm correction of the three Q5 p values | 3 tests |
| `Check` | every Excel result next to the R value | {len(ck)} |

Blue headers mark data as imported; orange headers mark columns you add; yellow cells are typed by you (the participant, ad, and AOI grid and the roster). **Analysis sample:** include = 1 ({len(ROSTER)} of {len(IDS)} participants; P18 and P21 were excluded by rule R2 in Workshop 7).

---

## 1. Import the analysis ready table (Data sheet)

`analysis_ready_long.csv` has one row per participant and stimulus (Baseline, Ad_A, Ad_B). Line 1 is the SIMULATED banner; the column names are on line 2.

1. **Data > From Text/CSV** (Excel 2016 and later), select the file, **Transform Data**; in Power Query **Home > Remove Rows > Remove Top Rows** (1), then **Home > Use First Row as Headers**, set `participant_id`, `stimulus`, and `order` to **Text**, **Close & Load**. (Older Excel: **File > Open**, Text Import Wizard, **Delimited**, **Start import at row: 2**, **Comma**.)
2. **You should see** {len(L_HEAD)} columns (`A` = participant_id to `{DC['recall_correct']}` = recall_correct) and {len(L_ROWS)} data rows. Cells that read `NA` are text: the value does not exist (for example liking on a Baseline row). Leave them; AVERAGEIFS, STDEV.S, and MEDIAN skip text.
3. Rename the sheet `Data`. Add the key in column `{KEY}` and fill it down to row {D_LAST}:

| Result | Cell | Formula | Value |
|---|---|---|---|
| Key of the first row | `Data!{cl('Data', 'key')}` | `{fx('Data', 'key')}` | `{lov('Data', 'key')}` |

**Common mistake:** importing with the banner as the header row. Every column is then named `# SIMULATED ...` or `Column1`, and the numbers come in as text.

---

## 2. Stack the three AOIs into one column (AOI_long sheet)

The table has three dwell columns (`{DC['logo_gaze_dwell_pct']}` Logo, `{DC['product_gaze_dwell_pct']}` Product, `{DC['headline_gaze_dwell_pct']}` Headline). AVERAGEIFS and a PivotTable need ONE value column and one column that says which AOI the value belongs to ("long" format). Type the grid in columns `A` to `C` (yellow: P01 Ad_A Logo, P01 Ad_A Product, P01 Ad_A Headline, P01 Ad_B Logo, ...; {A_LAST - 1} rows), then fill these two formulas down to row {A_LAST}:

| Column | Result | Cell | Formula |
|---|---|---|---|
| `D` | include, looked up from Data | `AOI_long!{cl('AOI_long', 'D')}` | `{fx('AOI_long', 'D')}` |
| `E` | gaze dwell of the AOI named in column C | `AOI_long!{cl('AOI_long', 'E')}` | `{fx('AOI_long', 'E')}` |

**Excel 365:** `=XLOOKUP(A2&"|"&B2,{R_KEY},{DR('include')})` does the same as the INDEX and MATCH pair in column `D`.

**Check the stacking** (cells on the right of `AOI_long`):

| Check | Cell | Formula | Value |
|---|---|---|---|
{row_('Rows (24 x 2 x 3)', 'AOI_long', 'n_rows', 0)}
{row_('Lookups that failed (#N/A)', 'AOI_long', 'n_na', 0)}
{row_('Rows with include = 1 (22 x 6)', 'AOI_long', 'n_incl', 0)}

---

## 3. Descriptive statistics (Descriptives sheet)

Report n, a center, and a spread for every measure before any test. Percent of time and dwell pile up at zero (a participant who never looked at the Logo, or never reached the Joy threshold), so the median, the quartiles, and the percent at zero say things the mean hides.

**The formulas for the first row** (gaze dwell, Ad_A, Logo, row {DESC_ROWS[dw_first]}); the other rows are the same with their own row number. `B` and `C` hold the ad and the AOI, so the formulas never contain a typed name:

| Result | Cell | Formula |
|---|---|---|
""" + "\n".join(f"| {lab} | `Descriptives!{cl('Descriptives', f'{k}|{dw_first}')}` | `{fx('Descriptives', f'{k}|{dw_first}')}` |" for k, lab, _ in desc_keys) + f"""

**Array formulas.** STDEV.S, MEDIAN, and QUARTILE.INC have no "IFS" version, so `IF(...)` builds the list of values that meet the conditions: `(range = B{DESC_ROWS[dw_first]})*(range = C{DESC_ROWS[dw_first]})*(include = 1)` is 1 only where all three conditions hold, and `IF` returns the dwell value there and FALSE elsewhere, which STDEV.S skips. In **Excel 365** type the formula and press **Enter**; in **Excel 2019 or older** press **Ctrl+Shift+Enter**, and Excel shows braces `{{...}}` around it (do not type the braces). **Common mistake:** pressing Enter in an older Excel: the formula then uses only one row and returns a wrong number without an error.

**Excel 365 alternative** without an array formula: `=STDEV.S(FILTER({AR('E')},({AR('B')}=B{DESC_ROWS[dw_first]})*({AR('C')}=C{DESC_ROWS[dw_first]})*({AR('D')}=1)))`.

**A. Gaze dwell by ad and AOI** (percent of the 30 s ad)

{desc_table([('Logo', ad, 'logo_gaze_dwell_pct') for ad in ADS] + [('Product', ad, 'product_gaze_dwell_pct') for ad in ADS] + [('Headline', ad, 'headline_gaze_dwell_pct') for ad in ADS])}

**B. Facial expressions** (percent of time at or above threshold 50)

{desc_table([(e, st, var) for e, var, _ in EXPR for st in STIMS])}

**C. GSR Peaks Per Minute** (valid GSR only, rule R3; the formulas add the condition `gsr_valid = 1`)

{desc_table([('GSR Peaks Per Minute', st, 'gsr_peaks_per_min') for st in STIMS])}

**D. Liking** (1 to 7)

{desc_table([('Liking', ad, 'liking') for ad in ADS])}

**Read the tables.** The Product held {val('Descriptives', 'mean|product_gaze_dwell_pct|Ad_B', 1)} to {val('Descriptives', 'mean|product_gaze_dwell_pct|Ad_A', 1)} percent of the ad on average and the Logo less than 2 percent; {val('Descriptives', 'zero|logo_gaze_dwell_pct|Ad_A', 0)} percent of participants never had their gaze on the Logo in each ad, so its median is below its mean. Brow Furrow was zero for most participants (its median and both quartiles are 0): a mean of {val('Descriptives', 'mean|browfurrow_pct|Ad_A', 2)} percent describes almost nobody.

---

## 4. The PivotTable (Pivot sheet)

The iMotions article Creating Excel Pivot Tables (local copy: Making Excel Pivot Tables) gives the steps for an AOI metrics table; on `AOI_long` they are:

1. Click any cell of `AOI_long`, then **Insert > PivotTable**, choose **New Worksheet**, **OK**.
2. In the **PivotTable Fields** pane drag `AOI` to **Rows**, `ad` to **Columns**, and `gaze dwell (% of ad)` to **Values**.
3. **You should see** "Sum of gaze dwell (% of ad)". Right click a number, **Value Field Settings**, choose **Average**, **OK**.
4. Drag `include (from Data)` to **Filters**, open the filter, tick only **1**, **OK**.
5. **You should see** the averages below, with the rows sorted A to Z (Headline, Logo, Product). Add a second copy of the dwell field to Values with **Count** to see n = 22 in every cell.

**Common mistakes:** leaving the value on **Sum** (it adds the participants' values together); forgetting the include filter (24 participants instead of 22); building the PivotTable from `Data` (three dwell columns cannot go into Columns as one field; that is why section 2 stacked them).

A PivotTable cannot be checked cell by cell, so the `Pivot` sheet builds the same table with AVERAGEIFS and COUNTIFS:

| Result | Cell | Formula | Value |
|---|---|---|---|
""" + "\n".join(row_(f"Average dwell, {aoi}, {ad}", "Pivot", f"m_{aoi}_{ad}", 2) for aoi, _ in AOIS for ad in ADS) + "\n" + \
    "\n".join(row_(f"Grand average, {aoi} (both ads)", "Pivot", f"g_{aoi}", 2) for aoi, _ in AOIS) + "\n" + \
    "\n".join(row_(f"Grand average, {ad} (all AOIs)", "Pivot", f"g_{ad}", 2) for ad in ADS) + "\n" + \
    row_("Count, Logo, Ad_A", "Pivot", "n_Logo_Ad_A", 0) + f"""

The grand averages equal the estimated marginal means of the Q4 mixed model in `SIMULATED_q4_marginal_means.csv`, because every participant has every cell. The PivotTable shows WHAT the averages are; whether the AOI pattern differs between the ads is the Q4 interaction test, which needs R (`analysis_output.txt`, STEP 5: F(2, 126) = {float(T['Q4_AD_X_AOI']['statistic']):.2f}, p {p3(T['Q4_AD_X_AOI']['p'])}).

---

## 5. Q5 with T.TEST (Q5 sheet)

Q5 asks whether Joy, Engagement, and Brow Furrow differ between Ad_A and Ad_B. Each participant saw both ads, so the tests are **paired**: each one works on the participant's own difference (Ad_A minus Ad_B). The roster in `Q5!A{R0}:A{Q5_LAST}` lists the {len(ROSTER)} participants with include = 1 (typed, yellow; the `Check` sheet confirms it). The values come from `Data` by key:

| Result | Cell | Formula |
|---|---|---|
| Joy, Ad_A, first participant | `Q5!{cl('Q5', 'Joy_a')}` | `{fx('Q5', 'Joy_a')}` |
| Joy, Ad_B, first participant | `Q5!{cl('Q5', 'Joy_b')}` | `{fx('Q5', 'Joy_b')}` |
| Difference | `Q5!{cl('Q5', 'Joy_d')}` | `{fx('Q5', 'Joy_d')}` |

Engagement and Brow Furrow use the same formulas in their own columns. **The results block for Joy** (the blocks for Engagement and Brow Furrow sit to its right and use their own columns):

{q5_formulas()}

**All three tests:**

{q5_table()}

**Read the results.** T.TEST returns only the p value; its arguments are (range 1, range 2, tails = 2, type = 1 for paired). The hand computed t and T.DIST.2T must give the same p, which confirms that the two ranges are aligned row by row. **Common mistake:** type = 2 or 3 (two independent groups) ignores the pairing and gives a different, wrong p. R adds what Excel has no function for: the Wilcoxon signed rank checks (Joy p {p3(tj['np_p'])}, Engagement p {p3(te['np_p'])}, Brow Furrow p {p3(tb['np_p'])}) and the exact noncentral t intervals of d_z (Joy [{float(tj['es_ci_low']):.2f}, {float(tj['es_ci_high']):.2f}], Engagement [{float(te['es_ci_low']):.2f}, {float(te['es_ci_high']):.2f}], Brow Furrow [{float(tb['es_ci_low']):.2f}, {float(tb['es_ci_high']):.2f}]). Brow Furrow was zero in both ads for {val('Q5', 'Brow Furrow_nzero', 0)} of {len(ROSTER)} participants: its t test rests on {len(ROSTER) - int(float(val('Q5', 'Brow Furrow_ntie', 0)))} nonzero differences, and the Wilcoxon test is the better guide.

---

## 6. The Holm correction (Holm sheet)

Three tests on the same question (do the ads differ in facial expression?) give three chances of a false positive. Holm's procedure (Holm, 1979) keeps the chance of at least one false positive in the family at .05: sort the p values from smallest to largest, multiply the smallest by m (the number of tests), the next by m minus 1, and so on, and never let an adjusted value fall below the one before it.

**The step table** (rows {S0} to {SN}; m in `Holm!{cl('Holm', 'm')}` = `{fx('Holm', 'm')}`):

| Step k | Cell | k-th smallest p | Value | Multiplier | p x multiplier | Value | Holm adjusted | Value |
|---|---|---|---|---|---|---|---|---|
{step_rows}

**Back to each test** (rows {H0} to {HN}):

| Result | Cell | Formula |
|---|---|---|
| p, linked from Q5 | `Holm!{cl('Holm', 'p_Joy')}` | `{fx('Holm', 'p_Joy')}` |
| Rank (1 = smallest) | `Holm!{cl('Holm', 'rank_Joy')}` | `{fx('Holm', 'rank_Joy')}` |
| Multiplier | `Holm!{cl('Holm', 'mult_Joy')}` | `{fx('Holm', 'mult_Joy')}` |
| p x multiplier | `Holm!{cl('Holm', 'px_Joy')}` | `{fx('Holm', 'px_Joy')}` |
| Holm adjusted p | `Holm!{cl('Holm', 'holm_Joy')}` | `{fx('Holm', 'holm_Joy')}` |
| Bonferroni p | `Holm!{cl('Holm', 'bonf_Joy')}` | `{fx('Holm', 'bonf_Joy')}` |
| Decision | `Holm!{cl('Holm', 'dec_Joy')}` | `{fx('Holm', 'dec_Joy')}` |

| Test | p | Rank | Multiplier | p x multiplier | Holm p | Bonferroni p | Decision |
|---|---|---|---|---|---|---|---|
{holm_rows}

**Read the table.** The order from the smallest p is {HOLM_ORDER}. The running maximum matters here: {HOLM_RAISED}. **Common mistake:** reporting p x multiplier without the running maximum, which can make a later test look stronger than an earlier one. None of the three differences survives the correction (and none was significant before it). That does NOT show that the ads produce the same expressions: across the three expressions the 95% intervals of d_z run from {min(float(t_['es_ci_low']) for t_ in (tj, te, tb)):.2f} to {max(float(t_['es_ci_high']) for t_ in (tj, te, tb)):.2f}, so they include both no difference and a moderate advantage for Ad_A; 22 participants cannot tell these apart.

The R function `p.adjust(p, method = "holm")` gives the same numbers in one line (`analysis_output.txt`, STEP 6).

---

## 7. Check your numbers (Check sheet)

Every result printed above sits on the `Check` sheet next to the value that `analysis_full.R` wrote to `SIMULATED_descriptives.csv`, `SIMULATED_all_tests.csv`, `SIMULATED_q4_marginal_means.csv`, or `analysis_output.txt`, with the difference and a tolerance of half a unit in the last printed decimal. {check_line}. The last two rows confirm the Q5 roster: every listed participant has include = 1, and the roster holds all {len(ROSTER)}.

R only: H1 to H3 (Workshop 7), the Q4 mixed model and its checks, the Q4 and Q6 paired tests, the order checks, the Wilcoxon checks, the exact d_z intervals, the Spearman correlations, and the power analysis.

---

## 8. Common mistakes, collected

1. Averaging over all 24 participants: every COUNTIFS, AVERAGEIFS, and IF condition needs `include = 1`.
2. Pressing Enter on an array formula in Excel 2019 or older (no braces): the result is wrong without an error.
3. Reporting only the mean of a measure that piles up at zero: add the median and the percent at zero.
4. A PivotTable left on **Sum**, or built from the wide `Data` sheet instead of the stacked `AOI_long` sheet.
5. T.TEST with type 2 or 3 for a paired design.
6. Holm without the running maximum, or Bonferroni when Holm was planned (Bonferroni is never more powerful).
7. Reading "not significant" as "no difference": check the confidence interval of the difference and of d_z.
8. Typing values by hand instead of linking them (`=Q5!{cl('Q5', 'Joy_p_tt')}`, not `{val('Q5', 'Joy_p_tt', 5)}`): a typed value does not update when the data change (Broman & Woo, 2018).

---

## Sources

- Broman, K. W., & Woo, K. H. (2018). Data organization in spreadsheets. *The American Statistician, 72*(1), 2-10. https://doi.org/10.1080/00031305.2017.1375989
- Holm, S. (1979). A simple sequentially rejective multiple test procedure. *Scandinavian Journal of Statistics, 6*(2), 65-70.
- Lakens, D. (2013). Calculating and reporting effect sizes to facilitate cumulative science: A practical primer for t-tests and ANOVAs. *Frontiers in Psychology, 4*, Article 863. https://doi.org/10.3389/fpsyg.2013.00863
- iMotions help center, Creating Excel Pivot Tables (local PDF: Making Excel Pivot Tables): https://help.imotions.com/lab/exports/additional-export-information/creating-excel-pivot-tables

## Items to verify on the lab computer

1. The Excel version on the lab computers (Excel 365 takes array formulas with Enter and has FILTER and XLOOKUP; Excel 2019 or older needs Ctrl+Shift+Enter and has neither).
2. The Windows regional settings (expected: dot decimal, comma list separator); with other settings the formulas use semicolons between arguments.
"""


def esc_pipes(line):
    """Inside a Markdown table, a | in a code span must be written \\| or it splits the cell."""
    if not line.startswith("|"):
        return line
    return re.sub(r"`[^`]*`", lambda m: m.group(0).replace("|", "\\|"), line)


md = "\n".join(esc_pipes(ln) for ln in md.split("\n"))
with open(MD_OUT, "w", encoding="utf-8") as fh:
    fh.write(md)
print("wrote", MD_OUT)
