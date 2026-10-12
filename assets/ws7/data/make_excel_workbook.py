#!/usr/bin/env python3
"""
make_excel_workbook.py
Builds SIMULATED_data_prep.xlsx (the Excel route of Workshop 7) from the SIMULATED export bundle with
openpyxl, and writes excel_walkthrough.md from the SAME cell registry, so that every cell reference in
the walkthrough is the address of a live formula in the workbook.

The Excel route covers the core data preparation steps:
  1. Text Import of one metrics file (GSRPeakMetrics.csv), plus the survey matrix, the data quality file,
     and the session log, each placed on its own sheet exactly as the import would leave it;
  2. TRIM and UPPER to fix the respondent keys, with EXACT and LEN to SEE the problem;
  3. INDEX and MATCH (XLOOKUP in Excel 365) to join the survey and the data quality file to the GSR metrics;
  4. the exclusion rule (R2) and the GSR rule (R3) as formulas;
  5. a pivot table summary (a formula twin of the PivotTable, so that it can be checked);
  6. H3 with T.TEST (paired), d_z, and its approximate confidence interval.

Verification after the workbook is saved:
  1. openpyxl reads the workbook back and checks that every registered cell holds the registered formula.
  2. If LibreOffice is installed, a copy is recalculated headless (in a temporary folder; the delivered
     workbook is not changed) and every row of the Check sheet must say OK.
  3. Every result printed in the walkthrough is the LibreOffice value of that cell, and it must agree with
     the value that prepare_data.R or analysis.R produced at the number of decimals printed.

Live formulas use INDEX and MATCH, which every Excel version and LibreOffice understand; the walkthrough
gives the XLOOKUP version of each lookup for Excel 365 and Excel 2021 (LibreOffice 24.2, used for the
check, has no XLOOKUP).

Run: python3 make_excel_workbook.py   (after make_simulated_data.py, prepare_data.R, and analysis.R)
"""

import csv
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

HERE = os.path.dirname(os.path.abspath(__file__))
EX = os.path.join(HERE, "exports")
OUT = os.path.join(HERE, "SIMULATED_data_prep.xlsx")
MD_OUT = os.path.join(HERE, "excel_walkthrough.md")
REFS = os.path.join(HERE, "..", "refs", "references.json")

BLUE, ORANGE, LIGHT = "0072B2", "E69F00", "F2F2F2"
hdr_font = Font(bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor=BLUE)
help_fill = PatternFill("solid", fgColor=ORANGE)
bold = Font(bold=True)
note_font = Font(italic=True, color="555555")
code_font = Font(name="Consolas", size=9, color="555555")
thin = Side(style="thin", color="BFBFBF")
box = Border(top=thin, bottom=thin, left=thin, right=thin)
ET_MIN, FACE_MIN, GSR_MIN = 90, 80, 0.8


# ---------------------------------------------------------------- reading the bundle (as the import would)
def lines_of(path):
    with open(path, encoding="utf-8") as fh:
        return fh.read().splitlines()


def header_row(lines, first_col=None):
    """1 based line number of the header: the line after #DATA, else the first line that starts with first_col."""
    first = [re.split(r"[,;\t]", ln)[0] for ln in lines]
    if "#DATA" in first:
        return first.index("#DATA") + 2
    return first.index(first_col) + 1


def table(path, delim, first_col=None, quote=True):
    L = lines_of(path)
    h = header_row(L, first_col)
    rdr = csv.reader(io.StringIO("\n".join(L[h - 1:])), delimiter=delim, quoting=csv.QUOTE_MINIMAL if quote else csv.QUOTE_NONE)
    rows = list(rdr)
    return h, rows[0], rows[1:], L[0]


def to_cell(v):
    """Numbers become numbers; NA stays the text NA (as Excel imports it); an empty cell stays empty."""
    if v == "":
        return None
    try:
        f = float(v)
        return int(f) if f.is_integer() and "." not in v else f
    except ValueError:
        return v


GSR_PATH = os.path.join(EX, "04_GSR_Metrics", "GSRPeakMetrics.csv")
SV_PATH = os.path.join(EX, "05_Survey", "MERGED_SURVEY_RESPONSE_MATRIX.txt")
DQ_PATH = os.path.join(EX, "06_Data_Quality", "SIMULATED Multimodal Ad Study_DataQuality.txt")
LOG_PATH = os.path.join(EX, "08_Session_Log", "SIMULATED_AdStudy_SessionLog.csv")
g_h, g_head, g_rows, banner = table(GSR_PATH, ",")
s_h, s_head, s_rows, _ = table(SV_PATH, "\t", "STUDY", quote=False)
d_h, d_head, d_rows, _ = table(DQ_PATH, "\t", "Respondent_Name")
l_h, l_head, l_rows, _ = table(LOG_PATH, ",", "Participant ID")
SEED = re.search(r"seed (\d+)", banner).group(1)


def read_csv_banner(name):
    with open(os.path.join(HERE, name), encoding="utf-8") as fh:
        raw = fh.read().splitlines()
    start = 1 if raw[0].startswith("#") else 0
    rows = list(csv.reader(raw[start:]))
    return rows[0], rows[1:]


th, trows = read_csv_banner("SIMULATED_test_results.csv")
T = {r[0]: dict(zip(th, r)) for r in trows}
dh, drows = read_csv_banner("SIMULATED_descriptives.csv")
DS = {(r[0], r[1]): dict(zip(dh, r)) for r in drows}
lh, lrows = read_csv_banner("analysis_ready_long.csv")
LONG = [dict(zip(lh, r)) for r in lrows]
with open(os.path.join(HERE, "prepare_output.txt"), encoding="utf-8") as fh:
    prep_txt = fh.read()
with open(os.path.join(HERE, "analysis_output.txt"), encoding="utf-8") as fh:
    ana_txt = fh.read()
m_alpha = re.search(r"Alpha WITHOUT recoding ATT_3R = (-?[0-9.]+) \(the mistake\); WITH recoding = ([0-9.]+) \(n = (\d+)\)", prep_txt)
assert m_alpha, "alpha line not found in prepare_output.txt"
ALPHA_WRONG, ALPHA_RIGHT = float(m_alpha.group(1)), float(m_alpha.group(2))


def fnum(x):
    return float(x)


# ---------------------------------------------------------------- registry of documented cells
REG, EXPECT = {}, {}


def reg(sheet, key, cell, formula, label=""):
    REG[(sheet, key)] = {"cell": cell, "formula": formula.replace("_xlfn.", ""), "label": label}


def put(ws, key, cell, formula, label="", fmt=None, expect=None, nd=3):
    ws[cell] = formula
    if fmt:
        ws[cell].number_format = fmt
    reg(ws.title, key, cell, formula, label)
    if expect is not None:
        EXPECT[(ws.title, key)] = (expect, nd)


def style_header(ws, row=1, first=1, last=None, fill=None):
    last = last or ws.max_column
    for c in range(first, last + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill, cell.border = hdr_font, fill or hdr_fill, box
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def show_formula(cell, formula):
    cell.value = formula.replace("_xlfn.", "")
    cell.data_type = "s"
    cell.font = code_font


def widths(ws, d):
    for col, w in d.items():
        ws.column_dimensions[col].width = w


wb = Workbook()
# ================================================================ README
ws = wb.active
ws.title = "README"
readme = [
    ("SIMULATED DATA PREPARATION WORKBOOK (Workshop 7)", bold),
    (f"SIMULATED DATA for teaching (seed {SEED}). No real person was recorded. Built by make_simulated_data.py, prepare_data.R, analysis.R, "
     "and make_excel_workbook.py.", None),
    ("Study: SIMULATED Multimodal Ad Study. 24 participants (P01 to P24), Baseline (60 s), Ad_A and Ad_B (30 s, counterbalanced), survey.", None),
    ("This workbook is the Excel route through the core steps of prepare_data.R and the H3 test of analysis.R.", None),
    ("", None),
    ("Sheets", bold),
    ("GSR_Import    GSRPeakMetrics.csv as the Text Import leaves it (header in row 1), plus helper columns P to U (orange)", None),
    ("SessionLog    the lab session log (the key table: recording name -> participant, Use recording), plus the key in column M", None),
    ("Survey        MERGED_SURVEY_RESPONSE_MATRIX.txt (tab separated), plus helper columns V to AC and Cronbach's alpha", None),
    ("DataQuality   the _DataQuality.txt file (tab separated), plus helper columns I to M", None),
    ("Keys          the problem keys made visible with LEN, CODE, EXACT, TRIM, UPPER", None),
    ("Merged        one row per participant and stimulus (72 rows): INDEX and MATCH joins, rules R2 and R3", None),
    ("Pivot         the summary a PivotTable gives (built with AVERAGEIFS and COUNTIFS so that it can be checked)", None),
    ("H3            GSR Peaks Per Minute, each ad versus Baseline: paired T.TEST, d_z, and its approximate 95% CI", None),
    ("Check         every Excel result next to the value R produced, with OK or CHECK", None),
    ("", None),
    ("Colors: blue headers = data as exported; orange headers = columns you add. Yellow cells on Merged hold the participant and stimulus grid.", None),
    ("Rules written before collection: R1 use only recordings with Use recording = Yes; R2 exclude a participant when Eye tracking_Quality < 90 "
     "or Affectiva AFFDEX_Quality < 80 on Ad_A or Ad_B; R3 GSR of a stimulus is missing when Signal Duration < 80 percent of Duration.", None),
    ("The walkthrough excel_walkthrough.md explains every formula with its cell address.", None),
]
for i, (txt, f) in enumerate(readme, start=1):
    c = ws.cell(row=i, column=1, value=txt)
    if f:
        c.font = f
ws.column_dimensions["A"].width = 150

# ================================================================ SessionLog
ws = wb.create_sheet("SessionLog")
for c, h in enumerate(l_head, start=1):
    ws.cell(row=1, column=c, value=h)
for r, row in enumerate(l_rows, start=2):
    for c, v in enumerate(row, start=1):
        ws.cell(row=r, column=c, value=to_cell(v) if c != 4 else v)
style_header(ws, last=len(l_head))
LOG_LAST = len(l_rows) + 1
LC = {h: get_column_letter(i + 1) for i, h in enumerate(l_head)}
KEYCOL_LOG = get_column_letter(len(l_head) + 1)          # M
ws[f"{KEYCOL_LOG}1"] = "Key = UPPER(TRIM(name))"
style_header(ws, first=len(l_head) + 1, last=len(l_head) + 1, fill=help_fill)
for r in range(2, LOG_LAST + 1):
    ws[f"{KEYCOL_LOG}{r}"] = f"=UPPER(TRIM({LC['iMotions respondent name']}{r}))"
reg("SessionLog", "key2", f"{KEYCOL_LOG}2", f"=UPPER(TRIM({LC['iMotions respondent name']}2))", "key of the first recording")
R_LOG_PID = f"SessionLog!${LC['Participant ID']}$2:${LC['Participant ID']}${LOG_LAST}"
R_LOG_KEY = f"SessionLog!${KEYCOL_LOG}$2:${KEYCOL_LOG}${LOG_LAST}"
R_LOG_USE = f"SessionLog!${LC['Use recording']}$2:${LC['Use recording']}${LOG_LAST}"
R_LOG_NAME = f"SessionLog!${LC['iMotions respondent name']}$2:${LC['iMotions respondent name']}${LOG_LAST}"
widths(ws, {get_column_letter(i + 1): w for i, w in enumerate([10, 14, 9, 11, 8, 6, 20, 11, 10, 14, 20, 9, 90, 16])})
LOG_ROW = {row[1]: i + 2 for i, row in enumerate(l_rows)}       # recording name -> sheet row

# ================================================================ GSR_Import
ws = wb.create_sheet("GSR_Import")
for c, h in enumerate(g_head, start=1):
    ws.cell(row=1, column=c, value=h)
for r, row in enumerate(g_rows, start=2):
    for c, v in enumerate(row, start=1):
        ws.cell(row=r, column=c, value=to_cell(v))
style_header(ws, last=len(g_head))
G_LAST = len(g_rows) + 1
GC = {h: get_column_letter(i + 1) for i, h in enumerate(g_head)}
G_HELP = ["Key = UPPER(TRIM())", "Participant", "Use recording", "Lookup key", "Signal share", "GSR valid (R3)"]
GH = {h: get_column_letter(len(g_head) + 1 + k) for k, h in enumerate(G_HELP)}
for h in G_HELP:
    ws[f"{GH[h]}1"] = h
style_header(ws, first=len(g_head) + 1, last=len(g_head) + len(G_HELP), fill=help_fill)
B, F_, G_, I_, J_, K_ = GC["Respondent Name"], GC["Type"], GC["Label"], GC["Duration (ms)"], GC["Signal Duration (ms)"], GC["Parent Stimulus"]
gf = {
    "Key = UPPER(TRIM())": "=UPPER(TRIM({B}{r}))",
    "Participant": '=IFERROR(INDEX(' + R_LOG_PID + ',MATCH({KEY}{r},' + R_LOG_KEY + ',0)),"NOT FOUND")',
    "Use recording": '=IFERROR(INDEX(' + R_LOG_USE + ',MATCH({KEY}{r},' + R_LOG_KEY + ',0)),"NOT FOUND")',
    "Lookup key": '=IF({USE}{r}="Yes",{PID}{r}&"|"&{F}{r}&"|"&{G}{r}&"|"&{K}{r},"DROP")',
    "Signal share": "=IF({I}{r}>0,{J}{r}/{I}{r},0)",
    "GSR valid (R3)": "=IF({SH}{r}>=" + str(GSR_MIN) + ",1,0)",
}
subs = {"B": B, "F": F_, "G": G_, "I": I_, "J": J_, "K": K_, "KEY": GH["Key = UPPER(TRIM())"], "PID": GH["Participant"],
        "USE": GH["Use recording"], "SH": GH["Signal share"]}
for r in range(2, G_LAST + 1):
    for h in G_HELP:
        ws[f"{GH[h]}{r}"] = gf[h].format(r=r, **subs)
        if h == "Signal share":
            ws[f"{GH[h]}{r}"].number_format = "0.000"
for h, k in zip(G_HELP, ("key", "pid", "use", "lkey", "share", "valid")):
    reg("GSR_Import", k, f"{GH[h]}2", gf[h].format(r=2, **subs), h)
R_G_LKEY = f"GSR_Import!${GH['Lookup key']}$2:${GH['Lookup key']}${G_LAST}"
R_G_PPM = f"GSR_Import!${GC['Peaks Per Minute']}$2:${GC['Peaks Per Minute']}${G_LAST}"
R_G_VALID = f"GSR_Import!${GH['GSR valid (R3)']}$2:${GH['GSR valid (R3)']}${G_LAST}"
R_G_PID = f"GSR_Import!${GH['Participant']}$2:${GH['Participant']}${G_LAST}"
widths(ws, {get_column_letter(i): 13 for i in range(1, len(g_head) + len(G_HELP) + 1)})
ws.column_dimensions["A"].width = 30
ws.column_dimensions[GH["Lookup key"]].width = 30
G_ROWS = {(row[g_head.index("Respondent Name")], row[g_head.index("Type")], row[g_head.index("Label")], row[g_head.index("Parent Stimulus")]): i + 2
          for i, row in enumerate(g_rows)}

# ================================================================ Survey
ws = wb.create_sheet("Survey")
for c, h in enumerate(s_head, start=1):
    ws.cell(row=1, column=c, value=h)
for r, row in enumerate(s_rows, start=2):
    for c, v in enumerate(row, start=1):
        ws.cell(row=r, column=c, value=to_cell(v))
style_header(ws, last=len(s_head))
S_LAST = len(s_rows) + 1
SC = {}
for i, h in enumerate(s_head):
    m = re.match(r'^(LABELID|LABELVALUE)_Questionnaire_([^,]+),', h)
    SC[(m.group(1), m.group(2)) if m else h] = get_column_letter(i + 1)
V = lambda q: SC[("LABELVALUE", q)]
S_HELP = ["Key = UPPER(TRIM())", "Participant", "Recall_A correct", "Recall_B correct", "ATT_3 recoded", "Attitude", "Sum (recoded)",
          "Sum (not recoded)"]
SH = {h: get_column_letter(len(s_head) + 1 + k) for k, h in enumerate(S_HELP)}
for h in S_HELP:
    ws[f"{SH[h]}1"] = h
style_header(ws, first=len(s_head) + 1, last=len(s_head) + len(S_HELP), fill=help_fill)
sf = {
    "Key = UPPER(TRIM())": "=UPPER(TRIM(B{r}))",
    "Participant": '=IFERROR(INDEX(' + R_LOG_PID + ',MATCH(' + SH["Key = UPPER(TRIM())"] + '{r},' + R_LOG_KEY + ',0)),"NOT FOUND")',
    "Recall_A correct": "=IF(" + V("Recall_A") + "{r}=1,1,0)",
    "Recall_B correct": "=IF(" + V("Recall_B") + "{r}=2,1,0)",
    "ATT_3 recoded": "=8-" + V("ATT_3R") + "{r}",
    "Attitude": "=AVERAGE(" + V("ATT_1") + "{r}," + V("ATT_2") + "{r}," + SH["ATT_3 recoded"] + "{r})",
    "Sum (recoded)": "=" + V("ATT_1") + "{r}+" + V("ATT_2") + "{r}+" + SH["ATT_3 recoded"] + "{r}",
    "Sum (not recoded)": "=" + V("ATT_1") + "{r}+" + V("ATT_2") + "{r}+" + V("ATT_3R") + "{r}",
}
for r in range(2, S_LAST + 1):
    for h in S_HELP:
        ws[f"{SH[h]}{r}"] = sf[h].format(r=r)
        if h == "Attitude":
            ws[f"{SH[h]}{r}"].number_format = "0.00"
for h, k in zip(S_HELP, ("key", "pid", "rec_a", "rec_b", "att3", "att", "sum", "sum_raw")):
    reg("Survey", k, f"{SH[h]}2", sf[h].format(r=2), h)
rng = lambda col: f"{col}2:{col}{S_LAST}"
AL = get_column_letter(len(s_head) + len(S_HELP) + 2)
AV = get_column_letter(len(s_head) + len(S_HELP) + 3)
ws[f"{AL}1"] = "Cronbach's alpha (k = 3)"
ws[f"{AL}1"].font = bold
ws[f"{AL}2"] = "With ATT_3R recoded"
put(ws, "alpha", f"{AV}2", f"=3/2*(1-(_xlfn.VAR.S({rng(V('ATT_1'))})+_xlfn.VAR.S({rng(V('ATT_2'))})+_xlfn.VAR.S({rng(SH['ATT_3 recoded'])}))"
                           f"/_xlfn.VAR.S({rng(SH['Sum (recoded)'])}))", "Cronbach's alpha, recoded", "0.000", ALPHA_RIGHT, 3)
ws[f"{AL}3"] = "Without recoding (the mistake)"
put(ws, "alpha_wrong", f"{AV}3", f"=3/2*(1-(_xlfn.VAR.S({rng(V('ATT_1'))})+_xlfn.VAR.S({rng(V('ATT_2'))})+_xlfn.VAR.S({rng(V('ATT_3R'))}))"
                                 f"/_xlfn.VAR.S({rng(SH['Sum (not recoded)'])}))", "Cronbach's alpha, not recoded", "0.000", ALPHA_WRONG, 3)
ws.column_dimensions[AL].width = 30
for i in range(1, len(s_head) + 1):
    ws.column_dimensions[get_column_letter(i)].width = 12 if i > 5 else 14
ws.column_dimensions["A"].width = 30
R_S_PID = f"Survey!${SH['Participant']}$2:${SH['Participant']}${S_LAST}"
R_S_LIKE = {"Ad_A": f"Survey!${V('Liking_A')}$2:${V('Liking_A')}${S_LAST}", "Ad_B": f"Survey!${V('Liking_B')}$2:${V('Liking_B')}${S_LAST}"}
R_S_REC = {"Ad_A": f"Survey!${SH['Recall_A correct']}$2:${SH['Recall_A correct']}${S_LAST}",
           "Ad_B": f"Survey!${SH['Recall_B correct']}$2:${SH['Recall_B correct']}${S_LAST}"}
S_ROW = {row[1]: i + 2 for i, row in enumerate(s_rows)}

# ================================================================ DataQuality
ws = wb.create_sheet("DataQuality")
for c, h in enumerate(d_head, start=1):
    ws.cell(row=1, column=c, value=h)
for r, row in enumerate(d_rows, start=2):
    for c, v in enumerate(row, start=1):
        ws.cell(row=r, column=c, value=to_cell(v))
style_header(ws, last=len(d_head))
D_LAST = len(d_rows) + 1
DC = {h: get_column_letter(i + 1) for i, h in enumerate(d_head)}
D_HELP = ["Key = UPPER(TRIM())", "Participant", "Use recording", "Lookup key", "LEN of the raw name"]
DH = {h: get_column_letter(len(d_head) + 1 + k) for k, h in enumerate(D_HELP)}
for h in D_HELP:
    ws[f"{DH[h]}1"] = h
style_header(ws, first=len(d_head) + 1, last=len(d_head) + len(D_HELP), fill=help_fill)
df_ = {
    "Key = UPPER(TRIM())": "=UPPER(TRIM(A{r}))",
    "Participant": '=IFERROR(INDEX(' + R_LOG_PID + ',MATCH(' + DH["Key = UPPER(TRIM())"] + '{r},' + R_LOG_KEY + ',0)),"NOT FOUND")',
    "Use recording": '=IFERROR(INDEX(' + R_LOG_USE + ',MATCH(' + DH["Key = UPPER(TRIM())"] + '{r},' + R_LOG_KEY + ',0)),"NOT FOUND")',
    "Lookup key": '=IF(' + DH["Use recording"] + '{r}="Yes",' + DH["Participant"] + '{r}&"|"&B{r},"DROP")',
    "LEN of the raw name": "=LEN(A{r})",
}
for r in range(2, D_LAST + 1):
    for h in D_HELP:
        ws[f"{DH[h]}{r}"] = df_[h].format(r=r)
for h, k in zip(D_HELP, ("key", "pid", "use", "lkey", "len")):
    reg("DataQuality", k, f"{DH[h]}2", df_[h].format(r=2), h)
R_D_LKEY = f"DataQuality!${DH['Lookup key']}$2:${DH['Lookup key']}${D_LAST}"
R_D_ET = f"DataQuality!${DC['Eye tracking_Quality']}$2:${DC['Eye tracking_Quality']}${D_LAST}"
R_D_FACE = f"DataQuality!${DC['Affectiva AFFDEX_Quality']}$2:${DC['Affectiva AFFDEX_Quality']}${D_LAST}"
widths(ws, {get_column_letter(i): 15 for i in range(1, len(d_head) + len(D_HELP) + 1)})
D_ROWS = {}
for i, row in enumerate(d_rows):
    D_ROWS.setdefault(row[0], i + 2)

# ================================================================ Keys
ws = wb.create_sheet("Keys")
ws["A1"] = "The problem keys, made visible. Each raw key is a live reference to the cell where it sits in the imported data."
ws["A1"].font = bold
khead = ["Where", "Raw key (reference)", "LEN", "CODE of the last character", "EXACT match in the session log",
         "MATCH finds it (ignores case)", "UPPER(TRIM())", "Participant", "Use recording"]
for c, h in enumerate(khead, start=1):
    ws.cell(row=3, column=c, value=h)
style_header(ws, row=3, last=len(khead))
p07_dq = D_ROWS["P07 "]
p07_sv = S_ROW["p07"]
p04_g = G_ROWS[("P04", "Stimulus", "Baseline", "Baseline")]
p04r_g = G_ROWS[("P04_rerun", "Stimulus", "Baseline", "Baseline")]
p01_g = G_ROWS[("P01", "Stimulus", "Baseline", "Baseline")]
key_rows = [("DataQuality, column A", f"=DataQuality!A{p07_dq}", "kdq"), ("Survey, column B", f"=Survey!B{p07_sv}", "ksv"),
            ("GSR_Import, column B", f"=GSR_Import!B{p04_g}", "kp04"), ("GSR_Import, column B", f"=GSR_Import!B{p04r_g}", "kp04r"),
            ("GSR_Import, column B (a normal key)", f"=GSR_Import!B{p01_g}", "kp01")]
KROW = {}
for i, (where, ref, key) in enumerate(key_rows):
    r = 4 + i
    KROW[key] = r
    ws[f"A{r}"] = where
    fs = {"B": ref, "C": f"=LEN(B{r})", "D": f"=CODE(RIGHT(B{r},1))", "E": f"=SUMPRODUCT(--EXACT(B{r},{R_LOG_NAME}))>0",
          "F": f"=ISNUMBER(MATCH(B{r},{R_LOG_NAME},0))", "G": f"=UPPER(TRIM(B{r}))",
          "H": f'=IFERROR(INDEX({R_LOG_PID},MATCH(G{r},{R_LOG_KEY},0)),"NOT FOUND")',
          "I": f'=IFERROR(INDEX({R_LOG_USE},MATCH(G{r},{R_LOG_KEY},0)),"NOT FOUND")'}
    for col, f in fs.items():
        put(ws, f"{key}_{col}", f"{col}{r}", f, f"{where} {col}")
ws["A10"] = ("Read row by row: LEN 4 and CODE 32 reveal the trailing space; EXACT is FALSE for p07 although MATCH finds it, because Excel lookups "
             "ignore upper and lower case (R joins do not). TRIM and UPPER give every row a key that the session log knows; the Use recording "
             "column then says which recording of P04 to keep.")
ws["A10"].font = note_font
widths(ws, {"A": 34, "B": 20, "C": 8, "D": 12, "E": 18, "F": 18, "G": 16, "H": 12, "I": 12})

# ================================================================ Merged
ws = wb.create_sheet("Merged")
mhead = ["participant_id", "stimulus", "GSR lookup key", "GSR Peaks Per Minute (as exported)", "GSR valid (R3)", "GSR Peaks Per Minute (analysis)",
         "DQ lookup key", "Eye tracking quality (%)", "Face quality (%)", "Eye tracking minimum, ads", "Face minimum, ads", "include (R2)",
         "liking", "recall_correct"]
for c, h in enumerate(mhead, start=1):
    ws.cell(row=1, column=c, value=h)
style_header(ws, last=2)
style_header(ws, first=3, last=len(mhead), fill=help_fill)
yellow = PatternFill("solid", fgColor="FFF2CC")
IDS = [f"P{i:02d}" for i in range(1, 25)]
STIMS = ["Baseline", "Ad_A", "Ad_B"]
M_LAST = 1 + len(IDS) * len(STIMS)
mf = {
    "C": '=A{r}&"|Stimulus|"&B{r}&"|"&B{r}',
    "D": "=INDEX(" + R_G_PPM + ",MATCH(C{r}," + R_G_LKEY + ",0))",
    "E": "=INDEX(" + R_G_VALID + ",MATCH(C{r}," + R_G_LKEY + ",0))",
    "F": '=IF(E{r}=1,D{r},"")',
    "G": '=A{r}&"|"&B{r}',
    "H": "=INDEX(" + R_D_ET + ",MATCH(G{r}," + R_D_LKEY + ",0))",
    "I": "=INDEX(" + R_D_FACE + ",MATCH(G{r}," + R_D_LKEY + ",0))",
    "J": '=MIN(INDEX(' + R_D_ET + ',MATCH(A{r}&"|Ad_A",' + R_D_LKEY + ',0)),INDEX(' + R_D_ET + ',MATCH(A{r}&"|Ad_B",' + R_D_LKEY + ',0)))',
    "K": '=MIN(INDEX(' + R_D_FACE + ',MATCH(A{r}&"|Ad_A",' + R_D_LKEY + ',0)),INDEX(' + R_D_FACE + ',MATCH(A{r}&"|Ad_B",' + R_D_LKEY + ',0)))',
    "L": "=IF(AND(J{r}>=" + str(ET_MIN) + ",K{r}>=" + str(FACE_MIN) + "),1,0)",
    "M": '=IF(B{r}="Ad_A",INDEX(' + R_S_LIKE["Ad_A"] + ',MATCH(A{r},' + R_S_PID + ',0)),IF(B{r}="Ad_B",INDEX(' + R_S_LIKE["Ad_B"] + ',MATCH(A{r},' + R_S_PID + ',0)),""))',
    "N": '=IF(B{r}="Ad_A",INDEX(' + R_S_REC["Ad_A"] + ',MATCH(A{r},' + R_S_PID + ',0)),IF(B{r}="Ad_B",INDEX(' + R_S_REC["Ad_B"] + ',MATCH(A{r},' + R_S_PID + ',0)),""))',
}
MROW = {}
r = 2
for pid in IDS:
    for st in STIMS:
        ws[f"A{r}"], ws[f"B{r}"] = pid, st
        ws[f"A{r}"].fill = ws[f"B{r}"].fill = yellow
        for col, f in mf.items():
            ws[f"{col}{r}"] = f.format(r=r)
        MROW[(pid, st)] = r
        r += 1
for col in mf:
    reg("Merged", col, f"{col}2", mf[col].format(r=2), mhead[ord(col) - 65])
for col, key in (("D", "D_ad"), ("M", "M_ad"), ("N", "N_ad")):
    rr = MROW[("P01", "Ad_A")]
    reg("Merged", key, f"{col}{rr}", mf[col].format(r=rr), mhead[ord(col) - 65] + " (P01, Ad_A)")
ws[f"P2"] = "Checks"
ws["P2"].font = bold
mchk = [("n_rows", "Rows (participant x stimulus)", f"=COUNTA(A2:A{M_LAST})", len(IDS) * 3, 0),
        ("n_na", "Lookups that failed (#N/A) in D, H, I", f"=SUMPRODUCT(--ISNA(D2:D{M_LAST}))+SUMPRODUCT(--ISNA(H2:H{M_LAST}))+SUMPRODUCT(--ISNA(I2:I{M_LAST}))", 0, 0),
        ("n_incl", "Participants with include = 1", f'=COUNTIFS(B2:B{M_LAST},"Baseline",L2:L{M_LAST},1)', 22, 0),
        ("n_gsr_invalid", "Stimulus rows with GSR not valid", f"=COUNTIF(E2:E{M_LAST},0)", 1, 0)]
for i, (key, lab, f, ex, nd) in enumerate(mchk):
    ws[f"P{3 + i}"] = lab
    put(ws, key, f"Q{3 + i}", f, lab, "0", ex, nd)
widths(ws, {"A": 13, "B": 10, "C": 26, "D": 14, "E": 9, "F": 14, "G": 14, "H": 11, "I": 11, "J": 12, "K": 12, "L": 9, "M": 8, "N": 9, "P": 40, "Q": 8})
ws.freeze_panes = "C2"
R_M = lambda col: f"Merged!${col}$2:${col}${M_LAST}"

# ================================================================ Pivot
ws = wb.create_sheet("Pivot")
ws["A1"] = "Pivot summary: analysis sample (include = 1), GSR valid only for GSR. Built with AVERAGEIFS and COUNTIFS; a PivotTable gives the same numbers."
ws["A1"].font = bold
phead = ["Stimulus", "n (GSR)", "Mean GSR Peaks Per Minute", "n (liking)", "Mean liking", "Brand recalled (proportion)"]
for c, h in enumerate(phead, start=1):
    ws.cell(row=3, column=c, value=h)
style_header(ws, row=3, last=len(phead))
PROW = {}
for i, st in enumerate(STIMS):
    rr = 4 + i
    PROW[st] = rr
    ws[f"A{rr}"] = st
    put(ws, f"n_{st}", f"B{rr}", f'=COUNTIFS({R_M("B")},A{rr},{R_M("L")},1,{R_M("E")},1)', f"n GSR {st}", "0",
        int(DS[("GSR Peaks Per Minute", st)]["n"]), 0)
    put(ws, f"m_{st}", f"C{rr}", f'=AVERAGEIFS({R_M("D")},{R_M("B")},A{rr},{R_M("L")},1,{R_M("E")},1)', f"mean GSR {st}", "0.000",
        fnum(DS[("GSR Peaks Per Minute", st)]["mean"]), 3)
    if st != "Baseline":
        put(ws, f"nl_{st}", f"D{rr}", f'=COUNTIFS({R_M("B")},A{rr},{R_M("L")},1)', f"n liking {st}", "0", int(DS[("Liking (1 to 7)", st)]["n"]), 0)
        put(ws, f"ml_{st}", f"E{rr}", f'=AVERAGEIFS({R_M("M")},{R_M("B")},A{rr},{R_M("L")},1)', f"mean liking {st}", "0.000",
            fnum(DS[("Liking (1 to 7)", st)]["mean"]), 3)
        put(ws, f"rc_{st}", f"F{rr}", f'=AVERAGEIFS({R_M("N")},{R_M("B")},A{rr},{R_M("L")},1)', f"recall {st}", "0.000",
            fnum(DS[("Brand recalled (proportion)", st)]["mean"]), 3)
ws["A8"] = ("To build the real PivotTable: click any cell of Merged, Insert > PivotTable > New Worksheet > OK; drag stimulus to Rows, "
            "GSR Peaks Per Minute (analysis) to Values; right click the value, Value Field Settings, change Sum to Average; drag include (R2) "
            "to Filters and keep 1.")
ws["A8"].font = note_font
widths(ws, {"A": 14, "B": 10, "C": 26, "D": 11, "E": 12, "F": 26})


# ================================================================ H3
class Block:
    """A column of labeled results; {key} in a formula becomes the address of that key's cell."""

    def __init__(self, ws, row, lc):
        self.ws, self.row, self.lc, self.ref = ws, row, lc, {}

    def title(self, text):
        self.ws.cell(row=self.row, column=self.lc, value=text).font = bold
        self.row += 1

    def put(self, key, label, formula, note="", fmt="0.0000", expect=None, nd=3, pref=""):
        f = formula
        for k, rr in self.ref.items():
            f = f.replace("{" + k + "}", f"{get_column_letter(self.lc + 1)}{rr}")
        self.ws.cell(row=self.row, column=self.lc, value=label)
        cell = f"{get_column_letter(self.lc + 1)}{self.row}"
        put(self.ws, pref + key, cell, f, label, fmt, expect, nd)
        if note:
            self.ws.cell(row=self.row, column=self.lc + 2, value=note).font = note_font
        self.ref[key] = self.row
        self.row += 1


ws = wb.create_sheet("H3")
ws["A1"] = ("H3: GSR Peaks Per Minute, each ad versus Baseline (paired). Rosters: participants with include = 1 and valid GSR in both stimuli "
            "(static lists, checked on the Check sheet).")
ws["A1"].font = bold
H3B = {}
long_ok = {}
for row in LONG:
    long_ok[(row["participant_id"], row["stimulus"])] = (row["include"] == "1", row["gsr_valid"] == "1")
for bi, (ad, col0, tid) in enumerate((("Ad_A", 1, "H3_AD_A"), ("Ad_B", 7, "H3_AD_B"))):
    roster = [p for p in IDS if long_ok[(p, "Baseline")] == (True, True) and long_ok[(p, ad)] == (True, True)]
    c0, c1, c2, c3 = (get_column_letter(col0 + k) for k in range(4))
    ws[f"{c0}3"] = f"{ad} versus Baseline"
    ws[f"{c0}3"].font = bold
    for k, h in enumerate(["participant_id", "Baseline", ad, f"{ad} minus Baseline"]):
        ws.cell(row=4, column=col0 + k, value=h)
    style_header(ws, row=4, first=col0, last=col0 + 1)
    style_header(ws, first=col0 + 2, last=col0 + 3, row=4, fill=help_fill)
    for i, p in enumerate(roster):
        rr = 5 + i
        ws[f"{c0}{rr}"] = p
        ws[f"{c0}{rr}"].fill = yellow
        ws[f"{c1}{rr}"] = f'=INDEX({R_M("F")},MATCH({c0}{rr}&"|Stimulus|Baseline|Baseline",{R_M("C")},0))'
        ws[f"{c2}{rr}"] = f'=INDEX({R_M("F")},MATCH({c0}{rr}&"|Stimulus|{ad}|{ad}",{R_M("C")},0))'
        ws[f"{c3}{rr}"] = f"={c2}{rr}-{c1}{rr}"
    last = 4 + len(roster)
    reg("H3", f"{ad}_base", f"{c1}5", ws[f"{c1}5"].value, f"Baseline for {roster[0]}")
    reg("H3", f"{ad}_ad", f"{c2}5", ws[f"{c2}5"].value, f"{ad} for {roster[0]}")
    reg("H3", f"{ad}_diff", f"{c3}5", ws[f"{c3}5"].value, "difference")
    r1, r2, rd = f"{c1}5:{c1}{last}", f"{c2}5:{c2}{last}", f"{c3}5:{c3}{last}"
    t = T[tid]
    blk = Block(ws, last + 3, col0)
    pf_ = f"{ad}_"
    blk.title(f"Paired t test, {ad} minus Baseline")
    blk.put("m1", "Baseline mean", f"=AVERAGE({r1})", expect=fnum(t["m1"]), nd=3, pref=pf_)
    blk.put("s1", "Baseline SD", f"=_xlfn.STDEV.S({r1})", "STDEV.S = sample SD", expect=fnum(t["sd1"]), nd=3, pref=pf_)
    blk.put("m2", f"{ad} mean", f"=AVERAGE({r2})", expect=fnum(t["m2"]), nd=3, pref=pf_)
    blk.put("s2", f"{ad} SD", f"=_xlfn.STDEV.S({r2})", expect=fnum(t["sd2"]), nd=3, pref=pf_)
    blk.put("n", "n (participants)", f"=COUNT({rd})", fmt="0", expect=int(float(t["n"])), nd=0, pref=pf_)
    blk.put("md", "Mean difference", f"=AVERAGE({rd})", expect=fnum(t["mean_diff"]), nd=3, pref=pf_)
    blk.put("sdd", "SD of the differences", f"=_xlfn.STDEV.S({rd})", expect=fnum(t["sd_diff"]), nd=3, pref=pf_)
    blk.put("se", "SE of the mean difference", "={sdd}/SQRT({n})", "SD of the differences / SQRT(n)", pref=pf_)
    blk.put("t", "t", "={md}/{se}", "mean difference / SE", expect=fnum(t["statistic"]), nd=3, pref=pf_)
    blk.put("df", "df", "={n}-1", fmt="0", expect=int(float(t["df"])), nd=0, pref=pf_)
    blk.put("p_tt", "p, two tailed (T.TEST, type 1 = paired)", f"=_xlfn.T.TEST({r2},{r1},2,1)", "tails = 2, type = 1", fmt="0.000000",
            expect=fnum(t["p"]), nd=5, pref=pf_)
    blk.put("p_hand", "p from t (T.DIST.2T)", "=_xlfn.T.DIST.2T(ABS({t}),{df})", "must equal the T.TEST cell", fmt="0.000000",
            expect=fnum(t["p"]), nd=5, pref=pf_)
    blk.put("tcrit", "Critical t (95%)", "=_xlfn.T.INV.2T(0.05,{df})", pref=pf_)
    blk.put("ci_lo", "Mean difference, 95% CI lower", "={md}-{tcrit}*{se}", expect=fnum(t["diff_ci_low"]), nd=3, pref=pf_)
    blk.put("ci_hi", "Mean difference, 95% CI upper", "={md}+{tcrit}*{se}", expect=fnum(t["diff_ci_high"]), nd=3, pref=pf_)
    blk.put("dz", "Cohen's d_z", "={md}/{sdd}", "mean difference / SD of the differences", expect=fnum(t["es"]), nd=3, pref=pf_)
    blk.put("se_dz", "SE of d_z (approximation)", "=SQRT(1/{n}+{dz}^2/(2*{n}))", pref=pf_)
    blk.put("dz_lo", "d_z 95% CI lower (approximation)", "={dz}-1.96*{se_dz}", "R also gives the exact noncentral t interval",
            expect=fnum(t["es_ci_low_approx"]), nd=3, pref=pf_)
    blk.put("dz_hi", "d_z 95% CI upper (approximation)", "={dz}+1.96*{se_dz}", expect=fnum(t["es_ci_high_approx"]), nd=3, pref=pf_)
    blk.put("npos", "Participants with an increase", f'=COUNTIF({rd},">0")', fmt="0", pref=pf_)
    H3B[ad] = {"roster": roster, "blk": blk, "c": (c0, c1, c2, c3), "last": last}
widths(ws, {"A": 34, "B": 12, "C": 12, "D": 16, "G": 34, "H": 12, "I": 12, "J": 16, "E": 4, "F": 4})
ws.column_dimensions["C"].width = 34
ws.column_dimensions["I"].width = 34

# ================================================================ Check
ws = wb.create_sheet("Check")
ws["A1"] = "Excel results next to the values that prepare_data.R and analysis.R produced (SIMULATED data)."
ws["A1"].font = bold
for c, h in enumerate(["Result", "Excel", "R", "Difference", "Tolerance", "Status"], start=1):
    ws.cell(row=3, column=c, value=h)
style_header(ws, row=3, last=6)
ck = []
for (sheet, key), (val, nd) in EXPECT.items():
    tol = 1e-9 if nd == 0 else 0.5 * 10 ** (-nd)
    ck.append((f"{sheet}: {REG[(sheet, key)]['label']}", f"={sheet}!{REG[(sheet, key)]['cell']}", float(val), tol))
# the static rosters: every listed participant must have include = 1 and valid GSR (sum of flags = 2 x n)
for ad in ("Ad_A", "Ad_B"):
    c0 = H3B[ad]["c"][0]
    last = H3B[ad]["last"]
    f = (f'=SUMPRODUCT(COUNTIFS({R_M("A")},{c0}5:{c0}{last},{R_M("B")},"{ad}",{R_M("L")},1,{R_M("E")},1))')
    ck.append((f"H3 roster {ad}: participants with include = 1 and valid GSR in {ad}", f"=SUMPRODUCT(COUNTIFS({R_M('A')},H3!{c0}5:{c0}{last},"
               f"{R_M('B')},\"{ad}\",{R_M('L')},1,{R_M('E')},1))", float(len(H3B[ad]["roster"])), 1e-9))
CK1 = 4
for k, (lab, f, rv, tol) in enumerate(ck):
    rr = CK1 + k
    ws.cell(row=rr, column=1, value=lab)
    ws.cell(row=rr, column=2, value=f).number_format = "0.000000"
    ws.cell(row=rr, column=3, value=rv).number_format = "0.000000"
    ws.cell(row=rr, column=4, value=f"=B{rr}-C{rr}").number_format = "0.000000"
    ws.cell(row=rr, column=5, value=tol).number_format = "0.000000"
    ws.cell(row=rr, column=6, value=f'=IF(ABS(D{rr})<=E{rr},"OK","CHECK")')
CK2 = CK1 + len(ck) - 1
ws.cell(row=CK2 + 2, column=1, value="Rows with status CHECK").font = bold
put(ws, "n_check", f"B{CK2 + 2}", f'=COUNTIF(F{CK1}:F{CK2},"CHECK")', "Rows with status CHECK", "0")
ws.cell(row=CK2 + 3, column=1, value=("Not reproduced in Excel: H1 (Fisher's exact test, the mixed logistic model), H2 (the linear mixed model), "
                                      "the Wilcoxon checks, the exact noncentral t intervals, and the Sensor Data consistency check. "
                                      "See prepare_output.txt and analysis_output.txt.")).font = note_font
widths(ws, {"A": 78, "B": 16, "C": 16, "D": 14, "E": 12, "F": 9})

wb.calculation = CalcProperties(fullCalcOnLoad=True)
for name in ("GSR_Import", "Survey", "DataQuality", "SessionLog"):
    wb[name].freeze_panes = "A2"
wb.save(OUT)
print("wrote", OUT)

# =============================================================================
# Verification 1: read back and compare every registered cell
# =============================================================================
wb2 = load_workbook(OUT)
problems = []
for (sheet, key), info in REG.items():
    v = wb2[sheet][info["cell"]].value
    v = str(v).replace("_xlfn.", "") if v is not None else None
    if v != info["formula"]:
        problems.append(f"{sheet}!{info['cell']} ({key}): workbook {v!r} registry {info['formula']!r}")
for k, (lab, f, rv, tol) in enumerate(ck):
    if wb2["Check"][f"B{CK1 + k}"].value != f:
        problems.append(f"Check!B{CK1 + k} differs")
if problems:
    print("VERIFICATION FAILED:\n" + "\n".join(problems[:40]))
    sys.exit(1)
n_formulas = sum(1 for w in wb2.worksheets for row in w.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
print(f"verified (openpyxl): {len(REG)} registered cells and {len(ck)} Check rows hold the registered formulas; formula cells in the workbook: {n_formulas}")

# =============================================================================
# Verification 2 and 3: LibreOffice recalculation
# =============================================================================
LO = {}
soffice = shutil.which("soffice") or shutil.which("libreoffice")
lo_note = "LibreOffice was not found, so the formula values were not recalculated here; the printed values are the R values."


def same_at(got, val, nd):
    got, val = float(got), float(val)
    return f"{got:.{nd}f}" == f"{val:.{nd}f}" or abs(got - val) <= 1e-9 * max(1.0, abs(val))


if soffice:
    tmp = tempfile.mkdtemp(prefix="ws7_lo_")
    shutil.copy(OUT, os.path.join(tmp, "in.xlsx"))
    subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", os.path.join(tmp, "out"),
                    os.path.join(tmp, "in.xlsx")], check=True, capture_output=True, timeout=600)
    wb3 = load_workbook(os.path.join(tmp, "out", "in.xlsx"), data_only=True)
    for (sheet, key), info in REG.items():
        LO[(sheet, key)] = wb3[sheet][info["cell"]].value
    bad = [(ck[k][0], wb3["Check"][f"B{CK1 + k}"].value, ck[k][2]) for k in range(len(ck)) if wb3["Check"][f"F{CK1 + k}"].value != "OK"]
    errs = [(w.title, c.coordinate, c.value) for w in wb3.worksheets for row in w.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("#") and c.value[:4] in ("#N/A", "#VAL", "#REF", "#NAM", "#DIV", "#NUM")]
    if bad or errs:
        print("LIBREOFFICE CHECK FAILED:", bad[:20], errs[:20])
        sys.exit(1)
    mism = [(s, k, LO[(s, k)], v) for (s, k), (v, nd) in EXPECT.items() if LO[(s, k)] is None or isinstance(LO[(s, k)], str)
            or not same_at(LO[(s, k)], v, nd)]
    if mism:
        print("VALUES DIFFER FROM R AT THE PRINTED DECIMALS:", mism)
        sys.exit(1)
    lo_ver = subprocess.run([soffice, "--version"], capture_output=True, text=True).stdout.strip()
    lo_note = (f"A copy was recalculated with {lo_ver} (headless): all {len(ck)} Check rows are OK, no cell shows an error value, and "
               f"the {len(EXPECT)} results printed here agree with R at the printed decimals.")
    print(f"verified (LibreOffice): {len(ck)} Check rows OK; no error values; {len(EXPECT)} printed results agree with R")
    shutil.rmtree(tmp, ignore_errors=True)
else:
    print(lo_note)


# =============================================================================
# The walkthrough, generated from REG
# =============================================================================
def cl(sheet, key):
    return REG[(sheet, key)]["cell"]


def fx(sheet, key):
    return REG[(sheet, key)]["formula"]


def val(sheet, key, nd=None):
    v, nd0 = EXPECT.get((sheet, key), (None, 3))
    nd = nd0 if nd is None else nd
    got = LO.get((sheet, key))
    if got is not None and not isinstance(got, str):
        v = got
    if v is None:
        return str(got) if got is not None else ""
    if isinstance(v, bool):
        return str(v).upper()
    return f"{float(v):.{nd}f}"


def lov(sheet, key):
    """Any LibreOffice value as text (TRUE or FALSE, text, or a number)."""
    v = LO.get((sheet, key))
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return "" if v is None else str(v)


def row_(label, sheet, key, nd=None):
    return f"| {label} | `{sheet}!{cl(sheet, key)}` | `{fx(sheet, key)}` | {val(sheet, key, nd)} |"


def ref_apa(key):
    with open(REFS, encoding="utf-8") as fh:
        refs = {r["key"]: r for r in json.load(fh)}
    r = refs[key]
    au = r["authors"]
    a = au[0] if len(au) == 1 else (", ".join(au[:-1]) + ", & " + au[-1])
    vol = f", *{r['volume']}*" if r.get("volume") else ""
    iss = f"({r['issue']})" if r.get("issue") else ""
    pg = f", {r['pages']}".replace("\u2013", "-") if r.get("pages") else ""
    doi = f" https://doi.org/{r['doi']}" if r.get("doi") else (f" {r['url']}" if r.get("url") else "")
    title = r["title"].replace("\u2014", ", ").replace("\u2013", "-")
    return f"- {a} ({r['year']}). {title}. *{r['container']}*{vol}{iss}{pg}.{doi}"


def p3(x):
    x = float(x)
    return "< .001" if x < 0.001 else "= " + f"{x:.3f}".lstrip("0")


gA, gB = T["H3_AD_A"], T["H3_AD_B"]
miss_b = [p for p in H3B["Ad_A"]["roster"] if p not in H3B["Ad_B"]["roster"]]
excl = [p for p in IDS if not long_ok[(p, "Baseline")][0]]
GHL = {k: GH[h] for k, h in zip(("key", "pid", "use", "lkey", "share", "valid"), G_HELP)}
SHL = {k: SH[h] for k, h in zip(("key", "pid", "rec_a", "rec_b", "att3", "att", "sum", "sum_raw"), S_HELP)}
DHL = {k: DH[h] for k, h in zip(("key", "pid", "use", "lkey", "len"), D_HELP)}
p05b = MROW[("P05", "Ad_B")]
_p05 = g_rows[G_ROWS[("P05", "Stimulus", "Ad_B", "Ad_B")] - 2]
p05_sig, p05_ppm = float(_p05[g_head.index("Signal Duration (ms)")]), _p05[g_head.index("Peaks Per Minute")]
_st = table(os.path.join(EX, "05_Survey", "SurveyStatistics.txt"), "\t", "Stimulus Name", quote=False)
recall_avg = [r for r in _st[2] if r[1] == "Recall_A"][0][3]
p07dq_len = lov("Keys", "kdq_C")
blkA, blkB = H3B["Ad_A"]["blk"], H3B["Ad_B"]["blk"]


def h3_table(ad):
    keys = [("m1", "Baseline mean"), ("s1", "Baseline SD"), ("m2", f"{ad} mean"), ("s2", f"{ad} SD"), ("n", "n"), ("md", "Mean difference"),
            ("sdd", "SD of the differences"), ("se", "SE of the mean difference"), ("t", "t"), ("df", "df"),
            ("p_tt", "p (T.TEST, paired)"), ("p_hand", "p from t (T.DIST.2T)"), ("tcrit", "Critical t"), ("ci_lo", "95% CI lower"),
            ("ci_hi", "95% CI upper"), ("dz", "Cohen's d_z"), ("se_dz", "SE of d_z"), ("dz_lo", "d_z CI lower (approx.)"),
            ("dz_hi", "d_z CI upper (approx.)"), ("npos", "Increases")]
    out = ["| Result | Cell | Formula | Value |", "|---|---|---|---|"]
    for k, lab in keys:
        key = f"{ad}_{k}"
        nd = 0 if k in ("n", "df", "npos") else (5 if k.startswith("p_") else 3)
        if (("H3", key)) not in EXPECT:
            v = LO.get(("H3", key))
            out.append(f"| {lab} | `H3!{cl('H3', key)}` | `{fx('H3', key)}` | {('' if v is None else f'{float(v):.{nd}f}')} |")
        else:
            out.append(row_(lab, "H3", key, nd))
    return "\n".join(out)


md = f"""# Excel walkthrough: from iMotions exports to one analysis ready table

**SIMULATED DATA.** Every number in this walkthrough comes from the SIMULATED export bundle that `make_simulated_data.py` generated (seed {SEED}). No real person was recorded, and the numbers must never be quoted as findings. The study is "SIMULATED Multimodal Ad Study": 24 participants (P01 to P24) at the lab desktop with the Smart Eye Aurora, the respondent camera with AFFDEX, and a Shimmer3 GSR+; slides Instructions, Baseline (60 s), Ad_A and Ad_B (30 s video ads, order AB or BA), and a survey.

This is the Excel route through the core of `prepare_data.R` and the H3 test of `analysis.R`. The finished workbook is `SIMULATED_data_prep.xlsx`. **Every cell address below is written as `Sheet!Cell` and was generated from the workbook itself:** `make_excel_workbook.py` writes the workbook and this file from one list of cells, reads the workbook back with openpyxl to confirm that each address holds the formula printed here, and recalculates a copy to confirm the results. {lo_note}

Companion files: `prepare_data.R` and `prepare_output.txt` (the R pipeline and its full output), `analysis.R` and `analysis_output.txt` (the tests), `README.md` (the bundle), `results_text.md` (model paragraphs), `figures/`.

---

## 0. What you are working with

| Sheet | What it holds | Rows |
|---|---|---|
| `GSR_Import` | `exports/04_GSR_Metrics/GSRPeakMetrics.csv` as the Text Import leaves it; helper columns `{GHL['key']}` to `{GHL['valid']}` (orange) | {len(g_rows)} |
| `SessionLog` | `exports/08_Session_Log/SIMULATED_AdStudy_SessionLog.csv`, one row per recording; helper column `{KEYCOL_LOG}` | {len(l_rows)} |
| `Survey` | `exports/05_Survey/MERGED_SURVEY_RESPONSE_MATRIX.txt` (tab separated); helper columns `{SHL['key']}` to `{SHL['sum_raw']}`, Cronbach's alpha in `{AV}2` and `{AV}3` | {len(s_rows)} |
| `DataQuality` | `exports/06_Data_Quality/SIMULATED Multimodal Ad Study_DataQuality.txt` (tab separated); helper columns `{DHL['key']}` to `{DHL['len']}` | {len(d_rows)} |
| `Keys` | the problem keys made visible | 5 |
| `Merged` | one row per participant and stimulus, joined with INDEX and MATCH | {M_LAST - 1} |
| `Pivot` | the PivotTable summary, built with AVERAGEIFS and COUNTIFS | 3 |
| `H3` | GSR Peaks Per Minute, each ad versus Baseline, with T.TEST | {len(H3B['Ad_A']['roster'])} and {len(H3B['Ad_B']['roster'])} |
| `Check` | every Excel result next to the R value | {len(ck)} |

Blue headers mark data as exported; orange headers mark columns you add; yellow cells are typed by you (the participant and stimulus grid).

**Rules written before collection (study protocol).** R1: use only the recordings that the session log marks Use recording = Yes. R2: exclude a participant from all analyses when Eye tracking_Quality is below {ET_MIN} percent or Affectiva AFFDEX_Quality is below {FACE_MIN} percent on Ad_A or Ad_B. R3: for GSR analyses only, a stimulus is missing when Signal Duration (ms) is below {int(GSR_MIN * 100)} percent of Duration (ms).

---

## 1. Text Import of one metrics file (GSRPeakMetrics.csv)

Row numbers in this walkthrough count the SIMULATED banner in line 1 of every file; a real export has no banner, so subtract 1. iMotions writes its metrics files with information ABOVE the table. In `GSRPeakMetrics.csv` line 1 is the SIMULATED banner (a real export has no banner), lines 2 to {g_h - 2} are the `#METADATA` block, line {g_h - 1} reads `#DATA`, and the column names are on **line {g_h}**. Double clicking the file makes Excel guess, and on a computer with other regional settings everything lands in column A (Opening Data Export Files in Excel article).

**Route A, Text Import Wizard (any Excel version).**
1. **File > Open > Browse**, set the file type to **All Files (\\*.\\*)**, select `GSRPeakMetrics.csv`, **Open**.
2. Step 1 of 3: choose **Delimited**, set **Start import at row: {g_h}**, File origin **65001 : Unicode (UTF-8)**, tick **My data has headers**. **You should see** `Study Name,Respondent Name,...` as the first preview line.
3. Step 2 of 3: tick **Comma** only. **You should see** the preview split into columns.
4. Step 3 of 3: select the `Respondent Name` column and set it to **Text** (so that nothing reformats a key); click **Advanced...** and set **Decimal separator** to `.` and **Thousands separator** to `,`; **Finish**.
5. **You should see** {len(g_head)} columns (`A` = Study Name to `O` = Average Peak Amplitude (microSiemens)) and {len(g_rows)} data rows, header in row 1. Cells that read `NA` are text: iMotions writes NA where a value does not exist (no peaks, or no signal). Leave them as they are.

**Route B, Get Data (Excel 2016 and later).** **Data > From Text/CSV**, select the file, **Transform Data**; in Power Query **Home > Remove Rows > Remove Top Rows** ({g_h - 1}), then **Home > Use First Row as Headers**, set `Respondent Name` to **Text**, **Close & Load**.

**The other three files** go on their own sheets the same way: `MERGED_SURVEY_RESPONSE_MATRIX.txt` with **Tab** as the delimiter, **Text qualifier: {{none}}**, start at row {s_h} (its headers contain commas and quotation marks, so a comma delimiter breaks them apart); `SIMULATED Multimodal Ad Study_DataQuality.txt` with **Tab**, start at row {d_h} (two `#Study ...` lines and the banner sit above the header); `SIMULATED_AdStudy_SessionLog.csv` with **Comma**, start at row {l_h}.

**Common mistakes.** (1) Starting the import at row 1: the `#METADATA` lines become data and every column turns to text. (2) Leaving the delimiter on Tab for a comma file: one column. (3) Saving the result back as .csv: formulas, helper columns, and extra sheets are lost; save as `.xlsx` and keep the raw export untouched in its folder. (4) Opening the student's copy `SurveyStatistics_student_copy.csv`: it was saved again on a laptop with German regional settings (semicolons, decimal commas). Never analyze a resaved copy; import the raw export.

---

## 2. The session log is the key table

Every iMotions export names people by the **respondent name** typed at the lab computer. The study is about **participants** (P01 to P24). The session log connects the two, one row per recording, and says which recording to use. P04 was recorded twice: `P04` (stopped 23 s into Baseline after the Aurora calibration failed; Use recording = No) and `P04_rerun` (Use recording = Yes).

Add one helper column to `SessionLog`, a clean version of the respondent name:

| Result | Cell | Formula | Value |
|---|---|---|---|
| Key of the first recording | `SessionLog!{cl('SessionLog', 'key2')}` | `{fx('SessionLog', 'key2')}` | {lov('SessionLog', 'key2')} |

Fill it down to row {LOG_LAST}. **You should see** `P04` and `P04_RERUN` as two different keys in rows {LOG_ROW['P04']} and {LOG_ROW['P04_rerun']}.

---

## 3. Fix the keys with TRIM and UPPER (Keys sheet)

Before fixing a key, **look at it**. The `Keys` sheet pulls five raw keys from the imported sheets by reference and tests each one:

| Raw key | Where | LEN | CODE of last character | EXACT in the session log | MATCH finds it | UPPER(TRIM()) | Participant | Use |
|---|---|---|---|---|---|---|---|---|
""" + "\n".join(
    f"| `\"{lov('Keys', k + '_B')}\"` | {w} | {lov('Keys', k + '_C')} | {lov('Keys', k + '_D')} | {lov('Keys', k + '_E')} | {lov('Keys', k + '_F')} | "
    f"{lov('Keys', k + '_G')} | {lov('Keys', k + '_H')} | {lov('Keys', k + '_I')} |" for w, _, k in key_rows) + f"""

The formulas in row {KROW['kdq']} (the other rows are the same with their own row number):

| Test | Cell | Formula |
|---|---|---|
| The raw key (a reference) | `Keys!{cl('Keys', 'kdq_B')}` | `{fx('Keys', 'kdq_B')}` |
| Number of characters | `Keys!{cl('Keys', 'kdq_C')}` | `{fx('Keys', 'kdq_C')}` |
| Code of the last character (32 = space) | `Keys!{cl('Keys', 'kdq_D')}` | `{fx('Keys', 'kdq_D')}` |
| Case sensitive match | `Keys!{cl('Keys', 'kdq_E')}` | `{fx('Keys', 'kdq_E')}` |
| What MATCH does | `Keys!{cl('Keys', 'kdq_F')}` | `{fx('Keys', 'kdq_F')}` |
| The standardized key | `Keys!{cl('Keys', 'kdq_G')}` | `{fx('Keys', 'kdq_G')}` |
| Participant from the session log | `Keys!{cl('Keys', 'kdq_H')}` | `{fx('Keys', 'kdq_H')}` |
| Use recording | `Keys!{cl('Keys', 'kdq_I')}` | `{fx('Keys', 'kdq_I')}` |

**Read the table.** `"P07 "` has {p07dq_len} characters and its last character is code 32, a space: invisible on screen, fatal for a lookup. `"p07"` fails the case sensitive EXACT test, yet MATCH finds it: **Excel lookups (MATCH, VLOOKUP, XLOOKUP) ignore upper and lower case, R joins do not.** A trailing space breaks both. `P04_rerun` is a correct name of the second recording, and the session log maps it to participant P04.

Now give every imported sheet a clean key and the participant (orange columns, filled down):

| Sheet | Result | Cell | Formula |
|---|---|---|---|
| `GSR_Import` | Key | `GSR_Import!{cl('GSR_Import', 'key')}` | `{fx('GSR_Import', 'key')}` |
| `GSR_Import` | Participant | `GSR_Import!{cl('GSR_Import', 'pid')}` | `{fx('GSR_Import', 'pid')}` |
| `GSR_Import` | Use recording | `GSR_Import!{cl('GSR_Import', 'use')}` | `{fx('GSR_Import', 'use')}` |
| `GSR_Import` | Lookup key | `GSR_Import!{cl('GSR_Import', 'lkey')}` | `{fx('GSR_Import', 'lkey')}` |
| `Survey` | Key | `Survey!{cl('Survey', 'key')}` | `{fx('Survey', 'key')}` |
| `Survey` | Participant | `Survey!{cl('Survey', 'pid')}` | `{fx('Survey', 'pid')}` |
| `DataQuality` | Key | `DataQuality!{cl('DataQuality', 'key')}` | `{fx('DataQuality', 'key')}` |
| `DataQuality` | Participant | `DataQuality!{cl('DataQuality', 'pid')}` | `{fx('DataQuality', 'pid')}` |
| `DataQuality` | Use recording | `DataQuality!{cl('DataQuality', 'use')}` | `{fx('DataQuality', 'use')}` |
| `DataQuality` | Lookup key | `DataQuality!{cl('DataQuality', 'lkey')}` | `{fx('DataQuality', 'lkey')}` |
| `DataQuality` | LEN of the raw name | `DataQuality!{cl('DataQuality', 'len')}` | `{fx('DataQuality', 'len')}` |

The **lookup key** joins the participant with the other key columns. In `GSR_Import` it is participant, Type, Label, and Parent Stimulus, because Label alone is not unique: the annotation `EndCard` appears once in Ad_A and once in Ad_B. Rows of the dropped recording get the key `DROP`, so no lookup can ever return them (rule R1). **You should see** `NOT FOUND` nowhere: filter the Participant columns for it. **You should see** `DROP` in exactly 2 rows of `GSR_Import` and 2 rows of `DataQuality` (the Instructions and Baseline rows of the aborted recording).

**Excel 365 and Excel 2021:** each INDEX and MATCH pair can be written as XLOOKUP, for example `=XLOOKUP({GHL['key']}2,{R_LOG_KEY},{R_LOG_PID},"NOT FOUND")` for the participant. The workbook uses INDEX and MATCH because every Excel version and LibreOffice understand it.

---

## 4. Join the survey and the data quality file to the GSR metrics (Merged sheet)

The `Merged` sheet is the analysis ready table: one row per participant and stimulus, {len(IDS)} participants times 3 stimuli = {M_LAST - 1} rows. Type the grid first (yellow columns `A` and `B`: P01 Baseline, P01 Ad_A, P01 Ad_B, P02 Baseline, ...), then fill these formulas down from row 2 to row {M_LAST}:

| Column | Result | Cell | Formula |
|---|---|---|---|
""" + "\n".join(f"| `{c}` | {mhead[ord(c) - 65]} | `Merged!{cl('Merged', c)}` | `{fx('Merged', c)}` |" for c in mf) + f"""

Values in the first ad row (P01, Ad_A, row {MROW[('P01', 'Ad_A')]}):

| Result | Cell | Formula | Value |
|---|---|---|---|
| GSR Peaks Per Minute (as exported) | `Merged!{cl('Merged', 'D_ad')}` | `{fx('Merged', 'D_ad')}` | {lov('Merged', 'D_ad')} |
| Liking of Ad_A (survey) | `Merged!{cl('Merged', 'M_ad')}` | `{fx('Merged', 'M_ad')}` | {lov('Merged', 'M_ad')} |
| Brand recalled, Ad_A (survey, scored in section 6) | `Merged!{cl('Merged', 'N_ad')}` | `{fx('Merged', 'N_ad')}` | {lov('Merged', 'N_ad')} |

**XLOOKUP version** (Excel 365) of column `D`: `=XLOOKUP(C2,{R_G_LKEY},{R_G_PPM})`; of column `M`: `=IF(B2="Ad_A",XLOOKUP(A2,{R_S_PID},{R_S_LIKE['Ad_A']}),IF(B2="Ad_B",XLOOKUP(A2,{R_S_PID},{R_S_LIKE['Ad_B']}),""))`.

**Check the join** (cells on the right of `Merged`):

| Check | Cell | Formula | Value |
|---|---|---|---|
{row_('Rows (participant x stimulus)', 'Merged', 'n_rows', 0)}
{row_('Lookups that failed (#N/A)', 'Merged', 'n_na', 0)}
{row_('Participants with include = 1', 'Merged', 'n_incl', 0)}
{row_('Stimulus rows with GSR not valid', 'Merged', 'n_gsr_invalid', 0)}

**You should see** 0 failed lookups. Try it the wrong way once, in a copy: build the data quality lookup key in `DataQuality!{DHL['lkey']}2` from the raw name in column `A` instead of the participant in column `{DHL['pid']}`. P07 (raw name `"P07 "`) and P04 (whose data quality rows are named `P04_rerun`) then return `#N/A`, and a filter on quality would quietly drop them.

**Rules R2 and R3 in the table.** Columns `J` and `K` take the LOWER of the two ads for each participant; column `L` applies rule R2. {", ".join(excl)} get include = 0. Column `E` applies rule R3 through the helper columns of `GSR_Import`:

| Result | Cell | Formula |
|---|---|---|
| Signal share | `GSR_Import!{cl('GSR_Import', 'share')}` | `{fx('GSR_Import', 'share')}` |
| GSR valid (R3) | `GSR_Import!{cl('GSR_Import', 'valid')}` | `{fx('GSR_Import', 'valid')}` |

P05's Shimmer lost its battery {p05_sig / 1000:.1f} s into Ad_B, so row {p05b} of `Merged` (P05, Ad_B) shows GSR valid = 0 and an empty cell in column `F`: the exported {p05_ppm} peaks per minute were computed on {p05_sig / 1000:.1f} s of signal. P05 stays in every other analysis.

---

## 5. A pivot table summary (Pivot sheet)

The iMotions article Creating Excel Pivot Tables gives the steps: click a cell of the data, **Insert > PivotTable > New Worksheet > OK**, drag `stimulus` to **Rows**, drag `GSR Peaks Per Minute (analysis)` to **Values**, right click the value, **Value Field Settings**, change **Sum** to **Average**, **OK**. Then drag `include (R2)` to **Filters** and keep only 1. **Common mistake:** the default is **Sum**, which adds the participants' values together; always switch to Average, and use Count to check n.

A PivotTable cannot be checked cell by cell, so the `Pivot` sheet builds the same table with AVERAGEIFS and COUNTIFS:

| Result | Cell | Formula | Value |
|---|---|---|---|
""" + "\n".join(row_(f"n (GSR), {st}", "Pivot", f"n_{st}", 0) + "\n" + row_(f"Mean GSR Peaks Per Minute, {st}", "Pivot", f"m_{st}", 3) for st in STIMS) + "\n" + \
    "\n".join(row_(f"Mean liking, {st}", "Pivot", f"ml_{st}", 3) + "\n" + row_(f"Brand recalled (proportion), {st}", "Pivot", f"rc_{st}", 3) for st in ("Ad_A", "Ad_B")) + f"""

**You should see** n = {lov('Pivot', 'n_Ad_B')} for Ad_B and {lov('Pivot', 'n_Ad_A')} for the other two stimuli: rule R3 removed P05's Ad_B value. Your PivotTable must show the same means.

---

## 6. Score the survey (Survey sheet)

The survey export holds the **label value** that was chosen, not a score. Two items need scoring before analysis:

| Result | Cell | Formula |
|---|---|---|
| Recall_A correct (Label1 = ALTAVO) | `Survey!{cl('Survey', 'rec_a')}` | `{fx('Survey', 'rec_a')}` |
| Recall_B correct (Label2 = BRENLOW) | `Survey!{cl('Survey', 'rec_b')}` | `{fx('Survey', 'rec_b')}` |
| ATT_3R reverse coded (8 minus the answer) | `Survey!{cl('Survey', 'att3')}` | `{fx('Survey', 'att3')}` |
| Attitude (mean of three items) | `Survey!{cl('Survey', 'att')}` | `{fx('Survey', 'att')}` |
| Sum of the three items (recoded) | `Survey!{cl('Survey', 'sum')}` | `{fx('Survey', 'sum')}` |
| Sum of the three items (not recoded) | `Survey!{cl('Survey', 'sum_raw')}` | `{fx('Survey', 'sum_raw')}` |

**Cronbach's alpha** = k / (k minus 1) x (1 minus the sum of the item variances / the variance of the sum), with k = 3 items:

| Result | Cell | Formula | Value |
|---|---|---|---|
{row_("Cronbach's alpha, ATT_3R recoded", 'Survey', 'alpha', 3)}
{row_("Cronbach's alpha, ATT_3R NOT recoded (the mistake)", 'Survey', 'alpha_wrong', 3)}

A negative alpha is the classic sign of a reverse worded item that was not recoded. **Common mistake:** averaging the Recall label values (SurveyStatistics.txt does exactly that: an "Average" of {recall_avg} for Recall_A means nothing, because the labels are brand names).

---

## 7. H3 with T.TEST (H3 sheet)

H3: GSR Peaks Per Minute are higher during each ad than during Baseline. Each participant gives a Baseline value and an ad value, so the test is **paired**. The rosters in `H3!A5:A{H3B['Ad_A']['last']}` (Ad_A, {len(H3B['Ad_A']['roster'])} participants) and `H3!G5:G{H3B['Ad_B']['last']}` (Ad_B, {len(H3B['Ad_B']['roster'])} participants; {", ".join(miss_b)} is missing because of rule R3) list the participants with include = 1 and valid GSR in both stimuli; the `Check` sheet confirms each list. The values come from `Merged` column `F`:

| Result | Cell | Formula |
|---|---|---|
| Baseline, first participant | `H3!{cl('H3', 'Ad_A_base')}` | `{fx('H3', 'Ad_A_base')}` |
| Ad_A, first participant | `H3!{cl('H3', 'Ad_A_ad')}` | `{fx('H3', 'Ad_A_ad')}` |
| Difference | `H3!{cl('H3', 'Ad_A_diff')}` | `{fx('H3', 'Ad_A_diff')}` |

**Ad_A versus Baseline**

{h3_table('Ad_A')}

**Ad_B versus Baseline**

{h3_table('Ad_B')}

T.TEST returns only the p value; its arguments are (range 1, range 2, tails = 2, type = 1 for paired). The hand computed t and T.DIST.2T must give the same p, which is a good check that the two ranges are aligned row by row. R also reports the Wilcoxon signed rank test (Ad_A p {p3(gA['np_p'])}, Ad_B p {p3(gB['np_p'])}), the exact noncentral t intervals for d_z (Lakens, 2013) (Ad_A [{float(gA['es_ci_low']):.2f}, {float(gA['es_ci_high']):.2f}], Ad_B [{float(gB['es_ci_low']):.2f}, {float(gB['es_ci_high']):.2f}]), and Holm adjusted p values (Ad_A p {p3(gA['p_holm'])}, Ad_B p {p3(gB['p_holm'])}); Excel has no function for these.

---

## 8. Check your numbers (Check sheet)

Every result printed above sits on the `Check` sheet next to the value that R produced, with the difference and a tolerance of half a unit in the last printed decimal. The cell `Check!{cl('Check', 'n_check')}` counts rows with status CHECK and must read 0. The last two rows confirm the H3 rosters: every listed participant has include = 1 and valid GSR in that ad.

H1 (Fisher's exact test and the mixed logistic model), H2 (the linear mixed model), the Wilcoxon checks, and the Sensor Data consistency check are R only; see `analysis_output.txt` and `prepare_output.txt`.

---

## 9. Common mistakes, collected

1. Double clicking an export instead of importing it: wrong delimiter or decimal mark, and `#` lines inside the data.
2. Trusting a fixed row number for the header: it moves between iMotions releases. Find the `#DATA` line (or the first column name) every time.
3. Joining on the raw respondent name: `"P07 "` with its trailing space never matches. TRIM and UPPER first, and LEN to see what you cannot see.
4. Forgetting that Excel lookups ignore case while R joins do not: the same file can join in Excel and fail in R.
5. Keeping both recordings of P04, or keeping the first one: the session log decides (rule R1).
6. Using Label alone as a key: `EndCard` exists in both ads.
7. Leaving the PivotTable on Sum, or forgetting the include filter.
8. Treating `NA` as 0: an AOI never fixated has no time to first fixation (NA stays NA), but its dwell time is 0.
9. Averaging the reverse worded item as it was answered (alpha goes negative) or averaging multiple choice label values.
10. Analyzing a copy that someone opened and saved again in Excel (`SurveyStatistics_student_copy.csv`). Spreadsheet programs can also silently turn codes into dates; keep the raw export read only and work in an `.xlsx` copy (Broman & Woo, 2018; Ziemann et al., 2016).

---

## Sources

{ref_apa('Broman2018')}
{ref_apa('Ziemann2016')}
{ref_apa('Lakens2013')}
- iMotions help center, Opening Data Export Files in Excel: https://help.imotions.com/lab/exports/additional-export-information/opening-data-export-files-in-excel
- iMotions help center, Creating Excel Pivot Tables: https://help.imotions.com/lab/exports/additional-export-information/creating-excel-pivot-tables
- iMotions help center, Export: R Processed GSR Files: https://help.imotions.com/lab/exports/r-processed-exports/export-r-processed-gsr-files
- iMotions help center, Export: Survey data: https://help.imotions.com/lab/exports/export-survey-data
- iMotions help center, Export: Data Quality Statistics: https://help.imotions.com/lab/exports/export-data-quality-statistics

## Items to verify on the lab computer

1. The header line of a real `GSR Peak Metrics` export in Lab 11.1: file name, `.csv` or `.txt`, separator, and the lines above the header (modeled here as a `#METADATA` block ending in `#DATA`).
2. The separator of `_DataQuality.txt` and of the survey files (modeled as tab), and the device names in the data quality headers.
3. The header pattern of `MERGED_SURVEY_RESPONSE_MATRIX.txt` for the lab's survey elements (modeled on the screenshot pattern `LABELVALUE_<slide>_<Question ID>,Question1 "<text>"`).
4. Which Excel version runs on the lab computers (XLOOKUP needs Excel 365 or 2021) and the Windows regional settings (expected: dot decimal, comma list separator).
"""
def esc_pipes(line):
    """Inside a Markdown table, a | in a code span must be written \\| or it splits the cell."""
    if not line.startswith("|"):
        return line
    return re.sub(r"`[^`]*`", lambda m: m.group(0).replace("|", "\\|"), line)


md = "\n".join(esc_pipes(ln) for ln in md.split("\n"))
with open(MD_OUT, "w", encoding="utf-8") as fh:
    fh.write(md)
print("wrote", MD_OUT)
