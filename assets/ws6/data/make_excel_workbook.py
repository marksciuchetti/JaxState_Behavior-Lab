#!/usr/bin/env python3
"""
make_excel_workbook.py
Builds SIMULATED_remote_analysis.xlsx from the SIMULATED CSV files with openpyxl, and writes
analysis_excel_walkthrough.md from the SAME cell registry, so that every cell reference in the
walkthrough is the address of the formula in the workbook.

Three verification steps run after the workbook is saved:
  1. openpyxl reads the workbook back and checks that every registered cell holds the registered
     formula (the addresses in the walkthrough are real).
  2. If LibreOffice is installed, a headless copy of the workbook is recalculated by LibreOffice
     (in a temporary folder; the delivered workbook is not changed) and every Check row must say OK.
  3. Every result printed in the walkthrough is the LibreOffice value of that cell, and it must
     agree with the value analysis.R wrote to SIMULATED_test_results.csv or SIMULATED_descriptives.csv
     at the number of decimals printed. The script stops with an error list if anything differs.

The workbook covers the Excel route of Workshop 6: the funnel, differential attrition by order, the
exclusion rule, calibration drift, H1, and H3. H2 (Friedman test, Kendall's W, Holm adjusted Wilcoxon
tests) and the logistic regression have no Excel function; their R results are quoted in the
walkthrough.

  README     what the workbook is, and that every number is SIMULATED
  Session    the session data as downloaded (columns A to J) with live helper columns K to O
             (Step number, Completed, Order, and the two numbers inside the "Eye tracking" text)
  Funnel     sessions entering and stopping at each funnel step, by order, by device; chi square tests
  Quality    the quality table (MODELED layout) as imported
  Screening  one row per completed respondent: the five quality values, the rule, the decision, drift
  FEA        the FEA Metrics Export (MODELED header) as imported
  Survey     the survey export (MODELED layout) as imported
  H1         Engagement, visual versus narrative video (paired), and the 3.0 degree sensitivity check
  H3         intention by order (Welch t test)
  Check      live Excel results next to the values analysis.R produced, with a status

All formulas are ordinary Excel formulas (no array formulas are needed).

Run: python3 make_excel_workbook.py   (after make_simulated_data.py and analysis.R)
"""

import csv
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "SIMULATED_remote_analysis.xlsx")
MD_OUT = os.path.join(HERE, "analysis_excel_walkthrough.md")
REFS = os.path.join(HERE, "..", "refs", "references.json")

BLUE, VERM, ORANGE, LIGHT = "0072B2", "D55E00", "E69F00", "F2F2F2"
hdr_font = Font(bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor=BLUE)
help_fill = PatternFill("solid", fgColor=ORANGE)
sub_fill = PatternFill("solid", fgColor=LIGHT)
param_fill = PatternFill("solid", fgColor="FFF2CC")
bold = Font(bold=True)
note_font = Font(italic=True, color="555555")
code_font = Font(name="Consolas", size=9, color="555555")
thin = Side(style="thin", color="BFBFBF")
box = Border(top=thin, bottom=thin, left=thin, right=thin)

ACC_MAX, ACC_STRICT, ETQ_MIN, FACE_MIN = 4.0, 3.0, 90.0, 80.0
NARR, VIS = "foodwaste_narrative", "foodwaste_visual"
STEPS = ["STEP 1 Testing upload", "STEP 2 Consent screen", "STEP 3 Screen recording setup",
         "STEP 4 Respondent camera setup", "STEP 5 Screen size", "STEP 6 Enter fullscreen",
         "STEP 7 Prepare positioning", "STEP 8 Positioning check", "STEP 9 Stimuli presentation*",
         "STEP 10 Study data upload"]


def read_csv(name):
    with open(os.path.join(HERE, name), encoding="utf-8") as fh:
        raw = fh.readlines()
    banner = raw[0].strip()
    rows = list(csv.reader([ln for ln in raw if not ln.startswith("#")]))
    return banner, rows[0], rows[1:]


def to_cell(v):
    """Numbers become numbers; NA stays the text NA (as iMotions writes it); an empty cell stays empty."""
    if v == "":
        return None
    if v == "NA":
        return v
    try:
        f = float(v)
        return int(f) if f.is_integer() and "." not in v else f
    except ValueError:
        return v


def style_header(ws, row=1, ncol=None, first_col=1, fill=None):
    ncol = ncol or ws.max_column
    for c in range(first_col, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill, cell.border = hdr_font, fill or hdr_fill, box
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def autowidth(ws, min_w=8, max_w=44):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            L = 12 if str(cell.value).startswith("=") else len(str(cell.value))
            widths[cell.column_letter] = max(widths.get(cell.column_letter, 0), L)
    for col, w in widths.items():
        ws.column_dimensions[col].width = max(min_w, min(max_w, w + 2))


def show_formula(cell, formula):
    """Write the formula text into a cell as TEXT (not evaluated), without the _xlfn. prefixes."""
    cell.value = formula.replace("_xlfn.", "")
    cell.data_type = "s"
    cell.font = code_font


# registry of documented cells: REG[(sheet, key)] = {"cell", "formula", "label"}
REG = {}
# expected values from analysis.R: EXPECT[(sheet, key)] = (value, decimals shown in the walkthrough)
EXPECT = {}


def reg(sheet, key, cell, formula, label=""):
    REG[(sheet, key)] = {"cell": cell, "formula": formula.replace("_xlfn.", ""), "label": label}


def put(ws, sheet, key, cell, formula, label="", fmt=None, expect=None, nd=3):
    ws[cell] = formula
    if fmt:
        ws[cell].number_format = fmt
    reg(sheet, key, cell, formula, label)
    if expect is not None:
        EXPECT[(sheet, key)] = (expect, nd)


# ---------------------------------------------------------------- read the files
banner, ses_h, ses_rows = read_csv("SIMULATED_session_data.csv")
SEED = re.search(r"seed (\d+)", banner).group(1)
_, q_h, q_rows = read_csv("SIMULATED_quality.csv")
_, fea_h, fea_rows = read_csv("SIMULATED_fea_metrics.csv")
_, sv_h, sv_rows = read_csv("SIMULATED_survey.csv")
_, th, trows = read_csv("SIMULATED_test_results.csv")
_, dh, drows = read_csv("SIMULATED_descriptives.csv")
T = {r[0]: dict(zip(th, r)) for r in trows}
DS = {(r[0], r[1], r[2]): dict(zip(dh, r)) for r in drows}
with open(os.path.join(HERE, "analysis_output.txt"), encoding="utf-8") as fh:
    out_txt = fh.read()

si, qi, fi, vi = ({h: i for i, h in enumerate(hh)} for hh in (ses_h, q_h, fea_h, sv_h))
N_SES, N_Q, N_FEA, N_SV = len(ses_rows), len(q_rows), len(fea_rows), len(sv_rows)
S2, Q2, F2, V2 = N_SES + 1, N_Q + 1, N_FEA + 1, N_SV + 1          # last data row of each sheet


def rng_(sheet, col, last):
    return f"{sheet}!${col}${2}:${col}${last}"


SCOL = {h: get_column_letter(i + 1) for i, h in enumerate(ses_h)}
SHELP = ["Step number", "Completed (1 = yes)", "Order", "Session accuracy (degrees)", "Session quality (%)"]
for k, h in enumerate(SHELP):
    SCOL[h] = get_column_letter(len(ses_h) + 1 + k)
R_ID, R_STATUS, R_DEV = rng_("Session", SCOL["Id"], S2), rng_("Session", SCOL["Status"], S2), rng_("Session", SCOL["Device type (modeled)"], S2)
R_STEP, R_DONE, R_ORD = (rng_("Session", SCOL[h], S2) for h in ("Step number", "Completed (1 = yes)", "Order"))
QCOL = {h: get_column_letter(i + 1) for i, h in enumerate(q_h)}
RQ_ID, RQ_SLIDE, RQ_ACC, RQ_ETQ, RQ_FACE = (rng_("Quality", QCOL[h], Q2) for h in
                                             ("Respondent Name", "Slide", "Eye tracking accuracy (degrees)", "Eye tracking quality (%)", "Face quality (%)"))
FCOL = {h: get_column_letter(i + 1) for i, h in enumerate(fea_h)}
RF_ID, RF_LAB, RF_EXP, RF_PCT = (rng_("FEA", FCOL[h], F2) for h in ("Respondent Name", "Interval label", "Expression", "Percent of time present"))
VCOL = {h: get_column_letter(i + 1) for i, h in enumerate(sv_h)}
RV_ID, RV_INT = (rng_("Survey", VCOL[h], V2) for h in ("Respondent Name", "Intention_reduce_food_waste_1to7"))

# ---------------------------------------------------------------- rosters (static lists; the Check sheet verifies them)
order_of = {r[si["Id"]]: ("Narrative first" if r[si["Block"]] == "Study flow 1" else "Visual first") for r in ses_rows}
done_ids = [r[si["Id"]] for r in ses_rows if r[si["Status"]] == "Processed"]
qv = {}
for r in q_rows:
    d = qv.setdefault(r[qi["Respondent Name"]], {})
    d[r[qi["Slide"]]] = r
def qnum(rid, slide, col):
    return float(qv[rid][slide][qi[col]])
scr = {}
for rid in done_ids:
    pre, post = qnum(rid, "Pre-study calibration", "Eye tracking accuracy (degrees)"), qnum(rid, "Post-study calibration", "Eye tracking accuracy (degrees)")
    etq = qnum(rid, "infographic", "Eye tracking quality (%)")
    fn, fvv = qnum(rid, NARR, "Face quality (%)"), qnum(rid, VIS, "Face quality (%)")
    fails = [pre > ACC_MAX, post > ACC_MAX, etq < ETQ_MIN, fn < FACE_MIN or fvv < FACE_MIN]
    scr[rid] = {"pre": pre, "post": post, "keep": not any(fails), "strict": (not any(fails)) and max(pre, post) <= ACC_STRICT,
                "fails": fails}
kept = [i for i in done_ids if scr[i]["keep"]]
strict = [i for i in kept if scr[i]["strict"]]
drop30 = [i for i in kept if not scr[i]["strict"]]
H1_ROSTER = strict + drop30
H3_ROSTER = [i for i in kept if order_of[i] == "Narrative first"] + [i for i in kept if order_of[i] == "Visual first"]
nK, nS, nD = len(kept), len(strict), len(done_ids)
nN = sum(1 for i in kept if order_of[i] == "Narrative first")
nV = nK - nN
m_kept = re.search(r"Kept respondents: (\d+) \(Narrative first (\d+), Visual first (\d+)\)", out_txt)
assert m_kept and (int(m_kept.group(1)), int(m_kept.group(2)), int(m_kept.group(3))) == (nK, nN, nV), "kept counts differ from analysis_output.txt"


def fnum(x):
    return float(x)


def dsv(measure, group, subgroup="All", field="n"):
    return float(DS[(measure, group, subgroup)][field])


wb = Workbook()

# ---------------------------------------------------------------- README
ws = wb.active
ws.title = "README"
readme = [
    "SIMULATED remote study workbook (teaching material): webcam eye tracking and webcam facial expression analysis",
    "",
    f"Every number in this workbook is SIMULATED. It was generated by make_simulated_data.py (seed {SEED}) in the style of the iMotions",
    "vendor demo (DEMO) Food Waste_webET_FEA_voice: every respondent saw two one minute food waste videos (foodwaste_narrative and",
    "foodwaste_visual) in random order, then an infographic with four tips (15 s), then a questionnaire. Remote Data Collection study built in",
    "iMotions Lab 11.1, run through my.imotions.com, downloaded, and processed in Lab (Post Processing > Affectiva, FEA thresholding at 50).",
    "No real person was recorded.",
    "",
    "Each order of the videos is its own study flow: Study flow 1 = narrative first, Study flow 2 = visual first (the Block column).",
    "",
    "Sheets:",
    f"  Session    Session data as downloaded (columns A to J, rows 2 to {S2}); helper columns {SCOL['Step number']} to {SCOL['Session quality (%)']} are live formulas.",
    "  Funnel     Sessions entering and stopping at each step of the respondent funnel; completion by order (chi square); by device.",
    "  Quality    Quality table (MODELED layout): accuracy at the two calibrations, eye tracking quality and face quality per slide.",
    "  Screening  The exclusion rule, one row per completed respondent; the thresholds are in yellow cells and can be changed.",
    "  FEA        FEA Metrics Export (MODELED header, as in Workshop 3): percent of time each expression was above its threshold.",
    "  Survey     Survey export (MODELED layout): intention (1 to 7), more persuasive video, tips recalled (0 to 4).",
    "  H1         Engagement, visual versus narrative video, paired t test; the 3.0 degree sensitivity check.",
    "  H3         Intention to reduce food waste by order, Welch t test (the randomization check).",
    "  Check      Live Excel results next to the values analysis.R produced, with an OK or CHECK status.",
    "",
    f"Exclusion rule, written before collection: exclude when the Pre-study OR Post-study calibration accuracy is above {ACC_MAX:.1f} degrees,",
    f"eye tracking quality on the infographic is below {ETQ_MIN:.0f} percent, or face quality is below {FACE_MIN:.0f} percent during either video.",
    f"Sensitivity: accuracy at or below {ACC_STRICT:.1f} degrees at both calibrations.",
    "Documented column names: Id, Started, Ended, Eye tracking, Est. face %, Block, Status (RDC respondent list) and the funnel step names.",
    "MODELED, verify on the lab computer: the session CSV layout and its three (modeled) columns, the quality table, the FEA header, the survey.",
    "",
    "Not in this workbook (no Excel function): H2 (Friedman test, Kendall's W, Holm adjusted Wilcoxon tests), Wilcoxon and Fisher tests,",
    "the logistic regression of the AOI hit on accuracy, exact noncentral t intervals. See analysis_output.txt.",
    "Companion files: analysis_excel_walkthrough.md (step by step), analysis.R, analysis_output.txt, results_text.md.",
]
for i, line in enumerate(readme, start=1):
    ws.cell(row=i, column=1, value=line)
ws["A1"].font = Font(bold=True, size=14, color=BLUE)
ws.column_dimensions["A"].width = 150

# ---------------------------------------------------------------- Session
ws = wb.create_sheet("Session")
ws.append(ses_h + SHELP)
style_header(ws, ncol=len(ses_h))
style_header(ws, ncol=len(ses_h) + len(SHELP), first_col=len(ses_h) + 1, fill=help_fill)
for k, r in enumerate(ses_rows):
    i = k + 2
    vals = []
    for h, v in zip(ses_h, r):
        if h in ("Started", "Ended") and v:
            vals.append(datetime.strptime(v, "%Y-%m-%d %H:%M:%S"))
        elif h == "Eye tracking":
            vals.append(v if v else None)
        else:
            vals.append(to_cell(v))
    ws.append(vals)
    for h in ("Started", "Ended"):
        ws[f"{SCOL[h]}{i}"].number_format = "yyyy-mm-dd hh:mm:ss"
    L, G, F_, D = SCOL["Last funnel step (modeled)"], SCOL["Status"], SCOL["Block"], SCOL["Eye tracking"]
    f_step = f'=IF({L}{i}="Completed",11,VALUE(MID({L}{i},6,FIND(" ",{L}{i},6)-6)))'
    f_done = f'=IF({G}{i}="Processed",1,0)'
    f_ord = f'=IF({F_}{i}="Study flow 1","Narrative first","Visual first")'
    f_acc = f'=IF({D}{i}="","",VALUE(LEFT({D}{i},FIND(",",{D}{i})-1)))'
    f_q = f'=IF({D}{i}="","",VALUE(SUBSTITUTE(MID({D}{i},FIND(",",{D}{i})+2,10),"%","")))'
    for h, f in zip(SHELP, (f_step, f_done, f_ord, f_acc, f_q)):
        ws[f"{SCOL[h]}{i}"] = f
    if i == 2:
        for key, h, f in (("step", "Step number", f_step), ("done", "Completed (1 = yes)", f_done), ("ord", "Order", f_ord),
                          ("acc", "Session accuracy (degrees)", f_acc), ("q", "Session quality (%)", f_q)):
            reg("Session", key, f"{SCOL[h]}2", f, h)
    if r[si["Id"]] == done_ids[0]:
        reg("Session", "acc_first_done", f"{SCOL['Session accuracy (degrees)']}{i}", f_acc, "first completed row")
        FIRST_DONE_ROW = i
autowidth(ws, max_w=30)
ws.freeze_panes = "B2"

# ---------------------------------------------------------------- Quality, FEA, Survey (as imported)
for name, hh, rows in (("Quality", q_h, q_rows), ("FEA", fea_h, fea_rows), ("Survey", sv_h, sv_rows)):
    ws = wb.create_sheet(name)
    ws.append(hh)
    style_header(ws, ncol=len(hh))
    for r in rows:
        ws.append([to_cell(v) for v in r])
    autowidth(ws, max_w=28)
    ws.freeze_panes = "C2"


class Block:
    """Writes label, value, formula text, and note columns for one block and records every cell in REG."""

    def __init__(self, ws, sheet, label_col, start_row, prefix):
        self.ws, self.sheet, self.lc, self.prefix = ws, sheet, label_col, prefix
        self.row = start_row
        self.ref = {}
        for k, name in enumerate(("Statistic", "Value", "Formula shown", "Note")):
            c = ws.cell(row=start_row - 1, column=label_col + k, value=name)
            c.font, c.fill = hdr_font, hdr_fill

    def title(self, text):
        c = self.ws.cell(row=self.row, column=self.lc, value=text)
        c.font, c.fill = bold, sub_fill
        self.row += 1

    def put(self, key, label, formula, note="", fmt="0.0000", expect=None, nd=3):
        f = formula.format(**{k: f"{get_column_letter(self.lc + 1)}{v}" for k, v in self.ref.items()})
        i = self.row
        self.ws.cell(row=i, column=self.lc, value=label)
        self.ws.cell(row=i, column=self.lc + 1, value=f).number_format = fmt
        show_formula(self.ws.cell(row=i, column=self.lc + 2), f)
        if note:
            self.ws.cell(row=i, column=self.lc + 3, value=note).font = note_font
        self.ref[key] = i
        reg(self.sheet, f"{self.prefix}_{key}", f"{get_column_letter(self.lc + 1)}{i}", f, label)
        if expect is not None:
            EXPECT[(self.sheet, f"{self.prefix}_{key}")] = (expect, nd)
        self.row += 1

    def blank(self):
        self.row += 1

    def addr(self, key):
        return f"{get_column_letter(self.lc + 1)}{self.ref[key]}"


def paired_block(blk, title, r1, r2, rd, nm1, nm2, t, unit=""):
    """Paired t test as a one sample t test on the differences (rd = r2 - r1); t = the R result row."""
    blk.title(title)
    blk.put("m1", f"{nm1} mean{unit}", f"=AVERAGE({r1})", "AVERAGE of the first column", expect=fnum(t["m1"]), nd=2)
    blk.put("s1", f"{nm1} SD", f"=_xlfn.STDEV.S({r1})", "STDEV.S = sample standard deviation", expect=fnum(t["sd1"]), nd=2)
    blk.put("m2", f"{nm2} mean{unit}", f"=AVERAGE({r2})", expect=fnum(t["m2"]), nd=2)
    blk.put("s2", f"{nm2} SD", f"=_xlfn.STDEV.S({r2})", expect=fnum(t["sd2"]), nd=2)
    blk.put("n", "n (respondents)", f"=COUNT({rd})", fmt="0", expect=int(float(t["n"])), nd=0)
    blk.put("md", f"Mean difference, {nm2} minus {nm1}", f"=AVERAGE({rd})", expect=fnum(t["mean_diff"]), nd=2)
    blk.put("sdd", "SD of the differences", f"=_xlfn.STDEV.S({rd})", "the SD that matters in a paired design", expect=fnum(t["sd_diff"]), nd=2)
    blk.put("se", "SE of the mean difference", "={sdd}/SQRT({n})", "SD of the differences / sqrt(n)")
    blk.put("t", "t", "={md}/{se}", "t = mean difference / SE", expect=fnum(t["statistic"]), nd=2)
    blk.put("df", "df", "={n}-1", "n minus 1", fmt="0", expect=int(float(t["df"])), nd=0)
    blk.put("p_tt", "p, two tailed (T.TEST type 1)", f"=_xlfn.T.TEST({r1},{r2},2,1)", "tails = 2, type = 1 means paired",
            fmt="0.000000E+00", expect=fnum(t["p"]), nd=-3)
    blk.put("p_hand", "p from the hand computed t", "=_xlfn.T.DIST.2T(ABS({t}),{df})", "must match the T.TEST cell", fmt="0.000000E+00",
            expect=fnum(t["p"]), nd=-3)
    blk.put("tcrit", "Critical t (95%)", "=_xlfn.T.INV.2T(0.05,{df})", "two tailed, alpha = .05")
    blk.put("ci_lo", "Mean difference 95% CI lower", "={md}-{tcrit}*{se}", expect=fnum(t["diff_ci_low"]), nd=2)
    blk.put("ci_hi", "Mean difference 95% CI upper", "={md}+{tcrit}*{se}", expect=fnum(t["diff_ci_high"]), nd=2)
    blk.put("dz", "Cohen's d_z", "={md}/{sdd}", "d_z = mean difference / SD of the differences", expect=fnum(t["es"]), nd=2)
    blk.put("se_dz", "SE of d_z (approximation)", "=SQRT(1/{n}+{dz}^2/(2*{n}))", "sqrt(1/n + d_z^2 / (2 n))")
    blk.put("dz_lo", "d_z 95% CI lower (approximation)", "={dz}-1.96*{se_dz}", "R also gives the exact noncentral t interval",
            expect=fnum(t["es_ci_low_approx"]), nd=2)
    blk.put("dz_hi", "d_z 95% CI upper (approximation)", "={dz}+1.96*{se_dz}", expect=fnum(t["es_ci_high_approx"]), nd=2)
    blk.put("r", "Correlation of the two columns", f"=CORREL({r1},{r2})", "how stable a respondent is from one column to the other")
    blk.put("npos", "Respondents with a positive difference", f'=COUNTIF({rd},">0")', fmt="0")
    blk.blank()


def welch_block(blk, title, ra, rb, t, nm_a, nm_b):
    blk.title(title)
    blk.put("m1", f"{nm_a} mean", f"=AVERAGE({ra})", expect=fnum(t["m1"]), nd=2)
    blk.put("s1", f"{nm_a} SD", f"=_xlfn.STDEV.S({ra})", expect=fnum(t["sd1"]), nd=2)
    blk.put("n1", f"{nm_a} n", f"=COUNT({ra})", fmt="0", expect=int(float(t["n1"])), nd=0)
    blk.put("m2", f"{nm_b} mean", f"=AVERAGE({rb})", expect=fnum(t["m2"]), nd=2)
    blk.put("s2", f"{nm_b} SD", f"=_xlfn.STDEV.S({rb})", expect=fnum(t["sd2"]), nd=2)
    blk.put("n2", f"{nm_b} n", f"=COUNT({rb})", fmt="0", expect=int(float(t["n2"])), nd=0)
    blk.put("diff", f"Difference, {nm_a} minus {nm_b}", "={m1}-{m2}", expect=fnum(t["mean_diff"]), nd=2)
    blk.put("se", "SE of the difference (Welch)", "=SQRT({s1}^2/{n1}+{s2}^2/{n2})", "each group keeps its own variance")
    blk.put("t", "Welch t", "={diff}/{se}", expect=fnum(t["statistic"]), nd=2)
    blk.put("df", "Welch df", "=({s1}^2/{n1}+{s2}^2/{n2})^2/(({s1}^2/{n1})^2/({n1}-1)+({s2}^2/{n2})^2/({n2}-1))",
            "Welch Satterthwaite formula; not a whole number", fmt="0.00", expect=fnum(t["df"]), nd=2)
    blk.put("p_tt", "p, two tailed (T.TEST type 3)", f"=_xlfn.T.TEST({ra},{rb},2,3)", "type 3 = two groups, unequal variances (Welch)",
            fmt="0.000000", expect=fnum(t["p"]), nd=3)
    blk.put("p_hand", "p from the hand computed t", "=_xlfn.T.DIST.2T(ABS({t}),{df})",
            "may round the df down to a whole number; report the T.TEST cell", fmt="0.000000")
    blk.put("tcrit", "Critical t (95%)", "=_xlfn.T.INV.2T(0.05,{df})", "may also round the df down")
    blk.put("ci_lo", "Difference 95% CI lower", "={diff}-{tcrit}*{se}", expect=fnum(t["diff_ci_low"]), nd=2)
    blk.put("ci_hi", "Difference 95% CI upper", "={diff}+{tcrit}*{se}", expect=fnum(t["diff_ci_high"]), nd=2)
    blk.put("sp", "Pooled SD (for the effect size)", "=SQRT((({n1}-1)*{s1}^2+({n2}-1)*{s2}^2)/({n1}+{n2}-2))", expect=fnum(t["sd_pooled"]), nd=2)
    blk.put("d", "Cohen's d", "={diff}/{sp}", "difference / pooled SD", expect=fnum(t["es"]), nd=2)
    blk.put("J", "Small sample correction J", "=1-3/(4*({n1}+{n2})-9)", "Hedges' approximation")
    blk.put("g", "Hedges' g", "={J}*{d}", expect=fnum(t["g"]), nd=2)
    blk.put("se_g", "SE of g (approximation)", "=SQRT(({n1}+{n2})/({n1}*{n2})+{g}^2/(2*({n1}+{n2})))")
    blk.put("g_lo", "g 95% CI lower (approximation)", "={g}-1.96*{se_g}", "R also gives the exact noncentral t interval",
            expect=fnum(t["g_ci_low_approx"]), nd=2)
    blk.put("g_hi", "g 95% CI upper (approximation)", "={g}+1.96*{se_g}", expect=fnum(t["g_ci_high_approx"]), nd=2)
    blk.blank()


# ---------------------------------------------------------------- Funnel
ws = wb.create_sheet("Funnel")
SH = "Funnel"
ws["A1"] = "The respondent funnel, rebuilt from the session data (one row per session that opened a link). SIMULATED data."
ws["A1"].font = Font(bold=True, size=12, color=BLUE)
ws["A2"] = ("Entered = sessions whose last step is this step or a later one (Step number on the Session sheet; 11 = completed). "
            "Stopped here = sessions that ended at this step. Step names are those of the iMotions Respondent funnel.")
ws["A2"].font = note_font
fcols = ["Step", "Step name", "Entered", "Stopped here", "Continued", "Percent of opened", "Percent of previous step"]
for c, h in enumerate(fcols, start=1):
    ws.cell(row=4, column=c, value=h)
style_header(ws, row=4, ncol=len(fcols))
FR1 = 5
N_OPEN_F = f"COUNTA({R_ID})"
for k, name in enumerate(STEPS):
    i = FR1 + k
    ws.cell(row=i, column=1, value=k + 1)
    ws.cell(row=i, column=2, value=name)
    put(ws, SH, f"ent{k + 1}", f"C{i}", f'=COUNTIF({R_STEP},">="&A{i})', f"Entered {name}", fmt="0",
        expect=dsv("Funnel: sessions entering the step", name), nd=0)
    put(ws, SH, f"stop{k + 1}", f"D{i}", f"=COUNTIF({R_STEP},A{i})", f"Stopped at {name}", fmt="0",
        expect=dsv("Funnel: sessions that stopped at the step", name), nd=0)
    put(ws, SH, f"cont{k + 1}", f"E{i}", f"=C{i}-D{i}", "Continued", fmt="0")
    put(ws, SH, f"pct{k + 1}", f"F{i}", f"=100*C{i}/{N_OPEN_F}", "Percent of opened", fmt="0.0",
        expect=dsv("Funnel: sessions entering the step", name, field="mean"), nd=1)
    put(ws, SH, f"prev{k + 1}", f"G{i}", f"=100*C{i}/{N_OPEN_F}" if k == 0 else f"=100*C{i}/C{i - 1}", "Percent of previous step", fmt="0.0")
FRL = FR1 + 9
i = FRL + 1
ws.cell(row=i, column=2, value="Completed (Status Processed)").font = bold
put(ws, SH, "done", f"C{i}", f'=COUNTIF({R_STATUS},"Processed")', "Completed (Status Processed)", fmt="0",
    expect=dsv("Funnel: completed (Status Processed)", "Completed"), nd=0)
put(ws, SH, "done_pct", f"F{i}", f"=100*C{i}/{N_OPEN_F}", "Completed, percent of opened", fmt="0.0",
    expect=dsv("Funnel: completed (Status Processed)", "Completed", field="mean"), nd=1)
put(ws, SH, "done_prev", f"G{i}", f"=100*C{i}/C{FRL}", "Completed, percent of STEP 10", fmt="0.0")
DONE_ROW = i
i += 1
ws.cell(row=i, column=2, value="Check: Continued after STEP 10 = Completed (must be 0)")
put(ws, SH, "chk_cont", f"C{i}", f"=E{FRL}-C{DONE_ROW}", "Continued after STEP 10 minus Completed", fmt="0", expect=0, nd=0)
i += 1
ws.cell(row=i, column=2, value="Status Abandoned")
put(ws, SH, "aband", f"C{i}", f'=COUNTIF({R_STATUS},"Abandoned")', "Status Abandoned", fmt="0")
i += 1
ws.cell(row=i, column=2, value="Status Errors")
put(ws, SH, "errors", f"C{i}", f'=COUNTIF({R_STATUS},"Errors")', "Status Errors", fmt="0")
i += 1
ws.cell(row=i, column=2, value="Sessions that opened a link")
put(ws, SH, "opened", f"C{i}", f"={N_OPEN_F}", "Sessions that opened a link", fmt="0", expect=N_SES, nd=0)

# by order and the chi square test of completion
OR1 = i + 3
ws.cell(row=OR1 - 2, column=1, value="Completion by order (differential attrition): all sessions").font = bold
ocols = ["Order", "Opened", "Reached STEP 9", "Completed", "Not completed", "Percent completed", "", "Expected: completed",
         "Expected: not completed"]
for c, h in enumerate(ocols, start=1):
    ws.cell(row=OR1 - 1, column=c, value=h if h else None)
style_header(ws, row=OR1 - 1, ncol=6)
style_header(ws, row=OR1 - 1, ncol=9, first_col=8)
for k, o in enumerate(("Narrative first", "Visual first")):
    i = OR1 + k
    ws.cell(row=i, column=1, value=o)
    put(ws, SH, f"op_{k}", f"B{i}", f'=COUNTIF({R_ORD},A{i})', f"Opened, {o}", fmt="0", expect=dsv("Opened the link, by order", o), nd=0)
    put(ws, SH, f"r9_{k}", f"C{i}", f'=COUNTIFS({R_ORD},A{i},{R_STEP},">=9")', f"Reached STEP 9, {o}", fmt="0")
    put(ws, SH, f"dn_{k}", f"D{i}", f"=COUNTIFS({R_ORD},A{i},{R_DONE},1)", f"Completed, {o}", fmt="0",
        expect=dsv("Completed, by order", o), nd=0)
    put(ws, SH, f"nd_{k}", f"E{i}", f"=B{i}-D{i}", f"Not completed, {o}", fmt="0")
    put(ws, SH, f"pc_{k}", f"F{i}", f"=100*D{i}/B{i}", f"Percent completed, {o}", fmt="0.0", expect=dsv("Completed, by order", o, field="mean"), nd=1)
TOT = OR1 + 2
ws.cell(row=TOT, column=1, value="Total").font = bold
for c in "BCDE":
    put(ws, SH, f"tot_{c}", f"{c}{TOT}", f"=SUM({c}{OR1}:{c}{OR1 + 1})", f"Total {c}", fmt="0")
for k in range(2):
    i = OR1 + k
    put(ws, SH, f"ex_d{k}", f"H{i}", f"=B{i}*$D${TOT}/$B${TOT}", "Expected completed = row total x column total / N", fmt="0.00")
    put(ws, SH, f"ex_n{k}", f"I{i}", f"=B{i}*$E${TOT}/$B${TOT}", "Expected not completed", fmt="0.00")
OBS, EXP = f"D{OR1}:E{OR1 + 1}", f"H{OR1}:I{OR1 + 1}"
blk = Block(ws, SH, 11, OR1, "att")
t = T["ATTRITION_ALL"]
blk.title("Chi square test of independence: order x completed (all sessions)")
blk.put("x2", "Chi square (Pearson, no continuity correction)", f"=SUMPRODUCT(({OBS}-{EXP})^2/{EXP})", "sum of (observed - expected)^2 / expected",
        expect=fnum(t["statistic"]), nd=2)
blk.put("df", "df", "=1", "(rows - 1) x (columns - 1)", fmt="0")
blk.put("p", "p (CHISQ.DIST.RT)", "=_xlfn.CHISQ.DIST.RT({x2},{df})", "right tail of the chi square distribution", fmt="0.000000", expect=fnum(t["p"]), nd=3)
blk.put("p_ct", "p (CHISQ.TEST on observed and expected)", f"=_xlfn.CHISQ.TEST({OBS},{EXP})", "must match the cell above", fmt="0.000000",
        expect=fnum(t["p"]), nd=3)
blk.put("phi", "phi (effect size)", f"=SQRT({{x2}}/B{TOT})", "sqrt(chi square / N); .1 small, .3 medium", expect=fnum(t["es"]), nd=2)
blk.put("diff", "Percent completed, Narrative first minus Visual first", f"=F{OR1}-F{OR1 + 1}", "percentage points", fmt="0.0")
blk.put("minexp", "Smallest expected count", f"=MIN({EXP})", "5 or more: the chi square approximation is fine", fmt="0.00")
ATT = blk

# among sessions that reached STEP 9
R9 = TOT + 13
ws.cell(row=R9 - 2, column=1, value="Completion by order among the sessions that reached STEP 9 (the videos started)").font = bold
ocols9 = ["Order", "Reached STEP 9", "Completed", "Not completed", "Percent completed", "", "", "Expected: completed", "Expected: not completed"]
for c, h in enumerate(ocols9, start=1):
    ws.cell(row=R9 - 1, column=c, value=h if h else None)
style_header(ws, row=R9 - 1, ncol=5)
style_header(ws, row=R9 - 1, ncol=9, first_col=8)
for k, o in enumerate(("Narrative first", "Visual first")):
    i = R9 + k
    ws.cell(row=i, column=1, value=o)
    put(ws, SH, f"s9_r_{k}", f"B{i}", f"=C{OR1 + k}", f"Reached STEP 9, {o}", fmt="0")
    put(ws, SH, f"s9_d_{k}", f"C{i}", f"=D{OR1 + k}", f"Completed, {o}", fmt="0")
    put(ws, SH, f"s9_n_{k}", f"D{i}", f"=B{i}-C{i}", f"Not completed after STEP 9, {o}", fmt="0")
    put(ws, SH, f"s9_p_{k}", f"E{i}", f"=100*C{i}/B{i}", f"Percent completed after STEP 9, {o}", fmt="0.0")
T9 = R9 + 2
ws.cell(row=T9, column=1, value="Total").font = bold
for c in "BCD":
    put(ws, SH, f"t9_{c}", f"{c}{T9}", f"=SUM({c}{R9}:{c}{R9 + 1})", f"Total {c} (STEP 9)", fmt="0")
for k in range(2):
    i = R9 + k
    put(ws, SH, f"e9_d{k}", f"H{i}", f"=B{i}*$C${T9}/$B${T9}", "Expected completed", fmt="0.00")
    put(ws, SH, f"e9_n{k}", f"I{i}", f"=B{i}*$D${T9}/$B${T9}", "Expected not completed", fmt="0.00")
OBS9, EXP9 = f"C{R9}:D{R9 + 1}", f"H{R9}:I{R9 + 1}"
blk = Block(ws, SH, 11, R9, "att9")
t9 = T["ATTRITION_STIMULI"]
blk.title("Chi square test: order x completed, sessions that reached STEP 9")
blk.put("x2", "Chi square (Pearson, no continuity correction)", f"=SUMPRODUCT(({OBS9}-{EXP9})^2/{EXP9})", expect=fnum(t9["statistic"]), nd=2)
blk.put("p", "p (CHISQ.DIST.RT)", "=_xlfn.CHISQ.DIST.RT({x2},1)", fmt="0.000000", expect=fnum(t9["p"]), nd=3)
blk.put("minexp", "Smallest expected count", f"=MIN({EXP9})", "below 5: report the Fisher exact test from R", fmt="0.00")
blk.put("diff", "Percent completed, Narrative first minus Visual first", f"=E{R9}-E{R9 + 1}", "percentage points", fmt="0.0")
ATT9 = blk

# by device
DV = T9 + 14
ws.cell(row=DV - 2, column=1, value="Completion by device (modeled column)").font = bold
for c, h in enumerate(["Device type (modeled)", "Opened", "Completed", "Percent completed"], start=1):
    ws.cell(row=DV - 1, column=c, value=h)
style_header(ws, row=DV - 1, ncol=4)
for k, dname in enumerate(("Laptop", "Desktop", "Phone", "Tablet")):
    i = DV + k
    ws.cell(row=i, column=1, value=dname)
    put(ws, SH, f"dv_o_{dname}", f"B{i}", f"=COUNTIF({R_DEV},A{i})", f"Opened, {dname}", fmt="0")
    put(ws, SH, f"dv_c_{dname}", f"C{i}", f"=COUNTIFS({R_DEV},A{i},{R_DONE},1)", f"Completed, {dname}", fmt="0")
    put(ws, SH, f"dv_p_{dname}", f"D{i}", f'=IF(B{i}=0,"",100*C{i}/B{i})', f"Percent completed, {dname}", fmt="0.0")
# chart of the funnel
ch = BarChart()
ch.type = "bar"
ch.style = 10
ch.title = "SIMULATED: sessions entering each funnel step"
ch.y_axis.title = "Sessions"
ch.x_axis.scaling.orientation = "maxMin"
ch.add_data(Reference(ws, min_col=3, min_row=FR1 - 1, max_row=DONE_ROW), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=FR1, max_row=DONE_ROW))
ch.legend = None
ch.height, ch.width = 9, 18
ws.add_chart(ch, f"P{FR1 - 1}")
for col, wdt in (("A", 16), ("B", 34), ("C", 12), ("D", 12), ("E", 12), ("F", 12), ("G", 12), ("H", 13), ("I", 13), ("J", 3),
                 ("K", 52), ("L", 12), ("M", 60), ("N", 48)):
    ws.column_dimensions[col].width = wdt

# ---------------------------------------------------------------- Screening
ws = wb.create_sheet("Screening")
SH = "Screening"
ws["A1"] = "The exclusion rule, one row per completed respondent (Status Processed). SIMULATED data. Yellow cells hold the thresholds."
ws["A1"].font = Font(bold=True, size=12, color=BLUE)
params = [("A2", "Accuracy limit (degrees)", "B2", ACC_MAX), ("C2", "Sensitivity limit (degrees)", "D2", ACC_STRICT),
          ("E2", "Eye tracking quality minimum (%)", "F2", ETQ_MIN), ("G2", "Face quality minimum (%)", "H2", FACE_MIN)]
for lc, lab, vc, v in params:
    ws[lc] = lab
    ws[lc].font = bold
    ws[vc] = v
    ws[vc].fill = param_fill
    ws[vc].border = box
scols = ["Respondent Name", "Order", "Pre-study accuracy", "Post-study accuracy", "Eye tracking quality, infographic",
         "Face quality, narrative", "Face quality, visual", "Fail: Pre-study accuracy", "Fail: Post-study accuracy",
         "Fail: eye tracking quality", "Fail: face quality", "Decision", "First reason", "Largest accuracy", "Decision at 3.0",
         "Drift (Post minus Pre)"]
for c, h in enumerate(scols, start=1):
    ws.cell(row=4, column=c, value=h)
style_header(ws, row=4, ncol=len(scols))
ws.row_dimensions[4].height = 46
P1 = 5
PL = P1 + nD - 1


def qsum(col_rng, slide, r):
    return f'=SUMIFS({col_rng},{RQ_ID},$A{r},{RQ_SLIDE},"{slide}")'


for k, rid in enumerate(done_ids):
    r = P1 + k
    ws.cell(row=r, column=1, value=rid)
    fs = {
        "ord": f"=INDEX({R_ORD},MATCH($A{r},{R_ID},0))",
        "pre": qsum(RQ_ACC, "Pre-study calibration", r), "post": qsum(RQ_ACC, "Post-study calibration", r),
        "etq": qsum(RQ_ETQ, "infographic", r), "fn": qsum(RQ_FACE, NARR, r), "fv": qsum(RQ_FACE, VIS, r),
        "f_pre": f"=IF(C{r}>$B$2,1,0)", "f_post": f"=IF(D{r}>$B$2,1,0)", "f_etq": f"=IF(E{r}<$F$2,1,0)",
        "f_face": f"=IF(OR(F{r}<$H$2,G{r}<$H$2),1,0)", "dec": f'=IF(SUM(H{r}:K{r})>0,"Exclude","Keep")',
        "first": f'=IF(H{r}=1,"Pre-study accuracy",IF(I{r}=1,"Post-study accuracy",IF(J{r}=1,"Eye tracking quality",IF(K{r}=1,"Face quality",""))))',
        "big": f"=MAX(C{r}:D{r})", "dec30": f'=IF(AND(L{r}="Keep",N{r}<=$D$2),"Keep","Exclude")', "drift": f"=D{r}-C{r}"}
    for c, key in enumerate(fs, start=2):
        ws.cell(row=r, column=c, value=fs[key])
        if k == 0:
            reg(SH, key, f"{get_column_letter(c)}{r}", fs[key], scols[c - 1])
RS_ORD, RS_DEC, RS_FIRST, RS_DEC30 = (f"$B${P1}:$B${PL}", f"$L${P1}:$L${PL}", f"$M${P1}:$M${PL}", f"$O${P1}:$O${PL}")
blk = Block(ws, SH, 18, P1 + 1, "cnt")
blk.title("Counts")
blk.put("n", "Completed respondents screened", f"=COUNTA(A{P1}:A{PL})", fmt="0", expect=nD, nd=0)
blk.put("excl", "Excluded", f'=COUNTIF({RS_DEC},"Exclude")', fmt="0", expect=nD - nK, nd=0)
blk.put("kept", "Kept", f'=COUNTIF({RS_DEC},"Keep")', fmt="0", expect=dsv("Funnel: kept after the exclusion rule", "Kept (4.0 degrees)"), nd=0)
reasons = {"pre": "Pre-study calibration accuracy above 4.0 degrees", "post": "Post-study calibration accuracy above 4.0 degrees",
           "etq": "Eye tracking quality on the infographic below 90 percent", "face": "Face quality below 80 percent during a video"}
labels = {"pre": "Pre-study accuracy", "post": "Post-study accuracy", "etq": "Eye tracking quality", "face": "Face quality"}
cols_fail = {"pre": "H", "post": "I", "etq": "J", "face": "K"}
for key in ("pre", "post", "etq", "face"):
    blk.put(f"first_{key}", f"First reason: {labels[key]}", f'=COUNTIF({RS_FIRST},"{labels[key]}")', fmt="0",
            expect=dsv("Exclusions: first reason", reasons[key]), nd=0)
for key in ("pre", "post", "etq", "face"):
    c = cols_fail[key]
    blk.put(f"any_{key}", f"Failed (any): {labels[key]}", f"=SUM({c}{P1}:{c}{PL})", fmt="0",
            expect=dsv("Exclusions: failed the criterion (any)", reasons[key]), nd=0)
blk.put("multi", "Failed more than one criterion", f"=SUMPRODUCT(--((H{P1}:H{PL}+I{P1}:I{PL}+J{P1}:J{PL}+K{P1}:K{PL})>1))", fmt="0")
blk.put("postonly", "Failed only at the Post-study calibration", f"=COUNTIFS(H{P1}:H{PL},0,I{P1}:I{PL},1)", "fine at the start, drifted", fmt="0")
blk.put("kept_n", "Kept, Narrative first", f'=COUNTIFS({RS_ORD},"Narrative first",{RS_DEC},"Keep")', fmt="0",
        expect=dsv("Kept, by order", "Narrative first"), nd=0)
blk.put("kept_v", "Kept, Visual first", f'=COUNTIFS({RS_ORD},"Visual first",{RS_DEC},"Keep")', fmt="0",
        expect=dsv("Kept, by order", "Visual first"), nd=0)
blk.put("kept30", "Kept at 3.0 degrees", f'=COUNTIF({RS_DEC30},"Keep")', fmt="0",
        expect=dsv("Funnel: kept after the exclusion rule", "Kept (3.0 degrees, sensitivity)"), nd=0)
blk.put("more30", "Further removed at 3.0 degrees", "={kept}-{kept30}", fmt="0", expect=nK - nS, nd=0)
drb = Block(ws, SH, 18, blk.row + 3, "dr")
paired_block(drb, "Calibration drift: Pre-study versus Post-study accuracy (all completed respondents)", f"C{P1}:C{PL}", f"D{P1}:D{PL}",
             f"P{P1}:P{PL}", "Pre-study", "Post-study", T["DRIFT_ALL"], " (degrees)")
drb.put("worse", "Respondents with a worse Post-study accuracy", f'=COUNTIF(P{P1}:P{PL},">0")', fmt="0")
SCR = drb
for col, wdt in (("A", 13), ("B", 15), ("C", 10), ("D", 10), ("E", 11), ("F", 11), ("G", 11), ("H", 9), ("I", 9), ("J", 9), ("K", 9),
                 ("L", 10), ("M", 20), ("N", 10), ("O", 10), ("P", 10), ("Q", 3), ("R", 50), ("S", 13), ("T", 60), ("U", 44)):
    ws.column_dimensions[col].width = wdt
ws.freeze_panes = "B5"

# ---------------------------------------------------------------- H1
ws = wb.create_sheet("H1")
SH = "H1"
ws["A1"] = "H1. Percent of time Engaged (AFFDEX Engagement above 50), foodwaste_visual versus foodwaste_narrative, kept respondents. SIMULATED data."
ws["A1"].font = Font(bold=True, size=12, color=BLUE)
H1A, H1L, H1S = 5, 5 + nK - 1, 5 + nS - 1
ws["A2"] = (f"One row per kept respondent. Rows {H1A} to {H1S} are also kept under the {ACC_STRICT:.1f} degree sensitivity rule; rows {H1S + 1} to "
            f"{H1L} are removed by it. Columns D and E look the value up in the FEA sheet (respondent, video, expression).")
ws["A2"].font = note_font
h1_cols = ["Respondent Name", "Order", "Decision at 3.0", "Engagement, narrative (%)", "Engagement, visual (%)", "Difference (visual minus narrative)"]
for c, h in enumerate(h1_cols, start=1):
    ws.cell(row=4, column=c, value=h)
style_header(ws, row=4, ncol=len(h1_cols))
ws.row_dimensions[4].height = 46


def fea_lookup(r, video, expr="Engagement"):
    return f'=SUMIFS({RF_PCT},{RF_ID},$A{r},{RF_LAB},"{video}",{RF_EXP},"{expr}")'


for k, rid in enumerate(H1_ROSTER):
    r = H1A + k
    ws.cell(row=r, column=1, value=rid)
    fs = {"ord": f"=INDEX({R_ORD},MATCH($A{r},{R_ID},0))", "dec30": f"=INDEX(Screening!$O${P1}:$O${PL},MATCH($A{r},Screening!$A${P1}:$A${PL},0))",
          "narr": fea_lookup(r, NARR), "vis": fea_lookup(r, VIS), "diff": f"=E{r}-D{r}"}
    for c, key in enumerate(fs, start=2):
        ws.cell(row=r, column=c, value=fs[key]).number_format = "0.00" if key in ("narr", "vis", "diff") else "General"
        if k == 0:
            reg(SH, key, f"{get_column_letter(c)}{r}", fs[key], h1_cols[c - 1])
blk = Block(ws, SH, 8, H1A + 1, "h1")
paired_block(blk, f"H1: Engagement, visual minus narrative (kept, n = {nK})", f"D{H1A}:D{H1L}", f"E{H1A}:E{H1L}", f"F{H1A}:F{H1L}",
             "Narrative", "Visual", T["H1"], " (% of time)")
H1_MAIN = blk
blk.put("z_n", "Respondents at exactly 0, narrative", f"=COUNTIF(D{H1A}:D{H1L},0)", "never above the threshold", fmt="0")
blk.put("z_v", "Respondents at exactly 0, visual", f"=COUNTIF(E{H1A}:E{H1L},0)", fmt="0")
blk.put("med_n", "Median, narrative", f"=MEDIAN(D{H1A}:D{H1L})", "mean above median = right skew", fmt="0.00",
        expect=dsv("Percent of time present: Engagement", NARR, field="median"), nd=2)
blk.put("med_v", "Median, visual", f"=MEDIAN(E{H1A}:E{H1L})", fmt="0.00", expect=dsv("Percent of time present: Engagement", VIS, field="median"), nd=2)
blk.put("med_d", "Median difference", f"=MEDIAN(F{H1A}:F{H1L})", fmt="0.00")
s30 = Block(ws, SH, 8, blk.row + 3, "s30")
paired_block(s30, f"Sensitivity: H1 at {ACC_STRICT:.1f} degrees (rows {H1A} to {H1S}, n = {nS})", f"D{H1A}:D{H1S}", f"E{H1A}:E{H1S}",
             f"F{H1A}:F{H1S}", "Narrative", "Visual", T["S30_H1"], " (% of time)")
H1_S30_REF = s30.ref
blk = Block(ws, SH, 8, s30.row + 3, "chk")
blk.title("Roster checks")
blk.put("chk_n", "Kept respondents on the Screening sheet", f'=COUNTIF(Screening!$L${P1}:$L${PL},"Keep")', "must equal n of the H1 block", fmt="0",
        expect=nK, nd=0)
blk.put("chk_s30", "Rows in the sensitivity block marked Exclude at 3.0 (must be 0)", f'=COUNTIF(C{H1A}:C{H1S},"Exclude")', fmt="0", expect=0, nd=0)
blk.put("chk_rest", "Rows below the sensitivity block marked Exclude at 3.0", f'=COUNTIF(C{H1S + 1}:C{H1L},"Exclude")', fmt="0", expect=nK - nS, nd=0)
H1_BLK = blk
for col, wdt in (("A", 13), ("B", 15), ("C", 10), ("D", 12), ("E", 12), ("F", 13), ("G", 3), ("H", 52), ("I", 14), ("J", 62), ("K", 50)):
    ws.column_dimensions[col].width = wdt
ws.freeze_panes = "B5"

# ---------------------------------------------------------------- H3
ws = wb.create_sheet("H3")
SH = "H3"
ws["A1"] = "H3. Intention to reduce food waste (1 to 7) by order, kept respondents: the randomization check. SIMULATED data."
ws["A1"].font = Font(bold=True, size=12, color=BLUE)
H3A, H3N, H3L = 5, 5 + nN - 1, 5 + nK - 1
ws["A2"] = (f"Rows {H3A} to {H3N}: Narrative first (Study flow 1); rows {H3N + 1} to {H3L}: Visual first (Study flow 2). "
            "Column C looks the answer up in the Survey sheet.")
ws["A2"].font = note_font
h3_cols = ["Respondent Name", "Order", "Intention (1 to 7)"]
for c, h in enumerate(h3_cols, start=1):
    ws.cell(row=4, column=c, value=h)
style_header(ws, row=4, ncol=len(h3_cols))
for k, rid in enumerate(H3_ROSTER):
    r = H3A + k
    ws.cell(row=r, column=1, value=rid)
    fs = {"ord": f"=INDEX({R_ORD},MATCH($A{r},{R_ID},0))", "int": f"=INDEX({RV_INT},MATCH($A{r},{RV_ID},0))"}
    for c, key in enumerate(fs, start=2):
        ws.cell(row=r, column=c, value=fs[key])
        if k == 0:
            reg(SH, key, f"{get_column_letter(c)}{r}", fs[key], h3_cols[c - 1])
blk = Block(ws, SH, 5, H3A + 1, "h3")
welch_block(blk, f"H3: intention, Narrative first versus Visual first (n = {nN} and {nV})", f"C{H3A}:C{H3N}", f"C{H3N + 1}:C{H3L}", T["H3"],
            "Narrative first", "Visual first")
blk.title("Roster checks")
blk.put("chk_a", "Visual first rows inside the Narrative first range (must be 0)", f'=COUNTIF(B{H3A}:B{H3N},"Visual first")', fmt="0", expect=0, nd=0)
blk.put("chk_b", "Narrative first rows inside the Visual first range (must be 0)", f'=COUNTIF(B{H3N + 1}:B{H3L},"Narrative first")', fmt="0",
        expect=0, nd=0)
H3_BLK = blk
for col, wdt in (("A", 13), ("B", 15), ("C", 11), ("D", 3), ("E", 52), ("F", 14), ("G", 70), ("H", 56)):
    ws.column_dimensions[col].width = wdt
ws.freeze_panes = "B5"

# ---------------------------------------------------------------- Check
ws = wb.create_sheet("Check")
ws["A1"] = "Live Excel results next to the values analysis.R produced (SIMULATED data). Status is OK when they agree within the tolerance."
ws["A1"].font = Font(bold=True, size=12, color=BLUE)
ck_cols = ["Quantity", "Excel (live formula)", "analysis.R", "Difference", "Tolerance", "Status"]
for c, h in enumerate(ck_cols, start=1):
    ws.cell(row=3, column=c, value=h)
style_header(ws, row=3, ncol=len(ck_cols))
ck_rows = []


def add_check(label, ref, r_value, tol, note=""):
    ck_rows.append((label, f"={ref}", float(r_value), tol, note))


COVERED = set()


def blk_checks(tid, sheet, ref_dict, lc, kind, prefix):
    t = T[tid]
    COVERED.update((sheet, f"{prefix}_{k}") for k in ref_dict)
    col = get_column_letter(lc + 1)
    if kind == "paired":
        items = (("m1", "m1", 1e-6), ("m2", "m2", 1e-6), ("n", "n", 1e-9), ("md", "mean_diff", 1e-6), ("sdd", "sd_diff", 1e-6),
                 ("t", "statistic", 1e-4), ("df", "df", 1e-9), ("p_tt", "p", 1e-9), ("p_hand", "p", 1e-9),
                 ("ci_lo", "diff_ci_low", 1e-6), ("ci_hi", "diff_ci_high", 1e-6), ("dz", "es", 1e-5),
                 ("dz_lo", "es_ci_low_approx", 1e-5), ("dz_hi", "es_ci_high_approx", 1e-5))
    else:
        items = (("m1", "m1", 1e-6), ("m2", "m2", 1e-6), ("n1", "n1", 1e-9), ("n2", "n2", 1e-9), ("diff", "mean_diff", 1e-6),
                 ("t", "statistic", 1e-4), ("df", "df", 1e-4), ("p_tt", "p", 1e-6), ("ci_lo", "diff_ci_low", 5e-4), ("ci_hi", "diff_ci_high", 5e-4),
                 ("sp", "sd_pooled", 1e-6), ("d", "es", 1e-5), ("g", "g", 1e-5), ("g_lo", "g_ci_low_approx", 1e-5), ("g_hi", "g_ci_high_approx", 1e-5))
    for key, rkey, tol in items:
        note = "Welch CI: the spreadsheet critical t may use a whole number df" if key in ("ci_lo", "ci_hi") and kind == "welch" else ""
        add_check(f"{tid}: {rkey} ({key})", f"{sheet}!{col}{ref_dict[key]}", t[rkey], tol, note)
    if kind == "paired":
        add_check(f"{tid}: d_z exact CI lower (noncentral t in R) versus the approximation", f"{sheet}!{col}{ref_dict['dz_lo']}", t["es_ci_low"], 0.03,
                  "method differs; looser tolerance")
        add_check(f"{tid}: d_z exact CI upper (noncentral t in R) versus the approximation", f"{sheet}!{col}{ref_dict['dz_hi']}", t["es_ci_high"], 0.03,
                  "method differs; looser tolerance")


blk_checks("H1", "H1", H1_MAIN.ref, 8, "paired", "h1")
blk_checks("S30_H1", "H1", H1_S30_REF, 8, "paired", "s30")
blk_checks("H3", "H3", H3_BLK.ref, 5, "welch", "h3")
blk_checks("DRIFT_ALL", "Screening", SCR.ref, 18, "paired", "dr")
for key, rkey in (("x2", "statistic"), ("p", "p"), ("p_ct", "p"), ("phi", "es")):
    add_check(f"ATTRITION_ALL: {rkey} ({key})", f"Funnel!{ATT.addr(key)}", T["ATTRITION_ALL"][rkey], 1e-6)
    COVERED.add(("Funnel", f"att_{key}"))
for key, rkey in (("x2", "statistic"), ("p", "p")):
    add_check(f"ATTRITION_STIMULI: {rkey} ({key})", f"Funnel!{ATT9.addr(key)}", T["ATTRITION_STIMULI"][rkey], 1e-6)
    COVERED.add(("Funnel", f"att9_{key}"))
for (sheet, key), (val, nd) in EXPECT.items():
    if (sheet, key) not in COVERED and sheet != "Check":
        tol = 1e-9 if nd == 0 else (0.5 * 10 ** (-nd) if nd > 0 else 1e-9)
        add_check(f"{sheet} {REG[(sheet, key)]['label']}", f"{sheet}!{REG[(sheet, key)]['cell']}", val, tol,
                  "from SIMULATED_descriptives.csv or counted by this script")
CK1 = 4
for k, (lab, f, rv, tol, note) in enumerate(ck_rows):
    r = CK1 + k
    ws.cell(row=r, column=1, value=lab)
    ws.cell(row=r, column=2, value=f).number_format = "0.000000"
    ws.cell(row=r, column=3, value=rv).number_format = "0.000000"
    ws.cell(row=r, column=4, value=f"=B{r}-C{r}").number_format = "0.000000"
    ws.cell(row=r, column=5, value=tol).number_format = "0.000000"
    ws.cell(row=r, column=6, value=f'=IF(ABS(D{r})<=E{r},"OK","CHECK")')
    if note:
        ws.cell(row=r, column=7, value=note).font = note_font
CK2 = CK1 + len(ck_rows) - 1
ws.cell(row=CK2 + 2, column=1, value="Rows with status CHECK").font = bold
ws.cell(row=CK2 + 2, column=2, value=f'=COUNTIF(F{CK1}:F{CK2},"CHECK")')
reg("Check", "n_check", f"B{CK2 + 2}", f'=COUNTIF(F{CK1}:F{CK2},"CHECK")', "Rows with status CHECK")
EXPECT[("Check", "n_check")] = (0, 0)
ws.cell(row=CK2 + 3, column=1, value=("Not reproducible in Excel: H2 (Friedman test, Kendall's W, Holm adjusted Wilcoxon tests), the Wilcoxon, Mann "
                                      "Whitney, and Fisher exact tests, the exact noncentral t intervals, the logistic regression of the Headline hit on "
                                      "accuracy, Shapiro Wilk tests, and the Spearman correlations. See analysis_output.txt.")).font = note_font
r = CK2 + 6
ws.cell(row=r - 1, column=1, value="analysis.R: SIMULATED_test_results.csv").font = bold
for c, h in enumerate(th, start=1):
    ws.cell(row=r, column=c, value=h)
style_header(ws, row=r, ncol=len(th))
for rw in trows:
    r += 1
    for c, v in enumerate(rw, start=1):
        ws.cell(row=r, column=c, value=to_cell(v))
autowidth(ws, max_w=34)
ws.column_dimensions["A"].width = 70
for col in "BCDEF":
    ws.column_dimensions[col].width = 18

wb.calculation = CalcProperties(fullCalcOnLoad=True)    # Excel computes every formula when the file opens
wb.save(OUT)
print("wrote", OUT)

# =============================================================================
# Verification 1: read the workbook back with openpyxl and compare every registered cell
# =============================================================================
wb2 = load_workbook(OUT)
problems = []
for (sheet, key), info in REG.items():
    v = wb2[sheet][info["cell"]].value
    v = str(v).replace("_xlfn.", "") if v is not None else None
    if v != info["formula"]:
        problems.append(f"{sheet}!{info['cell']} ({key}): workbook has {v!r}, registry has {info['formula']!r}")
for k, (lab, f, rv, tol, note) in enumerate(ck_rows):
    if wb2["Check"][f"B{CK1 + k}"].value != f:
        problems.append(f"Check!B{CK1 + k} differs")
if problems:
    print("VERIFICATION FAILED:")
    print("\n".join(problems[:40]))
    sys.exit(1)
print(f"verified (openpyxl): {len(REG)} registered cells and {len(ck_rows)} Check rows hold the registered formulas")
n_formulas = sum(1 for wsx in wb2.worksheets for row in wsx.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
print(f"formula cells in the workbook: {n_formulas}")

# =============================================================================
# Verification 2 and 3: recalculate a copy with LibreOffice and compare the values with analysis.R
# =============================================================================
LO = {}
soffice = shutil.which("soffice") or shutil.which("libreoffice")
lo_note = "LibreOffice was not found, so the formula values were not recalculated here."


def same_at(got, val, nd):
    """Equal at the printed precision: nd >= 0 decimals, or nd < 0 = that many significant digits (p values)."""
    got, val = float(got), float(val)
    if nd >= 0:
        # equal at nd decimals, allowing a value that sits exactly on a rounding tie (43.255 against 43.254999...)
        return f"{got:.{nd}f}" == f"{val:.{nd}f}" or abs(got - val) <= 1e-9 * max(1.0, abs(val))
    return f"{got:.{-nd - 1}e}" == f"{val:.{-nd - 1}e}" or abs(got - val) <= 1e-9 * max(abs(val), 1e-300)


if soffice:
    tmp = tempfile.mkdtemp(prefix="ws6_lo_")
    shutil.copy(OUT, os.path.join(tmp, "in.xlsx"))
    subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir",
                    os.path.join(tmp, "out"), os.path.join(tmp, "in.xlsx")], check=True, capture_output=True, timeout=600)
    wb3 = load_workbook(os.path.join(tmp, "out", "in.xlsx"), data_only=True)
    for (sheet, key), info in REG.items():
        LO[(sheet, key)] = wb3[sheet][info["cell"]].value
    statuses = [wb3["Check"][f"F{CK1 + k}"].value for k in range(len(ck_rows))]
    bad = [(ck_rows[k][0], wb3["Check"][f"B{CK1 + k}"].value, ck_rows[k][2]) for k, s in enumerate(statuses) if s != "OK"]
    if bad:
        print("LIBREOFFICE CHECK FAILED:")
        for b in bad[:40]:
            print("  ", b)
        sys.exit(1)
    mism = []
    for (sheet, key), (val, nd) in EXPECT.items():
        got = LO[(sheet, key)]
        if got is None or isinstance(got, str) or not same_at(got, val, nd):
            mism.append((sheet, key, got, val))
    if mism:
        print("VALUES DIFFER FROM analysis.R AT THE PRINTED DECIMALS:")
        for m_ in mism[:40]:
            print("  ", m_)
        sys.exit(1)
    lo_ver = subprocess.run([soffice, "--version"], capture_output=True, text=True).stdout.strip()
    lo_note = (f"A copy was recalculated with {lo_ver} (headless): all {len(ck_rows)} Check rows are OK, and the {len(EXPECT)} results printed "
               "in this walkthrough agree with analysis.R at the printed decimals.")
    print(f"verified (LibreOffice): {len(ck_rows)} Check rows OK; {len(EXPECT)} printed results agree with analysis.R")
    shutil.rmtree(tmp, ignore_errors=True)
else:
    print(lo_note)


# =============================================================================
# The walkthrough, generated from REG so that every cell reference is the real one
# =============================================================================
def cl(sheet, key):
    return REG[(sheet, key)]["cell"]


def fx(sheet, key):
    return REG[(sheet, key)]["formula"]


def val(sheet, key, nd=None):
    """Printed result: the LibreOffice value when available (verified above), else the analysis.R value."""
    v, nd0 = EXPECT.get((sheet, key), (None, 3))
    nd = nd0 if nd is None else nd
    if (sheet, key) in LO and LO[(sheet, key)] is not None and not isinstance(LO[(sheet, key)], str):
        v = LO[(sheet, key)]
    if v is None:
        return ""
    if nd < 0:
        return f"{float(v):.{-nd - 1}e}"
    return f"{float(v):.{nd}f}"


def row(label, sheet, key, nd=None):
    return f"| {label} | `{sheet}!{cl(sheet, key)}` | `{fx(sheet, key)}` | {val(sheet, key, nd)} |"


def row3(label, sheet, key):
    return f"| {label} | `{sheet}!{cl(sheet, key)}` | `{fx(sheet, key)}` |"


def rcell(sheet, key):
    return f"{sheet}!{cl(sheet, key)}"


def f2(x, nd=2):
    return f"{float(x):.{nd}f}"


def p3(x):
    x = float(x)
    return "< .001" if x < 0.001 else "= " + f"{x:.3f}".lstrip("0")


def ref_apa(key):
    """APA style reference line from refs/references.json (only keys that exist there)."""
    with open(REFS, encoding="utf-8") as fh:
        refs = {r["key"]: r for r in json.load(fh)}
    r = refs[key]
    au = r["authors"]
    a = au[0] if len(au) == 1 else (", ".join(au[:-1]) + ", & " + au[-1] if len(au) <= 20 else ", ".join(au[:19]) + ", ... " + au[-1])
    vol = f", *{r['volume']}*" if r.get("volume") else ""
    iss = f"({r['issue']})" if r.get("issue") else ""
    pg = f", {r['pages']}" if r.get("pages") else ""
    pg = pg.replace("\u2013", "-")
    doi = f" https://doi.org/{r['doi']}" if r.get("doi") else (f" {r['url']}" if r.get("url") else "")
    title = r["title"].replace("\u2014", ", ").replace("\u2013", "-")
    return f"- {a} ({r['year']}). {title}. *{r['container']}*{vol}{iss}{pg}.{doi}".replace("**", "")


h1, s1, h3, dr, at, at9 = T["H1"], T["S30_H1"], T["H3"], T["DRIFT_ALL"], T["ATTRITION_ALL"], T["ATTRITION_STIMULI"]
fr, fr_s = T["H2_FRIEDMAN"], T["S30_H2_FRIEDMAN"]
pairs = [T[k] for k in T if k.startswith("H2_PAIR_")]
pairs_s = [T[k] for k in T if k.startswith("S30_H2_PAIR_")]
hit = T["HIT_HEADLINE"]
e1, e2 = T["EXP_PREF"], T["EXP_RECALL"]
mob = [r for r in ses_rows if r[si["Device type (modeled)"]] in ("Phone", "Tablet")]
mob5 = sum(1 for r in mob if r[si["Last funnel step (modeled)"]] == STEPS[4])
m_mob = re.search(r"Sessions that stopped at STEP 5 Screen size: (\d+) \((\d+) phones or tablets", out_txt)
assert m_mob and int(m_mob.group(2)) == mob5, "STEP 5 phone count differs from analysis_output.txt"
first_done = done_ids[0]
fd_row = [r for r in ses_rows if r[si["Id"]] == first_done][0]
first_stop_row, first_stop = next((k + 2, r[si["Id"]]) for k, r in enumerate(ses_rows) if r[si["Status"]] != "Processed")
sig_pairs = [p["comparison"].split(": ")[-1].replace(", dwell time (ms)", "") for p in pairs if float(p["p_holm"]) < 0.05]
nsig_pairs = [p["comparison"].split(": ")[-1].replace(", dwell time (ms)", "") for p in pairs if float(p["p_holm"]) >= 0.05]
m_hidden = re.search(r"calibration is above 4.0: (\d+)\. Screening", out_txt)
assert m_hidden, "step 4 line not found in analysis_output.txt"
m_fisher9 = re.search(r"\(b\) Completion by order among.*?Fisher exact test: p = ([0-9.e-]+); odds ratio = ([0-9.]+), 95% CI \[([0-9.]+), ([0-9.]+)\]",
                      out_txt, re.S)
assert m_fisher9, "Fisher line not found"
m_bands = re.findall(r"^\s+(1\.5 or less|1\.5 to 2\.5|2\.5 to 3\.5|above 3\.5)\s+(\d+)\s+(\d+)\s+(\d+)\s+([0-9.]+)\s+([0-9.]+)$", out_txt, re.M)
assert len(m_bands) == 4, "band table not found"
mw_h3 = (fnum(h3["nonparam_stat"]), fnum(h3["nonparam_p"]))


def _and(xs):
    xs = list(xs)
    if not xs:
        return "none"
    return xs[0] if len(xs) == 1 else (" and ".join(xs) if len(xs) == 2 else ", ".join(xs[:-1]) + ", and " + xs[-1])


def pcell(x):
    x = float(x)
    return "< .001" if x < 0.001 else f"{x:.3f}".lstrip("0")


def r3(x):
    """A correlation in APA style: two decimals, no leading zero."""
    v = f"{float(x):.2f}"
    return v.replace("0.", ".", 1) if v.startswith(("0.", "-0.")) else v


NUMW = {0: "None", 1: "One", 2: "Two", 3: "Three", 4: "Four", 5: "Five", 6: "Six"}


def pair_line(p):
    return (f"| {p['comparison'].split(': ')[-1].replace(', dwell time (ms)', '')} | {f2(p['mean_diff'], 0)} | {f2(p['es'])} | "
            f"{f2(p['nonparam_stat'], 1).rstrip('0').rstrip('.')} | {pcell(p['p'])} | {pcell(p['p_holm'])} |")


md = f"""# Remote study analysis in Excel: step by step walkthrough

**SIMULATED DATA.** Every number in this walkthrough comes from the SIMULATED files that `make_simulated_data.py` generated (seed {SEED}). No real person was recorded and no webcam was switched on. The design copies the logic of the iMotions vendor demo "(DEMO) Food Waste_webET_FEA_voice" (help center article "Demo study: Food Waste (image, videos and talkaloud)"): two one minute food waste videos, a narrative one and an emotional visual one, in random order, then an infographic with four tips, then a questionnaire. Nothing from the demo's respondents was used.

This is the Excel route through the analysis in `analysis.R`, for students who do not use R. It covers the funnel, differential attrition, the exclusion rule, calibration drift, H1, and H3. The finished workbook is `SIMULATED_remote_analysis.xlsx`; the sheet names below refer to it. **Every cell address in this walkthrough is written as `Sheet!Cell` and was generated from the workbook itself.** `make_excel_workbook.py` writes the workbook and this file from one list of cells, reads the workbook back with openpyxl to confirm that every address holds the formula printed here, and recalculates a copy to confirm the results. {lo_note}

Companion files: `analysis.R` (R version), `analysis_output.txt` (its console output), `results_text.md` (model paragraphs), `figures/` (PNG figures).

---

## 0. What you are analyzing

**Study.** "SIMULATED Food Waste Remote Study", a **Remote Data Collection** (RDC) study built in iMotions Lab 11.1 with the sensors "Eye tracking (Web Camera)" and "Respondent camera", uploaded to my.imotions.com, and shared by link. Respondents took part at home on their own computers. The flow: `Pre-study calibration` (13 points, 39 s), `Instructions1` (30 s), the two videos (`foodwaste_narrative`, 60 s, and `foodwaste_visual`, 62 s), `infographic` (15 s, four tips), `Questionnaire` (iMotions survey), and `Post-study calibration` (13 points, 39 s). After collection the study zip was downloaded and loaded into iMotions Lab, the faces were scored with **Post Processing > Affectiva**, and the FEA thresholding notebook was run at a threshold of 50 (as in the demo article).

**Order.** Each order of the two videos is its own study flow, recorded in the documented `Block` column: `Study flow 1` = narrative first, `Study flow 2` = visual first. The flow was assigned at random before the link was opened (modeled: the invitation list was split at random into two halves). The demo randomized the two videos inside one block instead; that hides the order of every session whose data never uploaded, and a differential attrition check needs it.

**Webcam quality has two parts** (iMotions Online data quality box). **Accuracy** is "the mean distance between the actual and expected gaze positions calculated during calibration", in degrees of visual angle, so **smaller is better**. **Quality** is "the ratio of valid eye tracking samples in relation to the total number of samples collected"; face quality is the ratio of samples "where the face can be detected".

**Files.**

| File | What it is | Documented or modeled | Used for |
|---|---|---|---|
| `SIMULATED_session_data.csv` | One row per session that opened a link ({N_SES}) | `Id`, `Started`, `Ended`, `Eye tracking`, `Est. face %`, `Block`, `Status` are the RDC respondent list columns (Accessing RDC Study Data). **Modeled:** the CSV layout and the columns `Last funnel step (modeled)`, `Device type (modeled)`, `Browser (modeled)` | Funnel, attrition |
| `SIMULATED_quality.csv` | One row per completed respondent and slide ({N_Q} rows) | **Modeled layout** (the Quality metrics table of a study session page, with EYE TRACKING split into accuracy and quality) | Exclusion rule, drift |
| `SIMULATED_fea_metrics.csv` | FEA Metrics Export, one row per completed respondent, video, and expression ({N_FEA} rows) | **Modeled header** (the Workshop 3 layout) | H1 |
| `SIMULATED_survey.csv` | Survey export ({N_SV} rows) | **Modeled layout** | H3 |
| `SIMULATED_aoi_metrics.csv` | Individual AOI metrics on the infographic | Metric names from the AOI Metrics article; identifier columns modeled | H2 (R only) |

**Hypotheses.**
- H1 (within subjects): percent of time Engaged is higher for `foodwaste_visual` than for `foodwaste_narrative`.
- H2 (within subjects): dwell time differs among the four tip AOIs on the infographic (R only; section 10).
- H3 (between subjects; randomization check): intention to reduce food waste does not differ by which video came first.
- Data quality: the funnel, completion by order (differential attrition), calibration drift, and accuracy against the hit on the smallest AOI (R only).

**Decision rules fixed before collection.** Exclude a completed respondent when the Pre-study **or** the Post-study calibration accuracy is above {ACC_MAX:.1f} degrees (iMotions Online Study Setup: "Generally acceptable range: 0-4 dva"), when eye tracking quality on the infographic is below {ETQ_MIN:.0f} percent ("We recommend a quality score of at least 90%"), or when face quality is below {FACE_MIN:.0f} percent during either video ("at least 80%"). Sensitivity: accuracy at or below {ACC_STRICT:.1f} degrees at both calibrations (the Data collection and Analysis articles: "We recommend a threshold of 3.0"). All tests are two tailed with alpha = .05.

---

## 1. Import the files

iMotions exports are comma separated with a dot decimal. Double clicking a file uses your regional defaults, which can put everything in one column (Opening Data Export Files in Excel article). Use the import route instead.

**Route A, Get Data (Excel 2016 and later):**
1. **Data > From Text/CSV**, select `SIMULATED_session_data.csv`, click **Import**.
2. In the preview set **Delimiter: Comma** and **File Origin: 65001 Unicode (UTF-8)**.
3. Click **Transform Data**. Every simulated file starts with **one** comment line (`# SIMULATED ...`). In Power Query choose **Home > Remove Rows > Remove Top Rows** (1), then **Home > Use First Row as Headers**. Select the `Eye tracking` column and set **Data Type: Text** (Power Query may otherwise try to read "1.9, 96%" as something else), then **Close & Load**.
4. **You should see** {len(ses_h)} columns with `Id` in the first column and {N_SES} data rows, one per session. Repeat for the other files ({N_Q}, {N_FEA}, and {N_SV} data rows).

**Route B, Text Import Wizard (any version):** **File > Open > Browse**, set the file type to **All Files**, open the .csv, choose **Delimited**, start at row 2, tick **Comma**, set the `Eye tracking` column to **Text** in the last step, click **Advanced** and set the decimal symbol to a dot, then **Finish**.

**Common mistakes.** (1) The `Eye tracking` cell holds **two numbers in one text cell** ("1.9, 96%": accuracy in degrees, then eye tracking quality). It cannot be averaged until it is split (section 2). (2) Empty cells are real: a session that never finished has no `Ended` time and no quality values. Leave them empty; do not type 0. (3) The quality table writes `NA` in the accuracy column of every slide that is not a calibration; leave it as text.

In the workbook the files sit on the **Session**, **Quality**, **FEA**, and **Survey** sheets, header in row 1 and data from row 2. Session column letters: `{SCOL['Id']}` = Id, `{SCOL['Eye tracking']}` = Eye tracking, `{SCOL['Block']}` = Block, `{SCOL['Status']}` = Status, `{SCOL['Last funnel step (modeled)']}` = Last funnel step (modeled), `{SCOL['Device type (modeled)']}` = Device type (modeled).

---

## 2. Helper columns on the Session sheet

Add five columns to the right of the session data (`{SCOL['Step number']}` to `{SCOL['Session quality (%)']}`, headers in orange) and fill them down. The formulas below are the ones in row 2.

| Header | Cell | Formula | Meaning |
|---|---|---|---|
| Step number | `Session!{cl('Session', 'step')}` | `{fx('Session', 'step')}` | The number inside "STEP 5 Screen size" (characters 6 up to the next space); 11 for a session that completed |
| Completed (1 = yes) | `Session!{cl('Session', 'done')}` | `{fx('Session', 'done')}` | 1 when Status is Processed (the data arrived and were processed) |
| Order | `Session!{cl('Session', 'ord')}` | `{fx('Session', 'ord')}` | The flow translated into the order of the videos |
| Session accuracy (degrees) | `Session!{cl('Session', 'acc')}` | `{fx('Session', 'acc')}` | The text before the comma, as a number |
| Session quality (%) | `Session!{cl('Session', 'q')}` | `{fx('Session', 'q')}` | The text after ", " without the % sign |

The first completed session, `{first_done}`, is in row {FIRST_DONE_ROW}: its `Eye tracking` cell reads "{fd_row[si['Eye tracking']]}" and `{rcell('Session', 'acc_first_done')}` returns {fd_row[si['Eye tracking']].split(',')[0]}. A session that stopped early has an empty `Eye tracking` cell, so its last two helper cells stay empty (the first one is `{first_stop}` in row {first_stop_row}).

**Why step numbers instead of step names.** The ninth step is called "STEP 9 Stimuli presentation*" (with an asterisk, as the portal writes it). In `COUNTIF` and `COUNTIFS` an asterisk in the criterion is a **wildcard** that matches any text, so a criterion typed from step names can count the wrong rows. Numbers avoid the problem.

**Common mistake.** Screening respondents on the session level `Eye tracking` number. It is rounded to one decimal and (in this simulation) is the mean of the two calibrations, so it hides a bad calibration: {m_hidden.group(1)} completed respondents show 4.0 or less there although one calibration was above 4.0 (step 4 of `analysis_output.txt`). The rule uses the quality table instead (section 5).

---

## 3. The respondent funnel (Funnel sheet)

The portal draws the respondent funnel as a chart ("View information on where respondents dropped off. Sessions that are in progress are not included."). The **Funnel** sheet rebuilds it as a table. A session **entered** a step when its step number is that step or a later one, and it **stopped** at the step where it ended.

| Column | Cell (STEP 1) | Formula |
|---|---|---|
{row3('Entered', 'Funnel', 'ent1')}
{row3('Stopped here', 'Funnel', 'stop1')}
{row3('Continued', 'Funnel', 'cont1')}
{row3('Percent of opened', 'Funnel', 'pct1')}
{row3('Percent of previous step (STEP 2)', 'Funnel', 'prev2')}

**You should see** (simulated):

| Step | Entered | Stopped here | Percent of opened |
|---|---|---|---|
""" + "\n".join(f"| {STEPS[k]} | {val('Funnel', f'ent{k + 1}', 0)} | {val('Funnel', f'stop{k + 1}', 0)} | {val('Funnel', f'pct{k + 1}', 1)} |" for k in range(10)) + f"""
| **Completed (Status Processed)** | **{val('Funnel', 'done', 0)}** | | **{val('Funnel', 'done_pct', 1)}** |

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('Completed (Status Processed)', 'Funnel', 'done', 0)}
{row('Continued after STEP 10 minus Completed (must be 0)', 'Funnel', 'chk_cont', 0)}
{row('Status Abandoned', 'Funnel', 'aband', 0)}
{row('Status Errors', 'Funnel', 'errors', 0)}

**Reading the funnel.** Of {N_SES} sessions that opened a link, {val('Funnel', 'done', 0)} ({val('Funnel', 'done_pct', 1)} percent) completed. The largest single loss is at **STEP 5 Screen size** ({val('Funnel', 'stop5', 0)} sessions: {mob5} phones or tablets and {int(float(val('Funnel', 'stop5', 0))) - mob5} computers). Every phone or tablet session that reached STEP 5 stopped there ({len(mob) - mob5} of the {len(mob)} phone and tablet sessions had left earlier). This is modeled: the RDC article says "Participation can not be done on phones or tablets" but not at which step a phone is stopped. Nobody stopped at STEP 3 because screen recording was off in this study. STEP 9 holds both calibrations, both videos, the infographic, and the questionnaire; {val('Funnel', 'stop9', 0)} sessions stopped there. At STEP 10 the upload failed for {val('Funnel', 'stop10', 0)} sessions ({val('Funnel', 'errors', 0)} with Status Errors).

**By device** (rows {DV} to {DV + 3}, modeled column): `Funnel!B{DV}` = `{fx('Funnel', 'dv_o_Laptop')}`, `Funnel!C{DV}` = `{fx('Funnel', 'dv_c_Laptop')}`. Laptop {val('Funnel', 'dv_c_Laptop', 0)} of {val('Funnel', 'dv_o_Laptop', 0)} completed, Desktop {val('Funnel', 'dv_c_Desktop', 0)} of {val('Funnel', 'dv_o_Desktop', 0)}, Phone {val('Funnel', 'dv_c_Phone', 0)} of {val('Funnel', 'dv_o_Phone', 0)}, Tablet {val('Funnel', 'dv_c_Tablet', 0)} of {val('Funnel', 'dv_o_Tablet', 0)}. Put "desktop or laptop computer only" in the invitation.

**Chart.** The sheet has a bar chart of the Entered column (**Insert > Bar > Clustered Bar** on `Funnel!B{FR1 - 1}:C{DONE_ROW}`; right click the vertical axis, **Format Axis > Categories in reverse order**, so STEP 1 is on top).

**Common mistake.** Reporting only the final number. A funnel shows **where** people were lost, and each kind of loss has a different fix: consent wording (STEP 2), camera help (STEP 4), "computer only" in the invitation (STEP 5), a shorter study (STEP 9), the data upload alternative (STEP 10).

---

## 4. Differential attrition: did completion depend on the order?

If one order made more people quit (for example, the visual video upset some viewers when it came first), the people who finished each order would no longer be comparable, and an order difference could appear that randomization did not create (Zhou & Fishbach, 2016). Because the flow was assigned before the link was opened, the `Block` column gives the order of **every** session.

| Column | Cell (Narrative first) | Formula |
|---|---|---|
{row3('Opened', 'Funnel', 'op_0')}
{row3('Reached STEP 9', 'Funnel', 'r9_0')}
{row3('Completed', 'Funnel', 'dn_0')}
{row3('Percent completed', 'Funnel', 'pc_0')}
{row3('Expected completed (if order and completion were unrelated)', 'Funnel', 'ex_d0')}

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('Opened, Narrative first', 'Funnel', 'op_0', 0)}
{row('Opened, Visual first', 'Funnel', 'op_1', 0)}
{row('Completed, Narrative first', 'Funnel', 'dn_0', 0)}
{row('Completed, Visual first', 'Funnel', 'dn_1', 0)}
{row('Percent completed, Narrative first', 'Funnel', 'pc_0', 1)}
{row('Percent completed, Visual first', 'Funnel', 'pc_1', 1)}
{row('Chi square (Pearson, no continuity correction)', 'Funnel', 'att_x2', 2)}
{row('p (CHISQ.DIST.RT)', 'Funnel', 'att_p', 3)}
{row('p (CHISQ.TEST)', 'Funnel', 'att_p_ct', 3)}
{row('phi', 'Funnel', 'att_phi', 2)}
{row('Smallest expected count', 'Funnel', 'att_minexp', 2)}

`CHISQ.TEST(observed, expected)` needs both tables, so the expected counts are built first (row total x column total / N). R's `chisq.test` applies the Yates continuity correction to a 2 x 2 table by default; `analysis.R` reports the uncorrected value as the main result because it matches Excel, and prints the corrected one beside it (X2 = {re.search(r"Completion by order, all.*?Yates continuity correction \(R default\): X2 = ([0-9.]+)", out_txt, re.S).group(1)}).

**The sharper check.** Before STEP 9 nobody had seen a video, so only losses from STEP 9 on can depend on the order. Rows {R9} to {R9 + 1} repeat the test for the sessions that reached STEP 9:

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('Percent completed after STEP 9, Narrative first', 'Funnel', 's9_p_0', 1)}
{row('Percent completed after STEP 9, Visual first', 'Funnel', 's9_p_1', 1)}
{row('Chi square (Pearson)', 'Funnel', 'att9_x2', 2)}
{row('p (CHISQ.DIST.RT)', 'Funnel', 'att9_p', 3)}
{row('Smallest expected count', 'Funnel', 'att9_minexp', 2)}

The smallest expected count is below 5, so report the **Fisher exact test**, which Excel does not have: in R, p = {f2(m_fisher9.group(1), 3)}, odds ratio = {m_fisher9.group(2)}, 95% CI [{m_fisher9.group(3)}, {m_fisher9.group(4)}].

**You should see** completion rates within a few points of each other and p values well above .05 in both tests: no sign of differential attrition (simulated). A nonsignificant test is not proof of equal attrition (the test has little power for small losses), so report the counts by order as well.

---

## 5. The exclusion rule (Screening sheet)

The **Screening** sheet lists every completed respondent once ({nD} respondents in rows {P1} to {PL}). The thresholds sit in the yellow cells `Screening!B2` ({ACC_MAX:.1f} degrees), `Screening!D2` ({ACC_STRICT:.1f}, sensitivity), `Screening!F2` ({ETQ_MIN:.0f} percent), and `Screening!H2` ({FACE_MIN:.0f} percent); change one and every decision updates. The five quality values come from the **Quality** sheet with `SUMIFS` on respondent and slide (each pair occurs once, so the sum is the value).

| Column | Cell (first respondent) | Formula |
|---|---|---|
{row3('Order', 'Screening', 'ord')}
{row3('Pre-study accuracy', 'Screening', 'pre')}
{row3('Post-study accuracy', 'Screening', 'post')}
{row3('Eye tracking quality, infographic', 'Screening', 'etq')}
{row3('Face quality, narrative', 'Screening', 'fn')}
{row3('Face quality, visual', 'Screening', 'fv')}
{row3('Fail: Pre-study accuracy', 'Screening', 'f_pre')}
{row3('Fail: Post-study accuracy', 'Screening', 'f_post')}
{row3('Fail: eye tracking quality', 'Screening', 'f_etq')}
{row3('Fail: face quality', 'Screening', 'f_face')}
{row3('Decision', 'Screening', 'dec')}
{row3('First reason', 'Screening', 'first')}
{row3('Largest accuracy', 'Screening', 'big')}
{row3('Decision at 3.0', 'Screening', 'dec30')}
{row3('Drift (Post minus Pre)', 'Screening', 'drift')}

The counts (column S):

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('Completed respondents screened', 'Screening', 'cnt_n', 0)}
{row('Excluded', 'Screening', 'cnt_excl', 0)}
{row('Kept', 'Screening', 'cnt_kept', 0)}
{row('First reason: Pre-study accuracy', 'Screening', 'cnt_first_pre', 0)}
{row('First reason: Post-study accuracy', 'Screening', 'cnt_first_post', 0)}
{row('First reason: eye tracking quality', 'Screening', 'cnt_first_etq', 0)}
{row('First reason: face quality', 'Screening', 'cnt_first_face', 0)}
{row('Failed (any): Pre-study accuracy', 'Screening', 'cnt_any_pre', 0)}
{row('Failed (any): Post-study accuracy', 'Screening', 'cnt_any_post', 0)}
{row('Failed (any): eye tracking quality', 'Screening', 'cnt_any_etq', 0)}
{row('Failed (any): face quality', 'Screening', 'cnt_any_face', 0)}
{row('Failed more than one criterion', 'Screening', 'cnt_multi', 0)}
{row('Failed only at the Post-study calibration', 'Screening', 'cnt_postonly', 0)}
{row('Kept, Narrative first', 'Screening', 'cnt_kept_n', 0)}
{row('Kept, Visual first', 'Screening', 'cnt_kept_v', 0)}
{row('Kept at 3.0 degrees', 'Screening', 'cnt_kept30', 0)}
{row('Further removed at 3.0 degrees', 'Screening', 'cnt_more30', 0)}

**Two ways to count reasons.** "Failed (any)" counts every respondent who failed a criterion, so a respondent can appear twice and the column adds up to more than the number excluded. "First reason" gives each excluded respondent one reason, the first one failed in the order of the rule, so it adds up to the total. Report both, as the lab convention asks ("report respondents kept per stimulus and the number excluded for each reason").

**You should see** {val('Screening', 'cnt_excl', 0)} exclusions and {val('Screening', 'cnt_kept', 0)} kept respondents ({val('Screening', 'cnt_kept_n', 0)} Narrative first, {val('Screening', 'cnt_kept_v', 0)} Visual first). Of the excluded respondents, {val('Screening', 'cnt_postonly', 0)} calibrated well at the start and failed only at the end: a single calibration at the start would have kept them. The RDC Testing with Webcam Eye Tracking article calls a calibration at the start and another at the end "currently the best practice that we recommend for webcam eye tracking".

**Common mistake.** Reading accuracy like a score where bigger is better. Accuracy is an **error** in degrees: the rule excludes values **above** {ACC_MAX:.1f}, while the two quality rules exclude values **below** their minimum.

---

## 6. Calibration drift (Screening sheet, column S)

The same calibration ran at the start and at the end, so a paired comparison shows how much the gaze estimate worsened while the study ran (people move, lean back, or the light changes).

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('Pre-study mean (degrees)', 'Screening', 'dr_m1', 2)}
{row('Post-study mean (degrees)', 'Screening', 'dr_m2', 2)}
{row('n', 'Screening', 'dr_n', 0)}
{row('Mean difference, Post minus Pre', 'Screening', 'dr_md', 2)}
{row('SD of the differences', 'Screening', 'dr_sdd', 2)}
{row('t', 'Screening', 'dr_t', 2)}
{row('df', 'Screening', 'dr_df', 0)}
{row('p, two tailed (T.TEST type 1)', 'Screening', 'dr_p_tt', -3)}
{row('Mean difference 95% CI lower', 'Screening', 'dr_ci_lo', 2)}
{row('Mean difference 95% CI upper', 'Screening', 'dr_ci_hi', 2)}
{row("Cohen's d_z", 'Screening', 'dr_dz', 2)}
{row('Correlation of the two calibrations', 'Screening', 'dr_r', 2)}
{row('Respondents with a worse Post-study accuracy', 'Screening', 'dr_worse', 0)}

The p value is far below .001; write "p < .001", never "p = .000". `analysis.R` also gives the exact interval for d_z, [{f2(dr['es_ci_low'])}, {f2(dr['es_ci_high'])}], and the Wilcoxon signed rank test (V = {f2(dr['nonparam_stat'], 0)}, p < .001).

---

## 7. H1: Engagement, visual versus narrative video (H1 sheet)

The **H1** sheet has one row per kept respondent (rows {H1A} to {H1L}). The respondents kept under the {ACC_STRICT:.1f} degree sensitivity rule come first (rows {H1A} to {H1S}) and the {nK - nS} removed by it last (rows {H1S + 1} to {H1L}), so the sensitivity block only shortens the range.

| Column | Cell (first respondent) | Formula |
|---|---|---|
{row3('Order', 'H1', 'ord')}
{row3('Decision at 3.0', 'H1', 'dec30')}
{row3('Engagement, narrative (%)', 'H1', 'narr')}
{row3('Engagement, visual (%)', 'H1', 'vis')}
{row3('Difference (visual minus narrative)', 'H1', 'diff')}

`SUMIFS` with three conditions (respondent, `Interval label`, `Expression`) picks exactly one row of the FEA export. The test sits in column I (formula text in column J). A paired t test **is** a one sample t test on the differences:

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('Narrative mean (% of time)', 'H1', 'h1_m1', 2)}
{row('Narrative SD', 'H1', 'h1_s1', 2)}
{row('Visual mean (% of time)', 'H1', 'h1_m2', 2)}
{row('Visual SD', 'H1', 'h1_s2', 2)}
{row('n', 'H1', 'h1_n', 0)}
{row('Mean difference, visual minus narrative', 'H1', 'h1_md', 2)}
{row('SD of the differences', 'H1', 'h1_sdd', 2)}
{row('t', 'H1', 'h1_t', 2)}
{row('df', 'H1', 'h1_df', 0)}
{row('**p, two tailed (T.TEST type 1)**', 'H1', 'h1_p_tt', -3)}
{row('p from the hand computed t', 'H1', 'h1_p_hand', -3)}
{row('Mean difference 95% CI lower', 'H1', 'h1_ci_lo', 2)}
{row('Mean difference 95% CI upper', 'H1', 'h1_ci_hi', 2)}
{row("Cohen's d_z", 'H1', 'h1_dz', 2)}
{row('d_z 95% CI lower (approximation)', 'H1', 'h1_dz_lo', 2)}
{row('d_z 95% CI upper (approximation)', 'H1', 'h1_dz_hi', 2)}
{row('Respondents with a positive difference', 'H1', 'h1_npos', 0)}
{row('Respondents at exactly 0, narrative', 'H1', 'h1_z_n', 0)}
{row('Respondents at exactly 0, visual', 'H1', 'h1_z_v', 0)}
{row('Median, narrative', 'H1', 'h1_med_n', 2)}
{row('Median, visual', 'H1', 'h1_med_v', 2)}
{row('Median difference', 'H1', 'h1_med_d', 2)}

`T.TEST(array1, array2, tails, type)`: `tails = 2` for a two tailed test; `type = 1` for **paired** data. `type = 2` and `type = 3` treat the two columns as different people and are wrong here. **You should see** the hand computed p equal the `T.TEST` p (both far below .001, so report p < .001). For paired data the effect size is **d_z**, the mean difference in units of the SD of the differences; SE(d_z) = sqrt(1/n + d_z^2 / (2 n)) is the large sample approximation, and `analysis.R` also reports the exact interval from the noncentral t distribution, [{f2(h1['es_ci_low'])}, {f2(h1['es_ci_high'])}] (Lakens, 2013). Rule of thumb for d_z: 0.2 small, 0.5 medium, 0.8 large.

**Why the Wilcoxon check matters here.** Percent of time data pile up at zero and have a long right tail: {val('H1', 'h1_z_n', 0)} respondents were never above the threshold during the narrative video, and its mean ({val('H1', 'h1_m1', 2)}) is above its median ({val('H1', 'h1_med_n', 2)}). With n = {nK} the t test on the mean difference is robust to this (the differences themselves are close to symmetric), and the Wilcoxon signed rank test, which compares ranks, is the check. It has no Excel function; in R: V = {f2(h1['nonparam_stat'], 0)}, p < .001.

**Common mistake.** Reading "percent of time Engaged" as "percent of viewers engaged". It is, for each respondent, the share of the video's frames in which the Engagement score was above 50; the mean is the average of those shares.

---

## 8. Sensitivity: H1 at the {ACC_STRICT:.1f} degree threshold (H1 sheet)

Accuracy does not affect facial expression scores, but the rule removes respondents, so the sensitivity analysis asks whether H1 depends on who was removed. The second block uses rows {H1A} to {H1S}.

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('n', 'H1', 's30_n', 0)}
{row('Narrative mean (% of time)', 'H1', 's30_m1', 2)}
{row('Visual mean (% of time)', 'H1', 's30_m2', 2)}
{row('Mean difference, visual minus narrative', 'H1', 's30_md', 2)}
{row('t', 'H1', 's30_t', 2)}
{row('p (T.TEST type 1)', 'H1', 's30_p_tt', -3)}
{row('Mean difference 95% CI lower', 'H1', 's30_ci_lo', 2)}
{row('Mean difference 95% CI upper', 'H1', 's30_ci_hi', 2)}
{row("Cohen's d_z", 'H1', 's30_dz', 2)}

| Roster check | Cell | Formula | Result |
|---|---|---|---|
{row('Kept respondents on the Screening sheet', 'H1', 'chk_chk_n', 0)}
{row('Rows in the sensitivity block marked Exclude at 3.0 (must be 0)', 'H1', 'chk_chk_s30', 0)}
{row('Rows below the sensitivity block marked Exclude at 3.0', 'H1', 'chk_chk_rest', 0)}

**You should see** the same conclusion with {nK - nS} fewer respondents: d_z {f2(h1['es'])} with all kept respondents and {f2(s1['es'])} at {ACC_STRICT:.1f} degrees.

---

## 9. H3: intention by order, the randomization check (H3 sheet)

The **H3** sheet lists the kept respondents by order: Narrative first in rows {H3A} to {H3N}, Visual first in rows {H3N + 1} to {H3L}. `{rcell('H3', 'int')}` = `{fx('H3', 'int')}` looks up the answer.

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('Narrative first mean', 'H3', 'h3_m1', 2)}
{row('Narrative first SD', 'H3', 'h3_s1', 2)}
{row('Narrative first n', 'H3', 'h3_n1', 0)}
{row('Visual first mean', 'H3', 'h3_m2', 2)}
{row('Visual first SD', 'H3', 'h3_s2', 2)}
{row('Visual first n', 'H3', 'h3_n2', 0)}
{row('Difference, Narrative first minus Visual first', 'H3', 'h3_diff', 2)}
{row('Welch t', 'H3', 'h3_t', 2)}
{row('Welch df', 'H3', 'h3_df', 2)}
{row('**p, two tailed (T.TEST type 3)**', 'H3', 'h3_p_tt', 3)}
{row('Difference 95% CI lower', 'H3', 'h3_ci_lo', 2)}
{row('Difference 95% CI upper', 'H3', 'h3_ci_hi', 2)}
{row('Pooled SD', 'H3', 'h3_sp', 2)}
{row("Cohen's d", 'H3', 'h3_d', 2)}
{row("Hedges' g", 'H3', 'h3_g', 2)}
{row('g 95% CI lower (approximation)', 'H3', 'h3_g_lo', 2)}
{row('g 95% CI upper (approximation)', 'H3', 'h3_g_hi', 2)}

`type = 3` in `T.TEST` is Welch's test, which does not assume equal variances (Delacre et al., 2017). Excel's `T.DIST.2T` and `T.INV.2T` may round the Welch df down to a whole number, so the hand computed p and the interval can differ from R in the third decimal; report the `T.TEST` p. The exact interval for g from R is [{f2(h3['g_ci_low'])}, {f2(h3['g_ci_high'])}]; the Mann Whitney check gives W = {mw_h3[0]:.0f}, p {p3(mw_h3[1])}.

**What this test means.** The two groups differ only in which video came first, and the order was assigned at random. A test of something measured after both videos asks two things at once: did randomization give comparable groups, and did the order carry over to the questionnaire? A **nonsignificant** result is what randomization predicts. A significant one would mean either a chance imbalance (expected in 5 percent of such tests) or a real order effect; the analysis would then report intention separately by order or add order to the model, and the attrition check in section 4 would matter even more. **You should see** a small difference with a confidence interval that includes zero (simulated).

**Common mistake.** Calling a nonsignificant randomization check "proof" that the groups are equal. It only means the data give no reason to think they differ; the confidence interval shows how large a difference is still compatible with the data.

---

## 10. What Excel cannot do here: the R results

**H2: dwell time on the four tips** (`SIMULATED_aoi_metrics.csv`, `Dwell time (ms)`, with NA counted as 0 ms because the respondent never fixated the AOI). Within each respondent the four tips are ranked 1 to 4 and Friedman's test asks whether the mean ranks differ; Kendall's W = Friedman chi square / (n (k - 1)) runs from 0 (no agreement) to 1 (every respondent ranks the tips the same way). In R: Friedman chi square(3) = {f2(fr['statistic'])}, p < .001, n = {int(float(fr['n']))}, W = {f2(fr['es'], 3)}. Holm adjusted Wilcoxon signed rank comparisons (mean difference in ms, first AOI minus second):

| Pair | Mean difference (ms) | d_z | V | p | p (Holm) |
|---|---|---|---|---|---|
""" + "\n".join(pair_line(p) for p in pairs) + f"""

{NUMW[len(sig_pairs)]} of the six pairs differ after the adjustment ({_and(sig_pairs)}); the other pairs ({_and(nsig_pairs)}) do not. At {ACC_STRICT:.1f} degrees: chi square(3) = {f2(fr_s['statistic'])}, W = {f2(fr_s['es'], 3)}, {sum(1 for p in pairs_s if float(p['p_holm']) < 0.05)} pairs. A PivotTable of the AOI file (**Insert > PivotTable**, `AOI Name` in **Rows**, `Dwell time (ms)` in **Values** set to **Average**) gives descriptive means, but it skips the NA rows, so its means are for the respondents who fixated each tip, not for everyone (Creating Excel Pivot Tables article).

**Accuracy and the hit on the smallest AOI** (the Headline banner). Logistic regression of hit (at least one fixation) on the mean of the two calibration accuracies, for every completed respondent whose eye tracking and face quality passed (n = {int(float(hit['n']))}): odds ratio per degree = {f2(hit['es'])}, 95% CI [{f2(hit['es_ci_low'])}, {f2(hit['es_ci_high'])}], p {p3(hit['p'])}. Hit ratio by mean accuracy: """ + "; ".join(f"{b[4]} percent {'at ' + b[0].replace('or less', 'degrees or less') if 'less' in b[0] else (b[0] + ' degrees' if 'above' in b[0] else 'at ' + b[0] + ' degrees')} (n = {b[1]})" for b in m_bands) + f"""; the last band holds too few respondents for a stable percentage.

**Exploratory.** Engagement difference (visual minus narrative) with naming the visual video more persuasive: point biserial r = {r3(e1['es'])}, 95% CI [{r3(e1['es_ci_low'])}, {r3(e1['es_ci_high'])}], p {p3(e1['p'])} (in Excel, `CORREL` of the difference column and a 0/1 column gives the same r). Total tip dwell time with tips recalled: r = {r3(e2['es'])}, 95% CI [{r3(e2['es_ci_low'])}, {r3(e2['es_ci_high'])}], p {p3(e2['p'])}; Spearman rho = {r3(e2['nonparam_stat'])}.

Also R only: the Wilcoxon, Mann Whitney, and Fisher exact tests, Shapiro Wilk tests, and the exact noncentral t intervals.

---

## 11. Check your numbers (Check sheet)

The **Check** sheet puts each live Excel result next to the value `analysis.R` produced and prints OK or CHECK (status in column F, the count of CHECK rows in `{rcell('Check', 'n_check')}`; it should be {val('Check', 'n_check', 0)}). If you rebuild the analysis yourself and a row says CHECK, the usual causes are a wrong range (an excluded respondent left in a roster), the wrong `T.TEST` type, or a lookup on the wrong video or expression.

If you have the **Analysis ToolPak** add in, **Data > Data Analysis > t-Test: Paired Two Sample for Means** (H1, drift) and **t-Test: Two-Sample Assuming Unequal Variances** (H3) are good cross checks.

---

## 12. Common mistakes, collected

1. Opening the .csv by double clicking, or letting the `Eye tracking` text become a number or a date.
2. Using the step names with an asterisk in `COUNTIF` (the asterisk is a wildcard); use the step numbers.
3. Screening on the rounded session level accuracy instead of both calibrations (section 2).
4. Treating accuracy like a score: bigger accuracy numbers are worse.
5. Counting reasons only once or only "any": report both columns (section 5).
6. Testing attrition only among completed respondents; the order of every session is in `Block`, so use all sessions (section 4).
7. Using `T.TEST` type 2 or 3 for H1 or drift (paired data), or type 1 for H3 (two groups).
8. Reading a nonsignificant randomization check as proof of equal groups (section 9).
9. Treating the {N_FEA} FEA rows as independent observations: each respondent contributes ten rows, and the tests use one value per respondent and video.
10. Forgetting that every number here is simulated.

## Sources

- iMotions help center, Accessing RDC Study Data: https://help.imotions.com/lab/remote-data-collection/accessing-rdc-study-data
- iMotions help center, Data Collection - Behavioral Research Online (respondent funnel and data quality box): https://help.imotions.com/online/behavioral-research/data-collection-behavioral-research-online
- iMotions help center, Analysis - Behavioral Research Online: https://help.imotions.com/online/behavioral-research/analysis-behavioral-research-online
- iMotions help center, RDC Testing with Webcam Eye Tracking: https://help.imotions.com/lab/remote-data-collection/getting-started-with-online-data-collection/rdc-testing-with-webcam-eye-tracking
- iMotions help center, Participating in RDC Studies: https://help.imotions.com/lab/remote-data-collection/participating-in-rdc-studies
- iMotions help center, Demo study: Food Waste (image, videos and talkaloud): https://help.imotions.com/lab/getting-started/demo-studies/10-0-demo-studies/demo-study-food-waste-image-videos-and-talkaloud
- iMotions help center, R Notebook (FEA): Facial expression analysis thresholding and aggregation: https://help.imotions.com/lab/analysis/facial-expression-analysis/r-notebook-fea-facial-expression-analysis-thresholding-and-aggregation
- iMotions help center, AOI Metrics: https://help.imotions.com/lab/analysis/eye-tracking/areas-of-interest/aoi-metrics
- iMotions help center, Opening Data Export Files in Excel: https://help.imotions.com/lab/exports/additional-export-information/opening-data-export-files-in-excel
- iMotions help center, Creating Excel Pivot Tables: https://help.imotions.com/lab/exports/additional-export-information/creating-excel-pivot-tables
- iMotions Online Study Setup (local manual PDF, "Generally acceptable range: 0-4 dva")
{ref_apa('Delacre2017')}
{ref_apa('Lakens2013')}
{ref_apa('Zhou2016')}

## Items to verify on the lab computer

- The **session data download**: the CSV layout of "Download session data" for an RDC study, its column names (the manuals show the on screen list), the Id format, and which calibration the `Eye tracking` accuracy summarizes. Whether it holds any per session funnel step, device, or browser column; the three columns marked (modeled) are not documented.
- That an RDC study built in Lab 11.1 can hold **two flows** with one link each, and that the `Block` column names the flow.
- The **Status** of a session that stopped before the end (modeled here as Abandoned) and the step at which a **phone** is stopped (modeled as STEP 5); test with a phone during the pilot.
- Where the **per slide quality values** live for an RDC study (portal session page, Lab Exposure summary, or the data quality export); `SIMULATED_quality.csv` is a modeled table.
- The **FEA Metrics Export** header in Lab 11.1, whether the thresholding notebook offers Engagement (for example under "Other Measures thresholding"), and the denominator of "Percent of time present".
- The **survey export** layout for an RDC study ("the survey data export option").
"""
with open(MD_OUT, "w", encoding="utf-8") as fh:
    fh.write(md)
print("wrote", MD_OUT)
