#!/usr/bin/env python3
"""
make_excel_workbook.py
Builds SIMULATED_eeg_analysis.xlsx from the SIMULATED CSV files with openpyxl, and writes
analysis_excel_walkthrough.md from the SAME cell registry, so that every cell reference in the
walkthrough is the address of the formula in the workbook.

Three verification steps run after the workbook is saved:
  1. openpyxl reads the workbook back and checks that every registered cell holds the registered
     formula (the addresses in the walkthrough are real).
  2. If LibreOffice is installed, a headless copy of the workbook is recalculated by LibreOffice
     (in a temporary folder; the delivered workbook is not changed) and every Check row must say OK.
  3. Every result printed in the walkthrough is the LibreOffice value of that cell, and it must
     agree with the value analysis.R wrote to the SIMULATED_*.csv result files at the number of
     decimals printed. The script stops with an error list if anything differs.

The workbook mirrors analysis.R for students who work in Excel:

  README     what the workbook is, and that every number is SIMULATED
  Data       the EEG Frontal Asymmetry Metrics export as exported (columns A to L), with live helper
             columns M to Q (Version, Order, Role, Keep, Keep15)
  Default    the same export made with the notebook DEFAULT channels (teaching contrast)
  Artifact   the artifact report (modeled layout) with a recalculated percent and the gap samples
  Log        the session log (modeled layout)
  Survey     the questionnaire (modeled layout)
  Screening  one row per respondent: largest percent removed at F3 and F4, Red status, decisions
  Summary    raw FA by segment and version (COUNTIFS, AVERAGEIFS, STDEV.S(IF(...)))
  H1         key scene versus its baseline (Version A), the sensitivity checks, the default channels
  H2         end scene of Commercial 1, Version A versus Version B (Welch), sensitivity checks
  H3         Commercial_2 versus Commercial_1 change (paired), simple estimates, liking correlations
  Check      live Excel results next to the values analysis.R produced, with a status

All formulas are ordinary Excel formulas. The STDEV.S(IF(...)) and MAX(IF(...)) formulas are array
formulas; Excel 365 evaluates them with Enter, older Excel needs Ctrl+Shift+Enter (openpyxl writes
them as array formulas already).

Run: python3 make_excel_workbook.py   (after make_simulated_data.py and analysis.R)
"""

import csv
import os
import re
import shutil
import subprocess
import sys
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.data_source import NumDataSource, NumRef
from openpyxl.chart.error_bar import ErrorBars
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.formula import ArrayFormula

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "SIMULATED_eeg_analysis.xlsx")
MD_OUT = os.path.join(HERE, "analysis_excel_walkthrough.md")

NAVY, TEAL, GRAY, ACCENT, LIGHT = "1F3864", "1B7F79", "8C8C8C", "A6561A", "F2F2F2"
hdr_font = Font(bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor=NAVY)
sub_fill = PatternFill("solid", fgColor=LIGHT)
bold = Font(bold=True)
note_font = Font(italic=True, color="555555")
code_font = Font(name="Consolas", size=9, color="555555")
thin = Side(style="thin", color="BFBFBF")
box = Border(top=thin, bottom=thin, left=thin, right=thin)

EXCLUDE_PCT, STRICT_PCT = 25.0, 15.0


def read_csv(name):
    with open(os.path.join(HERE, name), encoding="utf-8") as fh:
        raw = fh.readlines()
    banner = raw[0].strip()
    rows = list(csv.reader([ln for ln in raw if not ln.startswith("#")]))
    return banner, rows[0], rows[1:]


def to_cell(v):
    """Numbers become numbers; NA stays the text NA (as iMotions writes it)."""
    if v in ("NA", ""):
        return v
    try:
        f = float(v)
        return int(f) if f.is_integer() and "." not in v else f
    except ValueError:
        return v


def style_header(ws, row=1, ncol=None, first_col=1):
    ncol = ncol or ws.max_column
    for c in range(first_col, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill, cell.border = hdr_font, hdr_fill, box
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def autowidth(ws, min_w=8, max_w=44):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            L = 12 if str(cell.value).startswith("=") or isinstance(cell.value, ArrayFormula) else len(str(cell.value))
            widths[cell.column_letter] = max(widths.get(cell.column_letter, 0), L)
    for col, w in widths.items():
        ws.column_dimensions[col].width = max(min_w, min(max_w, w + 2))


def show_formula(cell, formula):
    """Write the formula text into a cell as TEXT (not evaluated), without the _xlfn. prefixes."""
    cell.value = formula.replace("_xlfn.", "")
    cell.data_type = "s"
    cell.font = code_font


# registry of documented cells: REG[(sheet, key)] = {"cell", "formula", "label", "array"}
REG = {}
# expected values from analysis.R: EXPECT[(sheet, key)] = (value, decimals shown in the walkthrough)
EXPECT = {}


def reg(sheet, key, cell, formula, label="", array=False):
    REG[(sheet, key)] = {"cell": cell, "formula": formula.replace("_xlfn.", ""), "label": label, "array": array}


# ---------------------------------------------------------------- read the files
banner, fa_h, fa_rows = read_csv("SIMULATED_eeg_fa_metrics.csv")
SEED = re.search(r"seed (\d+)", banner).group(1)
_, dfa_h, dfa_rows = read_csv("SIMULATED_eeg_fa_metrics_default_channels.csv")
_, art_h, art_rows = read_csv("SIMULATED_eeg_artifact_report.csv")
_, log_h, log_rows = read_csv("SIMULATED_eeg_session_log.csv")
_, sv_h, sv_rows = read_csv("SIMULATED_eeg_survey.csv")
assert fa_h == dfa_h
fi = {h: i for i, h in enumerate(fa_h)}
ai = {h: i for i, h in enumerate(art_h)}
li = {h: i for i, h in enumerate(log_h)}
si = {h: i for i, h in enumerate(sv_h)}
N_FA, N_ART, N_LOG, N_SV = len(fa_rows), len(art_rows), len(log_rows), len(sv_rows)
D1, D2 = 2, N_FA + 1                     # Data rows 2 to 313
A2_ = N_ART + 1                          # Artifact rows 2 to 385
L2_ = N_LOG + 1                          # Log rows 2 to 49
S2_ = N_SV + 1                           # Survey rows 2 to 49


def rng_(sheet, col, last):
    return f"{sheet}!${col}$2:${col}${last}"


# Data sheet columns: A to L as exported, then helpers
DCOL = {h: get_column_letter(i + 1) for i, h in enumerate(fa_h)}
HELP = ["Version", "Order", "Role", "Keep", "Keep15"]
for k, h in enumerate(HELP):
    DCOL[h] = get_column_letter(len(fa_h) + 1 + k)
R_RESP, R_GROUP, R_TYPE, R_LABEL, R_PAR = (rng_("Data", DCOL[h], D2) for h in ("Respondent Name", "Group", "Type", "Label", "Parent Stimulus"))
R_FA = rng_("Data", DCOL["Mean Frontal Asymmetry Alpha"], D2)
R_VER, R_ORD, R_ROLE, R_KEEP, R_KEEP15 = (rng_("Data", DCOL[h], D2) for h in HELP)
DR_RESP, DR_FA, DR_ROLE = (rng_("Default", DCOL[h], D2) for h in ("Respondent Name", "Mean Frontal Asymmetry Alpha", "Role"))
ACOL = {h: get_column_letter(i + 1) for i, h in enumerate(art_h)}
RA_RESP, RA_SITE, RA_PCT = (rng_("Artifact", ACOL[h], A2_) for h in ("Respondent Name", "Electrode (assumed)", "Percent removed (%)"))
LCOL = {h: get_column_letter(i + 1) for i, h in enumerate(log_h)}
RL_RESP, RL_F3, RL_F4 = (rng_("Log", LCOL[h], L2_) for h in ("Respondent Name", "NIC2 quality F3 before recording", "NIC2 quality F4 before recording"))
SCOL = {h: get_column_letter(i + 1) for i, h in enumerate(sv_h)}
RS_RESP, RS_L1, RS_L2 = (rng_("Survey", SCOL[h], S2_) for h in ("Respondent Name", "Liking_Commercial_1_1to7", "Liking_Commercial_2_1to7"))

# ---------------------------------------------------------------- exclusions and rosters (static lists; the Check sheet verifies them)
ids = sorted({r[fi["Respondent Name"]] for r in fa_rows})
version = {r[fi["Respondent Name"]]: r[fi["Group"]][0] for r in fa_rows}
order = {r[fi["Respondent Name"]]: r[fi["Group"]][2:] for r in fa_rows}
maxpct = {i: max(float(r[ai["Percent removed (%)"]]) for r in art_rows if r[ai["Respondent Name"]] == i) for i in ids}
red = {r[li["Respondent Name"]]: "Red" in (r[li["NIC2 quality F3 before recording"]], r[li["NIC2 quality F4 before recording"]]) for r in log_rows}
excluded = sorted(i for i in ids if maxpct[i] > EXCLUDE_PCT or red[i])
kept = [i for i in ids if i not in excluded]
drop15 = sorted(i for i in kept if maxpct[i] > STRICT_PCT)
A_strict = [i for i in kept if version[i] == "A" and i not in drop15]
A_drop = [i for i in kept if version[i] == "A" and i in drop15]
B_strict = [i for i in kept if version[i] == "B" and i not in drop15]
B_drop = [i for i in kept if version[i] == "B" and i in drop15]
rosterA = A_strict + A_drop
roster = A_strict + A_drop + B_strict + B_drop
nA, nB, nA15, nB15 = len(rosterA), len(B_strict) + len(B_drop), len(A_strict), len(B_strict)
P1 = 5
PA_LAST = P1 + nA - 1                      # H1: last Version A row
PA15_LAST = P1 + nA15 - 1                  # H1: last Version A row kept at 15 percent
QA1, QA2 = P1, P1 + nA - 1                 # H2/H3: Version A rows
QA15 = P1 + nA15 - 1
QB1, QB2 = QA2 + 1, QA2 + nB
QB15 = QB1 + nB15 - 1
Q_LAST = QB2

# ---------------------------------------------------------------- results from analysis.R
_, th, trows = read_csv("SIMULATED_test_results.csv")
_, mh, mrows = read_csv("SIMULATED_model_results.csv")
_, dh, drows = read_csv("SIMULATED_descriptives.csv")
T = {r[0]: dict(zip(th, r)) for r in trows}
M = {(r[0], r[1]): dict(zip(mh, r)) for r in mrows}
DS = {(r[0], r[1], r[2]): dict(zip(dh, r)) for r in drows}
AN = "Mixed model FA change ~ commercial * version + position + (1 | respondent)"
EX = "Exploratory: whole commercial FA change with liking"


def fnum(x):
    return float(x)


wb = Workbook()

# ---------------------------------------------------------------- README
ws = wb.active
ws.title = "README"
readme = [
    "SIMULATED EEG frontal alpha asymmetry workbook (teaching material)",
    "",
    f"Every number in this workbook is SIMULATED. It was generated by make_simulated_data.py (seed {SEED}) in the style of the",
    "iMotions demo study DEMO - SonyBravia: a commercial with a key scene (Version A, Commercial_1A) against the same commercial with the",
    "key scene cut (Version B, Commercial_1B), between subjects; a second commercial (Commercial_2) for everyone, within subjects; two",
    "neutral 1 minute nature videos as baselines; the last 10 s of each commercial (EndScene). Device: Neuroelectrics Enobio 8 through",
    "NIC2 and Lab Streaming Layer, iMotions device name Enobio EEG, 500 Hz. No real person was recorded.",
    "",
    "FA = ln(alpha power at the RIGHT frontal electrode) - ln(alpha power at the LEFT frontal electrode). Higher = relatively more left",
    "frontal activity, read as approach motivation. ASSUMED channel map (verify in NIC2): Ch5 = F3 (left), Ch6 = F4 (right).",
    "",
    "Sheets:",
    f"  Data       EEG Frontal Asymmetry Metrics export as exported (columns A to L, rows {D1} to {D2}); helper columns {DCOL['Version']} to {DCOL['Keep15']} are live formulas:",
    "             Version, Order, Role (which part of the design the row is), Keep (exclusion rule), Keep15 (stricter sensitivity rule).",
    "  Default    The same export made with the notebook DEFAULT channels (Ch1 = Fz, Ch7 = P3), for the teaching contrast only.",
    "  Artifact   Artifact report (MODELED layout): percent of F3 and F4 samples removed per respondent and stimulus.",
    "  Log        Session log (MODELED layout), including the NIC2 quality status at F3 and F4 before recording.",
    "  Survey     Questionnaire (MODELED layout): liking, key scene recall, brand named correctly.",
    "  Screening  One row per respondent: largest percent removed at F3 and F4, Red status, the decision at 25 and at 15 percent.",
    "  Summary    Raw FA by segment and version (kept respondents).",
    "  H1         Key scene versus the baseline before Commercial_1A (Version A), sensitivity checks, and the default channel contrast.",
    "  H2         End scene of Commercial 1, baseline corrected, Version A versus Version B (Welch t test), sensitivity checks.",
    "  H3         Whole commercial change, Commercial_2 versus Commercial_1 (paired), simple position and version estimates, liking.",
    "  Check      Live Excel results next to the values analysis.R produced, with an OK or CHECK status.",
    "",
    "Exclusion rule, written before collection: more than 25 percent of the F3 or F4 samples removed in any baseline or commercial, or",
    f"a Red NIC2 quality status at F3 or F4 that could not be fixed (excluded: {', '.join(excluded)}).",
    "Each commercial is compared with the baseline that immediately precedes it (Baseline_1 for the first commercial, Baseline_2 for the second).",
    "A segment is identified by Label AND Parent Stimulus: the Label EndScene appears on all three commercials.",
    "MODELED, verify on the lab computer: the artifact report, session log, and survey layouts; the Type values; Parent Stimulus on",
    "stimulus rows; the single Algorithm Device Name column; the channel map.",
    "",
    "Companion files: analysis_excel_walkthrough.md (step by step), analysis.R (same analysis in R),",
    "analysis_output.txt (its console output), results_text.md (model paragraphs).",
]
for i, line in enumerate(readme, start=1):
    ws.cell(row=i, column=1, value=line)
ws["A1"].font = Font(bold=True, size=14, color=NAVY)
ws.column_dimensions["A"].width = 150


# ---------------------------------------------------------------- Data and Default
def role_formula(i):
    F, G, J = DCOL["Type"], DCOL["Label"], DCOL["Parent Stimulus"]
    return (f'=IF({F}{i}="Stimulus",IF(LEFT({G}{i},8)="Baseline",{G}{i},IF(LEFT({G}{i},12)="Commercial_1","C1","C2")),'
            f'IF({G}{i}="KeyScene","KeyScene",IF(LEFT({J}{i},12)="Commercial_1","EndScene_C1","EndScene_C2")))')


def keep_formula(i, limit):
    B = DCOL["Respondent Name"]
    return (f'=IF(COUNTIFS({RA_RESP},{B}{i},{RA_PCT},">{limit:g}")+COUNTIFS({RL_RESP},{B}{i},{RL_F3},"Red")'
            f'+COUNTIFS({RL_RESP},{B}{i},{RL_F4},"Red")>0,"No","Yes")')


def build_export_sheet(name, rows):
    ws = wb.create_sheet(name)
    ws.append(fa_h + HELP)
    style_header(ws)
    E = DCOL["Group"]
    for i, r in enumerate(rows, start=D1):
        ws.append([to_cell(v) for v in r])
        ws[f"{DCOL['Version']}{i}"] = f"=LEFT({E}{i},1)"
        ws[f"{DCOL['Order']}{i}"] = f"=MID({E}{i},3,10)"
        ws[f"{DCOL['Role']}{i}"] = role_formula(i)
        ws[f"{DCOL['Keep']}{i}"] = keep_formula(i, EXCLUDE_PCT)
        ws[f"{DCOL['Keep15']}{i}"] = keep_formula(i, STRICT_PCT)
        ws[f"{DCOL['Mean Frontal Asymmetry Alpha']}{i}"].number_format = "0.0000"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{DCOL['Keep15']}{D2}"
    autowidth(ws, max_w=22)
    for h in HELP:
        ws[f"{DCOL[h]}1"].fill = PatternFill("solid", fgColor=ACCENT)
    return ws


ws = build_export_sheet("Data", fa_rows)
for key, h in (("version", "Version"), ("order", "Order"), ("role", "Role"), ("keep", "Keep"), ("keep15", "Keep15")):
    reg("Data", key, f"{DCOL[h]}{D1}", ws[f"{DCOL[h]}{D1}"].value, h)
ws = build_export_sheet("Default", dfa_rows)
reg("Default", "role", f"{DCOL['Role']}{D1}", ws[f"{DCOL['Role']}{D1}"].value, "Role")

# ---------------------------------------------------------------- Artifact, Log, Survey
ws = wb.create_sheet("Artifact")
ws.append(art_h + ["Percent recalculated", "Gap samples (expected - recorded)"])
style_header(ws)
cH, cI, cJ = ACOL["Samples expected"], ACOL["Samples recorded"], ACOL["Samples removed"]
cRecalc, cGap = get_column_letter(len(art_h) + 1), get_column_letter(len(art_h) + 2)
for i, r in enumerate(art_rows, start=2):
    ws.append([to_cell(v) for v in r])
    ws[f"{cRecalc}{i}"] = f"=100*{cJ}{i}/{cI}{i}"
    ws[f"{cRecalc}{i}"].number_format = "0.00"
    ws[f"{cGap}{i}"] = f"={cH}{i}-{cI}{i}"
for c in (cRecalc, cGap):
    ws[f"{c}1"].fill = PatternFill("solid", fgColor=ACCENT)
ws.freeze_panes = "C2"
ws.auto_filter.ref = f"A1:{cGap}{A2_}"
autowidth(ws, max_w=22)
reg("Artifact", "recalc", f"{cRecalc}2", f"=100*{cJ}2/{cI}2", "Percent recalculated")
reg("Artifact", "gap", f"{cGap}2", f"={cH}2-{cI}2", "Gap samples")

ws = wb.create_sheet("Log")
ws.append(log_h)
style_header(ws)
for r in log_rows:
    ws.append([to_cell(v) for v in r])
ws.freeze_panes = "B2"
autowidth(ws, max_w=40)
ws = wb.create_sheet("Survey")
ws.append(sv_h)
style_header(ws)
for r in sv_rows:
    ws.append([to_cell(v) for v in r])
ws.freeze_panes = "C2"
autowidth(ws, max_w=26)

# ---------------------------------------------------------------- Screening
ws = wb.create_sheet("Screening")
ws["A1"] = "Exclusion screening, one row per respondent. SIMULATED data."
ws["A1"].font = Font(bold=True, size=12, color=NAVY)
ws["A2"] = ("Largest percent of samples removed at F3 and at F4 over the two baselines and two commercials (array formulas MAX(IF(...))), "
            "the Red count from the session log, and the decision under the 25 percent rule and the 15 percent sensitivity rule.")
ws["A2"].font = note_font
sc_cols = ["Respondent Name", "Version", "Largest percent removed, F3", "Largest percent removed, F4", "Largest of the two",
           "Red at F3 or F4 before recording (count)", "Decision (25 percent rule)", "Decision (15 percent rule)"]
ws.append([])
ws.append(sc_cols)
style_header(ws, row=4, ncol=len(sc_cols))
ws.row_dimensions[4].height = 46
SC1, SC2 = 5, 5 + len(ids) - 1
for k, rid in enumerate(ids):
    r = SC1 + k
    ws.cell(row=r, column=1, value=rid)
    ws.cell(row=r, column=2, value=f"=INDEX({R_VER},MATCH($A{r},{R_RESP},0))")
    ws.cell(row=r, column=3).value = ArrayFormula(f"C{r}", f'=MAX(IF(({RA_RESP}=$A{r})*({RA_SITE}="F3"),{RA_PCT}))')
    ws.cell(row=r, column=4).value = ArrayFormula(f"D{r}", f'=MAX(IF(({RA_RESP}=$A{r})*({RA_SITE}="F4"),{RA_PCT}))')
    ws.cell(row=r, column=5, value=f"=MAX(C{r}:D{r})")
    ws.cell(row=r, column=6, value=f'=COUNTIFS({RL_RESP},$A{r},{RL_F3},"Red")+COUNTIFS({RL_RESP},$A{r},{RL_F4},"Red")')
    ws.cell(row=r, column=7, value=f'=IF(OR(E{r}>{EXCLUDE_PCT:g},F{r}>0),"Exclude","Keep")')
    ws.cell(row=r, column=8, value=f'=IF(OR(E{r}>{STRICT_PCT:g},F{r}>0),"Exclude","Keep")')
    for c in range(1, 9):
        ws.cell(row=r, column=c).border = box
    for c in (3, 4, 5):
        ws.cell(row=r, column=c).number_format = "0.00"
reg("Screening", "ver", f"B{SC1}", f"=INDEX({R_VER},MATCH($A{SC1},{R_RESP},0))", "Version")
reg("Screening", "f3", f"C{SC1}", f'=MAX(IF(({RA_RESP}=$A{SC1})*({RA_SITE}="F3"),{RA_PCT}))', "Largest percent removed, F3", array=True)
reg("Screening", "f4", f"D{SC1}", f'=MAX(IF(({RA_RESP}=$A{SC1})*({RA_SITE}="F4"),{RA_PCT}))', "Largest percent removed, F4", array=True)
reg("Screening", "big", f"E{SC1}", f"=MAX(C{SC1}:D{SC1})", "Largest of the two")
reg("Screening", "red", f"F{SC1}", f'=COUNTIFS({RL_RESP},$A{SC1},{RL_F3},"Red")+COUNTIFS({RL_RESP},$A{SC1},{RL_F4},"Red")', "Red count")
reg("Screening", "dec25", f"G{SC1}", f'=IF(OR(E{SC1}>{EXCLUDE_PCT:g},F{SC1}>0),"Exclude","Keep")', "Decision (25 percent rule)")
reg("Screening", "dec15", f"H{SC1}", f'=IF(OR(E{SC1}>{STRICT_PCT:g},F{SC1}>0),"Exclude","Keep")', "Decision (15 percent rule)")
cnt0 = SC2 + 2
counts = [("n_ex25", "Respondents excluded (25 percent rule)", f'=COUNTIF(G{SC1}:G{SC2},"Exclude")', len(excluded)),
          ("n_ex15", "Respondents excluded (15 percent rule)", f'=COUNTIF(H{SC1}:H{SC2},"Exclude")', len(excluded) + len(drop15)),
          ("n_keepA", "Kept, Version A (25 percent rule)", f'=COUNTIFS(B{SC1}:B{SC2},"A",G{SC1}:G{SC2},"Keep")', nA),
          ("n_keepB", "Kept, Version B (25 percent rule)", f'=COUNTIFS(B{SC1}:B{SC2},"B",G{SC1}:G{SC2},"Keep")', nB),
          ("n_keepA15", "Kept, Version A (15 percent rule)", f'=COUNTIFS(B{SC1}:B{SC2},"A",H{SC1}:H{SC2},"Keep")', nA15),
          ("n_keepB15", "Kept, Version B (15 percent rule)", f'=COUNTIFS(B{SC1}:B{SC2},"B",H{SC1}:H{SC2},"Keep")', nB15),
          ("n_gap", "Artifact report rows with a stream gap (recorded < expected)", f'=COUNTIF(Artifact!${cGap}$2:${cGap}${A2_},">0")',
           sum(1 for r in art_rows if int(r[ai["Samples recorded"]]) < int(r[ai["Samples expected"]])))]
for k, (key, lab, f, val) in enumerate(counts):
    r = cnt0 + k
    ws.cell(row=r, column=1, value=lab).font = bold
    ws.cell(row=r, column=5, value=f)
    reg("Screening", key, f"E{r}", f, lab)
    if val is not None:
        EXPECT[("Screening", key)] = (val, 0)
for col, wdt in (("A", 52), ("B", 10), ("C", 16), ("D", 16), ("E", 14), ("F", 18), ("G", 18), ("H", 18)):
    ws.column_dimensions[col].width = wdt

# ---------------------------------------------------------------- Summary
ws = wb.create_sheet("Summary")
ws["A1"] = "Raw FA by segment and version, kept respondents (Keep = Yes). SIMULATED data."
ws["A1"].font = Font(bold=True, size=12, color=NAVY)
ws["A2"] = ("Role names the part of the design (C1 = Commercial_1A or Commercial_1B; EndScene_C1 = the EndScene annotation whose Parent Stimulus "
            "is Commercial_1A or Commercial_1B). Version * = both versions. Columns D and E use AVERAGEIFS and STDEV.S(IF(...)).")
ws["A2"].font = note_font
sm_cols = ["Role", "Version", "n", "Mean FA", "SD FA"]
ws.append([])
ws.append(sm_cols)
style_header(ws, row=4, ncol=len(sm_cols))
ROLE_SEG = {"Baseline_1": "Baseline_1", "Baseline_2": "Baseline_2", "C1": "Commercial_1 (1A or 1B)", "C2": "Commercial_2",
            "KeyScene": "KeyScene (Commercial_1A)", "EndScene_C1": "EndScene of Commercial_1", "EndScene_C2": "EndScene of Commercial_2"}
r = 5
SUM_FIRST = r
summary_rows = {}
for role in ROLE_SEG:
    for ver in ("A", "B", "*"):
        if role == "KeyScene" and ver != "A":
            continue
        ws.cell(row=r, column=1, value=role)
        ws.cell(row=r, column=2, value=ver)
        crit = f'{R_ROLE},$A{r},{R_VER},$B{r},{R_KEEP},"Yes"'
        f_n = f"=COUNTIFS({crit})"
        f_m = f"=AVERAGEIFS({R_FA},{crit})"
        if ver == "*":
            f_s = f'=_xlfn.STDEV.S(IF(({R_ROLE}=$A{r})*({R_KEEP}="Yes"),{R_FA}))'
        else:
            f_s = f'=_xlfn.STDEV.S(IF(({R_ROLE}=$A{r})*({R_VER}=$B{r})*({R_KEEP}="Yes"),{R_FA}))'
        ws.cell(row=r, column=3, value=f_n)
        ws.cell(row=r, column=4, value=f_m).number_format = "0.0000"
        ws.cell(row=r, column=5).value = ArrayFormula(f"E{r}", f_s)
        ws.cell(row=r, column=5).number_format = "0.0000"
        for c in range(1, 6):
            ws.cell(row=r, column=c).border = box
        d = DS[("FA (raw)", ROLE_SEG[role], "All" if ver == "*" else ver)]
        key = f"{role}_{'All' if ver == '*' else ver}"
        reg("Summary", f"n_{key}", f"C{r}", f_n, "n")
        reg("Summary", f"m_{key}", f"D{r}", f_m, "Mean FA")
        reg("Summary", f"s_{key}", f"E{r}", f_s, "SD FA", array=True)
        EXPECT[("Summary", f"n_{key}")] = (int(d["n"]), 0)
        EXPECT[("Summary", f"m_{key}")] = (fnum(d["mean"]), 3)
        EXPECT[("Summary", f"s_{key}")] = (fnum(d["sd"]), 3)
        summary_rows[key] = r
        r += 1
SUM_LAST = r - 1
ws[f"A{SUM_LAST + 2}"] = ("Array formulas (column E): Excel 365 evaluates them with Enter; Excel 2019 or older needs Ctrl+Shift+Enter after editing. "
                          "The text NA in the FA column (a segment with no clean window) is ignored by AVERAGEIFS and COUNTIFS counts the row; "
                          "no kept respondent has an NA.")
ws[f"A{SUM_LAST + 2}"].font = note_font
for col, wdt in (("A", 16), ("B", 10), ("C", 8), ("D", 12), ("E", 12)):
    ws.column_dimensions[col].width = wdt


# ---------------------------------------------------------------- blocks
class Block:
    """Writes label, value, formula text, and note columns for one block and records every cell in REG."""

    def __init__(self, ws, sheet, label_col, start_row, prefix):
        self.ws, self.sheet, self.lc, self.prefix = ws, sheet, label_col, prefix
        self.row = start_row
        self.ref = {}
        for k, name in enumerate(("Statistic", "Value", "Formula shown", "Note")):
            c = ws.cell(row=start_row - 1, column=label_col + k, value=name)
            c.font, c.fill = hdr_font, hdr_fill

    def title(self, text):
        c = self.ws.cell(row=self.row, column=self.lc, value=text)
        c.font, c.fill = bold, sub_fill
        self.row += 1

    def put(self, key, label, formula, note="", fmt="0.0000", expect=None, nd=3):
        f = formula.format(**{k: f"{get_column_letter(self.lc + 1)}{v}" for k, v in self.ref.items()})
        i = self.row
        self.ws.cell(row=i, column=self.lc, value=label)
        self.ws.cell(row=i, column=self.lc + 1, value=f).number_format = fmt
        show_formula(self.ws.cell(row=i, column=self.lc + 2), f)
        if note:
            self.ws.cell(row=i, column=self.lc + 3, value=note).font = note_font
        self.ref[key] = i
        reg(self.sheet, f"{self.prefix}_{key}", f"{get_column_letter(self.lc + 1)}{i}", f, label)
        if expect is not None:
            EXPECT[(self.sheet, f"{self.prefix}_{key}")] = (expect, nd)
        self.row += 1

    def blank(self):
        self.row += 1

    def addr(self, key):
        return f"{get_column_letter(self.lc + 1)}{self.ref[key]}"

    def next_start(self):
        return self.row + 2


def change_block(blk, title, r1, r2, rd, nm1, nm2, t, note_n=""):
    """One sample t test on the change (rd = r2 - r1), the same as a paired t test; t = the R result row."""
    blk.title(title)
    blk.put("m1", f"{nm1} mean (FA)", f"=AVERAGE({r1})", "AVERAGE of the reference column", expect=fnum(t["m1"]))
    blk.put("s1", f"{nm1} SD", f"=_xlfn.STDEV.S({r1})", "STDEV.S = sample standard deviation", expect=fnum(t["sd1"]))
    blk.put("m2", f"{nm2} mean (FA)", f"=AVERAGE({r2})", expect=fnum(t["m2"]))
    blk.put("s2", f"{nm2} SD", f"=_xlfn.STDEV.S({r2})", expect=fnum(t["sd2"]))
    blk.put("n", "n (respondents)", f"=COUNT({rd})", note_n, fmt="0", expect=int(t["n"]), nd=0)
    blk.put("md", f"Mean change, {nm2} minus {nm1}", f"=AVERAGE({rd})", expect=fnum(t["mean_diff"]))
    blk.put("sdd", "SD of the changes", f"=_xlfn.STDEV.S({rd})", "the SD that matters in a paired design", expect=fnum(t["sd_diff"]))
    blk.put("se", "SE of the mean change", "={sdd}/SQRT({n})", "SD of the changes / sqrt(n)")
    blk.put("t", "t", "={md}/{se}", "t = mean change / SE", expect=fnum(t["t"]), nd=2)
    blk.put("df", "df", "={n}-1", "n minus 1", fmt="0", expect=int(float(t["df"])), nd=0)
    blk.put("p_tt", "p, two tailed (T.TEST type 1)", f"=_xlfn.T.TEST({r1},{r2},2,1)", "tails = 2, type = 1 means paired",
            fmt="0.000000", expect=fnum(t["p"]), nd=5)
    blk.put("p_hand", "p from the hand computed t", "=_xlfn.T.DIST.2T(ABS({t}),{df})", "must match the T.TEST cell", fmt="0.000000",
            expect=fnum(t["p"]), nd=5)
    blk.put("p_one", "p, one tailed (teaching only)", "=_xlfn.T.DIST.RT({t},{df})", "right tail: the predicted direction is positive",
            fmt="0.000000", expect=fnum(t["p_one_tailed"]), nd=5)
    blk.put("tcrit", "Critical t (95%)", "=_xlfn.T.INV.2T(0.05,{df})", "two tailed, alpha = .05")
    blk.put("ci_lo", "Mean change 95% CI lower", "={md}-{tcrit}*{se}", expect=fnum(t["diff_ci_low"]))
    blk.put("ci_hi", "Mean change 95% CI upper", "={md}+{tcrit}*{se}", expect=fnum(t["diff_ci_high"]))
    blk.put("dz", "Cohen's d_z", "={md}/{sdd}", "d_z = mean change / SD of the changes", expect=fnum(t["es"]), nd=2)
    blk.put("se_dz", "SE of d_z (approximation)", "=SQRT(1/{n}+{dz}^2/(2*{n}))", "sqrt(1/n + d_z^2 / (2 n))")
    blk.put("dz_lo", "d_z 95% CI lower (approximation)", "={dz}-1.96*{se_dz}", "R also gives the exact noncentral t interval",
            expect=fnum(t["es_ci_low_approx"]), nd=2)
    blk.put("dz_hi", "d_z 95% CI upper (approximation)", "={dz}+1.96*{se_dz}", expect=fnum(t["es_ci_high_approx"]), nd=2)
    blk.put("r", "Correlation of the two columns", f"=CORREL({r1},{r2})", "how stable a respondent is from one segment to the other")
    blk.put("npos", "Respondents with a positive change", f'=COUNTIF({rd},">0")', fmt="0")
    blk.blank()


def zero_block(blk, title, rx, t):
    """Raw scores against zero (the sensitivity check without baseline correction)."""
    blk.title(title)
    blk.put("n", "n", f"=COUNT({rx})", fmt="0", expect=int(t["n"]), nd=0)
    blk.put("m", "Mean (FA)", f"=AVERAGE({rx})", expect=fnum(t["m2"]))
    blk.put("s", "SD", f"=_xlfn.STDEV.S({rx})", expect=fnum(t["sd2"]))
    blk.put("se", "SE of the mean", "={s}/SQRT({n})")
    blk.put("t", "t against 0", "={m}/{se}", expect=fnum(t["t"]), nd=2)
    blk.put("df", "df", "={n}-1", fmt="0", expect=int(float(t["df"])), nd=0)
    blk.put("p", "p, two tailed", "=_xlfn.T.DIST.2T(ABS({t}),{df})", fmt="0.000000", expect=fnum(t["p"]), nd=5)
    blk.put("d", "Cohen's d (one sample)", "={m}/{s}", expect=fnum(t["es"]), nd=2)
    blk.blank()


def welch_block(blk, title, ra, rb, t, note_n=""):
    blk.title(title)
    blk.put("m1", "Version A mean", f"=AVERAGE({ra})", expect=fnum(t["m1"]))
    blk.put("s1", "Version A SD", f"=_xlfn.STDEV.S({ra})", expect=fnum(t["sd1"]))
    blk.put("n1", "Version A n", f"=COUNT({ra})", note_n, fmt="0", expect=int(t["n1"]), nd=0)
    blk.put("m2", "Version B mean", f"=AVERAGE({rb})", expect=fnum(t["m2"]))
    blk.put("s2", "Version B SD", f"=_xlfn.STDEV.S({rb})", expect=fnum(t["sd2"]))
    blk.put("n2", "Version B n", f"=COUNT({rb})", fmt="0", expect=int(t["n2"]), nd=0)
    blk.put("diff", "Difference, Version A minus Version B", "={m1}-{m2}", expect=fnum(t["mean_diff"]))
    blk.put("se", "SE of the difference (Welch)", "=SQRT({s1}^2/{n1}+{s2}^2/{n2})", "each group keeps its own variance")
    blk.put("t", "Welch t", "={diff}/{se}", expect=fnum(t["t"]), nd=2)
    blk.put("df", "Welch df", "=({s1}^2/{n1}+{s2}^2/{n2})^2/(({s1}^2/{n1})^2/({n1}-1)+({s2}^2/{n2})^2/({n2}-1))",
            "Welch Satterthwaite formula; not a whole number", fmt="0.00", expect=fnum(t["df"]), nd=2)
    blk.put("p_tt", "p, two tailed (T.TEST type 3)", f"=_xlfn.T.TEST({ra},{rb},2,3)", "type 3 = two groups, unequal variances (Welch)",
            fmt="0.000000", expect=fnum(t["p"]), nd=5)
    blk.put("p_hand", "p from the hand computed t", "=_xlfn.T.DIST.2T(ABS({t}),{df})",
            "may round the df down to a whole number; report the T.TEST cell", fmt="0.000000")
    blk.put("p_one", "p, one tailed, predicted direction A > B (teaching only)", "=IF({t}>0,{p_tt}/2,1-{p_tt}/2)",
            "half the two tailed p when t points the predicted way; T.TEST with tails = 1 ignores the direction", fmt="0.000000",
            expect=fnum(t["p_one_tailed"]), nd=5)
    blk.put("tcrit", "Critical t (95%)", "=_xlfn.T.INV.2T(0.05,{df})", "may also round the df down")
    blk.put("ci_lo", "Difference 95% CI lower", "={diff}-{tcrit}*{se}", expect=fnum(t["diff_ci_low"]))
    blk.put("ci_hi", "Difference 95% CI upper", "={diff}+{tcrit}*{se}", expect=fnum(t["diff_ci_high"]))
    blk.put("sp", "Pooled SD (for the effect size)", "=SQRT((({n1}-1)*{s1}^2+({n2}-1)*{s2}^2)/({n1}+{n2}-2))", expect=fnum(t["sd_pooled"]))
    blk.put("d", "Cohen's d", "={diff}/{sp}", "difference / pooled SD", expect=fnum(t["es"]), nd=2)
    blk.put("J", "Small sample correction J", "=1-3/(4*({n1}+{n2})-9)", "Hedges' approximation")
    blk.put("g", "Hedges' g", "={J}*{d}", expect=fnum(t["g"]), nd=2)
    blk.put("se_g", "SE of g (approximation)", "=SQRT(({n1}+{n2})/({n1}*{n2})+{g}^2/(2*({n1}+{n2})))")
    blk.put("g_lo", "g 95% CI lower (approximation)", "={g}-1.96*{se_g}", "R also gives the exact noncentral t interval",
            expect=fnum(t["g_ci_low_approx"]), nd=2)
    blk.put("g_hi", "g 95% CI upper (approximation)", "={g}+1.96*{se_g}", expect=fnum(t["g_ci_high_approx"]), nd=2)
    blk.blank()


def corr_block(blk, title, rx, ry, m):
    blk.title(title)
    blk.put("n", "n", f"=COUNT({rx})", fmt="0")
    blk.put("r", "Pearson r (CORREL)", f"=CORREL({rx},{ry})", expect=fnum(m["estimate"]), nd=2)
    blk.put("t", "t for r", "={r}*SQRT(({n}-2)/(1-{r}^2))", "t = r sqrt((n - 2) / (1 - r^2))", expect=fnum(m["statistic"]), nd=2)
    blk.put("p", "p, two tailed", "=_xlfn.T.DIST.2T(ABS({t}),{n}-2)", "df = n - 2", fmt="0.000000", expect=fnum(m["p"]), nd=5)
    blk.put("lo", "r 95% CI lower (Fisher z)", "=TANH(ATANH({r})-_xlfn.NORM.S.INV(0.975)/SQRT({n}-3))", expect=fnum(m["ci_low"]), nd=2)
    blk.put("hi", "r 95% CI upper (Fisher z)", "=TANH(ATANH({r})+_xlfn.NORM.S.INV(0.975)/SQRT({n}-3))", expect=fnum(m["ci_high"]), nd=2)
    blk.blank()


def lookup(sheet_rng, r):
    return f"=INDEX({sheet_rng},MATCH($A{r},{R_RESP},0))"


def fa_of(r, role, src="Data"):
    rr, rf, ro = (R_RESP, R_FA, R_ROLE) if src == "Data" else (DR_RESP, DR_FA, DR_ROLE)
    return f'SUMIFS({rf},{rr},$A{r},{ro},"{role}")'


def bl_before(r, which, ord_col, src="Data"):
    """Baseline that precedes Commercial 1 (which = 1) or Commercial 2 (which = 2)."""
    first, second = ("Baseline_1", "Baseline_2") if which == 1 else ("Baseline_2", "Baseline_1")
    return f'=IF(${ord_col}{r}="C1first",{fa_of(r, first, src)},{fa_of(r, second, src)})'


# ---------------------------------------------------------------- H1
ws = wb.create_sheet("H1")
SH = "H1"
ws["A1"] = "H1. FA during the KeyScene minus FA during the baseline before Commercial_1A, Version A, kept respondents. SIMULATED data."
ws["A1"].font = Font(bold=True, size=12, color=NAVY)
ws["A2"] = (f"One row per kept Version A respondent; rows {P1} to {PA15_LAST} are also kept under the 15 percent rule, rows {PA15_LAST + 1} to "
            f"{PA_LAST} are dropped by it. Column D picks Baseline_1 or Baseline_2, whichever came right before Commercial_1A. Columns G to I "
            "repeat the computation on the Default sheet (notebook default channels Ch1 and Ch7).")
ws["A2"].font = note_font
h1_cols = ["Respondent Name", "Order", "Keep15", "Baseline before Commercial_1A (FA)", "KeyScene (FA)", "Change (KeyScene minus baseline)",
           "DEFAULT channels: baseline", "DEFAULT channels: KeyScene", "DEFAULT channels: change"]
ws.append([])
ws.append(h1_cols)
style_header(ws, row=4, ncol=len(h1_cols))
ws.row_dimensions[4].height = 48
for k, rid in enumerate(rosterA):
    r = P1 + k
    ws.cell(row=r, column=1, value=rid)
    ws.cell(row=r, column=2, value=lookup(R_ORD, r))
    ws.cell(row=r, column=3, value=lookup(R_KEEP15, r))
    ws.cell(row=r, column=4, value=bl_before(r, 1, "B"))
    ws.cell(row=r, column=5, value="=" + fa_of(r, "KeyScene"))
    ws.cell(row=r, column=6, value=f"=E{r}-D{r}")
    ws.cell(row=r, column=7, value=bl_before(r, 1, "B", "Default"))
    ws.cell(row=r, column=8, value="=" + fa_of(r, "KeyScene", "Default"))
    ws.cell(row=r, column=9, value=f"=H{r}-G{r}")
    for c in range(1, 10):
        ws.cell(row=r, column=c).border = box
        if c >= 4:
            ws.cell(row=r, column=c).number_format = "0.0000"
for key, col, lab in (("row_ord", "B", "Order"), ("row_k15", "C", "Keep15"), ("row_bl", "D", "Baseline before Commercial_1A"),
                      ("row_ks", "E", "KeyScene"), ("row_chg", "F", "Change"), ("row_dbl", "G", "DEFAULT baseline"),
                      ("row_dks", "H", "DEFAULT KeyScene"), ("row_dchg", "I", "DEFAULT change")):
    reg(SH, key, f"{col}{P1}", ws[f"{col}{P1}"].value, lab)
rA = lambda c, last: f"{c}{P1}:{c}{last}"
blk = Block(ws, SH, 11, 5, "h1")                     # label K, value L, formula M, note N
change_block(blk, f"H1: one sample t on the change (= paired t), Version A, n = {nA}", rA("D", PA_LAST), rA("E", PA_LAST), rA("F", PA_LAST),
             "Baseline", "KeyScene", T["H1"], "COUNT of the change column; must equal the check below")
H1_MAIN = blk
blk = Block(ws, SH, 11, blk.next_start(), "s15")
change_block(blk, f"Sensitivity: 15 percent threshold (rows {P1} to {PA15_LAST})", rA("D", PA15_LAST), rA("E", PA15_LAST), rA("F", PA15_LAST),
             "Baseline", "KeyScene", T["S15_H1"], "the respondents dropped at 15 percent sit below this range")
H1_S15 = blk
blk = Block(ws, SH, 11, blk.next_start(), "raw")
zero_block(blk, "Sensitivity: no baseline correction (raw KeyScene FA against 0)", rA("E", PA_LAST), T["RAW_H1"])
H1_RAW = blk
blk = Block(ws, SH, 11, blk.next_start(), "def")
change_block(blk, "Teaching contrast: the same test with the DEFAULT channels (Ch1 = Fz, Ch7 = P3)", rA("G", PA_LAST), rA("H", PA_LAST),
             rA("I", PA_LAST), "Baseline", "KeyScene", T["DEFAULT_H1"])
H1_DEF = blk
blk = Block(ws, SH, 11, blk.next_start(), "chk")
blk.title("Checks")
blk.put("n_data", "Kept Version A respondents with a KeyScene row in Data (Keep = Yes)", f'=COUNTIFS({R_ROLE},"KeyScene",{R_KEEP},"Yes")',
        "must equal n of the H1 block", fmt="0", expect=nA, nd=0)
blk.put("n_drop", "Roster rows dropped under the 15 percent rule (Keep15 = No)", f'=COUNTIF(C{P1}:C{PA_LAST},"No")', fmt="0",
        expect=len(A_drop), nd=0)
blk.put("n_drop_top", "Keep15 = No inside the rows of the 15 percent block", f'=COUNTIF(C{P1}:C{PA15_LAST},"No")', "must be 0", fmt="0",
        expect=0, nd=0)
blk.put("r_def", "Correlation of the default channel change with the F3/F4 change", f"=CORREL(F{P1}:F{PA_LAST},I{P1}:I{PA_LAST})")
H1_CHK = blk
for col, wdt in (("A", 17), ("B", 11), ("C", 9), ("D", 16), ("E", 12), ("F", 16), ("G", 14), ("H", 14), ("I", 14), ("J", 3), ("K", 58),
                 ("L", 12), ("M", 70), ("N", 56)):
    ws.column_dimensions[col].width = wdt
cr = blk.row + 2
ws.cell(row=cr, column=11, value="Chart table: mean FA with 95% CI half width (Version A)").font = bold
ws.cell(row=cr + 1, column=11, value="Segment")
ws.cell(row=cr + 1, column=12, value="Mean FA")
ws.cell(row=cr + 1, column=13, value="CI half width")
for k, (lab, mk, sk) in enumerate((("Baseline before Commercial_1A", "m1", "s1"), ("KeyScene", "m2", "s2"))):
    rr = cr + 2 + k
    ws.cell(row=rr, column=11, value=lab)
    ws.cell(row=rr, column=12, value=f"={H1_MAIN.addr(mk)}").number_format = "0.000"
    f_ci = f"=_xlfn.CONFIDENCE.T(0.05,{H1_MAIN.addr(sk)},{H1_MAIN.addr('n')})"
    ws.cell(row=rr, column=13, value=f_ci).number_format = "0.000"
    reg(SH, f"chart_ci_{k}", f"M{rr}", f_ci, f"CI half width {lab}")
chart = BarChart()
chart.type, chart.style = "col", 10
chart.title = "Mean FA, baseline and key scene, with 95% CI (SIMULATED)"
chart.y_axis.title = "FA = ln(F4 alpha) - ln(F3 alpha)"
chart.add_data(Reference(ws, min_col=12, min_row=cr + 1, max_row=cr + 3), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=11, min_row=cr + 2, max_row=cr + 3))
chart.series[0].graphicalProperties.solidFill = NAVY
chart.series[0].errBars = ErrorBars(errDir="y", errBarType="both", errValType="cust",
                                    plus=NumDataSource(numRef=NumRef(f=f"H1!$M${cr + 2}:$M${cr + 3}")),
                                    minus=NumDataSource(numRef=NumRef(f=f"H1!$M${cr + 2}:$M${cr + 3}")))
chart.legend = None
chart.height, chart.width = 8, 13
ws.add_chart(chart, f"O{cr}")
H1_CHART_ROW = cr

# ---------------------------------------------------------------- H2
ws = wb.create_sheet("H2")
SH = "H2"
ws["A1"] = "H2. Baseline corrected FA during the EndScene of Commercial 1, Version A versus Version B, kept respondents. SIMULATED data."
ws["A1"].font = Font(bold=True, size=12, color=NAVY)
ws["A2"] = (f"Version A in rows {QA1} to {QA2} (rows {QA15 + 1} to {QA2} are dropped by the 15 percent rule), Version B in rows {QB1} to {QB2} "
            f"(rows {QB15 + 1} to {QB2} dropped by the 15 percent rule). Change = EndScene of Commercial 1 minus the baseline before Commercial 1.")
ws["A2"].font = note_font
h2_cols = ["Respondent Name", "Version", "Order", "Keep15", "Baseline before Commercial 1 (FA)", "EndScene of Commercial 1 (FA)",
           "Change (EndScene minus baseline)"]
ws.append([])
ws.append(h2_cols)
style_header(ws, row=4, ncol=len(h2_cols))
ws.row_dimensions[4].height = 48
for k, rid in enumerate(roster):
    r = P1 + k
    ws.cell(row=r, column=1, value=rid)
    ws.cell(row=r, column=2, value=lookup(R_VER, r))
    ws.cell(row=r, column=3, value=lookup(R_ORD, r))
    ws.cell(row=r, column=4, value=lookup(R_KEEP15, r))
    ws.cell(row=r, column=5, value=bl_before(r, 1, "C"))
    ws.cell(row=r, column=6, value="=" + fa_of(r, "EndScene_C1"))
    ws.cell(row=r, column=7, value=f"=F{r}-E{r}")
    for c in range(1, 8):
        ws.cell(row=r, column=c).border = box
        if c >= 5:
            ws.cell(row=r, column=c).number_format = "0.0000"
for key, col, lab in (("row_ver", "B", "Version"), ("row_ord", "C", "Order"), ("row_k15", "D", "Keep15"),
                      ("row_bl", "E", "Baseline before Commercial 1"), ("row_es", "F", "EndScene of Commercial 1"), ("row_chg", "G", "Change")):
    reg(SH, key, f"{col}{P1}", ws[f"{col}{P1}"].value, lab)
blk = Block(ws, SH, 9, 5, "h2")                      # label I, value J, formula K, note L
welch_block(blk, f"H2: Welch t test, Version A (n = {nA}) versus Version B (n = {nB})", f"G{QA1}:G{QA2}", f"G{QB1}:G{QB2}", T["H2"],
            "must equal the check below")
H2_MAIN = blk
blk = Block(ws, SH, 9, blk.next_start(), "s15")
welch_block(blk, f"Sensitivity: 15 percent threshold (rows {QA1} to {QA15} and {QB1} to {QB15})", f"G{QA1}:G{QA15}", f"G{QB1}:G{QB15}", T["S15_H2"])
H2_S15 = blk
blk = Block(ws, SH, 9, blk.next_start(), "raw")
welch_block(blk, "Sensitivity: no baseline correction (raw EndScene FA, column F)", f"F{QA1}:F{QA2}", f"F{QB1}:F{QB2}", T["RAW_H2"])
H2_RAW = blk
blk = Block(ws, SH, 9, blk.next_start(), "chk")
blk.title("Checks")
blk.put("n_data", "Kept respondents with an EndScene_C1 row in Data (Keep = Yes)", f'=COUNTIFS({R_ROLE},"EndScene_C1",{R_KEEP},"Yes")',
        "must equal n1 + n2", fmt="0", expect=nA + nB, nd=0)
blk.put("ver_top", "Version B rows inside the Version A range", f'=COUNTIF(B{QA1}:B{QA2},"B")', "must be 0", fmt="0", expect=0, nd=0)
blk.put("ver_bot", "Version A rows inside the Version B range", f'=COUNTIF(B{QB1}:B{QB2},"A")', "must be 0", fmt="0", expect=0, nd=0)
blk.put("bl_diff", "Baseline before Commercial 1: Version A mean minus Version B mean", f"=AVERAGE(E{QA1}:E{QA2})-AVERAGE(E{QB1}:E{QB2})",
        "a chance difference that the raw comparison carries along")
H2_CHK = blk
for col, wdt in (("A", 17), ("B", 9), ("C", 10), ("D", 9), ("E", 16), ("F", 16), ("G", 16), ("H", 3), ("I", 58), ("J", 12), ("K", 74), ("L", 56)):
    ws.column_dimensions[col].width = wdt
cr = blk.row + 2
ws.cell(row=cr, column=9, value="Chart table: mean EndScene change with 95% CI half width").font = bold
ws.cell(row=cr + 1, column=9, value="Version")
ws.cell(row=cr + 1, column=10, value="Mean change")
ws.cell(row=cr + 1, column=11, value="CI half width")
for k, (lab, mk, sk, nk) in enumerate((("Version A", "m1", "s1", "n1"), ("Version B", "m2", "s2", "n2"))):
    rr = cr + 2 + k
    ws.cell(row=rr, column=9, value=lab)
    ws.cell(row=rr, column=10, value=f"={H2_MAIN.addr(mk)}").number_format = "0.000"
    f_ci = f"=_xlfn.CONFIDENCE.T(0.05,{H2_MAIN.addr(sk)},{H2_MAIN.addr(nk)})"
    ws.cell(row=rr, column=11, value=f_ci).number_format = "0.000"
    reg(SH, f"chart_ci_{k}", f"K{rr}", f_ci, f"CI half width {lab}")
chart = BarChart()
chart.type, chart.style = "col", 10
chart.title = "EndScene FA change by version, 95% CI (SIMULATED)"
chart.y_axis.title = "FA change from baseline"
chart.add_data(Reference(ws, min_col=10, min_row=cr + 1, max_row=cr + 3), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=9, min_row=cr + 2, max_row=cr + 3))
chart.series[0].graphicalProperties.solidFill = TEAL
chart.series[0].errBars = ErrorBars(errDir="y", errBarType="both", errValType="cust",
                                    plus=NumDataSource(numRef=NumRef(f=f"H2!$K${cr + 2}:$K${cr + 3}")),
                                    minus=NumDataSource(numRef=NumRef(f=f"H2!$K${cr + 2}:$K${cr + 3}")))
chart.legend = None
chart.height, chart.width = 8, 13
ws.add_chart(chart, f"M{cr}")
H2_CHART_ROW = cr

# ---------------------------------------------------------------- H3
ws = wb.create_sheet("H3")
SH = "H3"
ws["A1"] = "H3. Whole commercial FA change from baseline, Commercial_2 versus Commercial_1, all kept respondents. SIMULATED data."
ws["A1"].font = Font(bold=True, size=12, color=NAVY)
ws["A2"] = ("Same roster order as H2. Each commercial minus the baseline right before it; column J = Commercial_2 change minus Commercial_1 "
            "change. Columns K and L pull the liking ratings from the Survey sheet.")
ws["A2"].font = note_font
h3_cols = ["Respondent Name", "Version", "Order", "Baseline before C1", "Commercial_1 (FA)", "C1 change", "Baseline before C2",
           "Commercial_2 (FA)", "C2 change", "Difference (C2 change minus C1 change)", "Liking_Commercial_1_1to7", "Liking_Commercial_2_1to7"]
ws.append([])
ws.append(h3_cols)
style_header(ws, row=4, ncol=len(h3_cols))
ws.row_dimensions[4].height = 48
for k, rid in enumerate(roster):
    r = P1 + k
    ws.cell(row=r, column=1, value=rid)
    ws.cell(row=r, column=2, value=lookup(R_VER, r))
    ws.cell(row=r, column=3, value=lookup(R_ORD, r))
    ws.cell(row=r, column=4, value=bl_before(r, 1, "C"))
    ws.cell(row=r, column=5, value="=" + fa_of(r, "C1"))
    ws.cell(row=r, column=6, value=f"=E{r}-D{r}")
    ws.cell(row=r, column=7, value=bl_before(r, 2, "C"))
    ws.cell(row=r, column=8, value="=" + fa_of(r, "C2"))
    ws.cell(row=r, column=9, value=f"=H{r}-G{r}")
    ws.cell(row=r, column=10, value=f"=I{r}-F{r}")
    ws.cell(row=r, column=11, value=f"=INDEX({RS_L1},MATCH($A{r},{RS_RESP},0))")
    ws.cell(row=r, column=12, value=f"=INDEX({RS_L2},MATCH($A{r},{RS_RESP},0))")
    for c in range(1, 13):
        ws.cell(row=r, column=c).border = box
        if 4 <= c <= 10:
            ws.cell(row=r, column=c).number_format = "0.0000"
for key, col, lab in (("row_bl1", "D", "Baseline before C1"), ("row_c1", "E", "Commercial_1"), ("row_ch1", "F", "C1 change"),
                      ("row_bl2", "G", "Baseline before C2"), ("row_c2", "H", "Commercial_2"), ("row_ch2", "I", "C2 change"),
                      ("row_dif", "J", "Difference"), ("row_l1", "K", "Liking C1"), ("row_l2", "L", "Liking C2")):
    reg(SH, key, f"{col}{P1}", ws[f"{col}{P1}"].value, lab)
RJ, RB_, RC_ = f"J{P1}:J{Q_LAST}", f"B{P1}:B{Q_LAST}", f"C{P1}:C{Q_LAST}"
blk = Block(ws, SH, 14, 5, "h3")                     # label N, value O, formula P, note Q
change_block(blk, f"H3: paired t on the two changes, all kept respondents, n = {nA + nB}", f"F{P1}:F{Q_LAST}", f"I{P1}:I{Q_LAST}", RJ,
             "C1 change", "C2 change", T["H3"], "must equal the check below")
H3_MAIN = blk
blk.title("Simple estimates (Excel cannot fit the mixed model)")
mpos = M[(AN, "Simple position estimate ((C1first - C2first) / 2)")]
mcom = M[(AN, "Simple commercial estimate ((C1first + C2first) / 2)")]
mint = M[(AN, "Simple commercial x version estimate (A - B)")]
blk.put("o1", "Mean difference, C1first respondents (Commercial_2 seen second)", f'=AVERAGEIFS({RJ},{RC_},"C1first")',
        "difference = commercial effect + position effect")
blk.put("o2", "Mean difference, C2first respondents (Commercial_2 seen first)", f'=AVERAGEIFS({RJ},{RC_},"C2first")',
        "difference = commercial effect - position effect")
blk.put("no1", "n, C1first", f'=COUNTIF({RC_},"C1first")', fmt="0")
blk.put("no2", "n, C2first", f'=COUNTIF({RC_},"C2first")', fmt="0")
blk.put("pos", "Position (second minus first), simple estimate", "=({o1}-{o2})/2", "R mixed model: see analysis_output.txt step 9b",
        expect=fnum(mpos["estimate"]), nd=4)
blk.put("com", "Commercial (Commercial_2 minus Commercial_1), position balanced", "=({o1}+{o2})/2", expect=fnum(mcom["estimate"]), nd=4)
blk.put("vA", "Mean difference, Version A", f'=AVERAGEIFS({RJ},{RB_},"A")')
blk.put("vB", "Mean difference, Version B", f'=AVERAGEIFS({RJ},{RB_},"B")')
blk.put("int", "Commercial x version, simple estimate (A minus B)", "={vA}-{vB}", "R tests it in the mixed model", expect=fnum(mint["estimate"]), nd=4)
blk.blank()
blk = Block(ws, SH, 14, blk.next_start(), "c1")
corr_block(blk, "Exploratory: Commercial_1 change with Liking_Commercial_1_1to7", f"F{P1}:F{Q_LAST}", f"K{P1}:K{Q_LAST}", M[(EX, "r Commercial_1")])
H3_C1 = blk
blk = Block(ws, SH, 14, blk.next_start(), "c2")
corr_block(blk, "Exploratory: Commercial_2 change with Liking_Commercial_2_1to7", f"I{P1}:I{Q_LAST}", f"L{P1}:L{Q_LAST}", M[(EX, "r Commercial_2")])
H3_C2 = blk
blk = Block(ws, SH, 14, blk.next_start(), "chk")
blk.title("Checks")
blk.put("n_data", "Kept respondents with a Commercial_2 row in Data (Keep = Yes)", f'=COUNTIFS({R_ROLE},"C2",{R_KEEP},"Yes")',
        "must equal n of the H3 block", fmt="0", expect=nA + nB, nd=0)
blk.put("bl_same", "Rows where both changes use the same baseline (must be 0)", f"=SUMPRODUCT(--(D{P1}:D{Q_LAST}=G{P1}:G{Q_LAST}))", fmt="0",
        expect=0, nd=0)
H3_CHK = blk
for col, wdt in (("A", 17), ("B", 9), ("C", 10), ("D", 13), ("E", 13), ("F", 12), ("G", 13), ("H", 13), ("I", 12), ("J", 18), ("K", 13), ("L", 13),
                 ("M", 3), ("N", 60), ("O", 12), ("P", 70), ("Q", 52)):
    ws.column_dimensions[col].width = wdt

# ---------------------------------------------------------------- Check
ws = wb.create_sheet("Check")
ws["A1"] = "Live Excel results next to the values analysis.R produced (SIMULATED data). Status is OK when they agree within the tolerance."
ws["A1"].font = Font(bold=True, size=12, color=NAVY)
ws.append([])
ck_cols = ["Quantity", "Excel (live formula)", "analysis.R", "Difference", "Tolerance", "Status"]
ws.append(ck_cols)
style_header(ws, row=3, ncol=len(ck_cols))
ck_rows = []


def add_check(label, ref, r_value, tol, note=""):
    ck_rows.append((label, f"={ref}", float(r_value), tol, note))


blocks = [("H1", "H1", H1_MAIN, "change"), ("S15_H1", "H1", H1_S15, "change"), ("DEFAULT_H1", "H1", H1_DEF, "change"),
          ("H3", "H3", H3_MAIN, "change"), ("H2", "H2", H2_MAIN, "welch"), ("S15_H2", "H2", H2_S15, "welch"), ("RAW_H2", "H2", H2_RAW, "welch")]
for tid, sheet, b, kind in blocks:
    t = T[tid]
    if kind == "change":
        items = (("m1", "m1", 1e-6), ("m2", "m2", 1e-6), ("n", "n", 1e-9), ("md", "mean_diff", 1e-6), ("sdd", "sd_diff", 1e-6),
                 ("t", "t", 1e-4), ("df", "df", 1e-9), ("p_tt", "p", 1e-6), ("p_hand", "p", 1e-6), ("p_one", "p_one_tailed", 1e-6),
                 ("ci_lo", "diff_ci_low", 1e-6), ("ci_hi", "diff_ci_high", 1e-6), ("dz", "es", 1e-5),
                 ("dz_lo", "es_ci_low_approx", 1e-5), ("dz_hi", "es_ci_high_approx", 1e-5))
    else:
        items = (("m1", "m1", 1e-6), ("m2", "m2", 1e-6), ("n1", "n1", 1e-9), ("n2", "n2", 1e-9), ("diff", "mean_diff", 1e-6),
                 ("t", "t", 1e-4), ("df", "df", 1e-4), ("p_tt", "p", 1e-6), ("p_one", "p_one_tailed", 1e-6),
                 ("ci_lo", "diff_ci_low", 5e-4), ("ci_hi", "diff_ci_high", 5e-4), ("sp", "sd_pooled", 1e-6), ("d", "es", 1e-5),
                 ("g", "g", 1e-5), ("g_lo", "g_ci_low_approx", 1e-5), ("g_hi", "g_ci_high_approx", 1e-5))
    for key, rkey, tol in items:
        note = "Welch CI: the spreadsheet critical t may use a whole number df" if key in ("ci_lo", "ci_hi") and kind == "welch" else ""
        add_check(f"{tid}: {rkey} ({key})", f"{sheet}!{b.addr(key)}", t[rkey], tol, note)
    if kind == "change":
        add_check(f"{tid}: d_z exact CI lower (noncentral t in R) versus the approximation", f"{sheet}!{b.addr('dz_lo')}", t["es_ci_low"], 0.03,
                  "method differs; looser tolerance")
        add_check(f"{tid}: d_z exact CI upper (noncentral t in R) versus the approximation", f"{sheet}!{b.addr('dz_hi')}", t["es_ci_high"], 0.03,
                  "method differs; looser tolerance")
t = T["RAW_H1"]
for key, rkey, tol in (("n", "n", 1e-9), ("m", "m2", 1e-6), ("s", "sd2", 1e-6), ("t", "t", 1e-4), ("p", "p", 1e-6), ("d", "es", 1e-5)):
    add_check(f"RAW_H1: {rkey} ({key})", f"H1!{H1_RAW.addr(key)}", t[rkey], tol)
for key, term in (("pos", "Simple position estimate ((C1first - C2first) / 2)"), ("com", "Simple commercial estimate ((C1first + C2first) / 2)"),
                  ("int", "Simple commercial x version estimate (A - B)")):
    add_check(f"H3 {term}", f"H3!{H3_MAIN.addr(key)}", M[(AN, term)]["estimate"], 1e-6)
for b, term in ((H3_C1, "r Commercial_1"), (H3_C2, "r Commercial_2")):
    m = M[(EX, term)]
    for key, col, tol in (("r", "estimate", 1e-6), ("t", "statistic", 1e-4), ("p", "p", 1e-6), ("lo", "ci_low", 1e-6), ("hi", "ci_high", 1e-6)):
        add_check(f"{term}: {col}", f"H3!{b.addr(key)}", m[col], tol)
for key, r_ in summary_rows.items():
    role, ver = key.rsplit("_", 1)
    d = DS[("FA (raw)", ROLE_SEG[role], ver)]
    add_check(f"Summary {role} {ver}: n", f"Summary!C{r_}", d["n"], 1e-9)
    add_check(f"Summary {role} {ver}: mean", f"Summary!D{r_}", d["mean"], 1e-6)
    add_check(f"Summary {role} {ver}: SD", f"Summary!E{r_}", d["sd"], 1e-6)
for key, val in (("n_ex25", len(excluded)), ("n_ex15", len(excluded) + len(drop15)), ("n_keepA", nA), ("n_keepB", nB),
                 ("n_keepA15", nA15), ("n_keepB15", nB15)):
    add_check(f"Screening {REG[('Screening', key)]['label']}", f"Screening!{REG[('Screening', key)]['cell']}", val, 1e-9,
              "counted from the CSV files by this script")
with open(os.path.join(HERE, "analysis_output.txt"), encoding="utf-8") as fh:
    out_txt = fh.read()
m_kept = re.search(r"Respondents kept: (\d+) \(Version A (\d+), Version B (\d+)\)", out_txt)
assert m_kept and int(m_kept.group(2)) == nA and int(m_kept.group(3)) == nB, "kept counts differ from analysis_output.txt"
add_check("Data rows with Keep = No (R: rows of the excluded respondents)", f'COUNTIF({R_KEEP},"No")',
          sum(1 for r in fa_rows if r[fi["Respondent Name"]] in excluded), 1e-9)
CK1 = 4
for k, (lab, f, rv, tol, note) in enumerate(ck_rows):
    r = CK1 + k
    ws.cell(row=r, column=1, value=lab)
    ws.cell(row=r, column=2, value=f).number_format = "0.000000"
    ws.cell(row=r, column=3, value=rv).number_format = "0.000000"
    ws.cell(row=r, column=4, value=f"=B{r}-C{r}").number_format = "0.000000"
    ws.cell(row=r, column=5, value=tol).number_format = "0.000000"
    ws.cell(row=r, column=6, value=f'=IF(ABS(D{r})<=E{r},"OK","CHECK")')
    if note:
        ws.cell(row=r, column=7, value=note).font = note_font
CK2 = CK1 + len(ck_rows) - 1
ws.cell(row=CK2 + 2, column=1, value="Rows with status CHECK").font = bold
ws.cell(row=CK2 + 2, column=2, value=f'=COUNTIF(F{CK1}:F{CK2},"CHECK")')
reg("Check", "n_check", f"B{CK2 + 2}", f'=COUNTIF(F{CK1}:F{CK2},"CHECK")', "Rows with status CHECK")
reg("Check", "status_row", f"F{CK1}", f'=IF(ABS(D{CK1})<=E{CK1},"OK","CHECK")', "Status")
EXPECT[("Check", "n_check")] = (0, 0)
ws.cell(row=CK2 + 3, column=1, value=("Not reproducible in Excel: the Wilcoxon signed rank and Mann Whitney tests, the Shapiro Wilk tests, the exact "
                                      "noncentral t intervals, the mixed model with its Satterthwaite tests and ICC, the power analysis, the Spearman "
                                      "correlations, and the theta descriptives from the PSD export. See analysis_output.txt.")).font = note_font
r = CK2 + 6
for title, hdr_, body in (("analysis.R: tests (SIMULATED_test_results.csv)", th, trows),
                          ("analysis.R: mixed model, correlations, power note (SIMULATED_model_results.csv)", mh, mrows)):
    ws.cell(row=r - 1, column=1, value=title).font = bold
    for c, h in enumerate(hdr_, start=1):
        ws.cell(row=r, column=c, value=h)
    style_header(ws, row=r, ncol=len(hdr_))
    for rw in body:
        r += 1
        for c, v in enumerate(rw, start=1):
            ws.cell(row=r, column=c, value=to_cell(v))
    r += 3
autowidth(ws, max_w=34)
ws.column_dimensions["A"].width = 70
for col in "BCDEF":
    ws.column_dimensions[col].width = 18

wb.calculation = CalcProperties(fullCalcOnLoad=True)    # Excel computes every formula when the file opens
wb.save(OUT)
print("wrote", OUT)

# =============================================================================
# Verification 1: read the workbook back with openpyxl and compare every registered cell
# =============================================================================
wb2 = load_workbook(OUT)
problems = []
for (sheet, key), info in REG.items():
    v = wb2[sheet][info["cell"]].value
    if isinstance(v, ArrayFormula):
        v = v.text
    v = str(v).replace("_xlfn.", "") if v is not None else None
    if v != info["formula"]:
        problems.append(f"{sheet}!{info['cell']} ({key}): workbook has {v!r}, registry has {info['formula']!r}")
for k, (lab, f, rv, tol, note) in enumerate(ck_rows):
    if wb2["Check"][f"B{CK1 + k}"].value != f:
        problems.append(f"Check!B{CK1 + k} differs")
if problems:
    print("VERIFICATION FAILED:")
    print("\n".join(problems[:40]))
    sys.exit(1)
print(f"verified (openpyxl): {len(REG)} registered cells and {len(ck_rows)} Check rows hold the registered formulas")
n_formulas = sum(1 for wsx in wb2.worksheets for row in wsx.iter_rows() for c in row
                 if (isinstance(c.value, str) and c.value.startswith("=")) or isinstance(c.value, ArrayFormula))
print(f"formula cells in the workbook: {n_formulas}")

# =============================================================================
# Verification 2 and 3: recalculate a copy with LibreOffice and compare the values with analysis.R
# =============================================================================
LO = {}
soffice = shutil.which("soffice") or shutil.which("libreoffice")
lo_note = "LibreOffice was not found, so the formula values were not recalculated here."
if soffice:
    tmp = tempfile.mkdtemp(prefix="ws5_lo_")
    shutil.copy(OUT, os.path.join(tmp, "in.xlsx"))
    subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir",
                    os.path.join(tmp, "out"), os.path.join(tmp, "in.xlsx")], check=True, capture_output=True, timeout=300)
    wb3 = load_workbook(os.path.join(tmp, "out", "in.xlsx"), data_only=True)
    for (sheet, key), info in REG.items():
        LO[(sheet, key)] = wb3[sheet][info["cell"]].value
    statuses = [wb3["Check"][f"F{CK1 + k}"].value for k in range(len(ck_rows))]
    bad = [(ck_rows[k][0], wb3["Check"][f"B{CK1 + k}"].value, ck_rows[k][2]) for k, s in enumerate(statuses) if s != "OK"]
    if bad:
        print("LIBREOFFICE CHECK FAILED:")
        for b in bad[:40]:
            print("  ", b)
        sys.exit(1)
    mism = []
    for (sheet, key), (val, nd) in EXPECT.items():
        got = LO[(sheet, key)]
        if got is None or isinstance(got, str):
            mism.append((sheet, key, got, val))
            continue
        if f"{float(got):.{nd}f}" != f"{float(val):.{nd}f}":
            mism.append((sheet, key, got, val))
    if mism:
        print("VALUES DIFFER FROM analysis.R AT THE PRINTED DECIMALS:")
        for m_ in mism[:40]:
            print("  ", m_)
        sys.exit(1)
    lo_ver = subprocess.run([soffice, "--version"], capture_output=True, text=True).stdout.strip()
    lo_note = (f"A copy was recalculated with {lo_ver} (headless): all {len(ck_rows)} Check rows are OK, and the {len(EXPECT)} results printed "
               "in this walkthrough agree with analysis.R at the printed decimals.")
    print(f"verified (LibreOffice): {len(ck_rows)} Check rows OK; {len(EXPECT)} printed results agree with analysis.R")
    shutil.rmtree(tmp, ignore_errors=True)
else:
    print(lo_note)


# =============================================================================
# The walkthrough, generated from REG so that every cell reference is the real one
# =============================================================================
def cl(sheet, key):
    return REG[(sheet, key)]["cell"]


def fx(sheet, key):
    return REG[(sheet, key)]["formula"]


def val(sheet, key, nd=None):
    """Printed result: the LibreOffice value when available (verified above), else the analysis.R value."""
    if (sheet, key) in EXPECT:
        v, nd0 = EXPECT[(sheet, key)]
        nd = nd0 if nd is None else nd
    else:
        v = None
    if (sheet, key) in LO and LO[(sheet, key)] is not None and not isinstance(LO[(sheet, key)], str):
        v = LO[(sheet, key)]
    if v is None:
        return ""
    return f"{float(v):.{nd if nd is not None else 3}f}"


def row(label, sheet, key, nd=None, extra=""):
    return f"| {label} | `{sheet}!{cl(sheet, key)}` | `{fx(sheet, key)}` | {val(sheet, key, nd)}{extra} |"


def row3(label, sheet, key):
    return f"| {label} | `{sheet}!{cl(sheet, key)}` | `{fx(sheet, key)}` |"


def rcell(sheet, key):
    return f"{sheet}!{cl(sheet, key)}"


def f2(x, nd=2):
    return f"{float(x):.{nd}f}"


def p3(x):
    x = float(x)
    return "< .001" if x < 0.001 else "= " + f"{x:.3f}".lstrip("0")


def _and(xs):
    xs = list(xs)
    return xs[0] if len(xs) == 1 else (" and ".join(xs) if len(xs) == 2 else ", ".join(xs[:-1]) + ", and " + xs[-1])


h1, h2, h3 = T["H1"], T["H2"], T["H3"]
s1, s2, u1, u2, d1 = T["S15_H1"], T["S15_H2"], T["RAW_H1"], T["RAW_H2"], T["DEFAULT_H1"]
icc = M[(AN, "ICC adjusted (fitted model)")]["estimate"]
mm_int = M[(AN, "Commercial x version")]
mm_com = M[(AN, "Commercial (Commercial_2 minus Commercial_1)")]
mm_pos = M[(AN, "Position (second minus first)")]
pw_n = M[("Power note for H2", "n per version for 80 percent power (power.t.test, rounded up)")]["estimate"]
pw_now = M[("Power note for H2", "Power achieved with the kept sample")]["estimate"]
excl_info = {}
for rr_ in log_rows:
    rid = rr_[li["Respondent Name"]]
    if rid in excluded:
        excl_info[rid] = rr_[li["Incidents"]]


def cause_of(rid):
    inc = excl_info[rid]
    if red[rid]:
        return f"Red NIC2 status at F4 that could not be fixed (lost electrode contact); largest percent removed {maxpct[rid]:.2f}"
    if "LSL" in inc:
        return f"LSL stream gap and drift; {maxpct[rid]:.2f} percent removed"
    return f"heavy movement and blinking; {maxpct[rid]:.2f} percent removed"


excl_lines = "\n".join(f"- `{rid}` (Version {version[rid]}): {cause_of(rid)}." for rid in excluded)
d15_txt = _and([f"`{i}` (Version {version[i]}, {maxpct[i]:.2f} percent)" for i in drop15])
gap_resp = [r[ai["Respondent Name"]] for r in art_rows if int(r[ai["Samples recorded"]]) < int(r[ai["Samples expected"]])]
gap_txt = gap_resp[0] if gap_resp else "none"
m_ks1 = re.search(r"SD of the change from baseline: KeyScene ([0-9.]+) versus whole Commercial_1A ([0-9.]+)", out_txt)
m_ks2 = re.search(r"Correlation with the preceding baseline: KeyScene r = ([0-9.]+) versus whole Commercial_1A r = ([0-9.]+)", out_txt)
assert m_ks1 and m_ks2, "step 6 lines not found in analysis_output.txt"
ks_noise = (m_ks1.group(1), m_ks1.group(2), m_ks2.group(1), m_ks2.group(2))
red_fixed = [r[li["Respondent Name"]] for r in log_rows if "Red at the first check" in r[li["Incidents"]]]
NUMW = {0: "No", 1: "One", 2: "Two", 3: "Three", 4: "Four", 5: "Five", 6: "Six", 7: "Seven", 8: "Eight", 9: "Nine", 10: "Ten",
        11: "Eleven", 12: "Twelve"}
n_red_fixed_txt = (f"{NUMW.get(len(red_fixed), len(red_fixed))} other respondent{'s' if len(red_fixed) != 1 else ''} "
                   f"({_and(red_fixed)})" if red_fixed else "No other respondent")

md = f"""# EEG frontal alpha asymmetry analysis in Excel: step by step walkthrough

**SIMULATED DATA.** Every number in this walkthrough comes from the SIMULATED files that `make_simulated_data.py` generated (seed {SEED}). No real person was recorded. The design copies the logic of the iMotions demo study "DEMO - SonyBravia" (a commercial with a key scene against the same commercial without it, a second commercial for everyone, neutral 1 minute nature videos as baselines, and the last 10 seconds of each commercial); nothing about the content of the Sony commercials was used.

This is the same analysis as `analysis.R`, done in Excel for students who do not use R. The finished workbook is `SIMULATED_eeg_analysis.xlsx`; the sheet names below refer to it. **Every cell address in this walkthrough is written as `Sheet!Cell` and was generated from the workbook itself.** `make_excel_workbook.py` writes the workbook and this file from one list of cells, reads the workbook back with openpyxl to confirm that every address holds the formula printed here, and recalculates a copy to confirm the results. {lo_note}

Companion files: `analysis.R` (R version), `analysis_output.txt` (its console output), `results_text.md` (model paragraphs), `figures/` (PNG figures).

---

## 0. What you are analyzing

**Design.** {len(ids)} respondents (P01 to P{len(ids):02d}) were randomized to two versions of Commercial 1. **Version A** saw `Commercial_1A` (60 s) with a key scene at 32 to 38 s; **Version B** saw `Commercial_1B` (54 s), the same commercial with the key scene cut. Everyone also saw `Commercial_2` (60 s). Half of each version saw Commercial 1 first. The flow was `Baseline_1` (a 60 s neutral nature video), the first commercial, a rating slide, `Baseline_2` (another 60 s nature video), the second commercial, and a rating slide. **Each commercial is compared with the baseline that immediately precedes it.** Two stimulus annotations mark the analysis windows: `KeyScene` (32 to 38 s of `Commercial_1A`) and `EndScene` (the last 10 s of each commercial).

**Device.** One Neuroelectrics Enobio 8 streamed through NIC2 and Lab Streaming Layer (iMotions device name `Enobio EEG`, 500 Hz). iMotions reports the Enobio channels as Ch1 to Ch8; the scalp site of each channel is set in NIC2. This workbook **assumes** the lab manual montage order (Ch1 Fz, Ch2 Cz, Ch3 Pz, Ch4 Oz, Ch5 F3, Ch6 F4, Ch7 P3, Ch8 P4) and must be checked against the NIC2 protocol before real data are analyzed.

**The metric.** The EEG Frontal Asymmetry R Notebook computes, in each 1 s Welch window, **FA = ln(alpha power at the right frontal electrode) - ln(alpha power at the left frontal electrode)**, here ln(F4) - ln(F3). Alpha power goes **down** when a region is more active, so a higher FA means relatively more **left** frontal activity, which the notebook article reads as approach motivation; a lower FA means relatively more right frontal activity (withdrawal). Approach is not the same as liking. FA scores are read **relative to each other** (here, relative to a baseline), never as absolute values.

**Files.**

| File | What it is | Columns used here | Used for |
|---|---|---|---|
| `SIMULATED_eeg_fa_metrics.csv` | EEG Frontal Asymmetry Metrics export (R Processing), notebook run with left Ch5 and right Ch6 | `Respondent Name`, `Group`, `Type`, `Label`, `Parent Stimulus`, `Mean Frontal Asymmetry Alpha` | H1 to H3 |
| `SIMULATED_eeg_fa_metrics_default_channels.csv` | The same export made with the notebook **default** channels (left Ch1, right Ch7) | the same | Teaching contrast |
| `SIMULATED_eeg_artifact_report.csv` | Percent of F3 and F4 samples removed by the artifact threshold, per respondent and stimulus (**modeled layout**) | `Respondent Name`, `Electrode (assumed)`, `Samples expected`, `Samples recorded`, `Percent removed (%)` | Exclusion rule |
| `SIMULATED_eeg_session_log.csv` | Operator log (**modeled layout**) | `NIC2 quality F3 before recording`, `NIC2 quality F4 before recording`, `Incidents` | Exclusion rule |
| `SIMULATED_eeg_survey.csv` | Questionnaire (**modeled layout**) | `Liking_Commercial_1_1to7`, `Liking_Commercial_2_1to7` | Exploratory correlation |

A sixth file, `SIMULATED_eeg_psd_metrics.csv` (EEG Power Spectral Density Metrics export), is used in R for the theta descriptives and a consistency check; the workbook does not need it.

**Hypotheses.**
- H1 (within subjects, Version A): FA during the `KeyScene` minus FA during the baseline before `Commercial_1A` is greater than zero.
- H2 (between subjects): the baseline corrected FA during the `EndScene` of Commercial 1 is higher for Version A than for Version B.
- H3 (within subjects, all kept respondents): the baseline corrected whole commercial FA differs between `Commercial_2` and Commercial 1.
- Exploratory: whole commercial FA change and liking. Sensitivity: H1 and H2 with a 15 percent threshold and without baseline correction. Teaching contrast: H1 with the notebook default channels.

**Decision rules fixed before collection.** Exclude a respondent when more than 25 percent of the F3 or F4 samples were removed in any baseline or commercial, or when the session log records a Red NIC2 quality status at F3 or F4 that could not be fixed. All tests are two tailed with alpha = .05 (the one tailed p values in the blocks are for discussion only).

---

## 1. Import the files and build the Data sheet

iMotions exports are comma separated with a dot decimal. Double clicking a file uses your regional defaults, which can put everything in one column (Opening Data Export Files in Excel article). Use the import route instead.

**Route A, Get Data (Excel 2016 and later):**
1. **Data > From Text/CSV**, select `SIMULATED_eeg_fa_metrics.csv`, click **Import**.
2. In the preview set **Delimiter: Comma** and **File Origin: 65001 Unicode (UTF-8)**.
3. Click **Transform Data**. The two FA files start with **two** comment lines (`# SIMULATED ...` and `# METADATA ...`); the other files start with one. In Power Query choose **Home > Remove Rows > Remove Top Rows** (2 or 1), then **Home > Use First Row as Headers**, then **Close & Load**. (A real iMotions export has no comment lines; the simulated files carry them so nobody mistakes them for real data. Keep the METADATA line somewhere: it records the notebook settings.)
4. **You should see** {len(fa_h)} columns with `Study Name` in the first column and {N_FA} data rows: 7 rows for each Version A respondent (four stimuli, `KeyScene`, and two `EndScene` rows) and 6 for each Version B respondent.

**Route B, Text Import Wizard (any version):** **File > Open > Browse**, set the file type to **All Files**, open the .csv, choose **Delimited**, start at row 3 for the FA files (row 2 for the others), tick **Comma**, click **Advanced** and set the decimal symbol to a dot, then **Finish**.

**Common mistake.** The FA column can hold the text `NA` (a segment in which every 1 s window was removed by the artifact threshold). Leave it as `NA`; do not type 0. A zero FA would be a real value (equal left and right alpha).

In the workbook this is the **Data** sheet, header in row 1 and data in rows {D1} to {D2}. Column letters: `{DCOL['Respondent Name']}` = Respondent Name, `{DCOL['Group']}` = Group, `{DCOL['Type']}` = Type, `{DCOL['Label']}` = Label, `{DCOL['Parent Stimulus']}` = Parent Stimulus, `{DCOL['Mean Frontal Asymmetry Alpha']}` = Mean Frontal Asymmetry Alpha. The **Default** sheet has the same layout for the default channel export, and the **Artifact**, **Log**, and **Survey** sheets hold the other three files as imported.

---

## 2. Helper columns: Version, Order, Role, Keep, Keep15

Add five columns to the right of the export (`{DCOL['Version']}` to `{DCOL['Keep15']}`, headers in orange) and fill them down. The formulas below are the ones in row {D1}.

| Header | Cell | Formula | Meaning |
|---|---|---|---|
| Version | `Data!{cl('Data', 'version')}` | `{fx('Data', 'version')}` | A or B, the first letter of `Group` (for example `A_C1first`) |
| Order | `Data!{cl('Data', 'order')}` | `{fx('Data', 'order')}` | `C1first` or `C2first` |
| Role | `Data!{cl('Data', 'role')}` | `{fx('Data', 'role')}` | Which part of the design the row is: `Baseline_1`, `Baseline_2`, `C1`, `C2`, `KeyScene`, `EndScene_C1`, `EndScene_C2` |
| Keep | `Data!{cl('Data', 'keep')}` | `{fx('Data', 'keep')}` | The exclusion rule: No when any F3 or F4 row of this respondent in the artifact report is above 25 percent, or the log shows Red at F3 or F4 |
| Keep15 | `Data!{cl('Data', 'keep15')}` | `{fx('Data', 'keep15')}` | The same rule at 15 percent (sensitivity analysis only) |

**Why Role needs Parent Stimulus.** The label `EndScene` appears on all three commercials, so a row is identified by `Label` **and** `Parent Stimulus`. The Role formula sends an `EndScene` row whose parent starts with `Commercial_1` to `EndScene_C1` and the one on `Commercial_2` to `EndScene_C2`, and it gives `Commercial_1A` and `Commercial_1B` the same role, `C1`, because they are the two versions of the same commercial.

**Common mistake.** Averaging every `EndScene` row of a respondent, or filtering on `Label` alone. That mixes the two commercials.

---

## 3. The exclusion rule: Screening sheet

The **Screening** sheet lists every respondent once and shows why each one is kept or excluded. Columns C and D are **array formulas** (`MAX(IF(...))`): in Excel 365 press **Enter**; in Excel 2019 or older press **Ctrl+Shift+Enter**.

| Column | Cell (first respondent) | Formula |
|---|---|---|
| Version | `Screening!{cl('Screening', 'ver')}` | `{fx('Screening', 'ver')}` |
| Largest percent removed, F3 | `Screening!{cl('Screening', 'f3')}` | `{fx('Screening', 'f3')}` |
| Largest percent removed, F4 | `Screening!{cl('Screening', 'f4')}` | `{fx('Screening', 'f4')}` |
| Largest of the two | `Screening!{cl('Screening', 'big')}` | `{fx('Screening', 'big')}` |
| Red at F3 or F4 (count) | `Screening!{cl('Screening', 'red')}` | `{fx('Screening', 'red')}` |
| Decision (25 percent rule) | `Screening!{cl('Screening', 'dec25')}` | `{fx('Screening', 'dec25')}` |
| Decision (15 percent rule) | `Screening!{cl('Screening', 'dec15')}` | `{fx('Screening', 'dec15')}` |

The counts below the table:

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('Respondents excluded (25 percent rule)', 'Screening', 'n_ex25')}
{row('Respondents excluded (15 percent rule)', 'Screening', 'n_ex15')}
{row('Kept, Version A', 'Screening', 'n_keepA')}
{row('Kept, Version B', 'Screening', 'n_keepB')}
{row('Kept, Version A (15 percent rule)', 'Screening', 'n_keepA15')}
{row('Kept, Version B (15 percent rule)', 'Screening', 'n_keepB15')}
{row('Artifact report rows with a stream gap', 'Screening', 'n_gap')}

**You should see** three exclusions, each with a cause in the session log (**Log** sheet, column `Incidents`):

{excl_lines}

That leaves {nA} Version A and {nB} Version B respondents. The respondent with the stream gap ({gap_txt}) also shows `Samples recorded` below `Samples expected` on the **Artifact** sheet (column `{cGap}` holds the difference). {n_red_fixed_txt} had a Red status at the first impedance check that was fixed before recording; the log records Green for them, so the rule keeps them. Under the stricter 15 percent rule {d15_txt} are also dropped.

**Common mistake.** Applying the rule to the annotation rows or to a single channel. The rule looks at both frontal channels and at every baseline and commercial, because a respondent whose F3 or F4 signal is unreliable in any of them cannot give a trustworthy baseline corrected score.

---

## 4. Descriptive table (Summary sheet)

The **Summary** sheet gives n, mean, and SD of the raw FA for each role and version (`*` = both versions) among kept respondents. The formulas in row {SUM_FIRST} (`Baseline_1`, Version A):

| Statistic | Cell | Formula |
|---|---|---|
{row3('n', 'Summary', 'n_Baseline_1_A')}
{row3('Mean FA', 'Summary', 'm_Baseline_1_A')}
{row3('SD FA (array formula)', 'Summary', 's_Baseline_1_A')}

**You should see** raw FA values mostly below zero (as in the demo study chart): `Baseline_1` M = {val('Summary', 'm_Baseline_1_All')} and `Baseline_2` M = {val('Summary', 'm_Baseline_2_All')} over all kept respondents, with SDs of about {val('Summary', 's_Baseline_1_All', 2)}. The SD of the raw values is large because every respondent carries a stable asymmetry of their own; that stable part is what the baseline correction removes. A PivotTable gives the same numbers: **Insert > PivotTable**, **Keep** in **Filters** set to **Yes**, **Role** in **Rows**, **Version** in **Columns**, and **Mean Frontal Asymmetry Alpha** in **Values** set to **Average** (Creating Excel Pivot Tables article).

---

## 5. H1: key scene versus the baseline before Commercial_1A (Version A)

The **H1** sheet has one row per kept Version A respondent (rows {P1} to {PA_LAST}). The first respondents (rows {P1} to {PA15_LAST}) are also kept under the 15 percent rule, and the last {len(A_drop)} (rows {PA15_LAST + 1} to {PA_LAST}) are dropped by it, so the sensitivity block only shortens the range.

| Column | Cell (first respondent) | Formula |
|---|---|---|
{row3('Order', 'H1', 'row_ord')}
{row3('Keep15', 'H1', 'row_k15')}
{row3('Baseline before Commercial_1A', 'H1', 'row_bl')}
{row3('KeyScene', 'H1', 'row_ks')}
{row3('Change (KeyScene minus baseline)', 'H1', 'row_chg')}

The baseline formula is the heart of the design: for a `C1first` respondent the baseline right before `Commercial_1A` is `Baseline_1`; for a `C2first` respondent it is `Baseline_2`.

The test sits in column L (formula text in column M). A paired t test **is** a one sample t test on the change scores:

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('Baseline mean (FA)', 'H1', 'h1_m1')}
{row('KeyScene mean (FA)', 'H1', 'h1_m2')}
{row('n', 'H1', 'h1_n')}
{row('Mean change (KeyScene minus baseline)', 'H1', 'h1_md')}
{row('SD of the changes', 'H1', 'h1_sdd')}
{row('SE of the mean change', 'H1', 'h1_se', 4)}
{row('t', 'H1', 'h1_t')}
{row('df', 'H1', 'h1_df')}
{row('**p, two tailed (T.TEST type 1)**', 'H1', 'h1_p_tt')}
{row('p from the hand computed t', 'H1', 'h1_p_hand')}
{row('p, one tailed (teaching only)', 'H1', 'h1_p_one')}
{row('Mean change 95% CI lower', 'H1', 'h1_ci_lo')}
{row('Mean change 95% CI upper', 'H1', 'h1_ci_hi')}

`T.TEST(array1, array2, tails, type)`: `tails = 2` for a two tailed test; `type = 1` for **paired** data. `type = 2` and `type = 3` treat the two columns as different people and are wrong here. **You should see** the hand computed p equal the `T.TEST` p.

### Cohen's d_z with a 95 percent confidence interval

For paired data the effect size is **d_z**, the mean change in units of the SD of the changes (t = d_z x sqrt(n)).

| Step | Cell | Formula | Result |
|---|---|---|---|
{row("d_z", 'H1', 'h1_dz')}
{row('SE of d_z (approximation)', 'H1', 'h1_se_dz')}
{row('d_z 95% CI lower (approximation)', 'H1', 'h1_dz_lo')}
{row('d_z 95% CI upper (approximation)', 'H1', 'h1_dz_hi')}
{row('Correlation of the two columns', 'H1', 'h1_r', 2)}
{row('Respondents with a positive change', 'H1', 'h1_npos', 0)}

In words: SE(d_z) = sqrt(1/n + d_z^2 / (2 n)), the large sample approximation. `analysis.R` also reports the exact interval from the noncentral t distribution, [{f2(h1['es_ci_low'])}, {f2(h1['es_ci_high'])}] (Lakens, 2013); say which one you report. Rule of thumb for d_z: 0.2 small, 0.5 medium, 0.8 large. The Wilcoxon signed rank test, the check that the result does not depend on normality, has no Excel function; in R: V = {f2(h1['nonparam_stat'], 0) if float(h1['nonparam_stat']).is_integer() else h1['nonparam_stat']}, p {p3(h1['nonparam_p'])}.

**Why the key scene scatters more.** The 6 s key scene holds only 11 one second Welch windows, the 60 s baseline 119, so a key scene value averages out much less noise. The correlation of the two columns ({val('H1', 'h1_r', 2)}) is lower than the correlation of the whole `Commercial_1A` with the same baseline (r = {ks_noise[3]} in step 6 of `analysis_output.txt`), and the SD of the key scene changes ({ks_noise[0]}) is larger than the SD of the whole `Commercial_1A` changes ({ks_noise[1]}). A short window needs either more respondents or a larger effect.

The checks at the end of the column:

| Statistic | Cell | Formula | Result |
|---|---|---|---|
{row('Kept Version A respondents with a KeyScene row in Data', 'H1', 'chk_n_data')}
{row('Roster rows dropped under the 15 percent rule', 'H1', 'chk_n_drop')}
{row('Keep15 = No inside the 15 percent block (must be 0)', 'H1', 'chk_n_drop_top')}

**Common mistake.** Comparing the key scene with `Baseline_1` for everyone. For `C2first` respondents, `Baseline_1` came before `Commercial_2`, and several minutes and another commercial separate it from the key scene.

---

## 6. H2: end scene of Commercial 1, Version A versus Version B (Welch t test)

The **H2** sheet lists all {nA + nB} kept respondents: Version A in rows {QA1} to {QA2} and Version B in rows {QB1} to {QB2}. Within each version, the respondents dropped by the 15 percent rule sit at the bottom (A rows {QA15 + 1} to {QA2}, B rows {QB15 + 1} to {QB2}), so the sensitivity block uses rows {QA1} to {QA15} and {QB1} to {QB15}.

| Column | Cell (first respondent) | Formula |
|---|---|---|
{row3('Version', 'H2', 'row_ver')}
{row3('Baseline before Commercial 1', 'H2', 'row_bl')}
{row3('EndScene of Commercial 1', 'H2', 'row_es')}
{row3('Change', 'H2', 'row_chg')}

Different respondents saw the two versions, so this is a **two group** comparison. Welch's t test does not assume equal variances and is the recommended default (Delacre et al., 2017). The block sits in column J:

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('Version A mean change', 'H2', 'h2_m1')}
{row('Version A SD', 'H2', 'h2_s1')}
{row('Version A n', 'H2', 'h2_n1')}
{row('Version B mean change', 'H2', 'h2_m2')}
{row('Version B SD', 'H2', 'h2_s2')}
{row('Version B n', 'H2', 'h2_n2')}
{row('Difference, A minus B', 'H2', 'h2_diff')}
{row('SE of the difference (Welch)', 'H2', 'h2_se', 4)}
{row('Welch t', 'H2', 'h2_t')}
{row('Welch df', 'H2', 'h2_df')}
{row('**p, two tailed (T.TEST type 3)**', 'H2', 'h2_p_tt')}
{row('p from the hand computed t', 'H2', 'h2_p_hand', 4)}
{row('p, one tailed (teaching only)', 'H2', 'h2_p_one')}
{row('Difference 95% CI lower', 'H2', 'h2_ci_lo')}
{row('Difference 95% CI upper', 'H2', 'h2_ci_hi')}
{row('Pooled SD', 'H2', 'h2_sp')}
{row("Cohen's d", 'H2', 'h2_d')}
{row("Hedges' g", 'H2', 'h2_g')}
{row('g 95% CI lower (approximation)', 'H2', 'h2_g_lo')}
{row('g 95% CI upper (approximation)', 'H2', 'h2_g_hi')}

`T.TEST(..., 2, 3)`: `type = 3` is the unequal variance (Welch) test for two groups. The Welch df is not a whole number. Spreadsheet functions such as `T.DIST.2T` and `T.INV.2T` may round a fractional df down to a whole number (LibreOffice does), so the hand computed p and the interval can differ from R in the third or fourth decimal; **report the `T.TEST` p**, which matched R to six decimals in the recalculated copy. Hedges' g is Cohen's d times a small sample correction; `analysis.R` also gives the exact interval, g = {f2(h2['g'])}, 95% CI [{f2(h2['g_ci_low'])}, {f2(h2['g_ci_high'])}], and the Mann Whitney check, W = {h2['nonparam_stat'].rstrip('0').rstrip('.') if '.' in h2['nonparam_stat'] else h2['nonparam_stat']}, p {p3(h2['nonparam_p'])}.

**Reading a nonsignificant result.** The difference points the predicted way (Version A higher), but the 95 percent interval of the difference includes zero, so the data are compatible with no difference and with a difference as large as g = {f2(h2['g_ci_high'])}. That is not evidence that the key scene has no effect on the end scene. The power analysis in R (step 8 of `analysis_output.txt`) shows why: at the observed d = {f2(h2['es'])}, a new study would need about {int(float(pw_n))} respondents **per version** for 80 percent power, and this study had a power of about {f2(pw_now)}. Between subjects comparisons need far more respondents than within subjects ones, because the differences between people stay in the error term.

The checks at the end of the column:

| Statistic | Cell | Formula | Result |
|---|---|---|---|
{row('Kept respondents with an EndScene_C1 row in Data', 'H2', 'chk_n_data')}
{row('Version B rows inside the Version A range (must be 0)', 'H2', 'chk_ver_top')}
{row('Version A rows inside the Version B range (must be 0)', 'H2', 'chk_ver_bot')}

---

## 7. H3: Commercial_2 versus Commercial_1 (paired), and what Excel cannot do

The **H3** sheet uses the same roster order as H2. Each commercial is compared with the baseline right before it, so the two changes of a respondent use **different** baselines:

| Column | Cell (first respondent) | Formula |
|---|---|---|
{row3('Baseline before C1', 'H3', 'row_bl1')}
{row3('Commercial_1 (FA)', 'H3', 'row_c1')}
{row3('C1 change', 'H3', 'row_ch1')}
{row3('Baseline before C2', 'H3', 'row_bl2')}
{row3('Commercial_2 (FA)', 'H3', 'row_c2')}
{row3('C2 change', 'H3', 'row_ch2')}
{row3('Difference (C2 change minus C1 change)', 'H3', 'row_dif')}

The paired block sits in column O:

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('C1 change mean', 'H3', 'h3_m1')}
{row('C2 change mean', 'H3', 'h3_m2')}
{row('n', 'H3', 'h3_n')}
{row('Mean difference (C2 change minus C1 change)', 'H3', 'h3_md')}
{row('SD of the differences', 'H3', 'h3_sdd')}
{row('t', 'H3', 'h3_t')}
{row('df', 'H3', 'h3_df')}
{row('**p, two tailed (T.TEST type 1)**', 'H3', 'h3_p_tt')}
{row('Mean difference 95% CI lower', 'H3', 'h3_ci_lo')}
{row('Mean difference 95% CI upper', 'H3', 'h3_ci_hi')}
{row('d_z', 'H3', 'h3_dz')}
{row('d_z 95% CI lower (approximation)', 'H3', 'h3_dz_lo')}
{row('d_z 95% CI upper (approximation)', 'H3', 'h3_dz_hi')}

**Order and version.** Every respondent saw both commercials, so the one seen second may be affected by the first (carryover or fatigue). For `C1first` respondents the difference C2 change minus C1 change is the commercial effect **plus** the position effect (Commercial_2 came second); for `C2first` respondents it is the commercial effect **minus** the position effect. Half the gap between the two order groups estimates the position effect:

| Statistic | Cell | Formula | Result |
|---|---|---|---|
{row('Mean difference, C1first', 'H3', 'h3_o1', 4)}
{row('Mean difference, C2first', 'H3', 'h3_o2', 4)}
{row('Position (second minus first), simple estimate', 'H3', 'h3_pos')}
{row('Commercial effect, position balanced', 'H3', 'h3_com')}
{row('Commercial x version, simple estimate (A minus B)', 'H3', 'h3_int')}

`analysis.R` fits the linear mixed model `FA_change ~ commercial * version + position + (1 | respondent)` with lme4 (Bates et al., 2015) and lmerTest, which Excel cannot do. Its estimates differ slightly from the simple ones because the four version by order cells are unequal after the exclusions: commercial b = {f2(mm_com['estimate'], 4)} (p {p3(mm_com['p'])}), position b = {f2(mm_pos['estimate'], 4)} (p {p3(mm_pos['p'])}), commercial x version b = {f2(mm_int['estimate'], 4)} (p {p3(mm_int['p'])}), and the **intraclass correlation** ICC = {f2(icc, 3)}, the share of the remaining variance that lies between respondents. The ICC is low because the baseline correction already removed each respondent's stable asymmetry.

---

## 8. Exploratory: FA change and liking

Columns K and L of the H3 sheet pull the liking ratings from the **Survey** sheet with `INDEX` and `MATCH`. The correlation blocks use `CORREL`, then the usual t and Fisher z formulas:

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('Commercial_1: Pearson r', 'H3', 'c1_r')}
{row('Commercial_1: t for r', 'H3', 'c1_t')}
{row('Commercial_1: p, two tailed', 'H3', 'c1_p')}
{row('Commercial_1: r 95% CI lower (Fisher z)', 'H3', 'c1_lo')}
{row('Commercial_1: r 95% CI upper (Fisher z)', 'H3', 'c1_hi')}
{row('Commercial_2: Pearson r', 'H3', 'c2_r')}
{row('Commercial_2: p, two tailed', 'H3', 'c2_p')}

Approach motivation and liking are different things (approach can come with anger or curiosity), and a rating and a brain signal are different kinds of measure, so a small to moderate correlation is the expected result. The Spearman rank correlations are in R.

---

## 9. Sensitivity analyses

**Stricter threshold (15 percent).** The sensitivity blocks use the shorter row ranges described in sections 5 and 6:

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('H1, 15 percent: n', 'H1', 's15_n')}
{row('H1, 15 percent: mean change', 'H1', 's15_md')}
{row('H1, 15 percent: t', 'H1', 's15_t')}
{row('H1, 15 percent: p (T.TEST type 1)', 'H1', 's15_p_tt')}
{row('H1, 15 percent: d_z', 'H1', 's15_dz')}
{row('H2, 15 percent: difference', 'H2', 's15_diff')}
{row('H2, 15 percent: Welch t', 'H2', 's15_t')}
{row('H2, 15 percent: p (T.TEST type 3)', 'H2', 's15_p_tt')}
{row("H2, 15 percent: Hedges' g", 'H2', 's15_g')}

**Without baseline correction.** H1 then asks whether the raw `KeyScene` FA is above zero, and H2 compares the raw `EndScene` FA of the two versions:

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('H1 raw: mean KeyScene FA', 'H1', 'raw_m')}
{row('H1 raw: t against 0', 'H1', 'raw_t')}
{row('H1 raw: p, two tailed', 'H1', 'raw_p')}
{row('H2 raw: difference, A minus B', 'H2', 'raw_diff')}
{row('H2 raw: Welch t', 'H2', 'raw_t')}
{row('H2 raw: p (T.TEST type 3)', 'H2', 'raw_p_tt')}
{row('H2 raw: pooled SD', 'H2', 'raw_sp')}

FA has no meaningful zero point: each respondent carries a stable asymmetry, and most raw scores are negative. The raw key scene FA therefore says nothing about the key scene, and the raw end scene comparison carries a chance difference between the baselines of the two groups (`{rcell('H2', 'chk_bl_diff')}`, {val('H2', 'chk_bl_diff', 3) if ('H2', 'chk_bl_diff') in LO else ''}) and a larger SD (pooled {val('H2', 'raw_sp', 3)} against {val('H2', 'h2_sp', 3)} for the change scores). This is why the planned tests use baseline corrected scores.

---

## 10. Teaching contrast: the notebook default channels

The FA notebook's Enobio defaults are **left Ch1, right Ch7**. With the lab manual montage those are Fz (midline frontal) and P3 (left parietal), not a left and right frontal pair. The **Default** sheet holds the export made with those defaults, and columns G to I of the H1 sheet repeat H1 on it:

| Statistic | Cell | Formula | Result (simulated) |
|---|---|---|---|
{row('Default channels: baseline mean', 'H1', 'def_m1')}
{row('Default channels: KeyScene mean', 'H1', 'def_m2')}
{row('Default channels: mean change', 'H1', 'def_md')}
{row('Default channels: t', 'H1', 'def_t')}
{row('Default channels: p (T.TEST type 1)', 'H1', 'def_p_tt')}
{row('Default channels: d_z', 'H1', 'def_dz')}
{row('Correlation of the default channel change with the F3/F4 change', 'H1', 'chk_r_def', 2)}

**You should see** "FA" values near +0.7 in every row, a first warning sign: ln(P3 alpha) - ln(Fz alpha) is large and positive because alpha is stronger over the back of the head. The key scene effect disappears, and the default channel changes are unrelated to the F3/F4 changes. Nothing in the export flags the mistake; only the channel map does. Read the NIC2 protocol (the order in which the electrodes were dragged onto the montage) before setting the left and right electrode sets.

---

## 11. Charts with 95 percent confidence intervals

1. Make a small chart table: one row per mean, columns **Mean** (point to the mean cells of a block) and **CI half width** = `CONFIDENCE.T(0.05, SD cell, n cell)`. The workbook has one on the H1 sheet at `H1!K{H1_CHART_ROW + 2}:M{H1_CHART_ROW + 3}` (half widths in `{rcell('H1', 'chart_ci_0')}` and `{rcell('H1', 'chart_ci_1')}`) and one on the H2 sheet at `H2!I{H2_CHART_ROW + 2}:K{H2_CHART_ROW + 3}` (half widths in `{rcell('H2', 'chart_ci_0')}` and `{rcell('H2', 'chart_ci_1')}`).
2. Select the labels and the means, **Insert > Column chart > Clustered Column**.
3. Click the bars, **Chart Design > Add Chart Element > Error Bars > More Error Bars Options**, choose **Custom > Specify Value**, and select the CI half width cells for both **Positive** and **Negative**.

**Common mistake.** Reading two overlapping error bars on the H1 chart as "no difference". The paired test uses the interval of the **changes** (`{rcell('H1', 'h1_ci_lo')}` and `{rcell('H1', 'h1_ci_hi')}`), which is narrower because it removes the differences between respondents. Bars around negative means also hang below zero; FA is a difference score, so a bar chart starting at zero is less informative than the dot and line figures in `figures/`.

---

## 12. Check your numbers and know what Excel cannot do

The **Check** sheet puts each live Excel result next to the value `analysis.R` produced and prints OK or CHECK (status in column F, the count of CHECK rows in `{rcell('Check', 'n_check')}`; it should be 0). Several parts of the R analysis have no Excel equivalent:

- **Wilcoxon signed rank** (H1, H3) and **Mann Whitney** (H2) tests, the check that a conclusion does not depend on normality.
- **Exact noncentral t intervals** for d_z, d, and g.
- The **mixed model** with Satterthwaite tests and the **ICC**.
- The **power analysis** (`power.t.test`) for H2.
- **Shapiro Wilk** tests, QQ plots, and **Spearman** correlations.
- The **frontal midline theta** descriptives and the FA against PSD consistency check, which use the PSD export.

If you have the **Analysis ToolPak** add in, **Data > Data Analysis > t-Test: Paired Two Sample for Means** (H1, H3) and **t-Test: Two-Sample Assuming Unequal Variances** (H2) are good cross checks.

---

## 13. Common mistakes, collected

1. Opening the .csv by double clicking and getting one column, or forgetting that the FA files have two comment lines.
2. Running the FA notebook with its default channels without checking the NIC2 montage (section 10).
3. Identifying a segment by `Label` alone; `EndScene` exists on all three commercials (section 2).
4. Comparing a commercial with the wrong baseline; each commercial goes with the baseline right before it (section 5).
5. Leaving excluded respondents in a roster, or applying the 25 percent rule to one channel only.
6. Using `T.TEST` type 1 for the two versions (H2) or type 2 or 3 for the within subjects tests (H1, H3).
7. Reading the raw FA sign as approach or withdrawal. FA is interpreted relative to a baseline or to another stimulus.
8. Reading a nonsignificant H2 as "no effect". Report the interval and the power.
9. Treating {N_FA} rows as {N_FA} independent observations. Each respondent contributes several rows; the tests use one value per respondent.
10. Reading approach as liking. FA indexes approach or withdrawal motivation, not whether the viewer liked the commercial.

## Sources

- iMotions help center, R Notebook: EEG Frontal Asymmetry (FA): https://help.imotions.com/lab/analysis/eeg/r-notebook-eeg-frontal-asymmetry-fa
- iMotions help center, R Processed EEG Export Files: https://help.imotions.com/lab/analysis/eeg/r-processed-eeg-export-files
- iMotions help center, R Notebook: EEG Power Spectral Density (PSD): https://help.imotions.com/lab/analysis/eeg/r-notebook-eeg-power-spectral-density-psd
- iMotions help center, How to check EEG data quality: https://help.imotions.com/lab/sensor-setup/eeg/how-to-check-eeg-data-quality
- iMotions help center, Setting up NeuroElectrics Enobio: https://help.imotions.com/lab/sensor-setup/eeg/neuroelectrics/setting-up-neuroelectrics-enobio
- iMotions help center, Annotations: https://help.imotions.com/lab/analysis/my-library/annotations
- iMotions help center, Opening Data Export Files in Excel: https://help.imotions.com/lab/exports/additional-export-information/opening-data-export-files-in-excel
- iMotions help center, Creating Excel Pivot Tables: https://help.imotions.com/lab/exports/additional-export-information/creating-excel-pivot-tables
- Bates, D., Mächler, M., Bolker, B., & Walker, S. (2015). Fitting linear mixed-effects models using lme4. *Journal of Statistical Software, 67*(1), 1-48. https://doi.org/10.18637/jss.v067.i01
- Delacre, M., Lakens, D., & Leys, C. (2017). Why psychologists should by default use Welch's t-test instead of Student's t-test. *International Review of Social Psychology, 30*(1), 92-101. https://doi.org/10.5334/irsp.82
- Lakens, D. (2013). Calculating and reporting effect sizes to facilitate cumulative science: A practical primer for t-tests and ANOVAs. *Frontiers in Psychology, 4*, Article 863. https://doi.org/10.3389/fpsyg.2013.00863

## Items to verify on the lab computer

- The **channel map**: which 10-10 site each of Ch1 to Ch8 holds in the lab NIC2 protocol (this workbook assumes Ch5 = F3 and Ch6 = F4), and therefore which channels to enter as the left and right frontal electrode sets.
- The **artifact report**: the layout of the notebook report box "Percentage of data filtered (nb of samples filtered / nb total of samples)", whether it gives values per stimulus or for the whole recording, and whether the threshold is 120 or 100 microvolts in Lab 11.1. The simulated file is a modeled table.
- The **FA and PSD export details**: the Type values written for annotation rows ("Annotation" or "Annotation (Instance)"), what Parent Stimulus holds on a stimulus row, whether "Algorithm Device Name" is one column or two, the file type (.csv or .txt), and the PSD column order.
- How the notebook computes a **stimulus mean FA** (the mean of the 1 s FA values, as simulated here, or ln of the mean powers) and how it handles a segment in which every window was removed (NA in the simulation).
- The **session log** and **survey** layouts are lab designs, not iMotions exports.
- The sampling rate shown in the LSL widget (500 Hz expected) and whether the Enobio device profile was selected before recording.
"""
with open(MD_OUT, "w", encoding="utf-8") as fh:
    fh.write(md)
print("wrote", MD_OUT)
