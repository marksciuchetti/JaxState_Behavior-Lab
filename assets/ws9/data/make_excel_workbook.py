#!/usr/bin/env python3
"""
make_excel_workbook.py  (PLANNED STUDY, SIMULATED planning data; teaching only)
PLANNED AND SIMULATED. Builds SIMULATED_planning.xlsx (the Excel route of Workshop 9) with openpyxl and
writes excel_walkthrough.md from the SAME registry of cells, so that every cell address in the walkthrough
is the address of a live formula in the workbook. No real person was recorded; ALTAVO and BRENLOW are
invented brands; participant codes (P01, R01) and staff codes (S1 to S6) are placeholders.

The Excel route covers:
  1. SampleSize: the sample size of a paired comparison by the normal approximation with the small
     sample correction, for d_z = 0.30, 0.50, and the Workshop 7 and 8 estimates, with and without a Holm
     family of three, and the approximate power at 48 participants; then the approximate power of the
     brand test when ads are a random factor (why more ads beat more participants);
  2. Attrition: the sessions to book from the Workshop 7 exclusion rate, the reserve, the cap, the days,
     and the remote comparison from Workshop 6;
  3. LatinSquare: the Williams design for 12 ads built with MOD, the ad labels, and its balance checks;
  4. Assignment: the Latin square row of each participant code (from R) and its block check;
  5. Calendar: the session calendar with WORKDAY, slot times, codes, rows, and the staff pairs;
  6. Check: every Excel result next to the value plan_power.R produced.

Verification after the workbook is saved (as in Workshop 8):
  1. openpyxl reads the workbook back and checks that every registered cell holds the registered formula.
  2. LibreOffice (headless) recalculates a copy in a temporary folder (the delivered workbook is not
     changed); every row of the Check sheet must say OK and no cell may show an error value.
  3. Every result printed in the walkthrough is the LibreOffice value of that cell, and it must agree with
     the R value (at the printed decimals, or within the stated tolerance for an approximation).

Run: python3 make_excel_workbook.py   (after Rscript plan_power.R)
"""

import csv
import datetime as dt
import os
import re
import shutil
import subprocess
import sys
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "SIMULATED_planning.xlsx")
MD_OUT = os.path.join(HERE, "excel_walkthrough.md")

BLUE, ORANGE, GRAY = "0072B2", "E69F00", "8C8C8C"
hdr_font = Font(bold=True, color="FFFFFF")
hdr_fill = PatternFill("solid", fgColor=BLUE)
r_fill = PatternFill("solid", fgColor="E7E6E6")
yellow = PatternFill("solid", fgColor="FFF2CC")
bold = Font(bold=True)
note_font = Font(italic=True, color="555555")
thin = Side(style="thin", color="BFBFBF")
box = Border(top=thin, bottom=thin, left=thin, right=thin)
EXCEL_EPOCH = dt.date(1899, 12, 30)


# ---------------------------------------------------------------- inputs (the R results)
def read_banner_csv(name):
    with open(os.path.join(HERE, name), encoding="utf-8") as fh:
        raw = fh.read().splitlines()
    start = 1 if raw[0].startswith("#") else 0
    rows = list(csv.reader(raw[start:]))
    return raw[0], rows[0], [dict(zip(rows[0], r)) for r in rows[1:]]


banner_an, _, AN = read_banner_csv("SIMULATED_power_analytic.csv")
SEED = re.search(r"seed (\d+)", banner_an).group(1)
_, _, GRID = read_banner_csv("SIMULATED_power_grid.csv")
_, _, ASSUME = read_banner_csv("SIMULATED_planning_assumptions.csv")
_, SQ_HEAD, SQ = read_banner_csv("SIMULATED_latin_square.csv")
_, _, ORD = read_banner_csv("SIMULATED_order_schedule.csv")
_, _, SCH = read_banner_csv("SIMULATED_session_schedule.csv")
with open(os.path.join(HERE, "plan_output.txt"), encoding="utf-8") as fh:
    OUT_TXT = fh.read()


def grab(pattern, cast=float):
    m = re.search(pattern, OUT_TXT)
    assert m, f"not found in plan_output.txt: {pattern}"
    return [cast(g) for g in m.groups()]


R_SESS, R_MAIN, R_RES, R_CAP = grab(r"-> (\d+) sessions: (\d+) main \+ (\d+) reserve; hard cap (\d+) sessions", int)
R_REMOTE, = grab(r"would need about (\d+) people to open the link", int)
R_SLOTS, R_DAYS, R_OPEN = grab(r"Main collection: (\d+) slots on (\d+) weekdays .*?, (\d+) open", int)
R_FIRST_A, = grab(r"first brand ALTAVO for (\d+) and BRENLOW", int)
A15 = next(a for a in ASSUME if a["id"] == "A15")
A16 = next(a for a in ASSUME if a["id"] == "A16")
A3 = next(a for a in ASSUME if a["id"] == "A3")
WS6 = [int(x) for x in re.findall(r"(\d+) of (\d+) in Workshop 6; then (\d+) of (\d+) completed respondents were excluded, so (\d+) of", A16["how"])[0]]
WS7 = [int(x) for x in re.findall(r"(\d+) of (\d+) participants kept", A15["how"])[0]]
MAIN = [o for o in ORD if o["phase"] == "main"]
SCH_MAIN = [s for s in SCH if s["phase"] in ("main", "reserve", "open")]
T_ADS = len(SQ)
assert T_ADS == 12, "the LatinSquare sheet is laid out for 12 ads"
EFFECTS = [(a["effect"], float(a["d_z"])) for a in AN if float(a["alpha"]) == 0.05]
AN_SINGLE = {float(a["d_z"]): a for a in AN if float(a["alpha"]) == 0.05}
AN_HOLM = {float(a["d_z"]): a for a in AN if float(a["alpha"]) < 0.05}

# ---------------------------------------------------------------- registry of documented cells
REG, EXPECT = {}, {}


def reg(sheet, key, cell, formula, label=""):
    REG[(sheet, key)] = {"cell": cell, "formula": formula.replace("_xlfn.", ""), "label": label}


def put(ws, key, cell, formula, label="", fmt=None, expect=None, nd=3, tol=None):
    """Write a formula, register it, and (optionally) register the R value it must reproduce.
    tol = None: Excel and R must agree at nd printed decimals; tol = x: |Excel - R| <= x (approximations)."""
    ws[cell] = formula
    if fmt:
        ws[cell].number_format = fmt
    reg(ws.title, key, cell, formula, label)
    if expect is not None:
        EXPECT[(ws.title, key)] = (float(expect), nd, tol)


def hdr(ws, row, first, last, fill=None):
    for c in range(first, last + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill, cell.border = hdr_font, fill or hdr_fill, box
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def widths(ws, d):
    for col, w in d.items():
        ws.column_dimensions[col].width = w


def typed(ws, cell, value, fmt=None):
    ws[cell] = value
    ws[cell].fill = yellow
    if fmt:
        ws[cell].number_format = fmt


def from_r(ws, cell, value, fmt=None):
    ws[cell] = value
    ws[cell].fill = r_fill
    if fmt:
        ws[cell].number_format = fmt


wb = Workbook()
# ================================================================ README
ws = wb.active
ws.title = "README"
readme = [
    ("PLANNED STUDY / SIMULATED PLANNING DATA: the Excel route of Workshop 9 (study design, consent, and the IRB)", bold),
    (f"Built by make_excel_workbook.py from the CSV files that plan_power.R wrote (seed {SEED}). No real person was recorded; ALTAVO and BRENLOW are "
     "invented brands; participant codes (P01, R01) and staff codes (S1 to S6) are placeholders. Dates are placeholders.", None),
    ("The follow up study: 12 ads (6 per brand), every participant sees every ad in a Williams order, one participant at a time at the single "
     "Smart Eye Aurora seat; 48 participants analyzed.", None),
    ("", None),
    ("Sheets", bold),
    ("SampleSize   A. sample size of a paired comparison (normal approximation with the small sample correction), with and without a Holm "
     "family of three; B. approximate power of the brand test when ads are a random factor", None),
    ("Attrition    sessions to book from the Workshop 7 exclusion rate; reserve, cap, days; the remote comparison from Workshop 6", None),
    ("LatinSquare  the Williams design for 12 ads (MOD formulas), the ad labels, and the balance checks", None),
    ("Assignment   the Latin square row of each participant code (drawn in R with a fixed seed) and the block check", None),
    ("Calendar     one row per session slot: WORKDAY dates, slot times, codes, rows, and the staff pair", None),
    ("Check        every Excel result next to the value plan_power.R produced, with OK or CHECK", None),
    ("", None),
    ("Colors: blue headers = column names; yellow cells are typed by you (inputs); gray cells are values copied from R for comparison.", None),
    ("The walkthrough excel_walkthrough.md explains every formula with its cell address.", None),
]
for i, (txt, fnt) in enumerate(readme, start=1):
    c = ws.cell(row=i, column=1, value=txt)
    if fnt:
        c.font = fnt
ws.column_dimensions["A"].width = 150

# ================================================================ SampleSize
ws = wb.create_sheet("SampleSize")
ws["A1"] = "A. Sample size of a paired comparison (each participant gives one score in each condition)"
ws["A1"].font = bold
ws["A2"] = ("Normal approximation n = ((z_alpha + z_power) / d_z)^2, plus z_alpha^2 / 2 as the small sample correction; power at n from the same "
            "formula solved for the power. Compare with R's power.t.test on the Check sheet.")
ws["A2"].font = note_font
inputs = [("Alpha (two tailed)", 0.05, "0.00"), ("Power target", 0.80, "0.00"), ("Tests in the Holm family", 3, "0"), ("Participants for the power columns", 48, "0")]
for i, (lab, v, fm) in enumerate(inputs, start=4):
    ws[f"A{i}"] = lab
    typed(ws, f"B{i}", v, fm)
put(ws, "z_a", "B9", "=_xlfn.NORM.S.INV(1-B4/2)", "z for alpha (two tailed)", "0.0000")
put(ws, "z_h", "B10", "=_xlfn.NORM.S.INV(1-B4/B6/2)", "z at the strictest Holm step (alpha / 3)", "0.0000")
put(ws, "z_b", "B11", "=_xlfn.NORM.S.INV(B5)", "z for the power target", "0.0000")
ws["A9"], ws["A10"], ws["A11"] = "z for alpha (two tailed)", "z at the strictest Holm step", "z for the power target"
heads = ["Effect", "d_z", "n, normal approximation", "n, with the small sample correction", "n needed (round up)", "Holm: n, normal",
         "Holm: n, corrected", "Holm: n needed", "Power at B7 (approximate)", "Holm: power at B7 (approximate)"]
for c, h in enumerate(heads, start=1):
    ws.cell(row=13, column=c, value=h)
hdr(ws, 13, 1, len(heads))
SS = {}
for k, (lab, d) in enumerate(EFFECTS):
    r = 14 + k
    ws[f"A{r}"] = lab
    typed(ws, f"B{r}", d, "0.0000")
    s, h = AN_SINGLE[d], AN_HOLM[d]
    tag = f"d_z = {d:.3f}"
    put(ws, f"n_norm_{k}", f"C{r}", f"=(($B$9+$B$11)/B{r})^2", f"n, normal approximation, {tag}", "0.00")
    put(ws, f"n_corr_{k}", f"D{r}", f"=C{r}+$B$9^2/2", f"n, corrected, {tag} (R: power.t.test n)", "0.00", float(s["n_exact"]), 2, 0.25)
    put(ws, f"n_need_{k}", f"E{r}", f"=ROUNDUP(D{r},0)", f"n needed, {tag}", "0", float(s["n_needed"]), 0)
    put(ws, f"h_norm_{k}", f"F{r}", f"=(($B$10+$B$11)/B{r})^2", f"Holm: n, normal approximation, {tag}", "0.00")
    put(ws, f"h_corr_{k}", f"G{r}", f"=F{r}+$B$10^2/2", f"Holm: n, corrected, {tag} (R: power.t.test n)", "0.00", float(h["n_exact"]), 2, 0.25)
    put(ws, f"h_need_{k}", f"H{r}", f"=ROUNDUP(G{r},0)", f"Holm: n needed, {tag}", "0", float(h["n_needed"]), 0)
    put(ws, f"pow_{k}", f"I{r}", f"=_xlfn.NORM.S.DIST(B{r}*SQRT($B$7-$B$9^2/2)-$B$9,TRUE)", f"power at 48, {tag} (R: power.t.test)", "0.0000",
        float(s["power_n48"]), 4, 0.005)
    put(ws, f"hpow_{k}", f"J{r}", f"=_xlfn.NORM.S.DIST(B{r}*SQRT($B$7-$B$10^2/2)-$B$10,TRUE)", f"Holm: power at 48, {tag} (R: power.t.test)", "0.0000",
        float(h["power_n48"]), 4, 0.005)
    SS[k] = (lab, d, r)
# B. brand test with ads random
ws["A20"] = "B. The brand test when ads are a random factor: approximate power (balanced design, every participant rates every ad)"
ws["A20"].font = bold
binp = [("Brand difference to detect (scale points)", 0.50, "0.00"), ("SD of the ads within a brand (judgment)", 0.20, "0.00"),
        ("Residual SD of liking (Workshop 7)", float(A3["value"]), "0.000")]
for i, (lab, v, fm) in enumerate(binp, start=21):
    ws[f"A{i}"] = lab
    typed(ws, f"B{i}", v, fm)
bh = ["Participants", "Ads per brand", "df = 2 x ads - 2", "SE of the brand difference", "Noncentrality", "Critical t", "Approximate power"]
for c, h in enumerate(bh, start=1):
    ws.cell(row=25, column=c, value=h)
hdr(ws, 25, 1, len(bh))
GRIDF = {(int(float(g["n_participants"])), int(float(g["ads_per_brand"]))): float(g["power_formula"]) for g in GRID}
BT = []
for k, (n, m) in enumerate([(48, 2), (48, 3), (48, 4), (48, 6), (72, 2), (72, 3), (72, 4), (72, 6)]):
    r = 26 + k
    typed(ws, f"A{r}", n, "0")
    typed(ws, f"B{r}", m, "0")
    put(ws, f"bdf_{k}", f"C{r}", f"=2*B{r}-2", f"df, {n} participants, {m} ads per brand", "0")
    put(ws, f"bse_{k}", f"D{r}", f"=SQRT(2*($B$22^2+$B$23^2/A{r})/B{r})", f"SE, {n} participants, {m} ads per brand", "0.0000")
    put(ws, f"bncp_{k}", f"E{r}", f"=$B$21/D{r}", f"noncentrality, {n} x {m}", "0.000")
    put(ws, f"btc_{k}", f"F{r}", f"=_xlfn.T.INV.2T($B$4,C{r})", f"critical t, {m} ads per brand", "0.000")
    put(ws, f"bpow_{k}", f"G{r}", f"=1-_xlfn.NORM.S.DIST((F{r}*(1-1/(4*C{r}))-E{r})/SQRT(1+F{r}^2/(2*C{r})),TRUE)",
        f"approximate power, {n} participants, {m} ads per brand (R: exact formula)", "0.000", GRIDF[(n, m)], 3, 0.02)
    BT.append((n, m, r))
ws["A35"] = ("The power uses a normal approximation to the noncentral t distribution (Abramowitz and Stegun, 1964, chapter 26); R computes it exactly "
             "with pt(..., ncp). Read down column G: going from 48 to 72 participants changes little; going from 3 to 6 ads per brand changes a lot.")
ws["A35"].font = note_font
widths(ws, {"A": 38, "B": 12, "C": 16, "D": 18, "E": 14, "F": 14, "G": 16, "H": 14, "I": 16, "J": 18})

# ================================================================ Attrition
ws = wb.create_sheet("Attrition")
ws["A1"] = "Sessions to book, reserve, cap, and the remote comparison"
ws["A1"].font = bold
ainp = [("Participants to analyze (target)", R_MAIN), ("Workshop 7: participants excluded by rule R2", WS7[1] - WS7[0]), ("Workshop 7: participants recorded", WS7[1]),
        ("Rows of the Williams design (block size)", T_ADS), ("Cap as a multiple of the target", 1.25), ("Workshop 6: sessions that opened the link", WS6[1]),
        ("Workshop 6: completed", WS6[0]), ("Workshop 6: kept after the exclusion rule", WS6[4]), ("Session slots per day (one Aurora seat)", 5)]
for i, (lab, v) in enumerate(ainp, start=4):
    ws[f"A{i}"] = lab
    typed(ws, f"B{i}", v, "0.00" if isinstance(v, float) else "0")
ws["A14"] = "Results"
ws["A14"].font = bold
alab = [("rate", "B15", "=B5/B6", "Exclusion rate in the lab (Workshop 7)", "0.0000", float(A15["value"]), 4, None),
        ("exp", "B16", "=B4/(1-B15)", "Sessions expected to end with the target", "0.00", None, 2, None),
        ("book", "B17", "=ROUNDUP(B16,0)", "Sessions to book (main + reserve)", "0", R_SESS, 0, None),
        ("reserve", "B18", "=B17-B4", "Reserve sessions", "0", R_RES, 0, None),
        ("cap", "B19", "=B4*B8", "Hard cap on sessions", "0", R_CAP, 0, None),
        ("days", "B20", "=ROUNDUP(B17/B12,0)", "Weekdays of data collection", "0", R_DAYS, 0, None),
        ("slots", "B21", "=B20*B12", "Slots on those days", "0", R_SLOTS, 0, None),
        ("open", "B22", "=B21-B17", "Open slots for rebooking no shows", "0", R_OPEN, 0, None),
        ("r_done", "B25", "=B10/B9", "Remote: completed share of those who opened the link", "0.0000", WS6[0] / WS6[1], 4, None),
        ("r_excl", "B26", "=(B10-B11)/B10", "Remote: share of completed respondents excluded", "0.0000", (WS6[0] - WS6[4]) / WS6[0], 4, None),
        ("r_kept", "B27", "=B11/B9", "Remote: kept share of those who opened the link", "0.0000", WS6[4] / WS6[1], 4, None),
        ("r_need", "B28", "=ROUNDUP(B4/B27,0)", "Remote: people who must open the link to keep the target", "0", R_REMOTE, 0, None)]
ws["A24"] = "If the same study were run remotely (Workshop 6 rates)"
ws["A24"].font = bold
for key, cell, fm, lab, nf, ex, nd, tol in alab:
    ws[f"A{cell[1:]}"] = lab
    put(ws, key, cell, fm, lab, nf, ex, nd, tol)
widths(ws, {"A": 58, "B": 14})

# ================================================================ LatinSquare
ws = wb.create_sheet("LatinSquare")
ws["A1"] = "Williams design for 12 ads: row 1 = 0, 1, 11, 2, 10, 3, ...; row r adds r - 1 (mod 12). Labels: even symbols ALTAVO, odd symbols BRENLOW."
ws["A1"].font = bold
ws["A3"] = "Number of ads (t)"
ws["B3"] = T_ADS
ws["C3"] = "Fixed at 12 for this layout (the chosen design). The construction works for any even t."
ws["C3"].font = note_font
C = [get_column_letter(3 + j) for j in range(T_ADS)]          # C .. N
ws["B5"] = "Position"
for j, col in enumerate(C):
    ws[f"{col}5"] = j + 1
hdr(ws, 5, 2, 2 + T_ADS)
ws["B6"] = "Row 1 (symbols)"
for col in C:
    put(ws, f"first_{col}", f"{col}6", f"=IF({col}$5=1,0,IF(MOD({col}$5,2)=0,{col}$5/2,$B$3-({col}$5-1)/2))", f"first row, position in column {col}")
ws["A8"] = "Symbols"
for r in range(T_ADS):
    rr = 8 + r
    ws[f"B{rr}"] = r + 1
    for col in C:
        put(ws, f"sym_{col}{rr}", f"{col}{rr}", f"=MOD({col}$6+$B{rr}-1,$B$3)", f"symbol, row {r + 1}, column {col}")
ws["A22"] = "Ad labels"
ws["O21"] = "Brand sequence"
for r in range(T_ADS):
    rr, rs = 22 + r, 8 + r
    ws[f"B{rr}"] = r + 1
    for col in C:
        put(ws, f"lab_{col}{rr}", f"{col}{rr}", f'=IF(MOD({col}{rs},2)=0,"ALTAVO_"&({col}{rs}/2+1),"BRENLOW_"&(({col}{rs}+1)/2))', f"label, row {r + 1}, column {col}")
    put(ws, f"seq_{rr}", f"O{rr}", "=" + "&".join(f'IF(MOD({col}{rs},2)=0,"A","B")' for col in C), f"brand sequence, row {r + 1}")
# checks: once per row, once per position, each ordered pair adjacent once, ALTAVO per position
ws["A35"] = "Checks"
ws["A35"].font = bold
ws["B36"] = "Symbol ->"
for j, col in enumerate(C):
    ws[f"{col}36"] = j
for r in range(T_ADS):
    rr = 37 + r
    ws[f"B{rr}"] = f"row {r + 1}"
    for col in C:
        put(ws, f"cr_{col}{rr}", f"{col}{rr}", f"=COUNTIF($C{8 + r}:$N{8 + r},{col}$36)", "count of a symbol in a row")
put(ws, "row_min", "C50", "=MIN(C37:N48)", "each ad once per row: smallest count", "0", 1, 0)
put(ws, "row_max", "C51", "=MAX(C37:N48)", "each ad once per row: largest count", "0", 1, 0)
ws["B50"], ws["B51"] = "smallest count", "largest count"
ws["B53"] = "Symbol (rows) x position (columns)"
for s in range(T_ADS):
    rr = 54 + s
    ws[f"B{rr}"] = s
    for col in C:
        put(ws, f"cp_{col}{rr}", f"{col}{rr}", f"=COUNTIF({col}$8:{col}$19,$B{rr})", "count of a symbol in a position")
put(ws, "pos_min", "C67", "=MIN(C54:N65)", "each ad once per position: smallest count", "0", 1, 0)
put(ws, "pos_max", "C68", "=MAX(C54:N65)", "each ad once per position: largest count", "0", 1, 0)
ws["B67"], ws["B68"] = "smallest count", "largest count"
ws["B70"] = "Pairs: ad i (rows) immediately followed by ad j (columns)"
for j, col in enumerate(C):
    ws[f"{col}71"] = j
for i in range(T_ADS):
    rr = 72 + i
    ws[f"B{rr}"] = i
    for col in C:
        put(ws, f"pr_{col}{rr}", f"{col}{rr}", f"=SUMPRODUCT(($C$8:$M$19=$B{rr})*($D$8:$N$19={col}$71))", "count of an ordered pair")
put(ws, "pair_sum", "C85", "=SUM(C72:N83)", "adjacent pairs in all rows (12 x 11)", "0", T_ADS * (T_ADS - 1), 0)
put(ws, "pair_diag", "C86", "=" + "+".join(f"{C[i]}{72 + i}" for i in range(T_ADS)), "pairs of an ad with itself (diagonal)", "0", 0, 0)
put(ws, "pair_max", "C87", "=MAX(C72:N83)", "largest pair count", "0", 1, 0)
ws["B85"], ws["B86"], ws["B87"] = "sum of all pair counts", "diagonal (an ad after itself)", "largest pair count"
ws["B89"] = "ALTAVO ads per position"
for col in C:
    put(ws, f"alt_{col}", f"{col}89", f"=SUMPRODUCT(--(MOD({col}8:{col}19,2)=0))", f"ALTAVO ads in the position of column {col}", "0")
put(ws, "alt_min", "C90", "=MIN(C89:N89)", "ALTAVO ads per position: smallest", "0", T_ADS / 2, 0)
put(ws, "alt_max", "C91", "=MAX(C89:N89)", "ALTAVO ads per position: largest", "0", T_ADS / 2, 0)
ws["B90"], ws["B91"] = "smallest", "largest"
ws["B93"] = "The square from R (SIMULATED_latin_square.csv), for comparison"
for r, row in enumerate(SQ):
    rr = 94 + r
    ws[f"B{rr}"] = int(row["row"])
    for j, col in enumerate(C):
        from_r(ws, f"{col}{rr}", row[f"position_{j + 1:02d}"])
put(ws, "sq_mismatch", "C107", "=SUMPRODUCT(--(C22:N33<>C94:N105))", "labels that differ from the R square", "0", 0, 0)
ws["B107"] = "labels that differ from R"
widths(ws, {"A": 10, "B": 18, **{col: 11 for col in C}, "O": 16})

# ================================================================ Assignment
ws = wb.create_sheet("Assignment")
ws["A1"] = "Latin square row of each participant code: blocks of 12, every row once per block, random order within a block (drawn in R, seed fixed)"
ws["A1"].font = bold
for c, h in enumerate(["Participant code", "Block", "Latin square row", "First ad", "First brand"], start=1):
    ws.cell(row=3, column=c, value=h)
hdr(ws, 3, 1, 5)
for i, o in enumerate(MAIN):
    rr = 4 + i
    typed(ws, f"A{rr}", o["participant_code"])
    typed(ws, f"B{rr}", int(o["block"]), "0")
    typed(ws, f"C{rr}", int(o["latin_square_row"]), "0")
    put(ws, f"first_{rr}", f"D{rr}", f"=INDEX(LatinSquare!$C$22:$C$33,C{rr})", f"first ad of {o['participant_code']}")
    put(ws, f"fb_{rr}", f"E{rr}", f'=IF(LEFT(D{rr},6)="ALTAVO","ALTAVO","BRENLOW")', f"first brand of {o['participant_code']}")
AL = 3 + len(MAIN)
n_blocks = len(MAIN) // T_ADS
ws["G3"] = "Row (down) x block (across): participants"
for b in range(n_blocks):
    ws.cell(row=4, column=8 + b, value=b + 1)
for r in range(T_ADS):
    rr = 5 + r
    ws[f"G{rr}"] = r + 1
    for b in range(n_blocks):
        col = get_column_letter(8 + b)
        put(ws, f"blk_{col}{rr}", f"{col}{rr}", f"=COUNTIFS($B$4:$B${AL},{col}$4,$C$4:$C${AL},$G{rr})", "participants per row and block")
lastc = get_column_letter(7 + n_blocks)
put(ws, "blk_min", "H18", f"=MIN(H5:{lastc}16)", "every row once in every block: smallest count", "0", 1, 0)
put(ws, "blk_max", "H19", f"=MAX(H5:{lastc}16)", "every row once in every block: largest count", "0", 1, 0)
put(ws, "first_alt", "H20", f'=COUNTIF(E4:E{AL},"ALTAVO")', "participants who see an ALTAVO ad first", "0", R_FIRST_A, 0)
ws["G18"], ws["G19"], ws["G20"] = "smallest count", "largest count", "ALTAVO first"
widths(ws, {"A": 16, "B": 8, "C": 16, "D": 13, "E": 13, "G": 22})

# ================================================================ Calendar
ws = wb.create_sheet("Calendar")
ws["A1"] = "Session calendar for one Aurora seat (PLACEHOLDER dates: set them after IRB approval; check the JSU academic calendar for closures)"
ws["A1"].font = bold
cinp = [("First day of data collection", dt.date(2027, 2, 8), "yyyy-mm-dd"), ("Slots per day", None, "0"), ("First slot", dt.time(9, 0), "hh:mm"),
        ("Slot length (minutes)", 90, "0"), ("Main participants", None, "0"), ("Reserve sessions", None, "0")]
for i, (lab, v, fm) in enumerate(cinp, start=3):
    ws[f"A{i}"] = lab
    if v is not None:
        typed(ws, f"B{i}", v, fm)
put(ws, "c_slots", "B4", "=Attrition!B12", "slots per day (from Attrition)", "0")
put(ws, "c_main", "B7", "=Attrition!B4", "main participants (from Attrition)", "0")
put(ws, "c_res", "B8", "=Attrition!B18", "reserve sessions (from Attrition)", "0")
put(ws, "c_total", "B9", "=Attrition!B21", "slots in all (from Attrition)", "0", R_SLOTS, 0)
ws["A9"] = "Slots in all"
cal_h = ["k", "Day index", "Date", "Weekday", "Slot", "Arrive", "Recording starts", "Recording ends", "Participant leaves", "Room ready", "Phase",
         "Participant code", "Latin square row", "Shift", "Staff pair", "Lead operator", "Assistant"]
for c, h in enumerate(cal_h, start=1):
    ws.cell(row=11, column=c, value=h)
hdr(ws, 11, 1, len(cal_h))
for c, h in enumerate(["R: date", "R: arrive", "R: code", "R: row", "R: lead", "R: assistant"], start=20):
    ws.cell(row=11, column=c, value=h)
hdr(ws, 11, 20, 25, PatternFill("solid", fgColor=GRAY))
K1 = 12
KL = K1 + len(SCH_MAIN) - 1
for i, srow in enumerate(SCH_MAIN):
    r = K1 + i
    ws[f"A{r}"] = i + 1
    f = {"B": f"=INT((A{r}-1)/$B$4)", "C": f"=WORKDAY($B$3,B{r})", "D": f'=TEXT(C{r},"dddd")', "E": f"=MOD(A{r}-1,$B$4)+1",
         "F": f"=$B$5+(E{r}-1)*$B$6/1440", "G": f"=F{r}+25/1440", "H": f"=F{r}+45/1440", "I": f"=F{r}+60/1440", "J": f"=F{r}+$B$6/1440",
         "K": f'=IF(A{r}<=$B$7,"main",IF(A{r}<=$B$7+$B$8,"reserve","open"))',
         "L": f'=IF(K{r}="main","P"&TEXT(A{r},"00"),IF(K{r}="reserve","R"&TEXT(A{r}-$B$7,"00"),""))',
         "M": f'=IF(K{r}="main",INDEX(Assignment!$C$4:$C${AL},MATCH(L{r},Assignment!$A$4:$A${AL},0)),"")',
         "N": f"=2*B{r}+IF(E{r}>=4,1,0)", "O": f"=MOD(N{r},3)+1", "P": f'=CHOOSE(O{r},"S1","S3","S5")', "Q": f'=CHOOSE(O{r},"S2","S4","S6")'}
    fmts = {"C": "yyyy-mm-dd", "F": "hh:mm", "G": "hh:mm", "H": "hh:mm", "I": "hh:mm", "J": "hh:mm"}
    for col, fm in f.items():
        if r == K1:
            put(ws, f"cal_{col}", f"{col}{r}", fm, f"column {col} of the first slot", fmts.get(col))
        else:
            ws[f"{col}{r}"] = fm
            if col in fmts:
                ws[f"{col}{r}"].number_format = fmts[col]
    y, mo, d = (int(x) for x in srow["date"].split("-"))
    hh, mm = (int(x) for x in srow["arrive"].split(":"))
    from_r(ws, f"T{r}", dt.date(y, mo, d), "yyyy-mm-dd")
    from_r(ws, f"U{r}", dt.time(hh, mm), "hh:mm")
    from_r(ws, f"V{r}", "" if srow["participant_code"] in ("", "NA") else srow["participant_code"])
    from_r(ws, f"W{r}", "" if srow["latin_square_row"] in ("", "NA") else int(srow["latin_square_row"]))
    from_r(ws, f"X{r}", srow["lead_operator"])
    from_r(ws, f"Y{r}", srow["assistant"])
mis = [("mis_date", f"=SUMPRODUCT(--(C{K1}:C{KL}<>T{K1}:T{KL}))", "dates that differ from R"),
       ("mis_time", f"=SUMPRODUCT(--(ROUND(F{K1}:F{KL}*1440,0)<>ROUND(U{K1}:U{KL}*1440,0)))", "arrival times that differ from R"),
       ("mis_code", f'=SUMPRODUCT(--(L{K1}:L{KL}&""<>V{K1}:V{KL}&""))', "participant codes that differ from R"),
       ("mis_row", f'=SUMPRODUCT(--(M{K1}:M{KL}&""<>W{K1}:W{KL}&""))', "Latin square rows that differ from R"),
       ("mis_lead", f'=SUMPRODUCT(--(P{K1}:P{KL}<>X{K1}:X{KL}))', "lead operators that differ from R"),
       ("mis_asst", f'=SUMPRODUCT(--(Q{K1}:Q{KL}<>Y{K1}:Y{KL}))', "assistants that differ from R")]
for i, (key, fm, lab) in enumerate(mis):
    ws[f"S{3 + i}"] = lab
    put(ws, key, f"T{3 + i}", fm, lab, "0", 0, 0)
put(ws, "last_day", "T9", f"=MAX(C{K1}:C{KL})", "last day of data collection (date serial)", "yyyy-mm-dd",
    (dt.date(*(int(x) for x in SCH_MAIN[-1]["date"].split("-"))) - EXCEL_EPOCH).days, 0)
ws["S9"] = "last day of data collection"
widths(ws, {"A": 30, "B": 12, "C": 12, "D": 11, "E": 6, "F": 8, "G": 10, "H": 10, "I": 10, "J": 10, "K": 9, "L": 11, "M": 10, "N": 7, "O": 7,
            "P": 9, "Q": 9, "S": 38, "T": 12, "U": 9, "V": 9, "W": 7, "X": 7, "Y": 9})
ws.freeze_panes = "B12"

# ================================================================ Check
ws = wb.create_sheet("Check")
ws["A1"] = ("Excel results next to the values that plan_power.R produced (PLANNED / SIMULATED). Tolerance: half a unit in the last printed decimal, or "
            "the stated tolerance for an approximation.")
ws["A1"].font = bold
for c, h in enumerate(["Result", "Excel", "R", "Difference", "Tolerance", "Status"], start=1):
    ws.cell(row=3, column=c, value=h)
hdr(ws, 3, 1, 6)
ck = []
for (sheet, key), (val, nd, tol) in EXPECT.items():
    t = tol if tol is not None else (1e-9 if nd == 0 else 0.5 * 10 ** (-nd))
    ck.append((f"{sheet}: {REG[(sheet, key)]['label']}", f"={sheet}!{REG[(sheet, key)]['cell']}", float(val), t))
CK1 = 4
for k, (lab, fm, rv, tol) in enumerate(ck):
    rr = CK1 + k
    ws.cell(row=rr, column=1, value=lab)
    ws.cell(row=rr, column=2, value=fm).number_format = "0.000000"
    ws.cell(row=rr, column=3, value=rv).number_format = "0.000000"
    ws.cell(row=rr, column=4, value=f"=B{rr}-C{rr}").number_format = "0.000000"
    ws.cell(row=rr, column=5, value=tol).number_format = "0.000000"
    ws.cell(row=rr, column=6, value=f'=IF(ABS(D{rr})<=E{rr}+1E-12,"OK","CHECK")')
CK2 = CK1 + len(ck) - 1
ws.cell(row=CK2 + 2, column=1, value="Rows with status CHECK").font = bold
put(ws, "n_check", f"B{CK2 + 2}", f'=COUNTIF(F{CK1}:F{CK2},"CHECK")', "Rows with status CHECK", "0")
ws.cell(row=CK2 + 3, column=1, value=("Not reproduced in Excel: the simulations of STEP 2 and STEP 3 (Excel cannot fit a mixed model), the exact noncentral t "
                                      "power, and the random draw of the rows (R's random numbers). See plan_output.txt.")).font = note_font
widths(ws, {"A": 84, "B": 14, "C": 14, "D": 13, "E": 11, "F": 9})

wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
print("wrote", OUT)

# =============================================================================
# Verification 1: read back and compare every registered cell
# =============================================================================
wb2 = load_workbook(OUT)
problems = []
for (sheet, key), info in REG.items():
    v = wb2[sheet][info["cell"]].value
    v = str(v).replace("_xlfn.", "") if v is not None else None
    if v != info["formula"]:
        problems.append(f"{sheet}!{info['cell']} ({key}): workbook {v!r} registry {info['formula']!r}")
for k, (lab, fm, rv, tol) in enumerate(ck):
    if wb2["Check"][f"B{CK1 + k}"].value != fm:
        problems.append(f"Check!B{CK1 + k} differs")
if problems:
    print("VERIFICATION FAILED:\n" + "\n".join(problems[:40]))
    sys.exit(1)
n_formulas = sum(1 for w in wb2.worksheets for row in w.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
print(f"verified (openpyxl): {len(REG)} registered cells and {len(ck)} Check rows hold the registered formulas; formula cells in the workbook: {n_formulas}")

# =============================================================================
# Verification 2 and 3: LibreOffice recalculation
# =============================================================================
LO = {}
soffice = shutil.which("soffice") or shutil.which("libreoffice")
lo_note = "LibreOffice was not found, so the formula values were not recalculated here; the printed values are the R values."
LO_SUMMARY = None


def agrees(got, val, nd, tol):
    got, val = float(got), float(val)
    if tol is not None:
        return abs(got - val) <= tol + 1e-12
    return f"{got:.{nd}f}" == f"{val:.{nd}f}" or abs(got - val) <= 1e-9 * max(1.0, abs(val))


def as_number(v):
    if isinstance(v, dt.datetime):
        return (v.date() - EXCEL_EPOCH).days + (v - dt.datetime.combine(v.date(), dt.time())).total_seconds() / 86400
    if isinstance(v, dt.date):
        return (v - EXCEL_EPOCH).days
    if isinstance(v, dt.time):
        return (v.hour * 3600 + v.minute * 60 + v.second) / 86400
    return v


if soffice:
    tmp = tempfile.mkdtemp(prefix="ws9_lo_")
    shutil.copy(OUT, os.path.join(tmp, "in.xlsx"))
    subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", os.path.join(tmp, "out"),
                    os.path.join(tmp, "in.xlsx")], check=True, capture_output=True, timeout=600)
    wb3 = load_workbook(os.path.join(tmp, "out", "in.xlsx"), data_only=True)
    for (sheet, key), info in REG.items():
        LO[(sheet, key)] = as_number(wb3[sheet][info["cell"]].value)
    bad = [(ck[k][0], wb3["Check"][f"B{CK1 + k}"].value, ck[k][2]) for k in range(len(ck)) if wb3["Check"][f"F{CK1 + k}"].value != "OK"]
    errs = [(w.title, c.coordinate, c.value) for w in wb3.worksheets for row in w.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("#") and c.value[:4] in ("#N/A", "#VAL", "#REF", "#NAM", "#DIV", "#NUM")]
    if bad or errs:
        print("LIBREOFFICE CHECK FAILED:", bad[:20], errs[:20])
        sys.exit(1)
    mism = [(s, k, LO[(s, k)], v) for (s, k), (v, nd, tol) in EXPECT.items() if LO[(s, k)] is None or isinstance(LO[(s, k)], str)
            or not agrees(LO[(s, k)], v, nd, tol)]
    if mism:
        print("VALUES DIFFER FROM R:", mism)
        sys.exit(1)
    exact = [(s, k) for (s, k), (v, nd, tol) in EXPECT.items() if tol is None or tol == 0]
    approx = [(s, k) for (s, k), (v, nd, tol) in EXPECT.items() if tol not in (None, 0)]
    max_exact = max(abs(float(LO[x]) - EXPECT[x][0]) for x in exact)
    max_approx = max(abs(float(LO[x]) - EXPECT[x][0]) for x in approx)
    by_tol = {}
    for x in approx:
        by_tol.setdefault(EXPECT[x][2], []).append(abs(float(LO[x]) - EXPECT[x][0]))
    tol_name = {0.25: "participants for the corrected sample sizes", 0.02: "for the approximate power of the brand test", 0.005: "for the approximate paired power"}
    tol_txt = "; ".join(f"{t:g} {tol_name.get(t, '')} ({len(v)} values, largest difference {max(v):.4f})" for t, v in sorted(by_tol.items(), reverse=True))
    lo_ver = subprocess.run([soffice, "--version"], capture_output=True, text=True).stdout.strip()
    n_check_cell = wb3["Check"][REG[("Check", "n_check")]["cell"]].value
    weekday0 = wb3["Calendar"]["D12"].value
    LO_SUMMARY = (len(ck), n_check_cell, max_exact, max_approx, lo_ver, len(exact), len(approx))
    lo_note = (f"A copy was recalculated with {lo_ver} (headless): all {len(ck)} Check rows are OK (`Check!{REG[('Check', 'n_check')]['cell']}` "
               f"= {int(n_check_cell)}), and no cell shows an error value. The {len(exact)} exact results (counts, sample sizes, dates, codes) agree "
               f"with R exactly or at the printed decimals (largest difference {max_exact:.1e}); the {len(approx)} approximations agree within their "
               f"stated tolerances: {tol_txt}.")
    print(f"verified (LibreOffice): {len(ck)} Check rows OK; rows with CHECK = {int(n_check_cell)}; no error values; exact results: largest |Excel - R| = "
          f"{max_exact:.2e} ({len(exact)} values); approximations ({len(approx)} values): {tol_txt}; first weekday = {weekday0}")
    shutil.rmtree(tmp, ignore_errors=True)
else:
    print(lo_note)


# =============================================================================
# The walkthrough, generated from REG
# =============================================================================
def cl(sheet, key):
    return REG[(sheet, key)]["cell"]


def fx(sheet, key):
    return REG[(sheet, key)]["formula"]


def val(sheet, key, nd=3):
    v = LO.get((sheet, key))
    if v is None or isinstance(v, str):
        v = EXPECT.get((sheet, key), (v, nd, None))[0]
    if v is None:
        return ""
    if isinstance(v, str):
        return v
    return f"{float(v):.{nd}f}"


def rv(sheet, key, nd=3):
    return f"{EXPECT[(sheet, key)][0]:.{nd}f}"


def row4(label, sheet, key, nd=3, r_nd=None):
    r_nd = nd if r_nd is None else r_nd
    rcol = rv(sheet, key, r_nd) if (sheet, key) in EXPECT else ""
    return f"| {label} | `{sheet}!{cl(sheet, key)}` | `{fx(sheet, key)}` | {val(sheet, key, nd)} | {rcol} |"


ss_rows = []
for k, (lab, d, r) in SS.items():
    ss_rows.append(f"| {lab} | {d:.3f} | {val('SampleSize', f'n_norm_{k}', 2)} | {val('SampleSize', f'n_corr_{k}', 2)} | {rv('SampleSize', f'n_corr_{k}', 2)} | "
                   f"{val('SampleSize', f'n_need_{k}', 0)} | {rv('SampleSize', f'n_need_{k}', 0)} | {val('SampleSize', f'h_corr_{k}', 2)} | "
                   f"{rv('SampleSize', f'h_corr_{k}', 2)} | {val('SampleSize', f'h_need_{k}', 0)} | {rv('SampleSize', f'h_need_{k}', 0)} |")
pw_rows = []
for k, (lab, d, r) in SS.items():
    pw_rows.append(f"| {lab} | {val('SampleSize', f'pow_{k}', 4)} | {rv('SampleSize', f'pow_{k}', 4)} | {val('SampleSize', f'hpow_{k}', 4)} | "
                   f"{rv('SampleSize', f'hpow_{k}', 4)} |")
bt_rows = []
for k, (n, m, r) in enumerate(BT):
    bt_rows.append(f"| {n} | {m} | {val('SampleSize', f'bdf_{k}', 0)} | {val('SampleSize', f'bse_{k}', 4)} | {val('SampleSize', f'bncp_{k}', 3)} | "
                   f"{val('SampleSize', f'btc_{k}', 3)} | {val('SampleSize', f'bpow_{k}', 3)} | {rv('SampleSize', f'bpow_{k}', 3)} |")
att_rows = "\n".join(row4(lab, "Attrition", key, nd) for key, cell, fm, lab, nf, ex, nd, tol in alab if (("Attrition", key) in EXPECT or key == "exp"))
check_line = (f"`Check!{cl('Check', 'n_check')}` counts the rows with status CHECK and must read 0" +
              (f"; in the LibreOffice recalculation it reads {int(LO_SUMMARY[1])} for {LO_SUMMARY[0]} rows" if LO_SUMMARY else ""))
first_row_vals = ", ".join(val("LatinSquare", f"first_{col}", 0) for col in C)
seq1 = val("LatinSquare", "seq_22")
seq2 = val("LatinSquare", "seq_23")

md = f"""# Excel walkthrough: sample size, attrition, the Williams design, and the session calendar

**PLANNED STUDY, SIMULATED PLANNING DATA.** Every number in this walkthrough comes from `plan_power.R` (Workshop 9, seed {SEED}), which plans the follow up study that Workshop 8 called for. No real person was recorded, ALTAVO and BRENLOW are invented brands, participant and staff codes are placeholders, the dates are placeholders, and no number may be quoted as a finding.

This is the Excel route through the parts of `plan_power.R` that a spreadsheet can do: the analytic sample size (STEP 1), the reason why the number of ads matters (STEP 2, as a formula rather than a simulation), the attrition allowance (STEP 4), the Williams design (STEP 5), and the session calendar (STEP 6). The finished workbook is `SIMULATED_planning.xlsx`. **Every cell address below is written as `Sheet!Cell` and was generated from the workbook itself:** `make_excel_workbook.py` writes the workbook and this file from one list of cells, reads the workbook back with openpyxl to confirm that each address holds the formula printed here, and recalculates a copy to confirm the results. {lo_note}

Companion files: `plan_power.R` and `plan_output.txt` (the R planning script and its full output), the `SIMULATED_*.csv` files (the R values the `Check` sheet uses), `analysis_plan_prereg.md`, `planning_checklist.md`, `README.md`.

---

## 0. What you are working with

| Sheet | What it holds |
|---|---|
| `SampleSize` | A: sample size and power of a paired comparison for four values of d_z, with and without a Holm family of three. B: approximate power of the brand test when ads are a random factor |
| `Attrition` | sessions to book from the Workshop 7 exclusion rate; reserve, cap, days, open slots; the Workshop 6 remote comparison |
| `LatinSquare` | the Williams design for 12 ads, the ad labels, the brand sequence of each row, and the balance checks |
| `Assignment` | the Latin square row of each participant code (P01 to P{len(MAIN):02d}) and the block check |
| `Calendar` | one row per session slot ({len(SCH_MAIN)} slots): date, times, phase, code, row, staff pair |
| `Check` | every Excel result next to the R value ({len(ck)} rows) |

Yellow cells are typed by you (inputs). Gray cells hold values copied from R so that the `Check` sheet can compare them. Excel 2010 or later is needed for `NORM.S.INV`, `NORM.S.DIST`, and `T.INV.2T`.

---

## 1. Sample size of a paired comparison (SampleSize, part A)

A paired comparison means that every participant gives one score in each of two conditions, for example the mean liking of the six ALTAVO ads and the mean liking of the six BRENLOW ads. The effect size d_z is the mean of the participants' differences divided by the SD of those differences (Workshop 8). The ads are treated as fixed: the answer is about THESE ads (section 2 explains the difference).

Type the inputs in `SampleSize!B4:B7` (alpha .05, power .80, three tests in the Holm family, 48 participants for the power columns) and the four effect sizes in `SampleSize!B14:B17`. The three z values:

| Result | Cell | Formula | Value | R |
|---|---|---|---|---|
{row4('z for alpha (two tailed)', 'SampleSize', 'z_a', 4)}
{row4('z at the strictest Holm step (alpha / 3)', 'SampleSize', 'z_h', 4)}
{row4('z for the power target', 'SampleSize', 'z_b', 4)}

The formulas of the first effect row (fill them down to row 17):

| Result | Cell | Formula |
|---|---|---|
| n, normal approximation | `SampleSize!{cl('SampleSize', 'n_norm_0')}` | `{fx('SampleSize', 'n_norm_0')}` |
| n, with the small sample correction | `SampleSize!{cl('SampleSize', 'n_corr_0')}` | `{fx('SampleSize', 'n_corr_0')}` |
| n needed (always round UP) | `SampleSize!{cl('SampleSize', 'n_need_0')}` | `{fx('SampleSize', 'n_need_0')}` |
| Holm: n, normal approximation | `SampleSize!{cl('SampleSize', 'h_norm_0')}` | `{fx('SampleSize', 'h_norm_0')}` |
| Holm: n, corrected | `SampleSize!{cl('SampleSize', 'h_corr_0')}` | `{fx('SampleSize', 'h_corr_0')}` |
| Holm: n needed | `SampleSize!{cl('SampleSize', 'h_need_0')}` | `{fx('SampleSize', 'h_need_0')}` |
| Power at the n in B7 | `SampleSize!{cl('SampleSize', 'pow_0')}` | `{fx('SampleSize', 'pow_0')}` |
| Holm: power at the n in B7 | `SampleSize!{cl('SampleSize', 'hpow_0')}` | `{fx('SampleSize', 'hpow_0')}` |

**You should see** (Excel next to R's `power.t.test`, STEP 1 of `plan_output.txt`):

| Effect | d_z | n, normal | n, corrected | R n exact | n needed | R | Holm n, corrected | R Holm n exact | Holm n needed | R |
|---|---|---|---|---|---|---|---|---|---|---|
{chr(10).join(ss_rows)}

| Effect | Power at 48 (Excel) | R | Holm power at 48 (Excel) | R |
|---|---|---|---|---|
{chr(10).join(pw_rows)}

Why the correction: the normal approximation alone pretends that the SD of the differences is known, so it gives about two participants too few; adding half of the squared critical z brings it within a fraction of a participant of the exact t based answer (Guenther, 1981), and the rounded up numbers agree with R in every row. The power formula is the same equation solved for the power. **Common mistake:** rounding the sample size to the nearest whole number instead of UP. With 33.37 participants needed, 33 participants give less than 80 percent power.

**Read:** 34 participants give 80 percent power for one paired test at d_z = 0.50, and 45 at the strictest step of a Holm family of three, the numbers Workshop 8 reported. The secondary paired comparisons of the follow up study are planned as such a family, which is one reason the plan uses 48.

---

## 2. Why more ads beat more participants (SampleSize, part B)

The key planned test of the follow up study asks whether ALTAVO advertising is liked more than BRENLOW advertising IN GENERAL, not only for the twelve ads shown. Then the ads are a sample, and the brand difference must be judged against how much ads of the same brand differ. In a complete, balanced design (every participant rates every ad) the test has 2 x ads per brand - 2 degrees of freedom, and the standard error of the brand difference is

SE = SQRT(2 x (ad SD^2 + residual SD^2 / participants) / ads per brand).

More participants shrink only the second term; only more ads shrink the first. Type the brand difference (0.50), the ad SD (0.20, a planning judgment), and the residual SD from Workshop 7 in `SampleSize!B21:B23`, and the participants and ads per brand in columns A and B of rows 26 to 33. The formulas of row 26:

| Result | Cell | Formula |
|---|---|---|
| df | `SampleSize!{cl('SampleSize', 'bdf_0')}` | `{fx('SampleSize', 'bdf_0')}` |
| SE of the brand difference | `SampleSize!{cl('SampleSize', 'bse_0')}` | `{fx('SampleSize', 'bse_0')}` |
| Noncentrality (difference / SE) | `SampleSize!{cl('SampleSize', 'bncp_0')}` | `{fx('SampleSize', 'bncp_0')}` |
| Critical t | `SampleSize!{cl('SampleSize', 'btc_0')}` | `{fx('SampleSize', 'btc_0')}` |
| Approximate power | `SampleSize!{cl('SampleSize', 'bpow_0')}` | `{fx('SampleSize', 'bpow_0')}` |

**You should see** (R: the exact formula, column `power_formula` of `SIMULATED_power_grid.csv`; the simulated power of `plan_output.txt` agrees with it within Monte Carlo error):

| Participants | Ads per brand | df | SE | Noncentrality | Critical t | Power (Excel) | R (exact) |
|---|---|---|---|---|---|---|---|
{chr(10).join(bt_rows)}

Excel has no noncentral t distribution, so the power uses a normal approximation to it (Abramowitz & Stegun, 1964, chapter 26); the `Check` sheet accepts a difference of up to 0.02 from R's exact value. **Read:** with 3 ads per brand, going from 48 to 72 participants hardly moves the power; going from 3 to 6 ads per brand at 48 participants roughly doubles it. That is the reason the plan asks the client for six ads per brand.

---

## 3. Attrition, reserve, and the cap (Attrition sheet)

Type the inputs in `Attrition!B4:B12` (yellow): the target of {R_MAIN} analyzable participants, the Workshop 7 exclusions ({WS7[1] - WS7[0]} of {WS7[1]}), the 12 rows of the Williams design, the cap multiple 1.25, the Workshop 6 funnel ({WS6[1]} opened the link, {WS6[0]} completed, {WS6[4]} kept), and 5 slots per day.

| Result | Cell | Formula | Value | R |
|---|---|---|---|---|
{att_rows}

**Read:** in the lab, 53 sessions are expected to end with 48 usable participants, so 5 reserve sessions are booked; no shows are rebooked into the 2 open slots of the last day; data collection never runs beyond 60 sessions. The same 48 participants in a remote study would need about 88 people to open the link (Workshop 6), so moving this study online would need a new plan. **Common mistake:** multiplying the target by 1 plus the exclusion rate (48 x 1.083 = 52.0, so 52 sessions) instead of dividing it by the share that is kept (48 / 0.917 = 52.4, so 53). Of 52 sessions, 8.3 percent excluded leaves only 47.7 participants on average.

---

## 4. The Williams design (LatinSquare sheet)

A Latin square shows every ad once in every row and once in every position. A Williams design (Williams, 1949) is a Latin square in which, in addition, every ad is IMMEDIATELY followed by every other ad exactly once, so whatever one ad does to the next is balanced as well. For an even number of ads, one square is enough.

1. Type the positions 1 to 12 in `LatinSquare!C5:N5`. Row 1 of the square, in symbols 0 to 11, is 0, 1, 11, 2, 10, 3, ... (`LatinSquare!{cl('LatinSquare', 'first_C')}`, filled right to `N6`):

| Result | Cell | Formula |
|---|---|---|
| Row 1, position 1 | `LatinSquare!{cl('LatinSquare', 'first_C')}` | `{fx('LatinSquare', 'first_C')}` |
| Row r: add r - 1 and wrap at 12 | `LatinSquare!{cl('LatinSquare', 'sym_C8')}` | `{fx('LatinSquare', 'sym_C8')}` |
| Ad label (even symbol ALTAVO, odd BRENLOW) | `LatinSquare!{cl('LatinSquare', 'lab_C22')}` | `{fx('LatinSquare', 'lab_C22')}` |

   **You should see** row 1 = {first_row_vals}. Fill the symbol formula over `C8:N19` (row numbers 1 to 12 in `B8:B19`) and the label formula over `C22:N33`.
2. Why the labels go by even and odd symbols: which symbols belong to which brand does not change the balance of the ADS, but it decides the order of the BRANDS. `plan_power.R` tried all 462 ways to split the 12 symbols into two brands; only the even and odd split keeps every row to at most two ads of one brand in a row. Column `O` shows the brand sequence: row 1 is `{seq1}` and row 2 is `{seq2}`.
3. The checks (all must hold):

| Check | Cell | Formula | Value | Must be |
|---|---|---|---|---|
| Each ad once per row (smallest count) | `LatinSquare!{cl('LatinSquare', 'row_min')}` | `{fx('LatinSquare', 'row_min')}` | {val('LatinSquare', 'row_min', 0)} | 1 |
| Each ad once per row (largest count) | `LatinSquare!{cl('LatinSquare', 'row_max')}` | `{fx('LatinSquare', 'row_max')}` | {val('LatinSquare', 'row_max', 0)} | 1 |
| Each ad once per position (smallest) | `LatinSquare!{cl('LatinSquare', 'pos_min')}` | `{fx('LatinSquare', 'pos_min')}` | {val('LatinSquare', 'pos_min', 0)} | 1 |
| Each ad once per position (largest) | `LatinSquare!{cl('LatinSquare', 'pos_max')}` | `{fx('LatinSquare', 'pos_max')}` | {val('LatinSquare', 'pos_max', 0)} | 1 |
| Ordered pair count, ad i then ad j | `LatinSquare!{cl('LatinSquare', 'pr_C72')}` | `{fx('LatinSquare', 'pr_C72')}` | (a 12 x 12 table) | 0 on the diagonal, 1 elsewhere |
| Sum of all pair counts | `LatinSquare!{cl('LatinSquare', 'pair_sum')}` | `{fx('LatinSquare', 'pair_sum')}` | {val('LatinSquare', 'pair_sum', 0)} | 132 |
| Diagonal (an ad after itself) | `LatinSquare!{cl('LatinSquare', 'pair_diag')}` | `=C72+D73+...+N83` | {val('LatinSquare', 'pair_diag', 0)} | 0 |
| Largest pair count | `LatinSquare!{cl('LatinSquare', 'pair_max')}` | `{fx('LatinSquare', 'pair_max')}` | {val('LatinSquare', 'pair_max', 0)} | 1 |
| ALTAVO ads in position 1 | `LatinSquare!{cl('LatinSquare', 'alt_C')}` | `{fx('LatinSquare', 'alt_C')}` | {val('LatinSquare', 'alt_C', 0)} | 6 |
| Labels that differ from the R square | `LatinSquare!{cl('LatinSquare', 'sq_mismatch')}` | `{fx('LatinSquare', 'sq_mismatch')}` | {val('LatinSquare', 'sq_mismatch', 0)} | 0 |

   With 132 adjacent pairs in all, none on the diagonal, and no count above 1, each of the 132 ordered pairs of different ads occurs exactly once. **Common mistake:** a plain cyclic Latin square (1 2 3 ... 12, then 2 3 ... 12 1) is balanced for position but not for carryover: ad 2 then always follows ad 1.

---

## 5. Who gets which row (Assignment sheet)

`plan_power.R` assigned the 48 main participants in 4 blocks of 12: each block uses every row once, in a random order drawn with a fixed seed. If data collection stopped after any complete block, every row would still be used equally often. Excel cannot reproduce R's random numbers, so the codes, blocks, and rows are typed (yellow) in `Assignment!A4:C{AL}`. To draw your own order in Excel: list rows 1 to 12 for a block, put `=RAND()` next to each, copy and Paste Special > Values, and sort by that column (the values must be pasted, because RAND() changes on every recalculation).

| Result | Cell | Formula | Value | R |
|---|---|---|---|---|
| First ad of P01 | `Assignment!{cl('Assignment', 'first_4')}` | `{fx('Assignment', 'first_4')}` | {val('Assignment', 'first_4')} | |
{row4('Every row once in every block (smallest count)', 'Assignment', 'blk_min', 0)}
{row4('Every row once in every block (largest count)', 'Assignment', 'blk_max', 0)}
{row4('Participants who see an ALTAVO ad first', 'Assignment', 'first_alt', 0)}

Reserve participants (R01 to R{R_RES:02d}) are not in this table: each one takes the row of the participant it replaces, so that the rows stay balanced among the analyzed participants.

---

## 6. The session calendar (Calendar sheet)

The lab has one Smart Eye Aurora, so sessions run one at a time. A slot is 90 minutes: 15 for arrival and consent, 10 for the sensors and calibration, 20 for the recording, 15 for removing the sensors and the debriefing, and 30 for the backup, the session log, and cleaning (the lab procedure manual's times; see STEP 6 of `plan_output.txt`). Type the first day (a placeholder) in `Calendar!B3`, the first slot (09:00) in `B5`, and the slot length (90) in `B6`; `B4`, `B7`, and `B8` come from the Attrition sheet. Type k = 1 to {len(SCH_MAIN)} in column A and fill the formulas of row 12 down:

| Result | Cell | Formula |
|---|---|---|
| Day index (0 = first day) | `Calendar!{cl('Calendar', 'cal_B')}` | `{fx('Calendar', 'cal_B')}` |
| Date (weekdays only) | `Calendar!{cl('Calendar', 'cal_C')}` | `{fx('Calendar', 'cal_C')}` |
| Weekday | `Calendar!{cl('Calendar', 'cal_D')}` | `{fx('Calendar', 'cal_D')}` |
| Slot of the day (1 to 5) | `Calendar!{cl('Calendar', 'cal_E')}` | `{fx('Calendar', 'cal_E')}` |
| Arrival time | `Calendar!{cl('Calendar', 'cal_F')}` | `{fx('Calendar', 'cal_F')}` |
| Recording starts (25 min later) | `Calendar!{cl('Calendar', 'cal_G')}` | `{fx('Calendar', 'cal_G')}` |
| Room ready for the next participant | `Calendar!{cl('Calendar', 'cal_J')}` | `{fx('Calendar', 'cal_J')}` |
| Phase | `Calendar!{cl('Calendar', 'cal_K')}` | `{fx('Calendar', 'cal_K')}` |
| Participant code | `Calendar!{cl('Calendar', 'cal_L')}` | `{fx('Calendar', 'cal_L')}` |
| Latin square row | `Calendar!{cl('Calendar', 'cal_M')}` | `{fx('Calendar', 'cal_M')}` |
| Half day shift (morning = slots 1 to 3) | `Calendar!{cl('Calendar', 'cal_N')}` | `{fx('Calendar', 'cal_N')}` |
| Staff pair (1, 2, 3 in turn) | `Calendar!{cl('Calendar', 'cal_O')}` | `{fx('Calendar', 'cal_O')}` |
| Lead operator | `Calendar!{cl('Calendar', 'cal_P')}` | `{fx('Calendar', 'cal_P')}` |
| Assistant | `Calendar!{cl('Calendar', 'cal_Q')}` | `{fx('Calendar', 'cal_Q')}` |

Times are fractions of a day in Excel, so 25 minutes is `25/1440`. Format the date column as `yyyy-mm-dd` and the time columns as `hh:mm`. `WORKDAY` skips Saturdays and Sundays only; university closures must be added as its third argument (a range of holiday dates) once the JSU academic calendar for the term is known. **Common mistake:** formatting the time cells as numbers, which shows 0.375 instead of 09:00.

The columns `T` to `Y` (gray) hold R's schedule, and these cells count the differences:

| Check | Cell | Formula | Value | Must be |
|---|---|---|---|---|
{chr(10).join(f"| {REG[('Calendar', k)]['label']} | `Calendar!{cl('Calendar', k)}` | `{fx('Calendar', k)}` | {val('Calendar', k, 0)} | 0 |" for k, _, _ in mis)}
| Last day of data collection | `Calendar!{cl('Calendar', 'last_day')}` | `{fx('Calendar', 'last_day')}` | {SCH_MAIN[-1]['date']} | {SCH_MAIN[-1]['date']} |

---

## 7. Check your numbers (Check sheet)

Every result with an R value sits on the `Check` sheet next to that value, with the difference and a tolerance: half a unit in the last printed decimal for exact results (0 for counts, sample sizes, dates, and codes), 0.25 participants for the corrected sample sizes, 0.005 for the paired power, and 0.02 for the approximate power of the brand test. {check_line}. What Excel cannot do is listed at the bottom of the sheet: the mixed model simulations (STEP 2 and STEP 3), the exact noncentral t power, and R's random draw of the rows.

---

## Sources

- Abramowitz, M., & Stegun, I. A. (Eds.). (1964). *Handbook of mathematical functions with formulas, graphs, and mathematical tables* (National Bureau of Standards Applied Mathematics Series 55). U.S. Government Printing Office.
- Guenther, W. C. (1981). Sample size formulas for normal theory t tests. *The American Statistician, 35*(4), 243-244. https://doi.org/10.1080/00031305.1981.10479363
- Williams, E. J. (1949). Experimental designs balanced for the estimation of residual effects of treatments. *Australian Journal of Scientific Research, Series A, 2*(2), 149-168. https://doi.org/10.1071/CH9490149
"""
with open(MD_OUT, "w", encoding="utf-8") as fh:
    fh.write(md)
print("wrote", MD_OUT)
